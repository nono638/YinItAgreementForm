"""The run sheet as a refresher of a trial (the user, 2026-10-26: "to jog my memory of what happened, who worked on
what day, how many pages each person did"): By Reporter with each reporter's share of the pages and the days they
worked; Witnesses, every running head the takes printed, as printed, kept from day to day; Who Ordered, the case's
invoices and the firms with an agreement but no invoice yet, from the records, written again once Generate made the
invoices. Fictional names and cases."""
from openpyxl import load_workbook

from minute_filler.batch import fill_jobs, group, read_docs
from minute_filler.deliver import ledger_for
from minute_filler.firm_marks import SHADES
from minute_filler.records import Invoice
from minute_filler.runsheet import (BY_REPORTER, NO_ORDERS, ORDERS, WITNESSES, add_takes, refresh_orders,
                                    witness_rows)

from test_runsheet import TALK, TITLE_1, TITLE_2, job_of, s, transcript  # noqa: F401 (s: the fixture)

DAY_2 = [(320, "ds", "", TITLE_1), (321, "ds", "", TITLE_2), (322, "ds", "K. Park - Defendant - Direct", TALK),
         (323, "pr", "K. Park - Defendant - Cross", TALK)]


def two_days(tmp_path, s):
    """Day 1 (Pat and Dana: J. Doe direct and cross, M. Lee direct) and day 2 (K. Park) of Jane Roe's trial, each
    added to its run sheet by its own Generate; the run sheet's path."""
    one = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    a = one.runsheet_opts(s)
    add_takes(one.case, a, s, [])
    two = job_of(transcript(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", DAY_2, day="June 4, 2026"), s)
    b = two.runsheet_opts(s)
    add_takes(two.case, b, s, [])
    assert b.path == a.path and not b.created
    return a.path


def tab(path, name):
    """The rows of one tab of the run sheet, as values."""
    ws = load_workbook(path)[name]
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_by_reporter_has_each_reporters_share_and_days(tmp_path, s):
    """Pat wrote on both days, Dana on both: each reporter's days worked as dates, how many, and a share of the
    pages (a formula over the Pages column, so a row changed in Excel counts); the trial's days at the foot."""
    path = two_days(tmp_path, s)
    rows = tab(path, BY_REPORTER)
    assert rows[0] == ["Reporter", "Takes", "Pages", "Share", "Days", "Days worked"]
    days = {r[0]: (r[4], r[5]) for r in rows[1:-1]}
    assert set(days.values()) == {(2, "6/3/2026, 6/4/2026")}
    assert all(str(r[3]).startswith("=IF(N(C$") and f"C{i}/C$" in r[3] for i, r in enumerate(rows[1:-1], 2))
    assert rows[-1][0] == "Total" and rows[-1][4] == 2
    assert load_workbook(path)[BY_REPORTER]["D2"].number_format == "0%"


def test_witnesses_lists_every_head_as_printed_and_keeps_an_earlier_days(tmp_path, s):
    """Every running head the takes printed, as printed, once each, with its days and reporters: also those inside
    a take (J. Doe's cross, between his direct and 'Proceedings': in no Witness cell), and day 1's still there after
    day 2 rewrote the run sheet (the heads are kept on the hidden sheet)."""
    path = two_days(tmp_path, s)
    rows = tab(path, WITNESSES)
    assert rows[0][0] == "Running head (as printed)"
    got = {r[0]: (r[1], r[2]) for r in rows[1:]}
    assert got["J. Doe - Plaintiff - Cross"] == ("6/3/2026", "Pat")
    assert got["Proceedings"][0] == "6/3/2026" and set(got["Proceedings"][1].split(", ")) == {"Dana", "Pat"}
    assert got["K. Park - Defendant - Direct"] == ("6/4/2026", "Dana")
    assert list(got)[:2] == ["Proceedings", "J. Doe - Plaintiff - Direct"]  # (in the order of the takes)
    witness_cells = {c for r in tab(path, "Run Sheet")[4:] for c in r[8:10] if c}
    assert "J. Doe - Plaintiff - Cross" not in witness_cells  # (only the Witnesses tab has it)


def test_witnesses_are_listed_when_the_witness_cells_are_left_blank(tmp_path, s):
    """Settings.runsheet_witness "blank" leaves Witness Start and End for the user; the Witnesses tab still lists
    the heads read from the pages."""
    s.runsheet_witness = "blank"
    path = two_days(tmp_path, s)
    assert not {c for r in tab(path, "Run Sheet")[4:] for c in r[8:10] if c}
    assert "M. Lee - Defendant - Direct" in [r[0] for r in tab(path, WITNESSES)]


def test_a_take_without_kept_heads_has_its_witness_cells():
    """A row read back from a run sheet made before the heads were kept (or typed in) gives its Witness Start and
    End, each once."""
    from minute_filler.runsheet import Row
    from datetime import date
    rows = [Row(date(2026, 6, 3), "Pat", 4, 1, "A - Direct", "A - Direct"), Row(date(2026, 6, 4), "Dana", 3, 5,
                                                                               "A - Direct", "B - Cross")]
    assert witness_rows(rows) == [("A - Direct", [date(2026, 6, 3), date(2026, 6, 4)], ["Pat", "Dana"], 7),
                                  ("B - Cross", [date(2026, 6, 4)], ["Dana"], 3)]


def invoice(no, created, **kw):
    """A made-up invoice of Jane Roe's case."""
    return Invoice(no, created, "Jane Roe v. Sam Poe", index_no="712345/2021", **kw)


def test_who_ordered_lists_the_cases_invoices_and_agreements_from_the_records(tmp_path, s):
    """From the records: Advocate Law Group's invoice (the plaintiffs' first firm: its side and colour), oldest
    first, paid; Counsel & Ward with an agreement and no invoice yet (its days once each); another case's invoice
    left out. Written again (refresh_orders) once a new invoice is made, and not again when nothing changed."""
    lg = ledger_for(s)
    lg.add_invoice(invoice("2026-0001", "2026-06-05", firm="Advocate Law Group", bill_to="Samuel R. Advocate",
                           dates="6/3/2026", amounts={"Expedite": "98.00"}, billed_speed="Expedite", status="paid",
                           paid_date="2026-06-20", amount_paid="98.00"))
    lg.add_invoice(Invoice("2026-0002", "2026-06-05", "Lee v. Park", index_no="700001/2025", firm="Smith Law",
                           bill_to="Dana Smith", amounts={"Regular": "43.00"}, billed_speed="Regular"))
    for day in ("6/3/2026", "6/3/2026"):
        lg.log_activity("agreement", case_name="Jane Roe v. Sam Poe", index_no="712345/2021", firm="Counsel & Ward, LLP",
                        attorney="Alex B. Counsel", dates=day)
    path = two_days(tmp_path, s)
    rows = tab(path, ORDERS)
    assert rows[0][1:] == ["Firm", "Side", "Attorney", "Days", "Speed", "Invoice No.", "Invoice date", "Amount",
                           "Status"]
    assert rows[1][1:] == ["Advocate Law Group", "Plaintiff", "Samuel R. Advocate", "6/3/2026", "Expedite", "2026-0001",
                           "6/5/2026", "$98.00", "Paid 6/20/2026"]
    assert load_workbook(path)[ORDERS]["A2"].fill.start_color.rgb[-6:].lower() == SHADES["plaintiff"][0][1:]
    assert rows[2][1:] == ["Counsel & Ward, LLP", "Defense", "Alex B. Counsel", "6/3/2026", None, None, None, None,
                           "Agreement, no invoice yet"]
    assert len(rows) == 3  # (Lee v. Park's invoice is another case's)
    case = job_of(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", s).case
    assert refresh_orders(path, case, s) is False  # (the same: not saved again)
    lg.add_invoice(invoice("2026-0003", "2026-06-06", firm="Counsel & Ward, LLP", bill_to="Alex B. Counsel",
                           dates="6/3/2026, 6/4/2026", amounts={"Regular": "60.00", "Expedite": "90.00"}))
    assert refresh_orders(path, case, s) is True
    rows = tab(path, ORDERS)
    assert [r[6] for r in rows[1:]] == ["2026-0001", "2026-0003"]  # (Counsel & Ward billed now: no agreement row)
    assert rows[2][5:] == ["offered: Regular, Expedite", "2026-0003", "6/6/2026", "Regular $60.00; Expedite $90.00",
                           "Open"]
    assert tab(path, BY_REPORTER) and tab(path, WITNESSES)  # (the other tabs as they were)


def test_who_ordered_says_so_when_the_records_have_nothing(tmp_path, s):
    path = two_days(tmp_path, s)
    assert tab(path, ORDERS)[1][1] == NO_ORDERS
    case = job_of(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", s).case
    assert refresh_orders(path, case, s) is False


def test_a_run_sheet_open_in_excel_doesnt_stop_the_refresh(tmp_path, s, monkeypatch):
    """refresh_orders is an aside, after the invoices are made: a run sheet that can't be saved is logged and left
    as it was."""
    import minute_filler.runsheet as runsheet
    path = two_days(tmp_path, s)
    before = path.read_bytes()
    ledger_for(s).add_invoice(invoice("2026-0001", "2026-06-05", firm="Advocate Law Group", bill_to="Samuel R. Advocate",
                                      amounts={"Regular": "60.00"}, billed_speed="Regular"))

    def held(wb, p):
        raise PermissionError(f"{p.name} is open in another program")
    monkeypatch.setattr(runsheet, "_save", held)
    case = job_of(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", s).case
    assert refresh_orders(path, case, s) is False and path.read_bytes() == before


def test_generate_puts_the_invoices_it_just_made_on_who_ordered(tmp_path, s):
    """Generate all with the run sheet and the invoice: the run sheet is written first, then the invoice; its Who
    Ordered tab names the invoice all the same."""
    docs, _ = read_docs([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"))], s)
    jobs = group(docs, s)
    for a in jobs[0].case.attorneys:
        a.checked = "Advocate" in (a.firm or a.name)
    made = fill_jobs(jobs, s, outputs=["runsheet", "invoice"])
    sheet = next(p for p in made if p.suffix == ".xlsx")
    number = ledger_for(s).invoices()[0].invoice_no
    assert [r[6] for r in tab(sheet, ORDERS)[1:]] == [number]
