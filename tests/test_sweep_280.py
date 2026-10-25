"""The 2.8.0 sweep: a click on a text column's heading in Records sorts it (it used to crash the app); Add in Settings
-> Dates; the attorney details under the table vs a merge, another job shown, the status pill, rows added and
removed; the Help window closes with Settings and comes back when minimized; folded cards follow the speed; New
job's Undo vs a file dropped meanwhile, another rate sheet and the day; a new tick box's words follow its name; a
double-click on a tick box only ticks; the Records count after a tick, and when the records can't be read; Import
settings keeps this computer's own; the CSV copies get a box's column at once; Records opens on the year an invoice
made now gets; the Help results keep the one chosen in sight, and a result's line reads as the help does; text
dragged from the window's own boxes; the delivery date's note when the rate sheet gives the days; the turnaround
from the day you make it, across midnight and from a plain date; a delivery date emptied stays empty; a form at
another speed with the delivery date off; years 1 and 9999; a holiday weekend rule that can't be read; a damaged
word table; text starting with "=" in the workbook; a remembered update page that isn't the project's. All names
are made up."""
import os
import time
from datetime import date, datetime, time as dtime, timedelta

import pytest

from helpers import pat_settings, transcript_pdf
from minute_filler import workdays
from minute_filler.dates import us_date
from minute_filler.models import SRC_DEFAULT, SRC_USER, CaseInfo, FieldState
from minute_filler.records import Invoice, Ledger

CASE = "Jane Roe v. X.Y. Holding Corporation"


def invoice(no, firm="Smith Law", case=CASE, created="2026-06-04", **kw):
    return Invoice(invoice_no=no, created=created, case_name=case, index_no="712345/2021", dates="6/3/2026",
                   firm=firm, pages=100, amounts={"Regular": "100.00"}, billed_speed="Regular", **kw)


def pump(seconds=0.2):
    from PySide6.QtWidgets import QApplication
    end = time.time() + seconds
    while time.time() < end:
        QApplication.processEvents()
        time.sleep(0.01)


def wait(win, until, seconds=20):
    from PySide6.QtWidgets import QApplication
    end = time.time() + seconds
    while time.time() < end:
        QApplication.processEvents()
        if not win.runner._live and until():
            return
        time.sleep(0.01)
    raise AssertionError("timed out")


@pytest.fixture
def window(tmp_path, monkeypatch, make_window, qt):
    monkeypatch.setattr(qt.QMessageBox, "exec", lambda self: 0)
    s = pat_settings(initials="pr")
    s.use_ai = s.open_after = False
    s.welcomed = True
    s.output_dir = str(tmp_path / "out")
    return make_window(s)


def on_screen(win):
    """Shown and active: the details boxes under the attorney table can have the focus."""
    win.resize(1400, 900)
    win.show()
    win.activateWindow()
    pump()


def read_transcript(win, tmp_path, name="Roe.pdf", **kw):
    win.add_files([str(transcript_pdf(tmp_path / name, **kw))])
    wait(win, lambda: bool(win.cur.docs))


@pytest.fixture
def records(qt, tmp_path):
    """The Records window over four firms' invoices of one transcript, every year shown."""
    from minute_filler.deliver import ledger_for
    from minute_filler.gui.records_window import RecordsWindow
    from minute_filler.tracking import box_rows, default_boxes
    s = pat_settings()
    s.records_dir = str(tmp_path / "records")
    boxes = default_boxes()
    for b in boxes:
        b.on = b.key in ("paid", "transcribed", "sent")
    s.tracking_boxes = box_rows(boxes)
    lg = ledger_for(s)
    for n, (firm, case) in enumerate((("Smith Law", CASE), ("Counsel & Counsel", "jane Roe v. Sam Poe"),
                                      ("Roe & Co.", CASE), ("Doe PC", CASE)), 1):
        lg.add_invoice(invoice(f"2026-000{n}", firm, case))
    win = RecordsWindow(s)
    win.reload()
    win.i_year.setCurrentIndex(0)
    yield win, s, lg
    win.close()
    win.deleteLater()


def cell(win, no, key):
    from minute_filler.gui.records_window import col_index
    t = win.inv_table
    row = next(r for r in range(t.rowCount()) if t.item(r, 0).data(256) == no)
    return t.item(row, col_index(win.inv_cols, key))


# ------------------------------------------------------------------ Records
def test_a_click_on_a_text_columns_heading_sorts_it_whatever_the_capitals(records, qt):
    """2.7.0 crashed the app: the text columns' sort called back into itself until the stack ran out."""
    from PySide6.QtCore import Qt
    from minute_filler.gui.records_window import col_index
    win, _, _ = records
    t = win.inv_table
    case = col_index(win.inv_cols, "case")
    t.sortItems(case, Qt.AscendingOrder)
    assert [t.item(r, case).text() for r in range(t.rowCount())] == ["jane Roe v. Sam Poe"] + [CASE] * 3
    t.sortItems(case, Qt.DescendingOrder)
    assert t.item(t.rowCount() - 1, case).text() == "jane Roe v. Sam Poe"


def test_a_double_click_on_a_tick_box_doesnt_open_the_pdf(records, monkeypatch):
    win, _, _ = records
    opened = []
    monkeypatch.setattr(win, "_open_invoice", opened.append)
    win._double_clicked(cell(win, "2026-0001", "box:sent"))
    win._double_clicked(cell(win, "2026-0001", "paid"))
    assert opened == []
    win._double_clicked(cell(win, "2026-0001", "case"))
    assert len(opened) == 1


@pytest.mark.workdays
def test_records_show_an_invoice_dated_into_the_new_year(qt, tmp_path, monkeypatch):
    """Made on Dec 31 after the evening hour, an invoice is dated January and numbered in the new year: Records
    goes to that year once there is one (open or not), and keeps a year the user chose."""
    from minute_filler.deliver import ledger_for
    from minute_filler.gui.records_window import RecordsWindow
    year = date.today().year
    monkeypatch.setattr(workdays, "now", lambda: datetime(year, 12, 31, 18, 0))
    s = pat_settings()
    s.records_dir = str(tmp_path / "records")
    s.after_hours = True
    day = workdays.document_day(s).day
    assert day.year == year + 1
    lg = ledger_for(s)
    lg.add_invoice(invoice(f"{year}-0001", created=f"{year}-12-30"))
    win = RecordsWindow(s)
    try:
        win.reload()
        assert win.i_year.currentData() == year  # (nothing in the new year yet: the year's invoices)
        lg.add_invoice(invoice(f"{year + 1}-0001", created=day.isoformat()))
        win.reload()  # (Generate made it with Records open)
        assert win.i_year.currentData() == year + 1 and win.inv_table.rowCount() == 1
        assert win.a_year.currentData() == year  # (what was done, by the clock)
        win.i_year.setCurrentIndex(win.i_year.findData(0))  # (the user's choice: All years)
        win.reload()
        assert win.i_year.currentData() == 0 and win.inv_table.rowCount() == 2
    finally:
        win.close()
        win.deleteLater()


# ------------------------------------------------------------------ Settings' editors
def test_add_in_the_court_holidays_raises_nothing(qt, monkeypatch):
    """Each cell of the new row was a change to the table, which looked for the row's Next before it had one: the
    "Something went wrong" box twice."""
    import sys
    import traceback
    from minute_filler.courthouses import holiday_rows
    from minute_filler.gui.holiday_editor import NEXT, HolidayEditor
    seen = []

    def hook(t, e, tb):  # (only the editor's own: a window of an earlier test may still get a late answer)
        if any("holiday_editor" in f.filename for f in traceback.extract_tb(tb)):
            seen.append(f"{t.__name__}: {e}")

    monkeypatch.setattr(sys, "excepthook", hook)
    ed = HolidayEditor(holiday_rows())
    n = ed.table.rowCount()
    ed._add()
    qt.QApplication.processEvents()
    assert seen == [] and ed.table.rowCount() == n + 1 and ed.table.item(n, NEXT) is not None
    ed.table.item(n, 0).setText("Closing")
    ed.table.item(n, 1).setText("12/24/2026")
    assert seen == [] and ed.table.item(n, NEXT).text() != ""


def test_a_new_tick_boxs_words_follow_its_name(qt):
    from minute_filler.gui import tracking_editor as te
    from minute_filler.tracking import default_boxes, default_rules
    ed = te.TrackingEditor(default_boxes(), default_rules(), True)
    ed.box_table.setCurrentCell(3, te.NAME)
    ed._add_box()
    ed.box_table.item(4, te.NAME).setText("Proofread")
    b = ed.boxes()[4]
    assert (b.name, b.yes, b.no) == ("Proofread", "Proofread", "Not proofread")
    assert ed.box_table.item(4, te.NO).text() == "Not proofread"
    ed.box_table.item(4, te.NO).setText("To proofread")  # (words of the user's own stay theirs)
    ed.box_table.item(4, te.NAME).setText("Checked")
    b = ed.boxes()[4]
    assert (b.name, b.yes, b.no) == ("Checked", "Proofread", "To proofread")


# ------------------------------------------------------------------ the attorney details under the table
EMAIL = ("From: Dana Smith <dsmith@smithlaw.example>\nTo: Pat Reporter <preporter@example.com>\n"
         "Subject: minutes\n\nPlease send the minutes for Jane Roe v. X.Y. Holding Corporation, Index 712345/2021, "
         "of 6/3/2026, expedited.\nDana Smith, Smith Law, 555-0100\n")


def test_an_address_the_ai_finds_while_its_empty_box_has_the_focus_is_kept(window):
    """Left without typing, the box wrote its own (empty) text back over the AI's address."""
    from minute_filler.models import SRC_AI, Attorney, Extraction
    on_screen(window)
    window.add_text(EMAIL)
    wait(window, lambda: bool(window.cur.docs))
    window.att.setCurrentCell(0, 1)
    box = window.att_edits[3]
    box.setFocus()
    pump()
    assert box.hasFocus()
    window.cur.docs[0].ai = Extraction(attorneys=[Attorney(
        name="Dana Smith", firm="Smith Law", email="dsmith@smithlaw.example",
        address="1 Main Street\nKew Gardens, NY 11415", source=SRC_AI)])
    window._remerge(window.cur)
    window.rows["part"].edit.setFocus()  # (the box left)
    pump()
    assert window.case.attorneys[0].address == "1 Main Street\nKew Gardens, NY 11415"


def test_a_detail_box_only_clicked_into_keeps_what_the_ai_found(window):
    """The box with the focus kept its old (empty) text through the AI's answer, and the status pill, Ctrl+Enter
    or New job wrote it over the AI's address and fax. A value picked from a box's list counts as typed."""
    from minute_filler.models import SRC_AI, Attorney, Extraction
    on_screen(window)
    window.add_text(EMAIL)
    wait(window, lambda: bool(window.cur.docs))
    window.att.setCurrentCell(0, 1)
    for col in (3, 5):  # (the address, then the fax)
        box = window.att_edits[col]
        box.setFocus()
        pump()
        window.cur.docs[0].ai = Extraction(attorneys=[Attorney(
            name="Dana Smith", firm="Smith Law", email="dsmith@smithlaw.example",
            address="1 Main Street\nKew Gardens, NY 11415", fax="555-0109", source=SRC_AI)])
        window._remerge(window.cur)
        window._commit_att_details()  # (what a click on the status, Ctrl+Enter and New job do first)
        a = window.case.attorneys[0]
        assert (a.address, a.fax) == ("1 Main Street\nKew Gardens, NY 11415", "555-0109")
    firm = window.att_edits[2]
    firm.setFocus()
    pump()
    firm.setText("Smith Law PLLC")  # (as picking it from the box's list of past values does)
    window._commit_att_details()
    assert window.case.attorneys[0].firm == "Smith Law PLLC"


def test_a_detail_box_keeps_nothing_of_another_job(window, tmp_path):
    """Another job shown under a box with the focus: the box showed (and then wrote) the first job's firm. (Day by
    day: in a trial, a firm's details typed on one day are every day's, as one row is the firm's for the trial.)"""
    from PySide6.QtTest import QTest
    window.s.trial_orders = "day"
    on_screen(window)
    a = transcript_pdf(tmp_path / "Roe 6-3.pdf")
    b = transcript_pdf(tmp_path / "Roe 6-4.pdf", date="June 4, 2026")
    window.add_files([str(a), str(b)], batch=True)
    wait(window, lambda: len(window.jobs) == 2)
    first, second = window.jobs
    window._select_job(first)
    pump()
    window.att.setCurrentCell(1, 1)
    firm = window.att_edits[2]
    firm.setFocus()
    pump()
    before = second.case.attorneys[1].firm
    firm.selectAll()
    QTest.keyClicks(firm, "Counsel Firm of the first job")
    window._sync_from_ui()
    window.cur = second  # (as a read finishing does: another job shown, the focus where it was)
    window._show_job()
    assert firm.text() == before
    window.rows["part"].edit.setFocus()
    pump()
    assert second.case.attorneys[1].firm == before


def test_a_click_on_the_status_keeps_a_detail_being_typed(window, tmp_path):
    from PySide6.QtTest import QTest
    on_screen(window)
    read_transcript(window, tmp_path)
    assert window.cur.no_attorney_chosen()
    window.att.setCurrentCell(1, 1)
    firm = window.att_edits[2]
    firm.setFocus()
    pump()
    firm.selectAll()
    QTest.keyClicks(firm, "Counsel & Counsel LLP")
    window.status.clicked.emit()  # (the pill takes no focus)
    pump()
    assert window.case.attorneys[1].firm == "Counsel & Counsel LLP"


def test_the_details_follow_rows_removed_and_added(window, tmp_path):
    read_transcript(window, tmp_path)
    assert window.att.rowCount() == 2
    window.att.setCurrentCell(0, 1)
    window.att.selectRow(0)
    window._remove_att_rows()  # (the current cell moves to row 2 before row 1 goes: the details showed "row 2")
    pump()
    assert window._att_row == window.att.currentRow() == 0
    phone = window.att_edits[4]
    phone.setText("555-0199")
    phone.editingFinished.emit()
    assert window.case.attorneys[0].phone == "555-0199"
    window._add_att_row()
    assert window.att.currentRow() == window.att.rowCount() - 1 == window._att_row


# ------------------------------------------------------------------ the Help window
def test_the_help_of_settings_closes_with_it(qt, make_window, monkeypatch):
    from PySide6.QtWidgets import QApplication
    from minute_filler.gui.dialogs import SettingsDialog
    from minute_filler.gui.help_window import HelpWindow
    win = make_window(pat_settings())

    def fake_exec(self):
        self.setModal(True)
        self.show()
        QApplication.processEvents()
        self.show_help()  # (F1, or its Help button)
        assert self._help_window.isVisible()
        self.reject()
        QApplication.processEvents()
        return 0

    monkeypatch.setattr(SettingsDialog, "exec", fake_exec)
    win.open_settings()
    pump()
    assert not [w for w in QApplication.topLevelWidgets() if isinstance(w, HelpWindow) and w.isVisible()]


def test_f1_brings_back_a_minimized_help_window(qt, make_window):
    from minute_filler.gui.help_window import open_help
    win = make_window(pat_settings())
    help_win = open_help("start-steps", win)
    help_win.showMinimized()
    pump()
    assert open_help("check-index", win) is help_win
    assert not help_win.isMinimized()
    help_win.close()


def test_the_help_results_keep_the_one_chosen_in_sight(qt, make_window):
    from PySide6.QtCore import QEvent, Qt
    from PySide6.QtGui import QKeyEvent
    from minute_filler.gui.help_window import open_help
    win = make_window(pat_settings())
    help_win = open_help("", win)
    help_win.resize(700, 300)
    pump()
    help_win.search.setText("invoice")
    help_win._search()
    assert len(help_win.hits) > 8
    for _ in range(len(help_win.hits) - 1):
        help_win.eventFilter(help_win.search, QKeyEvent(QEvent.KeyPress, Qt.Key_Down, Qt.NoModifier))
    pump()
    view = help_win.results
    assert view.verticalScrollBar().value() > 0
    assert view.viewport().rect().intersects(view.cursorRect())
    help_win.close()


# ------------------------------------------------------------------ the main window
def test_a_folded_details_card_follows_the_speed(window, tmp_path):
    read_transcript(window, tmp_path)
    window._fold_card(window.cards["details"], True)
    boxes = window.inv_speed_boxes
    cur = window.case.get("delivery")
    other = next(n for n in boxes if n != cur)
    boxes[other].setChecked(True)
    boxes[cur].setChecked(False)  # (the speed shown is no longer offered: another is picked)
    pump()
    assert window.case.get("delivery") != cur
    assert window.case.get("delivery") in window.cards["details"].summary.full()


def test_undo_never_drops_a_file_dropped_after_new_job(window, tmp_path):
    read_transcript(window, tmp_path)
    window.new_job()
    assert window._undo is not None
    window.add_files([str(transcript_pdf(tmp_path / "Roe 6-4.pdf", date="June 4, 2026"))])
    window._undo_clear()  # (clicked before the file's read reported back)
    wait(window, lambda: bool(window.cur.docs))
    assert [d.ing.name for j in window.jobs for d in j.docs] == ["Roe 6-4.pdf"]


def test_undo_is_dropped_when_another_rate_sheet_is_picked(window, tmp_path):
    other = tmp_path / "sheets"
    other.mkdir()
    (other / "Other Rates.csv").write_text("Rate,Original,Copy,Email,Index\nExpedite,$7.70,$2.00,$2.00,$2.00\n"
                                           "Regular,$6.60,$2.00,$2.00,$2.00\n", encoding="utf-8")
    window.s.rate_sheets_dir = str(other)
    window._read_sheets_again()
    read_transcript(window, tmp_path)
    window.new_job()
    assert window._undo is not None
    window.sheet_box.setCurrentIndex(window.sheet_box.findData("Other Rates"))
    assert window._undo is None and window.undo_box.isHidden()


def test_jobs_brought_back_by_undo_are_dated_as_of_now(window, tmp_path, monkeypatch):
    read_transcript(window, tmp_path)
    window._redate()
    window.new_job()
    tomorrow = date.today() + timedelta(days=1)
    monkeypatch.setattr(workdays, "now", lambda: datetime.combine(tomorrow, dtime(10)))
    window._redate()  # (the minute's timer, past midnight, while Undo is offered)
    window._undo_clear()
    assert window.cur.case.fields["agreement_date"].value == us_date(tomorrow)


def test_the_records_count_follows_a_tick_made_just_before_coming_back(window, monkeypatch):
    from PySide6.QtCore import QEvent
    from minute_filler.tracking import NOT_PAID
    led = Ledger()
    led.add_invoice(invoice("2026-0001"))
    window.s.records_count = NOT_PAID
    monkeypatch.setattr(window, "isActiveWindow", lambda: True)
    window.changeEvent(QEvent(QEvent.ActivationChange))
    wait(window, lambda: window.rec_btn.text() == f"Records · {NOT_PAID}: 1")
    led.mark_paid("2026-0001", "Regular", "100.00", "2026-10-12")  # (ticked Paid in the Records window)
    window.changeEvent(QEvent(QEvent.ActivationChange))  # (back within 2 s)
    wait(window, lambda: window.rec_btn.text() == "Records", seconds=5)


def test_the_records_count_stays_when_the_records_cant_be_read(tmp_path, window, monkeypatch):
    from minute_filler import records
    from minute_filler.gui import workers
    from minute_filler.tracking import NOT_PAID
    db = tmp_path / "records.db"
    db.write_bytes(b"this is not a database" * 100)
    with pytest.raises(OSError):
        records.live_invoices(db)
    assert records.live_invoices(tmp_path / "none.db") == []
    window.s.records_count = NOT_PAID
    window._show_count(NOT_PAID, 3)
    logged = []
    monkeypatch.setattr(workers, "log_error", lambda *a: logged.append(a))

    def locked():
        raise OSError("the records could not be read")

    monkeypatch.setattr(records, "live_invoices", locked)
    window._count_records()  # (each time the window comes back, while a restore holds the records)
    wait(window, lambda: True)
    assert window.rec_btn.text() == f"Records · {NOT_PAID}: 3" and logged == []


def test_import_settings_keeps_this_computers_own(make_window, monkeypatch, tmp_path):
    """An update found while "Replace your settings?" waited was forgotten again."""
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    from minute_filler.settings import Settings
    s = pat_settings()
    s.check_updates = True
    win = make_window(s)
    exported = Settings()
    exported.profile.name = "Pat Reporter"
    exported.export_to(tmp_path / "x.json")
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (str(tmp_path / "x.json"), ""))

    def question(*a, **k):
        win._show_update("9.9.0", "https://github.com/nono638/YinItAgreementForm/releases/tag/v9.9.0")
        return QMessageBox.Yes

    monkeypatch.setattr(QMessageBox, "question", question)
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    monkeypatch.setattr(win, "_check_rate_sheet", lambda *a, **k: True)
    win.import_settings()
    assert win.s.update_version == "9.9.0" and Settings.load().update_version == "9.9.0"


def test_the_csv_copies_get_a_box_switched_on_while_records_is_closed(make_window, tmp_path):
    import csv
    from minute_filler.deliver import ledger_for
    s = pat_settings()
    s.records_dir = str(tmp_path / "rec")
    led = ledger_for(s)
    led.add_invoice(invoice("2026-0001"))
    win = make_window(s)

    def heads():
        with open(s.records_folder() / "invoices.csv", encoding="utf-8-sig") as f:
            return next(csv.reader(f))

    assert "Transcribed" not in heads()
    s.tracking_boxes = [dict(b, on=True) if b["key"] == "transcribed" else b for b in s.tracking_boxes]
    win._settings_changed()  # (Settings -> Records saved)
    wait(win, lambda: "Transcribed" in heads(), seconds=10)


def test_a_remembered_update_page_must_be_the_projects(make_window, monkeypatch):
    from minute_filler import update
    from minute_filler.gui import main_window
    opened = []
    monkeypatch.setattr(main_window, "open_url", opened.append)
    s = pat_settings()
    s.check_updates = True
    s.update_version, s.update_page = "9.9.0", "file:///C:/Windows/System32/calc.exe"  # (settings.json edited)
    win = make_window(s)
    win._show_known_update()
    win.update_btn.click()
    assert opened == [update.RELEASES_PAGE + "/latest"]


def test_text_dragged_from_the_windows_own_boxes_is_not_taken_as_new(window):
    from PySide6.QtCore import QMimeData, QPoint, Qt
    from PySide6.QtGui import QDragEnterEvent

    class FromABox(QDragEnterEvent):
        def source(self):
            return window.paste

    md = QMimeData()
    md.setText("Index 712345/2021")
    ev = FromABox(QPoint(5, 5), Qt.MoveAction | Qt.CopyAction, md, Qt.LeftButton, Qt.NoModifier)
    window.dragEnterEvent(ev)
    assert not ev.isAccepted()
    ev = QDragEnterEvent(QPoint(5, 5), Qt.MoveAction | Qt.CopyAction, md, Qt.LeftButton, Qt.NoModifier)
    window.dragEnterEvent(ev)  # (from another program: taken)
    assert ev.isAccepted()


# ------------------------------------------------------------------ dates
def test_the_delivery_dates_note_says_when_the_rate_sheet_gave_the_days(make_window, tmp_path, monkeypatch, qt):
    """A sheet filled in before 2.8.0 may still say 7 and 21, calendar days then: the note says where the number
    comes from, and the sheet's tooltip that they are business days."""
    monkeypatch.setattr(qt.QMessageBox, "exec", lambda self: 0)
    sheets = tmp_path / "sheets"
    sheets.mkdir()
    (sheets / "Pat Rates.csv").write_text("Rate,Original,Copy,Email,Index,Days\nExpedite,$5.40,$1.10,$1.10,$1.10,7\n"
                                          "Regular,$4.30,$1.00,$1.00,$1.00,21\n", encoding="utf-8")
    s = pat_settings(initials="pr")
    s.use_ai = s.open_after = False
    s.rate_sheets_dir, s.rate_sheet = str(sheets), "Pat Rates"
    win = make_window(s)
    read_transcript(win, tmp_path)
    note = win.date_notes["delivery_date"].text()
    assert note.startswith(f"{s.days_for(win.case.get('delivery'))} business days")
    assert note.endswith("(the rate sheet's Days column)")
    assert "21 business days" in win.sheet_box.itemData(win.sheet_box.currentIndex(), 3)  # (Qt.ToolTipRole)


@pytest.mark.workdays
def test_the_turnaround_from_the_day_you_make_it_counts_from_the_day_the_run_began(monkeypatch):
    """Previewed at 11:58 PM and saved at 12:01 AM: one delivery date, from the day it was previewed."""
    s = pat_settings()
    s.turnaround_from = "today"
    s.turnaround_business_days = False
    monkeypatch.setattr(workdays, "now", lambda: datetime(2026, 10, 13, 23, 58))
    day = workdays.document_day(s).day
    monkeypatch.setattr(workdays, "now", lambda: datetime(2026, 10, 14, 0, 1))
    assert workdays.delivery_day(s, 6, day)[0] == date(2026, 10, 19)  # (Tue + 6; from Wed it would be Tue 10/20)
    case = CaseInfo()
    case.fields["delivery"] = FieldState("Regular", SRC_DEFAULT, 0.7, ["Regular"])
    from minute_filler.merge import refresh_delivery_date
    s.days_regular = 6
    refresh_delivery_date(case, s, day)
    assert case.get("delivery_date") == "10/19/2026"


def test_a_delivery_date_emptied_by_the_user_stays_empty():
    """As in 2.7.0: Generate (which now dates the files once more) filled it in again, though the window showed it
    blank."""
    from minute_filler.merge import refresh_dates
    s = pat_settings()
    case = CaseInfo()
    case.fields["delivery"] = FieldState("Regular", SRC_DEFAULT, 0.7, ["Regular"])
    case.fields["delivery_date"] = FieldState("", SRC_USER, 1.0, [""])
    refresh_dates(case, s, date(2026, 10, 13))
    assert case.get("delivery_date") == "" and case.fields["delivery_date"].source == SRC_USER


def test_a_form_at_another_speed_has_no_delivery_date_when_that_is_off():
    from minute_filler.fill import speeds_case
    s = pat_settings()
    s.fill_delivery_date = False
    case = CaseInfo()
    case.fields["delivery"] = FieldState("Regular", SRC_DEFAULT, 0.7, ["Regular"])
    out = speeds_case(case, [("Expedite", "")], s, date(2026, 10, 13))
    assert out.get("delivery") == "Expedite" and out.get("delivery_date") == ""
    out = speeds_case(case, [("Expedite", "p. 1-10"), ("Regular", "p. 11-20")], s, date(2026, 10, 13))
    assert out.get("delivery_date") == ""
    # a date the case got before the setting was turned off is the case's own speed's: not carried, and gone
    from minute_filler.merge import refresh_dates
    from minute_filler.models import SRC_DERIVED
    case.fields["delivery_date"] = FieldState("10/28/2026", SRC_DERIVED, 0.8, ["10/28/2026"])
    assert speeds_case(case, [("Expedite", "")], s, date(2026, 10, 13)).get("delivery_date") == ""
    refresh_dates(case, s, date(2026, 10, 13))
    assert case.get("delivery_date") == ""
    case.fields["delivery_date"] = FieldState("11/2/2026", SRC_USER, 1.0, ["11/2/2026"])  # (typed: kept)
    assert speeds_case(case, [("Expedite", "")], s, date(2026, 10, 13)).get("delivery_date") == "11/2/2026"


@pytest.mark.workdays
def test_the_window_redates_at_midnight_when_the_turnaround_counts_from_the_day_you_make_it(window, tmp_path,
                                                                                            monkeypatch):
    """Monday 6 PM with the evening rule: dated Tuesday. At 12:05 AM the form's date is still Tuesday, but the
    turnaround now counts from Tuesday: the window used to keep Monday's delivery date while Generate saved
    Tuesday's."""
    s = window.s
    s.turnaround_from, s.after_hours = "today", True
    monkeypatch.setattr(workdays, "now", lambda: datetime(2026, 11, 9, 18, 0))
    read_transcript(window, tmp_path)
    window._redate(force=True)
    before = window.case.get("delivery_date")
    monkeypatch.setattr(workdays, "now", lambda: datetime(2026, 11, 10, 0, 5))
    window._redate()
    after = window.case.get("delivery_date")
    assert before and after != before and window.rows["delivery_date"].text() == after


def test_a_plain_date_counts_the_turnaround_from_today(monkeypatch):
    """Day.today (now made_on) hid date.today: a plain date's getattr found the class method, and counting from
    "the day you make it" crashed."""
    s = pat_settings()
    s.turnaround_from = "today"
    assert workdays.delivery_day(s, 5, date(2026, 10, 13))[0] > date.today()
    assert workdays.Day.today() == date.today()


def test_the_calendar_of_a_date_field_copes_with_years_1_and_9999(window):
    assert window._holidays_of(10000, 1) == {} and window._holidays_of(0, 12) == {}
    marks = window._holidays_of(9999, 12)  # (Christmas)
    assert marks and all((d.year, d.month) == (9999, 12) for d in marks)


def test_years_1_and_9999_dont_break_the_holidays():
    from minute_filler.courthouses import holiday_rows
    hol = workdays.Holidays(holiday_rows())
    assert hol.name(date(1, 1, 1)) and hol.name(date(9999, 12, 31)) == ""
    assert workdays.Holidays([{"name": "Eve", "when": "day after December 31"}]).name(date(9999, 12, 31)) == ""
    assert workdays.date_warning("12/31/9999", pat_settings()) == ""


def test_an_unreadable_holiday_weekend_rule_is_the_courthouses(monkeypatch):
    from minute_filler import settings as settings_mod
    real = settings_mod.current_courthouse()

    class Court:
        holiday_weekend = "none"

        def __getattr__(self, name):
            return getattr(real, name)

    monkeypatch.setattr(settings_mod, "current_courthouse", lambda: Court())
    assert settings_mod.Settings.from_dict({"holiday_weekend": "sideways"}).holiday_weekend == "none"


# ------------------------------------------------------------------ records, help
def test_text_starting_with_an_equals_sign_stays_text_in_the_workbook(tmp_path):
    from openpyxl import load_workbook
    led = Ledger(tmp_path / "r.db", boxes=[("c1234abcd", "=Judge copy")])
    led.add_invoice(invoice("2026-0001", notes="=> sent 10/3"))
    path = led.export_xlsx(tmp_path / "records.xlsx")
    ws = load_workbook(path)["Invoices"]
    texts = [c.value for row in ws.iter_rows() for c in row]
    assert "=Judge copy" in texts and "=> sent 10/3" in texts
    assert not [c for row in ws.iter_rows() for c in row if c.data_type == "f"]


def test_a_results_line_reads_as_the_help_does():
    """A result's line had a space before each comma after a bold word ("the Court , County") and could start in
    the middle of a word ("…ry date is worked out")."""
    import html
    import re
    from minute_filler import help_search
    from minute_filler.help_topics import topic
    assert "<b>Court</b>, County" in help_search.snippet(topic("set-defaults"), "the court")
    t = topic("dates-turnaround")
    line = html.unescape(re.sub(r"<[^>]+>", "", help_search.snippet(t, "the court is closed that day")))
    assert line.startswith("…") and f" {line[1:].split()[0]} " in f" {t.text} "


def test_a_word_table_of_the_wrong_shape_leaves_the_keywords(tmp_path):
    """A header that is JSON but not a table's raised TypeError in the loading thread: the crash box, each time
    the Help window opened."""
    from minute_filler import help_meaning
    from minute_filler.help_topics import TOPICS
    topics = list(TOPICS)
    try:
        for n, header in enumerate((b"[1, 2]", b'{"dims": null, "count": 1, "words_bytes": 1}')):
            bad = tmp_path / f"bad{n}.bin"
            bad.write_bytes(help_meaning.MAGIC + header + b"\nword\n")
            help_meaning._READY.pop(id(topics), None)
            assert help_meaning.load(topics, bad) is None
    finally:
        help_meaning._READY.pop(id(topics), None)
