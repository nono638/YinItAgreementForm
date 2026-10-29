"""Your pages and everyone's (the user, 2026-10-29: "Pages billed" read as every reporter's pages). The Who
ordered which pages window counts each run's pages four ways (excerpts.count_runs): Your pages billed, Your pages
transcribed, Total pages transcribed and Total pages billed, each hidden or shown under Columns ▾ (kept in
Settings); Only your pages hides the runs without any of the user's pages as it is ticked, and Clear filter shows
them again. The Who transcribed what card has Only yours, a box per reporter and Clear filter, and highlights the
user's takes (Settings -> Options, on at first). Pat Reporter ("pr") is the user, Dana Smith ("ds") the other
reporter; all names are made up."""
import os
import time

import pytest

from minute_filler.batch import group, read_docs
from minute_filler.excerpts import Run, case_days, count_runs
from minute_filler.models import Attorney
from minute_filler.settings import Settings

from helpers import pat_settings, transcript_pdf

PR, DS = "pr", "ds"
MIXED = [PR, PR, PR, DS, DS, DS, DS, PR, PR, DS]  # Pat 1-3 and 8-9, Dana 4-7 and 10


def alex():
    return Attorney(name="Alex B. Counsel", firm="Counsel & Counsel", checked=True)


def dana():
    return Attorney(name="Dana Smith", firm="Smith Law", checked=True)


@pytest.fixture
def s(tmp_path):
    s = pat_settings()
    s.output_dir = str(tmp_path / "out")
    s.records_dir = str(tmp_path / "records")
    return s


def job_with(tmp_path, s, initials, name="Roe 6-3-2026.pdf"):
    """The one job made from a transcript with these initials on its pages, Alex and Dana Smith's firms on it."""
    docs, errors = read_docs([str(transcript_pdf(tmp_path / name, len(initials), initials=initials))], s)
    assert not errors
    job, = group(docs, s)
    job.case.attorneys = [alex(), dana()]
    return job


def front_job(tmp_path, s):
    """A transcript of two reporters whose first 2 pages (its title pages) carry no initials."""
    from test_runsheet import TALK, TITLE_1, TITLE_2, transcript
    pages = [(310, "", "", TITLE_1), (311, "", "", TITLE_2)] + [
        (312 + i, x, "Proceedings", TALK) for i, x in enumerate([PR, PR, DS, DS, DS, DS])]
    docs, errors = read_docs([str(transcript(tmp_path / "Roe front 6-3-2026.pdf", pages))], s)
    job, = group(docs, s)
    job.case.attorneys = [alex(), dana()]
    return job


def three_runs(job):
    """Pages 1-5 ordered by both firms, 6-7 (Dana's) by nobody, 8-10 by Alex's firm."""
    a, b = alex().key(), dana().key()
    day, = case_days([job])
    return [Run(day, 1, 5, [a, b]), Run(day, 6, 7, []), Run(day, 8, 10, [a])]


def counted(runs):
    c = count_runs(runs)
    return [(c[id(r)].your_billed, c[id(r)].your_transcribed, c[id(r)].total_transcribed, c[id(r)].total_billed)
            for r in runs]


def test_each_run_counts_your_pages_and_everyones(tmp_path, s):
    job = job_with(tmp_path, s, MIXED)
    assert job.your_mask() == [True, True, True, False, False, False, False, True, True, False]
    # (your billed, your transcribed, total transcribed, total billed)
    assert counted(three_runs(job)) == [(3, 3, 5, 5), (0, 0, 2, 0), (2, 2, 3, 3)]


def test_your_pages_billed_follow_whose_pages(tmp_path, s):
    """Whose pages... bills Dana's pages only (made for her): none of the user's are billed, but they are still
    the user's transcribed pages."""
    job = job_with(tmp_path, s, MIXED)
    job.page_basis[job.docs[0].key()] = DS
    assert counted(three_runs(job)) == [(0, 3, 5, 5), (0, 0, 2, 0), (0, 2, 3, 3)]


def test_your_pages_unknown_until_whose_first_pages_are_said(tmp_path, s):
    """The first pages carry no initials on a transcript of two reporters: which are the user's can't be told
    (None, "?" in the window) until Whose pages... says whose they are."""
    job = front_job(tmp_path, s)
    day, = case_days([job])
    run = Run(day, 1, 8, [alex().key()])
    assert job.your_mask() is None and count_runs([run])[id(run)].your_transcribed is None
    job.front_owner[job.docs[0].key()] = PR
    assert count_runs([run])[id(run)].your_transcribed == 4


def test_a_job_of_two_dates_counts_every_reporters_pages(tmp_path, s):
    """Two dates read as one job (ordered whole: its one run is as long as the pages the user's invoice bills).
    Its totals were that run's length, the user's 10 pages, not the 20 the two reporters wrote."""
    s.batch_combine_dates = True
    docs, errors = read_docs([str(transcript_pdf(tmp_path / "D1.pdf", 10, initials=MIXED)),
                              str(transcript_pdf(tmp_path / "D2.pdf", 10, date="June 4, 2026", initials=MIXED))], s)
    assert not errors
    job, = group(docs, s)
    job.case.attorneys = [alex(), dana()]
    day, = case_days([job])
    assert not day.splittable and day.pages == 10
    run = Run(day, 1, day.pages, [alex().key()])
    assert counted([run]) == [(10, 10, 20, 20)]
    for d in job.transcripts():  # (Whose pages... bills Dana's: none of the user's pages are billed)
        job.page_basis[d.key()] = DS
    run = Run(case_days([job])[0], 1, 10, [alex().key()])
    assert counted([run]) == [(0, 10, 20, 20)]


# ------------------------------------------------------------------ the Who ordered window

@pytest.fixture
def excerpts(qt, tmp_path, s):
    from minute_filler.gui.excerpts import ExcerptsWindow
    from minute_filler.excerpts import store
    job = job_with(tmp_path, s, MIXED)
    runs = three_runs(job)
    store(runs[0].day, runs, [alex(), dana()], s)
    w = ExcerptsWindow(s, lambda: None)
    w.load([job])
    yield w
    w.close()
    w.deleteLater()


def cells(w, col):
    return [w.table.item(r, col).text() for r in range(w.table.rowCount())]


def shown_rows(w):
    return [r for r in range(w.table.rowCount()) if not w.table.isRowHidden(r)]


def test_the_window_has_the_four_counts(excerpts):
    from minute_filler.gui.excerpts import COUNTS, FIRST_FIRM
    w = excerpts
    first = FIRST_FIRM + len(w.firms)
    heads = [w.table.horizontalHeaderItem(first + k).text() for k in range(len(COUNTS))]
    assert heads == ["Your pages\nbilled", "Your pages\ntranscribed", "Total pages\ntranscribed",
                     "Total pages\nbilled"]
    assert cells(w, first) == ["3", "0 (nobody ordered it)", "2"]
    assert cells(w, first + 1) == ["3", "0", "2"]
    assert cells(w, first + 2) == ["5", "2", "3"]
    assert cells(w, first + 3) == ["5", "0", "3"]


def test_the_pages_nobody_ordered_are_every_reporters(excerpts):
    """Pages 6-7, Dana's, are ordered by nobody: the summary's Nobody line counts them, as Total pages
    transcribed does (it counted the user's pages only, none here, and wasn't shown)."""
    assert "Nobody: 2 pp. (not billed)" in excerpts.summary.text()


def test_a_count_hidden_under_columns_is_kept(excerpts):
    from minute_filler.gui.excerpts import FIRST_FIRM
    w = excerpts
    col = FIRST_FIRM + len(w.firms) + 2  # Total pages transcribed
    w.count_actions["total_transcribed"].setChecked(False)
    assert w.table.isColumnHidden(col) and Settings.load().excerpt_hidden_cols == ["total_transcribed"]
    w.load(w.group)  # (drawn again: still hidden)
    assert w.table.isColumnHidden(col)
    w.count_actions["total_transcribed"].setChecked(True)
    assert not w.table.isColumnHidden(col) and Settings.load().excerpt_hidden_cols == []


def test_only_your_pages_hides_the_runs_without_any_at_once(excerpts):
    w = excerpts
    assert shown_rows(w) == [0, 1, 2] and w.clear_filter.isHidden()
    w.only_mine.setChecked(True)
    assert shown_rows(w) == [0, 2] and not w.clear_filter.isHidden()
    assert w.filter_note.text() == "Showing 2 of 3 runs" and Settings.load().excerpt_only_mine
    w.load(w.group)  # (a change from the main window: the filter holds)
    assert shown_rows(w) == [0, 2]
    w.clear_filter.click()
    assert shown_rows(w) == [0, 1, 2] and w.clear_filter.isHidden() and not w.only_mine.isChecked()
    assert not Settings.load().excerpt_only_mine


def test_a_run_chosen_on_a_filtered_table_is_the_one_removed(excerpts):
    """Rows are hidden, not taken out: the run of the row chosen is still the one Remove run removes."""
    w = excerpts
    w.only_mine.setChecked(True)
    w.table.selectRow(2)  # pages 8-10
    w._remove()
    assert [(r.start, r.end) for r in w.runs] == [(1, 5), (6, 10)]


def test_a_run_hidden_by_the_filter_is_chosen_no_longer(excerpts):
    """The run chosen (pages 6-7, Dana's only) is hidden as Only your pages is ticked: Split run and the Delete
    key no longer act on it unseen, and a drawing again doesn't choose it again."""
    w = excerpts
    w.table.selectRow(1)
    w.only_mine.setChecked(True)
    assert w._chosen() is None and w.problem.text() == "Choose a run in the table first."
    w._remove()
    w._split()
    assert [(r.start, r.end) for r in w.runs] == [(1, 5), (6, 7), (8, 10)]
    w.table.selectRow(1)  # (as a drawing again did: _show kept the row chosen before)
    w._show()
    assert w._chosen() is None


def test_only_your_pages_is_disabled_when_none_can_be_told(qt, tmp_path, s):
    from minute_filler.gui.excerpts import UNKNOWN, ExcerptsWindow
    job = front_job(tmp_path, s)
    s.excerpt_only_mine = True  # (ticked on another job: it filters nothing here, so no Clear filter)
    w = ExcerptsWindow(s, lambda: None)
    w.load([job])
    assert not w.only_mine.isEnabled() and w.only_mine.toolTip() == UNKNOWN
    assert w.clear_filter.isHidden() and len(shown_rows(w)) == w.table.rowCount()
    w.deleteLater()


# ------------------------------------------------------------------ the Who transcribed what card

@pytest.fixture
def window(tmp_path, monkeypatch, make_window):
    from PySide6 import QtWidgets
    monkeypatch.setattr(QtWidgets.QMessageBox, "exec", lambda self: 0)
    s = pat_settings()
    s.use_ai = s.open_after = False
    s.output_dir = str(tmp_path / "out")
    s.records_dir = str(tmp_path / "records")
    w = make_window(s)
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    w.add_files([str(transcript_pdf(tmp_path / "Roe 6-3-2026.pdf", 6, initials=[PR, PR, DS, DS, DS, PR])),
                 str(transcript_pdf(tmp_path / "Roe 6-4-2026.pdf", 4, date="June 4, 2026",
                                    initials=[PR, DS, PR, PR]))])
    end = time.time() + 20
    while time.time() < end and (w.work or len(w.jobs) < 2 or not all(j.docs for j in w.jobs)):
        QtWidgets.QApplication.processEvents()
        time.sleep(0.01)
    w._update_status()
    return w


def reporters_shown(w):
    return [w.takes.item(r, 1).text() for r in range(w.takes.rowCount()) if not w.takes.isRowHidden(r)]


def test_only_yours_and_the_reporters_boxes_filter_the_takes(window):
    w = window
    assert reporters_shown(w) == ["Pat", "DS", "Pat", "Pat", "DS", "Pat"] and w.takes_clear.isHidden()
    full = w.takes.height()
    w.takes_mine.setChecked(True)
    assert reporters_shown(w) == ["Pat"] * 4 and not w.takes_clear.isHidden()
    assert not w.takes_who["DS"].isEnabled()  # (Only yours says whose are shown)
    assert w.takes.height() < full  # (as tall as the takes shown)
    assert w.takes_info.text().endswith("showing 4 of 6 takes") and Settings.load().takes_only_mine
    w.takes_mine.setChecked(False)
    w.takes_who["Pat"].setChecked(False)
    assert reporters_shown(w) == ["DS", "DS"] and not w.takes_clear.isHidden()
    w._show_takes()  # (drawn again, a read ending: what is unticked stays so)
    assert reporters_shown(w) == ["DS", "DS"]
    w.takes_mine.setChecked(True)
    w.takes_clear.click()
    assert len(reporters_shown(w)) == 6 and w.takes_clear.isHidden() and not w.takes_mine.isChecked()
    assert all(b.isChecked() and b.isEnabled() for b in w.takes_who.values())


def test_save_a_copy_saves_the_takes_shown(window, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from PySide6 import QtWidgets
    w = window
    w.takes_who["Pat"].setChecked(False)
    monkeypatch.setattr(QtWidgets.QFileDialog, "getSaveFileName",
                        lambda *a: (str(tmp_path / "w.xlsx"), "Excel workbook (*.xlsx)"))
    w._save_takes()
    ws = load_workbook(tmp_path / "w.xlsx").active
    assert [ws.cell(r, 2).value for r in (2, 3)] == ["DS", "DS"] and ws.cell(4, 1).value is None
    w.takes_who["DS"].setChecked(False)  # (no take shown: nothing to save, and the button says so)
    assert not w.takes_save.isEnabled()
    w.takes_clear.click()
    assert w.takes_save.isEnabled()


def test_your_takes_are_highlighted_unless_settings_say_not(window):
    from PySide6.QtCore import Qt
    w = window
    assert w.s.takes_highlight_mine  # (on at first)
    mine, theirs = w.takes.item(0, 1), w.takes.item(1, 1)
    assert mine.text() == "Pat" and mine.font().bold() and mine.background().style() != Qt.NoBrush
    assert theirs.background().style() == Qt.NoBrush and not theirs.font().bold()
    w.s.takes_highlight_mine = False
    w._show_takes()
    assert w.takes.item(0, 1).background().style() == Qt.NoBrush and not w.takes.item(0, 1).font().bold()
    w.s.takes_highlight_mine = True
    w._set_theme("dark")  # (tinted again in the dark theme's green, the takes not read again)
    from minute_filler.gui.theme import DARK
    assert w.takes.item(0, 1).background().color().name() == DARK["ok_bg"].lower()
    w._set_theme("light")


def test_settings_save_the_highlight_box(qt):
    from minute_filler.gui.dialogs import SettingsDialog
    assert Settings().takes_highlight_mine and not Settings().takes_only_mine
    assert Settings.from_dict({"settings_version": 17}).excerpt_hidden_cols == []
    dlg = SettingsDialog(pat_settings())
    assert dlg.o_takes_mine.isChecked()
    dlg.o_takes_mine.setChecked(False)
    dlg.accept()
    assert Settings.load().takes_highlight_mine is False
