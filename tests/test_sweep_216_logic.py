"""Regressions found by the 2.16 sweep, the logic half: a rate sheet's ".50" read as $50; the run sheet's blank
first page counted on from the row above, and a Pages sum on a run sheet of another app not read; another
reporter's invoices of a day typed in and split under Who ordered which pages… billed whole without a word; "your
pages billed" counting pages no invoice of the user's bills; a number typed above the transcript's count called
"total pages", and the index's reasons for days of both kinds; "of your 85 pages" for pages typed; page numbers
worked out where none was read; "alone" for a firm billed at two speeds; Restore bringing back rows deleted
before; a backup the user named deleted; the CSV export racing the copies; the exports' bare "Pages"; "Your pages
transcribed" counting every reporter's pages; and the build's pins. Fictional names and cases."""
import csv
import os
import re
import threading
import time
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from minute_filler import courthouses
from minute_filler import records as records_mod
from minute_filler import runsheet as R
from minute_filler.batch import group, joint_invoice_sets, read_docs
from minute_filler.excerpts import Run, case_days, count_runs
from minute_filler.invoice import DayOrder, FirmInvoice, InvoiceOpts, firm_invoices, index_reason
from minute_filler.invoice_calc import Share, extra_rate, quote_ordered
from minute_filler.invoice_math import about, who_pays
from minute_filler.models import Attorney, CaseInfo, FieldState, SRC_USER
from minute_filler.rates import RateSheet, Speed, load_sheet, parse_money
from minute_filler.records import Invoice, Ledger, report_html
from minute_filler.tracking import clean_boxes

from helpers import pat_settings, transcript_pdf

ROOT = Path(__file__).resolve().parents[1]
PR, DS = "pr", "ds"
MIXED = [PR, PR, PR, DS, DS, DS, DS, PR, PR, DS]  # Pat's 5 pages and Dana's 5 of one transcript


def alex():
    return Attorney(name="Alex B. Counsel", firm="Counsel & Counsel", checked=True)


def dana():
    return Attorney(name="Dana Smith", firm="Smith Law", checked=True)


@pytest.fixture
def s(tmp_path):
    s = pat_settings()
    s.output_dir, s.records_dir, s.save_math = str(tmp_path / "out"), str(tmp_path / "records"), "off"
    return s


def job_with(tmp_path, s, initials, name="Roe 6-3-2026.pdf"):
    docs, errors = read_docs([str(transcript_pdf(tmp_path / name, len(initials), initials=initials))], s)
    assert not errors, errors
    job, = group(docs, s)
    job.case.attorneys = [alex(), dana()]
    return job


def typed(job, n):
    job.case.fields["est_pages"] = FieldState(str(n), SRC_USER, 1.0, [str(n)])


# 1 ---------------------------------------------------------------------------------------------------------------
def test_a_rate_written_with_its_dot_first_is_cents(tmp_path):
    """".50" on a rate sheet is fifty cents: it was read as $50.00 a page ("$.75" as $75), with no warning."""
    assert parse_money(".50") == Decimal("0.50") and parse_money("$.75") == Decimal("0.75")
    assert parse_money("$1,234.5") == Decimal("1234.50") and parse_money("about 4.30 a page") == Decimal("4.30")
    path = tmp_path / "Pat's rates.csv"
    path.write_text("Rate,Original,Copy,Email,Index,Days\nRegular,$4.30,.50,.50,.50,15\n", encoding="utf-8")
    regular = load_sheet(path).speeds[0]
    assert (regular.original, regular.copy) == ("4.30", "0.50")
    assert extra_rate(regular, "email") == extra_rate(regular, "index") == Decimal("0.50")  # (as priced)


# 2 ---------------------------------------------------------------------------------------------------------------
def _case():
    case = CaseInfo()
    case.set("case_name", "Jane Roe v. Sam Poe")
    case.set("index_no", "712345/2021")
    return case


def test_a_take_without_a_first_page_is_not_given_the_pages_after_the_row_above(tmp_path):
    """A take whose first page wasn't read (a scan: Starting Page No. left blank) was read back as starting
    after the row above, so it seemed to cover Pat's later pages 51-60 of the next day: that take was skipped as
    there already, and the blank cell got "=H6+1" on the next write."""
    path = tmp_path / "Roe run sheet.xlsx"
    d1, d2 = date(2026, 6, 3), date(2026, 6, 4)
    rows = [R.Row(d1, "Dana", 25, 1, initials="ds"), R.Row(d1, "Pat", 25, 26, initials="pr"),
            R.Row(d2, "Dana", 10, None, initials="ds", note=R.NOT_FOUND)]
    R._write_ours(path, _case(), R.RunSheetOpts(rows=rows), exists=False)
    back = R._read_ours(load_workbook(path)[R.SHEET])
    assert [r.start for r in back] == [1, 26, None]
    opts = R.RunSheetOpts(rows=[R.Row(d2, "Pat", 10, 51, initials="pr")])
    R._write_ours(path, _case(), opts, exists=True)
    assert (opts.added, opts.skipped) == (1, 0)
    ws = load_workbook(path)[R.SHEET]
    starts = [ws.cell(r, R.COL["start"]).value for r in range(R.FIRST, R.FIRST + 4)]
    # (the blank cell stays blank; Pat's 51 follows the row the last page came from, Pat's 50)
    end = R.L["end"]
    assert starts == [1, f"={end}{R.FIRST}+1", None, f"={end}{R.FIRST + 1}+1"]


# 3 ---------------------------------------------------------------------------------------------------------------
def test_a_pages_sum_on_another_apps_run_sheet_is_read(tmp_path):
    """Pages typed as a sum ("=25-1+1") on a run sheet this app didn't make read as no pages: the same take was
    added again."""
    path = tmp_path / "theirs.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Reporter", "Pages", "Starting Page No.", "Ending Page No."])
    ws.append([datetime(2026, 6, 3), "Dana", "=25-1+1", 1, None])
    ws.append([datetime(2026, 6, 3), "Pat", 25, 26, None])
    wb.save(path)
    opts = R.RunSheetOpts(rows=[R.Row(date(2026, 6, 3), "Dana", 25, 1, initials="ds")])
    R._append_theirs(path, opts)
    assert (opts.added, opts.skipped) == (0, 1)


def test_a_run_sheet_is_saved_through_a_file_of_its_own_copy_and_thread(tmp_path):
    """Its temporary file was "<name>.saving.xlsx" for every copy of the app: two saving at once wrote and
    deleted one another's."""
    seen = []

    class Book:
        def save(self, p):
            seen.append(Path(p).name)
            Path(p).write_bytes(b"x")

    R._save(Book(), tmp_path / "Roe run sheet.xlsx")
    assert str(os.getpid()) in seen[0] and str(threading.get_ident()) in seen[0]
    assert (tmp_path / "Roe run sheet.xlsx").read_bytes() == b"x"


# 4 ---------------------------------------------------------------------------------------------------------------
def test_another_reporters_invoices_of_a_day_typed_in_and_split_are_held(tmp_path, s):
    """Pages typed 6, Who ordered which pages… rows on those 6 (Alex 1-3, Dana 4-6) and Whose pages… billing
    Pat and DS: DS's invoices split DS's pages of the 10 total pages, which those rows don't fit, so both firms
    were billed all of DS's pages, split, without a word. The invoice is held until the user checks."""
    job = job_with(tmp_path, s, MIXED)
    job.page_basis = {job.transcripts()[0].key(): ["me", "ds"]}
    typed(job, 6)
    job.portions = [(3, [alex().key()]), (6, [dana().key()])]
    assert job.billed_by("ds").portions_problem()  # (DS's set: its rows don't end on its pages)
    assert job.hold_kind() == "portions"
    assert "Who ordered which pages… needs checking" in job.invoice_hold()
    assert "6 pages typed" in job.invoice_hold() and "10 total pages" in job.invoice_hold()
    job.page_basis = {job.transcripts()[0].key(): ["me"]}  # (Pat's own set alone fits)
    assert job.portions_problem() == ""


# 5 ---------------------------------------------------------------------------------------------------------------
def test_your_pages_billed_is_none_when_whose_pages_bills_other_reporters_only(tmp_path, s):
    """Pages typed 40 and Whose pages… billing DS only: no invoice of Pat's bills any page, but the Who ordered
    window said 40 of your pages billed."""
    job = job_with(tmp_path, s, MIXED)
    typed(job, 40)
    job.page_basis[job.docs[0].key()] = DS
    assert job.bill_reporters() == ["ds"]
    day, = case_days([job])
    run = Run(day, 1, day.pages, [alex().key()])
    assert count_runs([run])[id(run)].your_billed == 0


# 6 ---------------------------------------------------------------------------------------------------------------
def test_a_number_typed_above_the_transcripts_count_is_pages_typed():
    """A day typed 90 whose transcript has 83 total pages is judged on the 90 typed: it said "90 total pages"
    (CLAUDE.md: a number typed is pages typed). Days of both kinds read properly: not "every day has 50 pages
    typed or total or more"."""
    s = pat_settings()
    s.invoice_include_index, s.invoice_index_rule = True, "any"
    one = InvoiceOpts(90, days=[("6/3/2026", 90)], orders=[DayOrder("6/3/2026", 90, total=83)], typed=True)
    assert index_reason(one, s) == "90 pages typed, 50 or more"
    assert index_reason(InvoiceOpts(90, typed=True, total_pages=83), s) == "90 pages typed, 50 or more"
    both = InvoiceOpts(143, days=[("6/1", 60), ("6/2", 40)],
                       orders=[DayOrder("6/1", 60, total=0), DayOrder("6/2", 40, total=83)])
    s.invoice_index_rule = "each"
    assert index_reason(both, s) == "every day has 50 pages or more (typed or total)"
    s.invoice_index_rule = "total"
    assert index_reason(both, s) == "143 pages over 2 days (typed and total), 50 or more"


# 7 ---------------------------------------------------------------------------------------------------------------
def test_a_day_typed_in_says_pages_typed(tmp_path, s):
    """Who ordered which pages… placed a run of a day typed in as "1–20 of your 85 pages": CLAUDE.md's word
    for a number typed is pages typed."""
    job = job_with(tmp_path, s, [""] * 30)
    typed(job, 85)
    day, = case_days([job])
    assert day.place(1, 20) == "1–20 of 85 pages typed"


# 8 ---------------------------------------------------------------------------------------------------------------
def test_a_job_of_several_days_shows_no_number_for_a_page_without_one(tmp_path, s):
    """A job of two days shows the numbers printed on its transcripts, as read: a page of the second whose
    number wasn't read was given one counted on from the page before. Its stretch now stops there."""
    s.batch_combine_dates = True
    docs, errors = read_docs([str(transcript_pdf(tmp_path / "D1.pdf", 30)),
                              str(transcript_pdf(tmp_path / "D2.pdf", 20, date="June 4, 2026"))], s)
    assert not errors
    second = next(d for d in docs if d.ing.name.startswith("D2"))
    second.ing.marks[5] = replace(second.ing.marks[5], number=None)
    jobs = group(docs, s)
    day, = case_days(jobs)
    assert not day.splittable
    span = day.span(1, day.pages)
    assert span.startswith("101–130, ") and span != "101–130, 101–120"
    assert jobs[0].page_numbers().count(None) >= 1
    assert None not in jobs[0]._numbers(count_on=True)  # (a single day's runs still count on: printed_pages)


# 9 ---------------------------------------------------------------------------------------------------------------
def test_a_firm_billed_at_two_speeds_says_how_it_ordered_each(tmp_path):
    """Day 1 at Daily alone (40 pages), day 2 at Regular shared by 2 (50 pages): Who pays what said "alone",
    from the first Original line only."""
    def sp(name, o, c):
        return Speed(name, o, c, None, {"Email": c, "Index": c})
    sheet = RateSheet("Sample", None, [sp("Regular", "4.30", "1.00"), sp("Daily", "6.50", "1.25")])
    q = quote_ordered([Share(40, 1, False, "Daily"), Share(50, 2, False, "Regular", (("Dana Smith", "Regular"),))],
                      sheet, courthouses.split_rule())
    f = FirmInvoice(Attorney(name="Alex B. Counsel"), CaseInfo(), InvoiceOpts(90, 2), [q])
    assert who_pays([f])[0].how.startswith("40 alone + 50 shared by 2")


# 10, 11 ----------------------------------------------------------------------------------------------------------
@pytest.fixture
def ledger(tmp_path):
    return Ledger(tmp_path / "db" / "records.db")


def _invoice(no="2026-0001", **kw):
    return Invoice(no, "2026-09-10", case_name="Jane Roe v. Sam Poe", index_no="712345/2021", dates="9/10/2026",
                   firm="Counsel & Counsel", pages=120, amounts={"Regular": "300.00"}, billed_speed="Regular", **kw)


def test_restoring_an_invoice_leaves_a_row_deleted_before_it_in_the_trash(ledger, monkeypatch):
    """A file row of the invoice the user deleted on its own, a day before the invoice was deleted, came back
    with the invoice: only the rows deleted with it do."""
    ledger.add_invoice(_invoice())
    ledger.log_activity("invoice", case_name="Jane Roe v. Sam Poe", invoice_no="2026-0001", file_path="x.pdf")
    ledger.log_activity("invoice", case_name="Jane Roe v. Sam Poe", invoice_no="2026-0001", file_path="y.pdf")
    first = ledger.activity(kind="invoice")[0]

    class Day1(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2026, 10, 1, 9, 0, 0)

    class Day2(Day1):
        @classmethod
        def now(cls, tz=None):
            return datetime(2026, 10, 2, 9, 0, 0)

    monkeypatch.setattr(records_mod, "datetime", Day1)
    ledger.delete(activity=[first.id])
    monkeypatch.setattr(records_mod, "datetime", Day2)
    ledger.delete(invoices=["2026-0001"])
    ledger.restore(invoices=["2026-0001"])
    left = ledger.activity(kind="invoice")
    assert len(left) == 1 and left[0].id != first.id
    assert ledger.invoice("2026-0001") is not None


def test_a_backup_folder_keeps_a_file_the_user_named(ledger, tmp_path):
    """"records keep me.db", put in the backup folder by the user, was pruned as the oldest copy."""
    folder = tmp_path / "backups"
    folder.mkdir()
    (folder / "records keep me.db").write_text("Pat's own copy")
    ledger.add_invoice(_invoice())  # (records with nothing in them aren't copied)
    for i in range(8):
        assert ledger.backup(folder, force=True, now=datetime(2026, 10, 1, 12, 0, i)) is not None
    assert (folder / "records keep me.db").exists()
    assert len(list(folder.glob("records 2026-10-01 *.db"))) == records_mod.FORCED_KEPT  # (the app's own: pruned)


# 12 --------------------------------------------------------------------------------------------------------------
def test_the_csv_export_waits_for_the_copies_being_written(ledger, tmp_path):
    """Records' Export CSV… (the window's thread) didn't take the lock Generate all's copies (mirror) hold, and
    both wrote the same temporary file: errors, or one's half put in the other's place."""
    ledger.add_invoice(_invoice())
    folder = tmp_path / "export"
    done = threading.Event()
    with records_mod._MIRROR_LOCK:
        t = threading.Thread(target=lambda: (ledger.export_csv(folder), done.set()))
        t.start()
        time.sleep(0.3)
        assert not done.is_set()  # (it waits for the copies)
    t.join(10)
    assert done.is_set() and (folder / "invoices.csv").exists()
    assert not list(folder.glob(".*.tmp"))
    ledger.mirror_dir = tmp_path / "mirror"
    ledger.mirror()  # (re-entrant: mirror holds the lock around export_csv)
    assert (tmp_path / "mirror" / "invoices.csv").exists()


# 13 --------------------------------------------------------------------------------------------------------------
def test_the_exports_say_whose_pages(ledger, tmp_path):
    """The CSV copies, the Excel export and the report said "Pages", "My pages" and "Transcript pages": they
    use the Records window's words. A tick box named "Pages" before still gets its " (2)"."""
    ledger.add_invoice(_invoice(my_pages=60, transcript_pages=120))
    ledger.log_activity("agreement", case_name="Jane Roe v. Sam Poe", pages=120, my_pages=60, transcript_pages=120)
    folder = tmp_path / "export"
    ledger.export_csv(folder)
    with open(folder / "invoices.csv", encoding="utf-8-sig") as f:
        row = next(csv.DictReader(f))
    assert (row["Your pages billed"], row["Your pages transcribed"], row["Total pages transcribed"]) == (
        "120", "60", "120")
    assert not {"Pages", "My pages", "Transcript pages"} & set(row)
    with open(folder / "activity.csv", encoding="utf-8-sig") as f:
        heads = next(csv.reader(f))
    assert {"Pages on the file", "Your pages transcribed", "Total pages transcribed"} <= set(heads)
    assert "Pages" not in heads
    assert "<th>Pages</th>" not in report_html(ledger.invoices(), "Pat's invoices")
    assert "<th>Your pages billed</th>" in report_html(ledger.invoices(), "Pat's invoices")
    box = clean_boxes([{"key": "c1a2b3", "name": "Pages", "on": True}])
    assert next(b for b in box if b.key == "c1a2b3").name == "Pages (2)"
    assert not hasattr(records_mod, "invoice_row")  # (unused, and it said the old headings)


# 14 --------------------------------------------------------------------------------------------------------------
def test_your_pages_transcribed_of_a_transcript_billed_whole_is_your_own(tmp_path, s):
    """Whose pages… → The whole transcript bills all 10 pages, of which Pat wrote 5: the records' "Your pages
    transcribed" said 10."""
    job = job_with(tmp_path, s, MIXED)
    assert job.your_pages() == 5
    job.page_basis = {job.transcripts()[0].key(): "*"}
    (case, opts), = joint_invoice_sets([job])
    assert opts.pages == 10 and opts.my_pages == 5
    f = firm_invoices(case, s, opts)[0]
    assert (f.opts.pages, f.opts.my_pages, f.opts.total_pages) == (10, 5, 10)
    assert "wrote" not in about(f)  # (the math as before: every page is billed, so nothing of who wrote them)
    job.page_basis = {job.transcripts()[0].key(): "me"}  # (Pat's own pages: the same as billed)
    (case, opts), = joint_invoice_sets([job])
    assert opts.my_pages == 5 == opts.pages


# 16 --------------------------------------------------------------------------------------------------------------
def test_the_build_pins_what_it_imports_and_reads_the_version_by_the_spec():
    """httpx is imported by extract_llm.py itself and pyinstaller-hooks-contrib decides what the build takes
    in: both pinned. exe_check.py imports minute_filler only after APPDATA is a folder of its own (else its
    records go into the user's), and the spec reads the version by its own folder."""
    req = (ROOT / "requirements.txt").read_text(encoding="utf-8")
    assert re.search(r"^httpx==\d", req, re.M) and re.search(r"^pyinstaller-hooks-contrib==\d", req, re.M)
    check = (ROOT / ".claude" / "skills" / "update-the-app" / "scripts" / "exe_check.py").read_text(encoding="utf-8")
    assert check.index('os.environ[_name]') < check.index("from minute_filler import")
    assert check.index('os.environ[_name]') < check.index("from test_runsheet import")
    spec = (ROOT / "YinItAgreementForm.spec").read_text(encoding="utf-8")
    assert 'Path(SPECPATH, "minute_filler", "__init__.py").read_text(encoding="utf-8")' in spec
