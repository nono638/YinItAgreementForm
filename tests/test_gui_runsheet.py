"""The run sheet from the window (off screen): Generate and Generate all ask where the takes go when the case
has one (once per trial), one run sheet counted per case, a run sheet open in Excel, the Settings → Run sheet tab
(initials and reporter lines written different ways, the Witness columns) and its folder, the run sheet in the
case's Trial Runs folder and "Open the run sheet's folder", and work that ends while a question is on screen.
Also a take corrected in Who transcribed what (double-click or right-click, the "Are you sure?" box, Undo), and a
past job's corrections coming back with it."""
import os
from copy import deepcopy

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
QtWidgets = pytest.importorskip("PySide6.QtWidgets")

from openpyxl import load_workbook  # noqa: E402

from helpers import pat_settings  # noqa: E402
from test_gui_batch import wait  # noqa: E402
from test_runsheet import transcript  # noqa: E402


@pytest.fixture
def window(tmp_path, monkeypatch, make_window):
    monkeypatch.setattr(QtWidgets.QMessageBox, "exec", lambda self: 0)  # no dialogs waiting for a click
    s = pat_settings()
    s.use_ai = s.open_after = False
    s.output_dir = str(tmp_path / "out")
    s.runsheet_dir = str(tmp_path / "sheets")
    s.outputs = ["runsheet"]
    return make_window(s)


def takes_on(path):
    """The Reporter column of a run sheet, one name per take."""
    ws =load_workbook(path)["Run Sheet"]
    return [ws.cell(r, 3).value for r in range(5, ws.max_row + 1)]


def test_generate_starts_a_run_sheet_then_asks(window, tmp_path, monkeypatch):
    from minute_filler.gui.dialogs import RunSheetDialog
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"))])
    wait(window, lambda: len(window.cur.docs) == 1)
    assert window.output_boxes["runsheet"].isChecked() and "takes" in window.output_boxes["runsheet"].toolTip()
    asked = []
    monkeypatch.setattr(RunSheetDialog, "exec", lambda self: asked.append(self) or 0)  # Cancel
    window.fill()  # no run sheet yet: nothing to ask, and the attorneys are not asked about for a run sheet
    sheet, = (tmp_path / "sheets").glob("*.xlsx")
    assert takes_on(sheet) == ["Pat", "Dana", "Pat", "Dana"] and not asked and window.cur.saved == [sheet]

    window.fill()  # now there is one: asked, and Cancel makes nothing
    assert len(asked) == 1 and list((tmp_path / "sheets").glob("*.xlsx")) == [sheet]
    monkeypatch.setattr(RunSheetDialog, "exec", lambda self: RunSheetDialog.Accepted)  # add to it
    window.fill()
    assert len(takes_on(sheet)) == 4 and window.cur.saved == []  # the same takes are not added twice

    window.s.runsheet_existing = "new"
    window.fill()
    assert len(list((tmp_path / "sheets").glob("*.xlsx"))) == 2


def test_run_sheet_settings(qt, tmp_path, monkeypatch):
    """Initials and reporter names are tidied on Save, a line that can't be read is named, and they show
    again the same way."""
    from minute_filler.gui.dialogs import SettingsDialog
    warned = []
    monkeypatch.setattr(QtWidgets.QMessageBox, "warning", lambda *a: warned.append(a[2]))
    s = pat_settings()
    dlg = SettingsDialog(s)
    dlg.p_initials.setText("P.R.")
    dlg.r_names.setPlainText("DS = Dana\nkl: Kim\nnonsense\n")
    dlg.r_existing.setCurrentIndex(dlg.r_existing.findData("add"))
    dlg.accept()
    assert s.profile.initials == "pr" and s.reporter_names() == {"ds": "Dana", "kl": "Kim"}
    assert len(warned) == 1 and "nonsense" in warned[0]  # the rest is saved; the line left out is named
    assert s.runsheet_existing == "add"
    again = SettingsDialog(s)
    assert again.r_names.toPlainText() == "ds = Dana\nkl = Kim"


def test_generate_all_asks_once_per_trial(window, tmp_path, monkeypatch):
    from minute_filler.gui.dialogs import RunSheetDialog
    from test_runsheet import TALK, TITLE_1
    first = transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf")
    window.add_files([str(first)])
    wait(window, lambda: len(window.cur.docs) == 1)
    window.fill()  # the run sheet of the trial so far
    second = transcript(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1), (321, "pr", "", TALK)],
                        day="June 4, 2026")
    third = transcript(tmp_path / "Transcript 6-5-2026 Roe v Poe.pdf", [(322, "pr", "", TITLE_1), (323, "ds", "", TALK)],
                       day="June 5, 2026")
    window.new_job()
    window.add_files([str(second), str(third)])
    wait(window, lambda: len(window.jobs) == 2 and all(j.docs for j in window.jobs))
    asked = []
    monkeypatch.setattr(RunSheetDialog, "exec", lambda self: asked.append(self) or RunSheetDialog.Accepted)
    window.fill_all_jobs()
    wait(window, lambda: not window.filling)
    sheet, = (tmp_path / "sheets").glob("*.xlsx")
    assert len(asked) == 1 and takes_on(sheet) == ["Pat", "Dana", "Pat", "Dana", "Dana", "Pat", "Pat", "Dana"]


# ------------------------------------------------------------------ found in the run sheet sweep

def two_days(tmp_path, window):
    """The case's run sheet so far (June 3), then a batch of two more days (June 4 and 5): two jobs."""
    from test_runsheet import TALK, TITLE_1
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"))])
    wait(window, lambda: len(window.cur.docs) == 1)
    window.fill()
    second = transcript(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1), (321, "pr", "", TALK)],
                        day="June 4, 2026")
    third = transcript(tmp_path / "Transcript 6-5-2026 Roe v Poe.pdf", [(322, "pr", "", TITLE_1), (323, "ds", "", TALK)],
                       day="June 5, 2026")
    window.new_job()
    window.add_files([str(second), str(third)])
    wait(window, lambda: len(window.jobs) == 2 and all(j.docs for j in window.jobs))
    sheet, = (tmp_path / "sheets").glob("*.xlsx")
    return sheet


def test_files_follow_an_answer_that_came_in_during_the_question(window, tmp_path, monkeypatch):
    """The AI's answer (or a read) that ends while the run sheet question is on screen changes the job: the
    files are made from the job as it is then, as the window shows it."""
    from minute_filler.gui.dialogs import RunSheetDialog
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"))])
    wait(window, lambda: len(window.cur.docs) == 1)
    window.fill()  # the case has a run sheet now: the next Generate asks

    def answer_comes_in(dlg):
        case = deepcopy(window.cur.case)
        case.set("case_name", "Jane Roe v. Terry Hill")
        window.cur.case = case
        return RunSheetDialog.Accepted
    monkeypatch.setattr(RunSheetDialog, "exec", answer_comes_in)
    monkeypatch.setattr(RunSheetDialog, "choice", lambda self: "")  # start a new run sheet
    window.fill()
    names = [p.name for p in (tmp_path / "sheets").glob("*.xlsx")]
    assert len(names) == 2 and any("Terry Hill" in n for n in names)


def test_jobs_joined_while_asking_about_the_run_sheet(window, tmp_path, monkeypatch):
    """A job that is gone when the question closes is left out of Generate all; the rest is made."""
    from minute_filler.gui.dialogs import RunSheetDialog
    sheet = two_days(tmp_path, window)

    def joined(dlg):  # a read that ended meanwhile put the second job's documents into the first
        window.jobs.remove(window.jobs[1])
        return RunSheetDialog.Accepted
    monkeypatch.setattr(RunSheetDialog, "exec", joined)
    window.fill_all_jobs()
    wait(window, lambda: not window.filling)
    assert takes_on(sheet) == ["Pat", "Dana", "Pat", "Dana", "Dana", "Pat"]


def test_ctrl_enter_while_typing_in_the_attorney_table(window, monkeypatch):
    """Generate takes the attorney cell still being typed in as typed."""
    told = []
    monkeypatch.setattr(QtWidgets.QMessageBox, "information", lambda *a: told.append(a[2]))
    window.show()
    window.activateWindow()
    window._add_att_row()  # the cursor is in the new row's Name cell
    QtWidgets.QApplication.processEvents()
    editor = QtWidgets.QApplication.focusWidget()
    assert isinstance(editor, QtWidgets.QLineEdit)
    editor.setText("Alex B. Counsel")
    window.s.outputs = []
    window.fill()  # Ctrl+Enter: nothing ticked to make, but what was typed is kept
    assert [a.name for a in window.cur.case.attorneys] == ["Alex B. Counsel"]
    assert "run sheet" in told[0]


def test_generate_all_counts_one_run_sheet_per_case(window, tmp_path):
    two_days(tmp_path, window)
    window._refresh_job_labels()
    assert window.fill_all_btn.text().endswith("  (1 file)")


def test_batch_message_says_why_the_run_sheet_was_not_saved(window, tmp_path, monkeypatch):
    """A run sheet open in Excel: the batch's message says so plainly, without the Python error name."""
    import minute_filler.runsheet as rs
    from test_runsheet import TALK, TITLE_1
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf")),
                      str(transcript(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1),
                                     (321, "pr", "", TALK)], day="June 4, 2026"))])
    wait(window, lambda: len(window.jobs) == 2 and all(j.docs for j in window.jobs))

    def locked(wb, path):
        raise PermissionError(f"{path.name} is open in another program - close it in Excel and try again")
    monkeypatch.setattr(rs, "_save", locked)
    shown = []
    monkeypatch.setattr(window, "_saved_box", lambda title, text, *a, **kw: shown.append(text))
    window.fill_all_jobs()
    wait(window, lambda: not window.filling)
    assert "Excel" in shown[0].split("\n\n")[1]  # the heading names the run sheet too
    assert "Run Sheet.xlsx is open in another program - close it in Excel" in shown[0]
    assert "PermissionError" not in shown[0]


def test_reading_gets_its_own_copy_of_the_settings(window, monkeypatch):
    """The background read works on a copy of the settings: Settings saved meanwhile can't change them under it."""
    import minute_filler.gui.main_window as mw
    seen = []
    real = mw.read_loaders
    monkeypatch.setattr(mw, "read_loaders", lambda loaders, names, s, *a: seen.append(s) or real(loaders, names, s, *a))
    window.add_text("Please send the minutes in Smith v Jones, Index No. 712222/2024, 5/22/2026.")
    wait(window, lambda: bool(window.cur.docs))
    assert seen and seen[0] is not window.s and seen[0].profile.name == "Pat Reporter"


def test_reporter_lines_written_other_ways(qt, monkeypatch):
    """"D. S. = Dana", "kl Kim", "mt - Maria" and a tab are all read; "Dana Smith - ds" (name first) is named."""
    from minute_filler.gui.dialogs import SettingsDialog
    warned = []
    monkeypatch.setattr(QtWidgets.QMessageBox, "warning", lambda *a: warned.append(a[2]))
    s = pat_settings()
    dlg = SettingsDialog(s)
    dlg.r_names.setPlainText("D. S. = Dana\nkl Kim\nmt - Maria\npr\tPat\nDana Smith - ds\n")
    dlg.accept()
    assert s.reporter_names() == {"ds": "Dana", "kl": "Kim", "mt": "Maria", "pr": "Pat"}
    assert "Dana Smith - ds" in warned[0]


def test_the_run_sheet_goes_in_the_cases_folder_and_its_link_opens_it(make_window, tmp_path, monkeypatch):
    """Without a folder of its own (Settings → Options → Folders), Generate puts the run sheet in the case's own
    folder with its other files. "Open the run sheet's folder" opens where it is; before there is one, the nearest
    folder that is there (a case folder is not made for it)."""
    import minute_filler.gui.main_window as mw
    from minute_filler.runsheet import old_folder
    monkeypatch.setattr(QtWidgets.QMessageBox, "exec", lambda self: 0)
    opened = []
    monkeypatch.setattr(mw, "open_path", opened.append)
    s = pat_settings()
    s.use_ai = s.open_after = False
    s.output_dir = str(tmp_path / "out")
    s.outputs = ["runsheet"]
    (tmp_path / "out").mkdir()
    window = make_window(s)
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"))])
    wait(window, lambda: len(window.cur.docs) == 1)
    assert window.rs_folder.text() == '<a href="open">Open the run sheet\'s folder</a>'
    window._open_run_sheets_folder()
    case = tmp_path / "out" / "712345-2021 6-3-2026 Jane Roe v Sam Poe"
    assert opened == [tmp_path / "out"] and not case.exists()
    window.fill()
    assert window.cur.folder == case.name
    sheet, = (case / "Trial Runs").glob("*.xlsx")  # (the case folder's Trial Runs folder: the user, 2026-10-26)
    assert window.cur.saved == [sheet] and not old_folder().exists() and not list(case.glob("*.xlsx"))
    window._open_run_sheets_folder()
    assert opened[-1] == case / "Trial Runs"


def test_a_take_is_corrected_in_who_transcribed_what(make_window, tmp_path, monkeypatch):
    """Dana's first take really ends on page 315 (its initials misread as Pat's): double-click it, type 315 and say
    yes in the "Are you sure?" box. The card shows the take corrected (✎), Yours in the job list follows, and Undo
    puts the takes back as the initials say. Cancel changes nothing; the last take can't be moved earlier."""
    from minute_filler.gui.main_window import JOB_YOURS
    asked, answer = [], {"yes": True}

    def exec_box(self):
        asked.append(self.text())
        yes = next((b for b in self.buttons() if self.buttonRole(b) == QtWidgets.QMessageBox.AcceptRole), None)
        if yes is not None and answer["yes"]:
            yes.click()
        return 0
    monkeypatch.setattr(QtWidgets.QMessageBox, "exec", exec_box)
    monkeypatch.setattr(QtWidgets.QMessageBox, "information", lambda *a: asked.append(a[2]))
    typed = []
    monkeypatch.setattr(QtWidgets.QInputDialog, "getInt", lambda *a, **k: typed.append(a) or (315, True))
    s = pat_settings()
    s.use_ai = s.open_after = False
    s.output_dir = str(tmp_path / "out")
    s.outputs = ["runsheet"]
    window = make_window(s)
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"))])
    wait(window, lambda: len(window.cur.docs) == 1 and bool(window._take_rows))
    takes = lambda: [(r.reporter, r.start, r.pages) for r in window._take_rows]
    yours = lambda: window._job_item(window.cur).text(JOB_YOURS)
    assert takes() == [("Pat", 310, 2), ("Dana", 312, 3), ("Pat", 315, 3), ("Dana", 318, 2)] and yours() == "5"
    answer["yes"] = False
    window.takes.cellDoubleClicked.emit(1, 0)
    assert "starts on page 312 and ends on page 314" in typed[-1][2]
    assert "Dana's take on 6/3/2026 will end on page 315 instead of 314: page 315 becomes Dana's (now Pat's)." \
        in asked[-1] and "Your pages that day: 5 → 4." in asked[-1] and "Are you sure?" in asked[-1]
    assert takes()[1] == ("Dana", 312, 3) and not window.cur.docs[0].fixes  # Cancel: nothing changed
    answer["yes"] = True
    window.takes.cellDoubleClicked.emit(1, 0)
    assert takes() == [("Pat", 310, 2), ("Dana", 312, 4), ("Pat", 316, 2), ("Dana", 318, 2)] and yours() == "4"
    assert window.takes.item(1, 2).text() == "pp. 312–315 ✎" and window.cur.docs[0].fixes == {5: "ds"}
    assert window.takes.item(2, 2).text() == "pp. 316–317 ✎"  # (Pat's take starts later: changed too)
    assert window.takes.item(0, 2).text() == "pp. 310–311" and window.takes.item(3, 2).text() == "pp. 318–319"
    sheet = window.cur.runsheet_opts(s).rows
    assert [r.pages for r in sheet] == [2, 4, 2, 2]
    window._fix_take(window._take_rows[3])  # the last take
    assert "last take" in asked[-1] and takes()[3] == ("Dana", 318, 2)
    window._undo_take_fixes(window.cur, window.cur.docs[0])
    assert takes()[1] == ("Dana", 312, 3) and yours() == "5" and not window.cur.docs[0].fixes
    assert window.takes.item(1, 2).text() == "pp. 312–314"


def test_correcting_a_take_acts_on_the_take_clicked_and_says_what_comes_of_it(make_window, tmp_path, monkeypatch):
    """Right-click: a read that ended while the menu was open drew the card again, and the row number then named
    another take. Dana's take made to reach Pat's next take: Dana's next take joins it, and the box says where it
    ends then. A read that moved the transcript to another job while the box was open lost the yes."""
    from minute_filler.batch import Job
    asked, typed = [], []

    def exec_box(self):
        asked.append(self.text())
        moved.get("do", lambda: None)()
        yes = next((b for b in self.buttons() if self.buttonRole(b) == QtWidgets.QMessageBox.AcceptRole), None)
        if yes is not None:
            yes.click()
        return 0
    moved = {}
    monkeypatch.setattr(QtWidgets.QMessageBox, "exec", exec_box)
    monkeypatch.setattr(QtWidgets.QInputDialog, "getInt", lambda *a, **k: typed.append(a) or (317, True))
    s = pat_settings()
    s.use_ai = s.open_after = False
    s.output_dir = str(tmp_path / "out")
    s.outputs = ["runsheet"]
    window = make_window(s)
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"))])
    wait(window, lambda: len(window.cur.docs) == 1 and bool(window._take_rows))

    import minute_filler.gui.main_window as mw

    class Menu(mw.QMenu):  # (shown, a menu would wait for a click)
        def exec(self, *a):
            window._take_rows = window._take_rows[::-1]  # (drawn again meanwhile: row 1 is another take now)
            self.actions()[0].trigger()
    monkeypatch.setattr(mw, "QMenu", Menu)
    doc = window.cur.docs[0]

    def move():  # a read that ends while the box is open puts the transcript in a job of its own
        window.cur.docs.remove(doc)
        window.jobs.append(Job(docs=[doc]))
    moved["do"] = move
    window._takes_menu(window.takes.visualItemRect(window.takes.item(1, 0)).center())
    assert "Dana's take on 6/3/2026 starts on page 312" in typed[-1][2]
    assert ("will end on page 319 instead of 314: pages 315–317 become Dana's (now Pat's), and the next take of "
            "Dana's, up to page 319, joins it.") in asked[-1]
    assert doc.fixes == {5: "ds", 6: "ds", 7: "ds"}  # (corrected, in the job it is in now)


def test_a_past_jobs_corrections_come_back_and_undo_stays_undone(make_window, tmp_path, monkeypatch):
    """Opened again from the records, the job's documents get the takes as they were corrected; Undo then cleared
    them, and the next merge put them back."""
    import json
    from minute_filler.batch import job_origin
    monkeypatch.setattr(QtWidgets.QMessageBox, "question", lambda *a, **k: QtWidgets.QMessageBox.Yes)

    def exec_box(self):
        yes = next((b for b in self.buttons() if self.buttonRole(b) == QtWidgets.QMessageBox.AcceptRole), None)
        if yes is not None:
            yes.click()
        return 0
    monkeypatch.setattr(QtWidgets.QMessageBox, "exec", exec_box)
    s = pat_settings()
    s.use_ai = s.open_after = False
    s.output_dir = str(tmp_path / "out")
    s.outputs = ["runsheet"]
    window = make_window(s)
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"))])
    wait(window, lambda: len(window.cur.docs) == 1)
    window.cur.docs[0].correct({5: "ds"})
    origin = json.dumps(job_origin(window.cur))
    assert window.open_past_job(origin) is True
    wait(window, lambda: len(window.cur.docs) == 1)
    doc = window.cur.docs[0]
    assert doc.fixes == {5: "ds"} and window.cur.your_pages() == 4
    window._undo_take_fixes(window.cur, doc)
    window._remerge()
    assert doc.fixes == {} and window.cur.your_pages() == 5


def test_the_run_sheets_folder_of_a_job_that_names_no_case(make_window, tmp_path, monkeypatch):
    """A blank job: the link opens the nearest folder that is there of where its files would go, and makes no case
    folder ("No index number") for it."""
    import minute_filler.gui.main_window as mw
    opened = []
    monkeypatch.setattr(mw, "open_path", opened.append)
    s = pat_settings()
    s.use_ai = s.open_after = False
    s.output_dir = str(tmp_path / "out")
    window = make_window(s)
    window._open_run_sheets_folder()
    assert opened == [tmp_path] and not (tmp_path / "out").exists()
    s.runsheet_dir = str(tmp_path / "Sheets")  # a folder of their own: made and opened
    window._open_run_sheets_folder()
    assert opened[-1] == tmp_path / "Sheets" and (tmp_path / "Sheets").is_dir()


def test_the_witness_columns_setting(qt):
    """Settings -> Run sheet -> Witness Start / End: the first and last head printed in the take (the default), or
    blank."""
    from minute_filler.gui.dialogs import SettingsDialog
    s = pat_settings()
    assert s.runsheet_witness == "headers"
    dlg = SettingsDialog(s)
    assert dlg.r_witness.currentData() == "headers"
    dlg.r_witness.setCurrentIndex(dlg.r_witness.findData("blank"))
    dlg.accept()
    assert s.runsheet_witness == "blank"
    s.save()
    from minute_filler.settings import Settings
    assert Settings.load().runsheet_witness == "blank"


def test_open_the_run_sheets_folder_before_there_is_one(qt, tmp_path, monkeypatch):
    import minute_filler.gui.dialogs as dialogs
    opened = []
    monkeypatch.setattr(dialogs, "open_path", opened.append)
    dlg = dialogs.SettingsDialog(pat_settings())
    edit = dlg.o_dirs["runsheet"]  # Options -> Folders
    edit.setText(str(tmp_path / "Run Sheets"))
    dlg._open_folder(edit, "runsheet")
    assert (tmp_path / "Run Sheets").is_dir() and opened == [tmp_path / "Run Sheets"]
