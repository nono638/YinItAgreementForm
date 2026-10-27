"""Regressions found by the 2.14.0 review: the run sheet's kept running heads lost or frozen after an edit in
Excel, Who Ordered coloured by one day's firms, wiped by a records read that failed, showing a paid invoice's
offered speed and another case's invoices; the case subfolders folding a folder of an output's own, the math
without a Math folder, a document in one of a case's own folders, and the math window's start folder; a firm of
another day losing its tooltip, and the square on the dark theme. Fictional names and cases."""
from datetime import date

from openpyxl import load_workbook

from minute_filler.batch import Job, out_dir_for, save_to_for, saved_folders, trial_marks
from minute_filler.deliver import ledger_for
from minute_filler.firm_marks import shade_of
from minute_filler.models import Attorney
from minute_filler.records import Invoice
from minute_filler.runsheet import (BY_REPORTER, META, ORDERS, SHEET, WITNESSES, Row, RunSheetOpts, _read_meta,
                                    add_takes, order_rows, refresh_orders)

from helpers import ROE, make_case, pat_settings
from test_runsheet import s  # noqa: F401 (the fixture)

D1, D2, D3 = date(2026, 6, 3), date(2026, 6, 4), date(2026, 6, 5)


def tab(path, name):
    return [[c.value for c in row] for row in load_workbook(path)[name].iter_rows()]


def heads(path):
    return [r[0] for r in tab(path, WITNESSES)[1:]]


def roe(*attorneys):
    return make_case(dict(ROE), list(attorneys))


def att(name, firm, party):
    return Attorney(name=name, firm=firm, party=party)


# ------------------------------------------------------------------ the kept running heads

def test_the_heads_inside_a_take_survive_a_page_count_corrected_in_excel(s):
    """Correcting an earlier take's pages in Excel moves the first page of every take after it (a formula): the
    heads kept for those takes are found again by their Witness cells, not lost (E, inside Dana's take)."""
    p = add_takes(roe(), RunSheetOpts(rows=[Row(D1, "Pat", 30, 1, "A", "C", head_list=["A", "B", "C"]),
                                            Row(D1, "Dana", 20, 31, "D", "F", head_list=["D", "E", "F"])]), s, [])
    wb = load_workbook(p)
    ws = wb[SHEET]
    ws["F5"].value = 29  # (Pat's pages)
    wb.save(p)
    add_takes(roe(), RunSheetOpts(rows=[Row(D2, "Pat", 5, 50, "Z", "Z", head_list=["Z"])]), s, [])
    assert heads(p) == ["A", "B", "C", "D", "E", "F", "Z"]


def test_typed_rows_are_read_from_their_cells_each_time(s):
    """Two rows typed in Excel for one day and reporter (the same day, first page and reporter) each keep their
    own Witness cell, and a cell corrected by hand later is what the Witnesses tab lists: typed rows' cells aren't
    frozen onto the hidden sheet as if read from a transcript."""
    p = add_takes(roe(), RunSheetOpts(rows=[Row(D1, "Pat", 30, 1, "A", "C", head_list=["A", "B", "C"])]), s, [])
    wb = load_workbook(p)
    ws = wb[SHEET]
    for r, text in ((6, "X - Direct"), (7, "Y - Cross")):
        ws.cell(r, 1).value, ws.cell(r, 3).value, ws.cell(r, 9).value = D2, "Dana", text
    wb.save(p)
    add_takes(roe(), RunSheetOpts(rows=[Row(D3, "Pat", 10, 31, "Z", "Z", head_list=["Z"])]), s, [])
    assert {"X - Direct", "Y - Cross"} <= set(heads(p))
    assert not [v for k, v in _read_meta(load_workbook(p)).items() if k == "heads" for x in v if "Cross\"" in x]
    wb = load_workbook(p)
    ws = wb[SHEET]
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 9).value == "Y - Cross":
            ws.cell(r, 9).value = "Y - Cross (corrected)"
    wb.save(p)
    add_takes(roe(), RunSheetOpts(rows=[Row(date(2026, 6, 8), "Pat", 5, 41, "R", "R", head_list=["R"])]), s, [])
    got = heads(p)
    assert "Y - Cross (corrected)" in got and "Y - Cross" not in got and "X - Direct" in got
    assert got[:3] == ["A", "B", "C"]  # (the transcript's heads still kept)


def test_a_witness_cell_corrected_by_hand_beats_the_kept_heads(s):
    """The app wrote Witness Start "A"; the user corrected it to "A. Doe": the Witnesses tab follows the cells."""
    p = add_takes(roe(), RunSheetOpts(rows=[Row(D1, "Pat", 30, 1, "A", "C", head_list=["A", "B", "C"])]), s, [])
    wb = load_workbook(p)
    wb[SHEET].cell(5, 9).value = "A. Doe - Direct"
    wb.save(p)
    add_takes(roe(), RunSheetOpts(rows=[Row(D2, "Pat", 5, 31, "Z", "Z", head_list=["Z"])]), s, [])
    assert heads(p) == ["A. Doe - Direct", "C", "Z"]


def test_a_take_whose_last_head_repeats_an_earlier_one_keeps_its_heads(s):
    """Witness End is the last head printed ("Proceedings" again, or "PROCEEDINGS" printed first and "Proceedings"
    last), which the take's heads, each once, have earlier: still the take's own, its heads kept."""
    p = add_takes(roe(), RunSheetOpts(rows=[
        Row(D1, "Pat", 30, 1, "Proceedings", "Proceedings", head_list=["Proceedings", "J. Doe - Direct"]),
        Row(D1, "Dana", 20, 31, "PROCEEDINGS", "Proceedings", head_list=["PROCEEDINGS", "M. Lee - Cross"])]), s, [])
    add_takes(roe(), RunSheetOpts(rows=[Row(D2, "Pat", 5, 51, "Z", "Z", head_list=["Z"])]), s, [])
    assert heads(p) == ["Proceedings", "J. Doe - Direct", "M. Lee - Cross", "Z"]


def test_a_hidden_sheet_cut_down_to_one_column_is_still_read(s):
    """Its rows are one cell each: read as keys without values, not a crash."""
    p = add_takes(roe(), RunSheetOpts(rows=[Row(D1, "Pat", 30, 1)]), s, [])
    wb = load_workbook(p)
    wb[META].delete_cols(2)
    wb.save(p)
    assert _read_meta(load_workbook(p))["case"] == [""]


def test_no_reporter_named_is_no_circular_total(s):
    """A scan's takes have no reporter: By Reporter's totals are 0, not a SUM of their own cell."""
    p = add_takes(roe(), RunSheetOpts(rows=[Row(D1, "", 30, 1)]), s, [])
    total = tab(p, BY_REPORTER)[1]
    assert total[0] == "Total" and total[1] == 0 and total[2] == 0


# ------------------------------------------------------------------ Who Ordered

def agree(s, a, day):
    ledger_for(s).log_activity("agreement", case_name=ROE["case_name"], index_no=ROE["index_no"], firm=a.firm,
                               attorney=a.name, dates=day)


def colours(path):
    return {r[1].value: (r[0].fill.start_color.rgb[-6:].lower() if r[0].fill.fill_type else None, r[2].value)
            for r in load_workbook(path)[ORDERS].iter_rows(min_row=2)}


def test_who_ordered_colours_the_firms_as_the_window_does_over_the_trials_days(s):
    """Day 1 names Smith Law and Counsel & Counsel, day 2 Park PLLC and Counsel & Counsel: after day 2's Generate
    each firm has the colour and side the window gives it over both days (trial_marks), Smith Law's too, and
    generating day 1 again changes nothing."""
    smith, counsel, park = (att("Dana Smith", "Smith Law", "Plaintiff"),
                            att("Alex B. Counsel", "Counsel & Counsel", "Defendant"),
                            att("Kim Park", "Park PLLC", "Defendant"))
    for a, day in ((smith, "6/3/2026"), (counsel, "6/3/2026"), (park, "6/4/2026")):
        agree(s, a, day)
    day1, day2 = roe(smith, counsel), make_case({**ROE, "dates": "6/4/2026"}, [park, counsel])
    p = add_takes(day1, RunSheetOpts(rows=[Row(D1, "Pat", 30, 1)]), s, [])
    add_takes(day2, RunSheetOpts(rows=[Row(D2, "Pat", 30, 31)]), s, [])
    j1, j2 = Job(), Job()
    j1.case, j2.case = day1, day2
    window = {a.firm: shade_of(m)[1:] for a in day1.attorneys + day2.attorneys
              if (m := trial_marks(j2, [j1, j2]).get(a.key()))}
    got = colours(p)
    assert {k: v[0] for k, v in got.items()} == window
    assert got["Smith Law"][1] == "Plaintiff"
    assert refresh_orders(p, day1, s) is False  # (no colours flipped back)


def test_generate_all_colours_a_firm_of_a_later_day_of_one_run(s):
    """fill_jobs refreshes Who Ordered with every day's attorneys of the run sheet, not the first day's only."""
    smith, park = att("Dana Smith", "Smith Law", "Plaintiff"), att("Kim Park", "Park PLLC", "Defendant")
    agree(s, park, "6/4/2026")
    p = add_takes(roe(smith), RunSheetOpts(rows=[Row(D1, "Pat", 30, 1)]), s, [])
    assert colours(p)["Park PLLC"] == (None, None)
    day2 = make_case({**ROE, "dates": "6/4/2026"}, [park])
    assert refresh_orders(p, roe(smith), s, [roe(smith), day2]) is True
    from minute_filler.firm_marks import mark_of, marks
    assert colours(p)["Park PLLC"] == (shade_of(mark_of(park, marks([smith, park])))[1:], "Defense")


def test_a_day_generated_after_a_later_one_gives_its_firms_their_colours_first(s):
    """Day 2 (Park PLLC) generated before day 1 (Counsel & Counsel): the window gives day 1's firm the first
    defense colour, over both days, and so does Who Ordered once day 1 is on the run sheet."""
    counsel, park = att("Alex B. Counsel", "Counsel & Counsel", "Defendant"), att("Kim Park", "Park PLLC", "Defendant")
    agree(s, counsel, "6/3/2026")
    agree(s, park, "6/4/2026")
    day1, day2 = roe(counsel), make_case({**ROE, "dates": "6/4/2026"}, [park])
    p = add_takes(day2, RunSheetOpts(rows=[Row(D2, "Pat", 30, 31)]), s, [])
    add_takes(day1, RunSheetOpts(rows=[Row(D1, "Pat", 30, 1)]), s, [])
    j1, j2 = Job(), Job()
    j1.case, j2.case = day1, day2
    window = {a.firm: shade_of(m)[1:] for a in (counsel, park) if (m := trial_marks(j1, [j1, j2]).get(a.key()))}
    assert {k: v[0] for k, v in colours(p).items()} == window


def test_a_records_read_that_fails_leaves_who_ordered_as_it_was(s, monkeypatch):
    import minute_filler.deliver as deliver
    agree(s, att("Dana Smith", "Smith Law", "Plaintiff"), "6/3/2026")
    p = add_takes(roe(), RunSheetOpts(rows=[Row(D1, "Pat", 30, 1)]), s, [])
    before = tab(p, ORDERS)
    assert before[1][1] == "Smith Law"

    def locked(_s):
        raise OSError("the records are locked")
    monkeypatch.setattr(deliver, "ledger_for", locked)
    add_takes(roe(), RunSheetOpts(rows=[Row(D2, "Pat", 30, 31)]), s, [])
    assert tab(p, ORDERS) == before and len(tab(p, SHEET)) > 5  # (the day added all the same)
    assert refresh_orders(p, roe(), s) is False


def invoice(no, **kw):
    return Invoice(no, "2026-06-05", ROE["case_name"], index_no=ROE["index_no"], **kw)


def test_a_party_corrected_after_day_one_is_the_firms_on_who_ordered(s):
    """Day 1 was read without Smith Law's party; it was corrected before day 2 (the window corrects it on every
    day): Who Ordered gives Smith Law the plaintiffs' side and colour."""
    agree(s, att("Dana Smith", "Smith Law", "Plaintiff"), "6/3/2026")
    p = add_takes(roe(att("Dana Smith", "Smith Law", "")), RunSheetOpts(rows=[Row(D1, "Pat", 30, 1)]), s, [])
    assert colours(p)["Smith Law"][1] == "Other"
    fixed = make_case({**ROE, "dates": "6/4/2026"}, [att("Dana Smith", "Smith Law", "Attorneys for Plaintiff")])
    add_takes(fixed, RunSheetOpts(rows=[Row(D2, "Pat", 30, 31)]), s, [])
    assert colours(p)["Smith Law"] == ("ea580c", "Plaintiff")


def test_a_run_sheets_days_are_named_by_the_earliest(s):
    """Generate all with day 2 listed first: the run sheet's case is day 1's, with its attorneys (the colours of
    Who Ordered follow its firms first)."""
    from minute_filler.batch import unit_case
    j1, j2 = Job(), Job()
    j1.case = roe(att("Alex B. Counsel", "Counsel & Counsel", "Defendant"))
    j2.case = make_case({**ROE, "dates": "6/4/2026"}, [att("Kim Park", "Park PLLC", "Defendant")])
    case = unit_case([j2, j1])
    assert [a.firm for a in case.attorneys] == ["Counsel & Counsel"] and "6/4/2026" in case.get("dates")


def test_a_paid_invoice_with_no_amount_recorded_shows_none(s):
    ledger_for(s).add_invoice(invoice("2026-0001", firm="Smith Law", bill_to="Dana Smith", billed_speed="Regular",
                                      amounts={"Regular": "60.00"}, status="paid", paid_speed="Daily"))
    row, = order_rows(ROE["case_name"], [ROE["index_no"]], roe(), s)
    assert row[5] == "Daily" and not row[8]


def test_who_ordered_shows_what_was_paid_and_what_is_offered(s):
    """A paid invoice: the speed and amount paid. An invoice offering two speeds: both, as offered, whichever is
    its billed speed."""
    lg = ledger_for(s)
    lg.add_invoice(invoice("2026-0001", firm="Smith Law", bill_to="Dana Smith", billed_speed="Regular",
                           amounts={"Regular": "60.00", "Expedite": "90.00"}, status="paid", paid_speed="Expedite",
                           amount_paid="90.00", paid_date="2026-06-20"))
    lg.add_invoice(invoice("2026-0002", firm="Park PLLC", bill_to="Kim Park", billed_speed="Regular",
                           amounts={"Regular": "60.00", "Expedite": "90.00"}))
    rows = order_rows(ROE["case_name"], [ROE["index_no"]], roe(), s)
    assert rows[0][5] == "Expedite" and rows[0][8] == "$90.00" and rows[0][9] == "Paid 6/20/2026"
    assert rows[1][5] == "offered: Regular, Expedite"


def test_who_ordered_leaves_out_a_case_of_another_index_number_and_the_same_caption(s):
    lg = ledger_for(s)
    lg.add_invoice(Invoice("2026-0001", "2026-06-05", ROE["case_name"], index_no="799999/2019", firm="Smith Law",
                           bill_to="Dana Smith", amounts={"Regular": "60.00"}, billed_speed="Regular"))
    lg.add_invoice(Invoice("2026-0002", "2026-06-05", ROE["case_name"], firm="Park PLLC", bill_to="Kim Park",
                           amounts={"Regular": "60.00"}, billed_speed="Regular"))  # (no index: by its caption)
    assert [r[1] for r in order_rows(ROE["case_name"], [ROE["index_no"]], roe(), s)] == ["Park PLLC"]


def test_an_agreement_logged_by_the_attorneys_name_is_billed_by_the_firms_invoice(s):
    agree(s, att("Alex B. Counsel", "", "Defendant"), "6/3/2026")
    ledger_for(s).add_invoice(invoice("2026-0003", firm="Counsel & Counsel", bill_to="Alex B. Counsel",
                                      amounts={"Regular": "60.00"}, billed_speed="Regular"))
    rows = order_rows(ROE["case_name"], [ROE["index_no"]], roe(), s)
    assert [r[-1] for r in rows] == ["Open"]


# ------------------------------------------------------------------ the case subfolders

def test_a_folder_of_an_outputs_own_named_invoices_is_named_as_it_is(tmp_path):
    s = pat_settings()
    s.case_folders = True
    own = tmp_path / "Billing" / "Invoices"
    s.output_dirs["invoice"] = str(own)
    case = tmp_path / "out" / "712345-2021"
    assert saved_folders([own / "Invoice 2026-0001.pdf", own / "Invoice 2026-0001 - the math.pdf"], s) == [own]
    assert saved_folders([case / "MOFR" / "MOFR.pdf", case / "Trial Runs" / "Run Sheet.xlsx"], s) == [case]


def test_without_a_math_folder_the_math_goes_next_to_the_invoices(tmp_path):
    from dataclasses import replace
    from minute_filler import courthouses
    s = pat_settings()
    s.case_folders = True
    was = courthouses.current()
    try:
        courthouses.use(replace(was, math_subfolder=""))
        assert s.folder_for("math", tmp_path / "712345-2021") == tmp_path / "712345-2021" / "Invoices"
    finally:
        courthouses.use(was)
    assert s.folder_for("math", tmp_path / "712345-2021") == tmp_path / "712345-2021" / "Math"


def test_a_case_folder_made_beside_another_is_the_cases_too(tmp_path):
    """"712345-2021 (2)" (made beside a 712345-2021 there already: pdfout.free_folder): a document in it, or in its
    Invoices folder, saves there, not in another 712345-2021 inside it."""
    from types import SimpleNamespace
    s = pat_settings()
    s.case_folders, s.output_dir = True, ""
    job = Job()
    job.case = roe()
    job.folder = "712345-2021"
    for where in ("712345-2021 (2)", "712345-2021 (2)/Invoices"):
        job.docs = [SimpleNamespace(path=str(tmp_path / where / "order.pdf"))]
        assert out_dir_for(job, s) == tmp_path / "712345-2021 (2)", where


def test_a_document_in_one_of_its_case_folders_own_folders_saves_to_that_case_folder(tmp_path):
    """An order e-mailed and saved in "...\\712345-2021\\Invoices": its files go in 712345-2021's folders, not in a
    second 712345-2021 inside Invoices."""
    from types import SimpleNamespace
    s = pat_settings()
    s.case_folders, s.output_dir = True, ""
    job = Job()
    job.case = roe()
    job.docs = [SimpleNamespace(path=str(tmp_path / "712345-2021" / "Invoices" / "order.pdf"))]
    job.folder = "712345-2021"
    assert save_to_for(job, s) == tmp_path / "712345-2021"
    assert out_dir_for(job, s) == tmp_path / "712345-2021"
    s.case_folders = False  # (a folder of that name with case folders off: it is just Save to)
    assert save_to_for(job, s) == tmp_path / "712345-2021" / "Invoices"


def test_generate_all_starts_the_maths_save_in_the_cases_math_folder(tmp_path, monkeypatch, make_window, qt):
    """The agreements have a folder of their own (saved first): the math window still starts in the invoiced
    case's Math folder, not in <the agreements' folder>\\Math."""
    from test_extras import wait
    from test_show_the_math import math_window
    win = math_window(tmp_path, monkeypatch, make_window, qt, preview=False, days=2)
    win.s.output_dirs["agreement"] = str(tmp_path / "Forms")
    win.s.save_math, win.s.show_math = "off", True
    shown = []
    monkeypatch.setattr(type(win), "_show_math", lambda self, made, folder, saved=None: shown.append(folder))
    monkeypatch.setattr(type(win), "_saved_box", lambda self, *a, **k: None)
    win.fill_all_jobs()
    wait(win, lambda: not win.filling and ledger_for(win.s).invoices())
    assert len(shown) == 1 and shown[0].name == "Math" and shown[0].parent.parent == tmp_path / "out", shown


# ------------------------------------------------------------------ the window

def test_a_firm_of_another_day_keeps_its_tooltip_with_its_side(qt, tmp_path, make_window, monkeypatch):
    from helpers import transcript_pdf
    from minute_filler.gui.main_window import VISITOR_TIP, _VISITOR
    from test_job_list import wait
    monkeypatch.setattr(qt.QMessageBox, "exec", lambda self: 0)
    s = pat_settings()
    s.use_ai = s.open_after = False
    s.output_dir, s.records_dir = str(tmp_path / "out"), str(tmp_path / "records")
    s.outputs = ["agreement", "invoice"]
    win = make_window(s)
    win.add_files([str(transcript_pdf(tmp_path / "d1.pdf", pages=30)),
                   str(transcript_pdf(tmp_path / "d2.pdf", pages=24, date="June 4, 2026"))], batch=True)
    wait(win, lambda: len(win.jobs) == 2 and all(j.docs for j in win.jobs))
    d1, d2 = sorted(win.jobs, key=lambda j: j.case.get("dates"))
    win._select_job(d2)  # (first: it writes the table back into the day shown before)
    d1.case.attorneys = [att("Dana Smith", "Smith Law", "Plaintiff"), att("Kim Park", "Park PLLC", "Defendant")]
    d2.case.attorneys = [att("Dana Smith", "Smith Law", "Plaintiff")]
    win._show_job()
    rows = [r for r in range(win.att.rowCount()) if win.att.item(r, 0).data(_VISITOR)]
    assert rows
    tip = win.att.item(rows[0], 1).toolTip()
    assert tip.startswith(VISITOR_TIP) and tip.endswith("A defense firm"), tip
    ours = next(r for r in range(win.att.rowCount()) if not win.att.item(r, 0).data(_VISITOR))
    assert win.att.item(ours, 1).toolTip() == "A plaintiffs' firm"


def test_the_square_has_a_lighter_edge_for_the_dark_theme(qt):
    from PySide6.QtGui import QImage
    from minute_filler.firm_marks import DEFENSE, Mark
    from minute_filler.gui.badges import badge_icon, tag_color
    navy = Mark(DEFENSE, 3)
    img = badge_icon(navy, 14).pixmap(14, 14).toImage().convertToFormat(QImage.Format_ARGB32)
    assert img.pixelColor(7, 7).name() == tag_color(navy).name()
    assert img.pixelColor(7, 0).lightness() > tag_color(navy).lightness()


def test_the_square_keeps_its_colour_on_a_selected_row(qt):
    """Qt tints an icon on a selected row unless it is given one for that: the square's colour is all it says."""
    from PySide6.QtGui import QIcon
    from minute_filler.firm_marks import PLAINTIFF, Mark
    from minute_filler.gui.badges import badge_icon, tag_color
    orange = Mark(PLAINTIFF, 0)
    selected = badge_icon(orange, 14).pixmap(14, 14, QIcon.Selected).toImage()
    assert selected.pixelColor(7, 7).name() == tag_color(orange).name()
