"""The run sheet: takes read from the reporters' initials on each transcript page, written to Excel, and added
to the case's run sheet when there is one, ours or one made elsewhere (whose original is kept as a "(before
YinIt)" copy, or the old "(before DjinnIt)" one). Its Witness columns get the running heads as printed, or are
blank; a new run sheet goes in the case's own folder, and the case's run sheet is found there, where the records
say or in the old run sheets folder. Takes the user corrected (takes.moved_end, batch.Doc.fixes) and the records
that keep them. Who transcribed what of several cases (batch.takes_of). Fictional names and cases."""
from copy import copy
from datetime import date

import pymupdf
import pytest
from openpyxl import Workbook, load_workbook

from minute_filler.batch import (case_dir_for, case_folders, fill_jobs, group, origin_fixes, out_dir_for, read_docs,
                                 takes_of)
from minute_filler.deliver import generate
from minute_filler.runsheet import (FIRST, Row, RunSheetOpts, add_takes, choose, find_sheets, old_folder, read_info,
                                    recorded_folders, run_sheet_summary, same_name, sheet_folders)
from minute_filler.takes import (PageMark, find_takes, fix_marks, initials_of, is_index_page, moved_end,
                                 page_numbers, reporter_label, scan_page, scan_pdf, title_reporters)

from helpers import pat_settings

TITLE_1 = """ 1    SUPREME COURT OF THE STATE OF NEW YORK
      COUNTY OF QUEENS:  CIVIL TERM:  PART 14
 2    JANE ROE,                          Index No. {index}
 3                       Plaintiff,
 4             -against-
 5    {defendant},
 6                       Defendant.
 7    {day}
 8    B E F O R E :  HONORABLE MARIA T. ALVAREZ
 9    A P P E A R A N C E S :
10    ADVOCATE LAW GROUP
      Attorneys for the Plaintiff
11    BY: SAMUEL R. ADVOCATE, ESQ.
12    (Title continues on next page.)"""

TITLE_2 = """ 1    COUNSEL & WARD, LLP
      Attorneys for the Defendant
 2    BY: ALEX B. COUNSEL, ESQ.
 3
 4                     PAT REPORTER
 5                     DANA SMITH
                       Senior Court Reporters"""

SWORN = """ 1              THE CLERK:  Please state your name.
 2           J O H N   D O E, called as a witness by and on behalf
 3    of the Plaintiff, after having been first duly sworn,
 4    testified as follows:"""

EXCUSED = """ 1              THE COURT:  Thank you.  You may step down.
 2              (Whereupon, the witness stepped down.)"""

TALK = """ 1              THE COURT:  Good morning.
 2              MS. ADVOCATE:  Good morning, Your Honor.
 3              THE COURT:  Let's continue."""

# (page number, initials, running head, text): two title pages by Pat, three pages by Dana with a witness
# sworn in, three by Pat where he steps down, then Dana's take with the next witness - its last page
# lacks the initials - and the word index after the transcript. (Who is on the stand is not read from the
# pages: the run sheet's Witness columns get the running heads as printed.)
DAY_1 = [
    (310, "pr", "", TITLE_1), (311, "pr", "", TITLE_2),
    (312, "ds", "Proceedings", TALK), (313, "ds", "Proceedings", SWORN), (314, "ds", "J. Doe - Plaintiff - Direct", TALK),
    (315, "pr", "J. Doe - Plaintiff - Direct", TALK), (316, "pr", "J. Doe - Plaintiff - Cross", EXCUSED),
    (317, "pr", "Proceedings", TALK),
    (318, "ds", "M. Lee - Defendant - Direct", TALK), (319, "", "M. Lee - Defendant - Direct", TALK),
    (None, "", "Roe v. Poe", "Min-U-Script\nabout (2) 312:4;315:9\nafter (1) 313:3"),
]


def transcript(path, pages=DAY_1, index="712345/2021", defendant="SAM POE", day="June 3, 2026"):
    """Writes a fictional transcript PDF: one page per (number, initials, head, text) in pages, the initials
    at the foot. index, defendant and day fill the title page. Returns path."""
    doc =pymupdf.open()
    for number, initials, head, text in pages:
        page = doc.new_page(width=612, height=792)
        if head:
            page.insert_text((150, 36), head, fontsize=10, fontname="cour")
        if number is not None:
            page.insert_text((300, 63), str(number), fontsize=10, fontname="cour")
        text = text.format(index=index, defendant=defendant, day=day)
        for i, line in enumerate(text.splitlines()):
            page.insert_text((40, 100 + 22 * i), line, fontsize=10, fontname="cour")
        if initials:
            page.insert_text((520, 740), initials, fontsize=10, fontname="cour")
    doc.save(path)
    return path


@pytest.fixture
def s(tmp_path):
    s = pat_settings()
    s.output_dir = str(tmp_path / "out")
    s.records_dir = str(tmp_path / "records")
    s.runsheet_dir = str(tmp_path / "sheets")
    return s


def job_of(path, s):
    """The one job made from a transcript."""
    docs, errors = read_docs([str(path)], s)
    assert not errors
    job, = group(docs, s)
    return job


def sheet_rows(path):
    """The rows of our run sheet's takes, columns A to K."""
    ws =load_workbook(path)["Run Sheet"]
    return [[ws.cell(r, c).value for c in range(1, 12)] for r in range(FIRST, ws.max_row + 1)]


# ------------------------------------------------------------------ reading the pages

def test_takes_follow_the_initials(tmp_path):
    with pymupdf.open(transcript(tmp_path / "t.pdf")) as doc:
        marks = scan_pdf(doc)
    assert len(marks) == 10  # the word index is not part of the transcript
    assert [m.number for m in marks] == list(range(310, 320)) and marks[0].initials == "pr" and not marks[9].initials
    takes = find_takes(marks, title_count=2)
    assert [(t.initials, t.start, t.pages, t.title_only, t.at) for t in takes] == \
        [("pr", 310, 2, True, 0), ("ds", 312, 3, False, 2), ("pr", 315, 3, False, 5), ("ds", 318, 2, False, 8)]
    assert [(t.heads, t.last_head) for t in takes] == [
        ([], ""), (["Proceedings", "J. Doe - Plaintiff - Direct"], "J. Doe - Plaintiff - Direct"),
        (["J. Doe - Plaintiff - Direct", "J. Doe - Plaintiff - Cross", "Proceedings"], "Proceedings"),
        (["M. Lee - Defendant - Direct"], "M. Lee - Defendant - Direct")]


def test_names_for_the_initials():
    names = title_reporters("BY: ALEX B. COUNSEL, ESQ.\n\n 4   PAT REPORTER\n 5   DANA SMITH\n   Senior Court Reporters")
    assert names == ["PAT REPORTER", "DANA SMITH"]
    assert initials_of("Pat Q. Reporter") == {"pr", "pqr"}
    assert reporter_label("ds", names, {}, "Pat Reporter") == "Dana"
    assert reporter_label("pr", [], {}, "Pat Reporter") == "Pat"       # the user's own, from My info
    assert reporter_label("px", [], {}, "Pat Reporter", own_initials="PX") == "Pat"
    assert reporter_label("ds", names, {"ds": "Dana S."}, "Pat Reporter") == "Dana S."  # Settings win
    assert reporter_label("kl", names, {}, "Pat Reporter") == "KL"      # nobody known: the initials


def test_the_word_index_is_not_billed(tmp_path, s):
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    assert job.docs[0].ing.page_count == 11 and job.transcript_pages() == 10
    job.page_basis[job.docs[0].key()] = "*"  # the whole transcript (Pat wrote 5 of its pages: see test_own_pages)
    assert job.invoice_pages() == 10
    note = job.docs[0].regex.fields["est_pages"][0].note
    assert "transcript pages 310-319" in note and "not counting 1 page(s) after the transcript (word index)" in note


# ------------------------------------------------------------------ the run sheet

def test_a_new_run_sheet(tmp_path, s):
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    opts = job.runsheet_opts(s)
    paths = generate(job.case, s, tmp_path / "out", ["runsheet"], runsheet=opts)
    assert paths == [opts.path] and opts.created and opts.added == 4
    assert opts.path.parent == tmp_path / "sheets" and opts.path.name == "June 2026 712345-2021 Jane Roe v. Sam Poe - Run Sheet.xlsx"
    rows = sheet_rows(opts.path)
    # Witness Start and End: the first and last heads printed in the take (Settings.runsheet_witness)
    assert [(r[0].date(), r[2], r[5], r[8], r[9], r[10]) for r in rows] == [
        (date(2026, 6, 3), "Pat", 2, None, None, "title page only"),
        (date(2026, 6, 3), "Dana", 3, "Proceedings", "J. Doe - Plaintiff - Direct", None),
        (date(2026, 6, 3), "Pat", 3, "J. Doe - Plaintiff - Direct", "Proceedings", None),
        (date(2026, 6, 3), "Dana", 2, "M. Lee - Defendant - Direct", "M. Lee - Defendant - Direct", None)]
    assert rows[0][6] == 310 and rows[1][6] == "=H5+1"  # the pages carry on from the row above
    assert rows[1][3].startswith('=IF(A6="","",IF(K6="title page only",0.5,COUNTIFS(')
    wb = load_workbook(opts.path)
    assert wb["YinIt"].sheet_state == "hidden" and [r.value for r in wb["By Reporter"]["A"]][:3] == \
        ["Reporter", "Pat", "Dana"]
    info = read_info(opts.path)
    assert info.ours and info.index_nos == ["712345/2021"] and info.case_name == "Jane Roe v. Sam Poe"
    assert "Started the run sheet" in run_sheet_summary(opts) and "4 takes" in run_sheet_summary(opts)


def test_the_next_day_is_added_in_order(tmp_path, s):
    day2 = [(320, "ds", "", TITLE_1), (321, "ds", "Proceedings", TALK), (322, "pr", "Proceedings", TALK)]
    second = job_of(transcript(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", day2, day="June 4, 2026"), s)
    first = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    a = second.runsheet_opts(s)
    add_takes(second.case, a, s, [])
    b = first.runsheet_opts(s)
    add_takes(first.case, b, s, [])  # "ask" with nobody to ask: the same index number is added to
    assert b.path == a.path and not b.created and b.added == 4
    rows = sheet_rows(a.path)
    assert [(r[0].day, r[2], r[5]) for r in rows] == [(3, "Pat", 2), (3, "Dana", 3), (3, "Pat", 3), (3, "Dana", 2),
                                                       (4, "Dana", 2), (4, "Pat", 1)]
    assert rows[4][6] == "=H8+1"  # day 2 starts at page 320, right after day 1's 319
    # the same transcript again: nothing is added twice
    again = first.runsheet_opts(s)
    add_takes(first.case, again, s, [])
    assert (again.added, again.skipped) == (0, 4) and len(sheet_rows(a.path)) == 6
    assert "already has these takes" in run_sheet_summary(again)


def test_another_index_number_of_the_same_trial(tmp_path, s):
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    add_takes(job.case, job.runsheet_opts(s), s, [])
    other = job_of(transcript(tmp_path / "Transcript 6-5-2026 Roe v Poe.pdf", index="700999/2022",
                              day="June 5, 2026"), s)
    found = find_sheets(other.case, [tmp_path / "sheets"])
    assert [f.why for f in found] == ["name"] and found[0].reason() == "same case name, index 712345/2021"
    assert choose(other.case, s, [tmp_path / "sheets"], None) is None  # "ask" with nobody to ask: a new one
    s.runsheet_existing = "add"
    assert choose(other.case, s, [tmp_path / "sheets"], None) == found[0].path
    opts = other.runsheet_opts(s)
    add_takes(other.case, opts, s, [])
    assert opts.path == found[0].path and read_info(opts.path).index_nos == ["712345/2021", "700999/2022"]
    unrelated = job_of(transcript(tmp_path / "Transcript 6-5-2026 Roe v Hill.pdf", index="700111/2022",
                                  defendant="TERRY HILL"), s)
    assert find_sheets(unrelated.case, [tmp_path / "sheets"]) == []
    assert same_name("Jane Roe, as Administrator of the Estate of Sam Roe v. Sam Poe, M.D., et al.",
                     "September 2026 712345-2021 ROE v Poe et al. - Page Sheet")


def test_adding_to_a_run_sheet_made_elsewhere(tmp_path, s):
    """A run sheet downloaded from Google Sheets: rows go at its end, its formulas are copied down."""
    path = tmp_path / "sheets" / "June 2026 712345-2021 ROE v Poe - Page Sheet.xlsx"
    path.parent.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(["Write here", "Don't write here", "Write here", "", "", "Write here", "", "", "", "", ""])
    ws.append(["Date", "Weekday", "Reporter", "Day's Take", "Running Take", "Pages written", "Starting Page No.",
               "Ending Page no.", "Witness Start", "Witness End", "Note"])
    ws.append([date(2026, 6, 2), '=TEXT(A3,"dddd")', "Pat", 1, 1, 10, 300, "=G3+F3-1", None, None, "opening"])
    ws.append([None, '=IF(A4="","-",TEXT(A4,"dddd"))', None, None, None, None, "=H3+1", "=IF(F4=\"\",FALSE,G4+F4-1)"])
    wb.save(path)
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    found = find_sheets(job.case, [path.parent])
    assert [(f.path, f.ours, f.why) for f in found] == [(path, False, "index")]
    opts = job.runsheet_opts(s)
    add_takes(job.case, opts, s, [])
    assert opts.path == path and opts.added == 4
    ws = load_workbook(path).active
    assert [ws.cell(r, 3).value for r in range(3, 8)] == ["Pat", "Pat", "Dana", "Pat", "Dana"]
    assert ws["A4"].value.date() == date(2026, 6, 3) and ws["F4"].value == 2
    assert ws["G4"].value == "=H3+1" and ws["G5"].value == "=H4+1"  # their formula, copied down
    assert ws["B5"].value == '=IF(A5="","-",TEXT(A5,"dddd"))' and ws["K3"].value == "opening"
    assert (path.parent / "June 2026 712345-2021 ROE v Poe - Page Sheet (before YinIt).xlsx").exists()


def test_a_run_sheet_djinnit_kept_a_copy_of_isnt_copied_again(tmp_path, s):
    """The original of a run sheet the old app added to is its "(before DjinnIt)" copy: a "(before YinIt)" copy
    of the sheet as DjinnIt left it would be a second, misleading one."""
    path = tmp_path / "sheets" / "June 2026 712345-2021 ROE v Poe - Page Sheet.xlsx"
    path.parent.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Weekday", "Reporter", "Day's Take", "Running Take", "Pages written", "Starting Page No.",
               "Ending Page no.", "Witness Start", "Witness End", "Note"])
    ws.append([date(2026, 6, 2), None, "Pat", 1, 1, 10, 300, 309, None, None, ""])
    wb.save(path)
    kept = path.with_name(f"{path.stem} (before DjinnIt).xlsx")
    wb.save(kept)
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    opts = job.runsheet_opts(s)
    opts.target = str(path)
    add_takes(job.case, opts, s, [])
    assert opts.path == path and opts.added
    assert kept.exists() and not path.with_name(f"{path.stem} (before YinIt).xlsx").exists()


def test_a_batch_puts_every_day_on_one_run_sheet(tmp_path, s):
    folder = tmp_path / "in"
    folder.mkdir()
    transcript(folder / "Transcript 6-3-2026 Roe v Poe.pdf")
    transcript(folder / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1), (321, "pr", "", TALK)],
               day="June 4, 2026")
    (folder / "note.txt").write_text("Please send the minutes in Smith v Jones, Index No. 712222/2024, 5/22/2026.")
    docs, _ = read_docs([str(p) for p in sorted(folder.iterdir())], s)
    jobs = group(docs, s)
    assert len(jobs) == 3
    done = fill_jobs(jobs, s, outputs=["runsheet"])
    sheets = list((tmp_path / "sheets").glob("*.xlsx"))
    assert len(sheets) == 1 and done == sheets and len(sheet_rows(sheets[0])) == 6
    smith = next(j for j in jobs if "Smith" in j.title())
    assert smith.error == "no run sheet: no transcript PDF among the inputs" and not smith.saved


def test_a_run_sheet_goes_in_the_cases_own_folder(tmp_path, s):
    """Run sheets went to Documents\\YinIt Run Sheets, away from each case's other files. They go in the case's
    own folder now (the user, 2026-10-23: "they are particular to each case"), every day of a trial on one, and a
    later day finds it there before its own folder is chosen (batch.case_dir_for)."""
    s.runsheet_dir = ""
    folder = tmp_path / "in"
    folder.mkdir()
    transcript(folder / "Transcript 6-3-2026 Roe v Poe.pdf")
    transcript(folder / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1), (321, "pr", "", TALK)],
               day="June 4, 2026")
    docs, _ = read_docs([str(p) for p in sorted(folder.iterdir())], s)
    done = fill_jobs(group(docs, s), s, outputs=["runsheet"])
    case = tmp_path / "out" / "712345-2021"
    assert [p.parent for p in done] == [case / "Trial Runs"] and len(sheet_rows(done[0])) == 6
    assert not old_folder().exists()
    third = job_of(transcript(folder / "Transcript 6-5-2026 Roe v Poe.pdf", [(330, "pr", "", TITLE_1),
                                                                          (331, "ds", "", TALK)], day="June 5, 2026"), s)
    assert case_dir_for(third, s) == case and out_dir_for(third, s) == tmp_path / "out"  # (not chosen yet)
    assert [f.path for f in find_sheets(third.case, sheet_folders(s, case_dir_for(third, s), []))] == done
    opts = third.runsheet_opts(s)
    generate(third.case, s, case, ["runsheet"], runsheet=opts)
    assert opts.path == done[0] and not opts.created and len(sheet_rows(done[0])) == 8


def test_a_run_sheet_in_the_old_run_sheets_folder_is_still_added_to(tmp_path, s):
    """A trial whose run sheet an earlier version started in Documents\\YinIt Run Sheets keeps it: the next day's
    takes go on it there, not on a new one in the case's Trial Runs folder."""
    first = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    s.runsheet_dir = str(old_folder())  # (where it was kept)
    a = first.runsheet_opts(s)
    add_takes(first.case, a, s, [])
    s.runsheet_dir = ""
    second = job_of(transcript(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1),
                                                                             (321, "pr", "", TALK)], day="June 4, 2026"), s)
    b = second.runsheet_opts(s)
    generate(second.case, s, case_dir_for(second, s), ["runsheet"], runsheet=b)
    assert b.path == a.path and b.path.parent == old_folder() and not b.created and len(sheet_rows(a.path)) == 6
    assert not list(tmp_path.rglob("out/**/*.xlsx"))


def test_a_run_sheet_in_the_case_folder_itself_is_still_added_to(tmp_path, s):
    """A new run sheet goes in the case folder's Trial Runs folder (the user, 2026-10-26); one 2.13 saved in the
    case's folder itself is still found there and gets the next day's takes, with no second one in Trial Runs."""
    first = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    case = case_dir_for(first, s)
    s.runsheet_dir = str(case)  # (where 2.13 put it)
    a = first.runsheet_opts(s)
    add_takes(first.case, a, s, [])
    assert a.path.parent == case
    s.runsheet_dir = ""
    second = job_of(transcript(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1),
                                                                             (321, "pr", "", TALK)], day="June 4, 2026"), s)
    b = second.runsheet_opts(s)
    generate(second.case, s, case_dir_for(second, s), ["runsheet"], runsheet=b)
    assert b.path == a.path and not b.created and len(sheet_rows(a.path)) == 6
    assert not (case / "Trial Runs").exists()


def test_a_scan_gets_one_row(tmp_path, s):
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    job.docs[0].ing.marks = []  # as for a scanned transcript: no text to find the initials in
    rows = job.runsheet_opts(s).rows
    # the row has the pages the invoice bills: the transcript's own, the word index on the 11th page left out
    assert len(rows) == 1 and rows[0].pages == job.transcript_pages() == 10
    assert rows[0].note == "reporter's initials not found"


def test_open_in_excel(tmp_path, s, monkeypatch):
    """A run sheet open in Excel can't be replaced: the message says to close it, and no half-saved copy stays."""
    job =job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    opts = job.runsheet_opts(s)
    add_takes(job.case, opts, s, [])
    import minute_filler.runsheet as rs

    def locked(src, dst):
        raise PermissionError(13, "The process cannot access the file")
    monkeypatch.setattr(rs.os, "replace", locked)
    again = RunSheetOpts([Row(date(2026, 6, 4), "Pat", 5, 320)], str(opts.path))
    with pytest.raises(PermissionError, match="close it in Excel"):
        add_takes(job.case, again, s, [])
    assert len(sheet_rows(opts.path)) == 4 and not list(opts.path.parent.glob("*.saving.xlsx"))


# ------------------------------------------------------------------ found in the run sheet sweep

D3, D4 = date(2026, 6, 3), date(2026, 6, 4)


def test_the_witness_columns_are_the_heads_as_printed_or_blank(tmp_path, s):
    """Witness Start and End were the witness read from the pages, which went wrong with heads laid out another way
    ("DIRECT - DEFENDANT - SAM POE"): they are the first and last heads printed in the take now, or blank
    for the user to type in (Settings -> Run sheet: the user, 2026-10-23). A character a workbook can't hold (a
    NUL in the PDF's text) is left out: once in a cell, Excel refused the file."""
    pages = DAY_1[:2] + [(312, "ds", "PROCEEDINGS", TALK), (313, "ds", "DIRECT - DEFENDANT - SAM POE", TALK),
                         (314, "pr", "CROSS - DEFENDANT - SAM POE", TALK), (315, "pr", "", TALK)]
    job = job_of(transcript(tmp_path / "t.pdf", pages), s)
    opts = job.runsheet_opts(s)
    add_takes(job.case, opts, s, [])
    assert [(r[8], r[9]) for r in sheet_rows(opts.path)] == [
        (None, None), ("PROCEEDINGS", "DIRECT - DEFENDANT - SAM POE"),
        ("CROSS - DEFENDANT - SAM POE", "CROSS - DEFENDANT - SAM POE")]  # (the last page has no head)
    s.runsheet_witness = "blank"
    assert {(r.witness_start, r.witness_end) for r in job.runsheet_opts(s).rows} == {("", "")}
    take, = find_takes([PageMark(1, "pr", "PROCEEDINGS\x00 "), PageMark(2, "pr", "\x07")], title_count=0)
    assert take.heads == ["PROCEEDINGS"] and take.last_head == "PROCEEDINGS"


def test_a_take_corrected_by_the_user(tmp_path, s):
    """A page's initials misread (or missing) put pages in the wrong take; the user moves where a take ends
    (Who transcribed what -> Correct where this take ends): the takes, Yours, the pages billed and the run sheet
    follow, and the records of the files made from then on keep the correction (job_origin), for the documents
    read again when the job is opened again (origin_fixes; the window gives them back: test_gui_runsheet)."""
    from minute_filler.batch import job_origin
    job = job_of(transcript(tmp_path / "t.pdf"), s)
    doc = job.docs[0]
    assert job.your_pages() == 5 and [(r.reporter, r.start, r.pages, r.at) for r in job.runsheet_opts(s).rows] == [
        ("Pat", 310, 2, 0), ("Dana", 312, 3, 2), ("Pat", 315, 3, 5), ("Dana", 318, 2, 8)]
    marks = doc.marks()
    assert page_numbers([PageMark(None), PageMark(5), PageMark(None)]) == [None, 5, 6]
    # Dana's first take really ends on 315 (its last page misread as Pat's): one page moves to Dana
    fix = moved_end(marks, 2, 5, title_count=2)
    assert fix == {5: "ds"}
    doc.fixes = {**doc.fixes, **fix}
    assert [(r.reporter, r.start, r.pages) for r in job.runsheet_opts(s).rows] == [
        ("Pat", 310, 2), ("Dana", 312, 4), ("Pat", 316, 2), ("Dana", 318, 2)]
    assert job.your_pages() == 4 and job.reporter_pages() == {"pr": 4, "ds": 6}
    assert doc.ing.marks[5].initials == "pr"  # (the pages as read are left as they were)
    # moved earlier: its pages after the new end go to the next take, which starts sooner
    assert moved_end(doc.marks(), 2, 3, title_count=2) == {4: "pr", 5: "pr"}
    # Pat's second take reaching to the end swallows Dana's last one
    assert moved_end(doc.marks(), 6, 9, title_count=2) == {8: "pr", 9: "pr"}
    for at, end in ((8, 7), (8, 10), (1, 5)):  # the last take earlier, past the last page, no take starting there
        with pytest.raises(ValueError):
            moved_end(doc.marks(), at, end, title_count=2)
    assert moved_end(doc.marks(), 8, 9, title_count=2) == {}  # (where it ends already)
    assert fix_marks(marks, None) is marks
    # kept with the records, read back as the document's corrections
    origin = job_origin(job)
    assert origin["job"]["page_fixes"] == {doc.key(): {"5": "ds"}} and origin_fixes(origin) == {doc.key(): {5: "ds"}}
    assert origin_fixes({"page_fixes": {"a.pdf": {"5": "ds", "x": "pr", "6": 3}}, "joint": True}) == {"a.pdf": {5: "ds"}}
    assert origin_fixes({"job": {"page_fixes": "bad"}}) == {} and origin_fixes({}) == {}
    assert origin_fixes({"page_fixes": {"a.pdf": {"²": "ds"}}}) == {"a.pdf": {}}  # (a digit int() can't read)
    assert "page_fixes" not in job_origin(job_of(tmp_path / "t.pdf", s))["job"]  # (none: the record as it always was)


def test_extending_a_take_leaves_the_next_takes_unsigned_pages_where_they_are(tmp_path, s):
    """Pat's take 315-317 made to end on 318: page 319, with no initials of its own, followed page 318 to Pat (Dana's
    last take gone, Yours 5 -> 7). It stays Dana's: pinned to the take it is in."""
    job = job_of(transcript(tmp_path / "t.pdf"), s)
    doc = job.docs[0]
    fix = moved_end(doc.marks(), 5, 8, title_count=2)
    assert fix == {8: "pr", 9: "ds"}
    doc.correct(fix)
    assert [(r.reporter, r.start, r.pages) for r in job.runsheet_opts(s).rows][2:] == [("Pat", 315, 4), ("Dana", 319, 1)]
    assert job.your_pages() == 6 and job.reporter_pages() == {"pr": 6, "ds": 4}


def test_a_correction_that_leaves_one_reporter_gives_them_the_title_pages_everywhere(tmp_path, s):
    """Whose pages... gave the title pages (no initials) to Pat; a correction then left Dana alone on the transcript.
    The bill counted the title pages as Dana's (one reporter: theirs) while the takes still said Pat's."""
    pages = [(310, "", "", TITLE_1), (311, "", "", TITLE_2), (312, "ds", "Proceedings", TALK),
             (313, "ds", "Proceedings", TALK), (314, "pr", "Proceedings", TALK), (315, "pr", "Proceedings", TALK)]
    job = job_of(transcript(tmp_path / "t.pdf", pages), s)
    doc = job.docs[0]
    job.front_owner = {doc.key(): "pr"}
    assert [(r.reporter, r.pages) for r in job.runsheet_opts(s).rows] == [("Pat", 2), ("Dana", 2), ("Pat", 2)]
    doc.correct(moved_end(doc.marks(), 2, 5, title_count=2, front="pr"))
    assert doc.reporters() == ["ds"] and job.reporter_pages() == {"ds": 6}
    assert [(r.reporter, r.start, r.pages) for r in job.runsheet_opts(s).rows] == [("Dana", 310, 6)]


def test_a_head_that_starts_with_equals_is_text_in_the_workbooks(tmp_path, s):
    """A running head printed "== DIRECT EXAMINATION ==" went into the run sheet as a (broken) formula, and Save a
    copy's CSV let Excel read "- CROSS -" as one."""
    from minute_filler.gui.main_window import save_takes
    pages = DAY_1[:2] + [(312, "ds", "== DIRECT EXAMINATION ==", TALK), (313, "pr", "- CROSS -", TALK)]
    job = job_of(transcript(tmp_path / "t.pdf", pages), s)
    opts = job.runsheet_opts(s)
    add_takes(job.case, opts, s, [])
    cell = load_workbook(opts.path)["Run Sheet"].cell(FIRST + 1, 9)
    assert cell.value == "== DIRECT EXAMINATION ==" and cell.data_type == "s"
    more = job_of(transcript(tmp_path / "t2.pdf", [(320, "ds", "", TITLE_1), (321, "ds", "PROCEEDINGS", TALK)],
                             day="June 4, 2026"), s)
    add_takes(more.case, more.runsheet_opts(s), s, [])  # (written again: still text, not taken for a formula typed)
    cell = load_workbook(opts.path)["Run Sheet"].cell(FIRST + 1, 9)
    assert cell.value == "== DIRECT EXAMINATION ==" and cell.data_type == "s"
    rows = job.runsheet_opts(s).rows
    saved = save_takes(rows, tmp_path / "copy.xlsx")
    assert not [c for r in load_workbook(saved).active.iter_rows() for c in r if c.data_type == "f"]
    text = save_takes(rows, tmp_path / "copy.csv", csv_file=True).read_text(encoding="utf-8-sig")
    assert "'== DIRECT EXAMINATION ==" in text and "'- CROSS -" in text


def test_a_page_number_on_the_heads_line_is_not_part_of_the_head():
    """A transcript printing "101 PROCEEDINGS" at the top listed every page's head once each in the card."""
    take, = find_takes([PageMark(101, "pr", "101 PROCEEDINGS"), PageMark(102, "pr", "PROCEEDINGS 102"),
                        PageMark(103, "pr", "PROCEEDINGS")], title_count=0)
    assert take.heads == ["PROCEEDINGS"] and take.last_head == "PROCEEDINGS"
    # any other number is the head's own: only the page's number goes
    take, = find_takes([PageMark(7, "pr", "PROCEEDINGS - DAY 2"), PageMark(8, "pr", "VOLUME 3"),
                        PageMark(9, "pr", "JANE ROE v. SAM POE - JUNE 3, 2026")], title_count=0)
    assert take.heads == ["PROCEEDINGS - DAY 2", "VOLUME 3", "JANE ROE v. SAM POE - JUNE 3, 2026"]


def test_two_cases_heard_the_same_day_are_kept_apart_in_who_transcribed_what(tmp_path, s):
    """Days of several cases ticked: two cases of the same date had their takes mixed by page number, with nothing
    to tell them apart."""
    from minute_filler.gui.main_window import take_cells
    a = job_of(transcript(tmp_path / "a.pdf"), s)
    b = job_of(transcript(tmp_path / "b.pdf", index="700111/2022", defendant="SAM HILL"), s)
    rows = takes_of([a, b], s)
    assert [r.case for r in rows] == ["Jane Roe v. Sam Poe"] * 4 + ["Jane Roe v. Sam Hill"] * 4
    assert take_cells(rows[0])[0] == "Jane Roe v. Sam Poe · 6/3/2026"
    assert {r.case for r in takes_of([a], s)} == {""}


def test_a_later_day_finds_the_trials_run_sheet_where_the_records_say(tmp_path, s):
    """Run sheets go with each case's files: day 2's transcript in another folder (Save to blank: next to it) looked
    only in its own case folder and started a second run sheet. The records say where day 1's is. And with a file of
    the case folder's name in Save to, the case's files go beside it ("712345-2021 (2)"), where day 2 looks too."""
    s.runsheet_dir, s.output_dir = "", ""
    for name in ("Day1", "Day2"):
        (tmp_path / name).mkdir()
    one = job_of(transcript(tmp_path / "Day1" / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    a = one.runsheet_opts(s)
    generate(one.case, s, case_dir_for(one, s), ["runsheet"], runsheet=a)
    assert a.path.parent == tmp_path / "Day1" / "712345-2021" / "Trial Runs"
    two = job_of(transcript(tmp_path / "Day2" / "Transcript 6-4-2026 Roe v Poe.pdf",
                            [(320, "ds", "", TITLE_1), (321, "pr", "", TALK)], day="June 4, 2026"), s)
    assert case_dir_for(two, s) == tmp_path / "Day2" / "712345-2021" and recorded_folders(two.case, s) == [a.path.parent]
    b = two.runsheet_opts(s)
    generate(two.case, s, case_dir_for(two, s), ["runsheet"], runsheet=b)
    assert b.path == a.path and not b.created and len(sheet_rows(a.path)) == 6
    s.output_dir = str(tmp_path / "out")
    (tmp_path / "out").mkdir()
    (tmp_path / "out" / "712345-2021").write_text("not a folder", encoding="utf-8")
    [f] = case_folders([two], s, ["runsheet"])
    assert case_dir_for(two, s) == f.path == tmp_path / "out" / "712345-2021 (2)"


def test_a_trials_records_keep_its_corrected_takes(tmp_path, s):
    """A trial's run sheet (like its forms and its joint invoice) is recorded for all its days at once: the takes
    corrected on them were in none of those records, and the trial opened again billed the pages as misread."""
    import json
    from minute_filler.deliver import ledger_for
    folder = tmp_path / "in"
    folder.mkdir()
    transcript(folder / "Transcript 6-3-2026 Roe v Poe.pdf")
    transcript(folder / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1), (321, "pr", "", TALK)],
               day="June 4, 2026")
    docs, _ = read_docs([str(p) for p in sorted(folder.iterdir())], s)
    jobs = group(docs, s)
    day1 = min(jobs, key=lambda j: j.docs[0].key())
    day1.docs[0].correct({5: "ds"})
    fill_jobs(jobs, s, outputs=["runsheet"])
    sheet, = ledger_for(s).activity(kind="runsheet")
    origin = json.loads(sheet.origin)
    assert origin["joint"] is True and origin_fixes(origin) == {day1.docs[0].key(): {5: "ds"}}


def test_who_transcribed_what_shows_the_running_heads_as_printed(tmp_path, s):
    """The card's last column is each take's running heads as printed, each once, comma-separated: no witness read
    from them (the user, 2026-10-23: heads with the examination first, "DIRECT - DEFENDANT - SAM POE", were
    shown as "Direct took the stand"). A title page's top line is no running head."""
    from minute_filler.gui.main_window import TAKE_COLS, take_cells
    pages = DAY_1[:2] + [(312, "ds", "PROCEEDINGS", TALK), (313, "ds", "DIRECT - DEFENDANT - SAM POE", TALK),
                         (314, "ds", "DIRECT - DEFENDANT - SAM POE", TALK),
                         (315, "pr", "CROSS - DEFENDANT - SAM POE", TALK), (316, "pr", "Proceedings", TALK),
                         (317, "pr", "", TALK), (318, "pr", "PROCEEDINGS", TALK)]
    with pymupdf.open(transcript(tmp_path / "t.pdf", pages)) as doc:
        takes = find_takes(scan_pdf(doc), title_count=2)
    assert [t.heads for t in takes] == [[], ["PROCEEDINGS", "DIRECT - DEFENDANT - SAM POE"],
                                        ["CROSS - DEFENDANT - SAM POE", "Proceedings"]]
    rows = job_of(tmp_path / "t.pdf", s).runsheet_opts(s).rows
    assert TAKE_COLS[-1] == "Headers"
    assert [take_cells(r)[-1] for r in rows] == ["", "PROCEEDINGS, DIRECT - DEFENDANT - SAM POE",
                                                 "CROSS - DEFENDANT - SAM POE, Proceedings"]
    take, = find_takes([PageMark(1, "pr", "SUPREME COURT OF THE STATE OF NEW YORK"), PageMark(2, "pr", "Proceedings")],
                       title_count=1)
    assert take.heads == ["Proceedings"]


def test_reporters_on_the_title_page():
    # an attorneys' block right above, a name on the reporter's line, letters after a name
    assert title_reporters("ADVOCATE LAW GROUP\nAttorneys for the Plaintiff\nDANA SMITH\nSenior Court Reporter") \
        == ["DANA SMITH"]
    assert title_reporters("BY: ALEX B. COUNSEL, ESQ.\n\nPAT REPORTER, Senior Court Reporter") == ["PAT REPORTER"]
    assert title_reporters("COUNSEL & WARD, LLP\n\nDANA SMITH, RPR\nOfficial Court Reporter") == ["DANA SMITH"]
    # the block of names ends at a blank line
    assert title_reporters("JOHN DOE\n\nDANA SMITH\nSenior Court Reporter") == ["DANA SMITH"]
    assert title_reporters("PAT REPORTER\nSenior Court Reporter\n\nDANA SMITH\nSenior Court Reporter") == \
        ["PAT REPORTER", "DANA SMITH"]


def _numbered_page(doc, top: str | None, foot: str | None, text: str = "Q. And then?") -> None:
    """A page with its number at the top or the foot, and the line numbers 1 and 25 alone at the left edge."""
    page = doc.new_page(width=612, height=792)
    if top:
        page.insert_text((534, 30), top, fontsize=10, fontname="cour")
    page.insert_text((40, 80), "1", fontsize=10, fontname="cour")      # the first line's number, in the top margin
    page.insert_text((80, 300), text, fontsize=10, fontname="cour")
    page.insert_text((40, 720), "25", fontsize=10, fontname="cour")    # the last line's, in the bottom margin
    if foot:
        page.insert_text((300, 770), foot, fontsize=10, fontname="cour")
    page.insert_text((520, 745), "pr", fontsize=10, fontname="cour")


def test_page_number_at_the_foot(tmp_path):
    doc = pymupdf.open()
    _numbered_page(doc, None, "312")   # at the foot only: the "1" at the top is the first line's number
    _numbered_page(doc, None, "313")
    _numbered_page(doc, "3", None)     # at the top (a transcript starting at page 1): the "25" is a line number
    marks = [scan_page(p) for p in doc]
    assert [m.number for m in marks] == [312, 313, 3]


def test_the_word_index_is_found_by_its_references():
    condensed = "Min-U-Script\nJane Roe v. Sam Poe\nPage 5\nQ. Were you there at 9:00?\nA. Yes."
    assert not is_index_page(condensed)  # a condensed transcript prints Min-U-Script on every page
    assert not is_index_page("7:00 7:00 8:00 9:38 10:04 7:00 7:00 8:00 7:00 9:45")  # times, not references
    assert is_index_page("about (2) 312:4;315:9\nafter (3) 313:3;316:12;318:20\nagain (3) 314:5;317:8;319:2")
    assert is_index_page("Min-U-Script\nabout (2) 312:4;315:9\nafter (1) 313:3")


def test_a_condensed_transcript_is_counted_whole(tmp_path):
    pages = [(n, i, h, t + "\nMin-U-Script") for n, i, h, t in DAY_1[:-1]] + [DAY_1[-1]]
    with pymupdf.open(transcript(tmp_path / "t.pdf", pages)) as doc:
        assert len(scan_pdf(doc)) == 10


def _add_day_1(tmp_path, s):
    """Starts the run sheet with DAY_1's four takes; returns the job and the run sheet's path."""
    job =job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    opts = job.runsheet_opts(s)
    add_takes(job.case, opts, s, [])
    return job, opts.path


def test_the_starting_page_follows_the_row_with_pages(tmp_path, s):
    job, path = _add_day_1(tmp_path, s)
    wb = load_workbook(path)
    wb["Run Sheet"]["A9"], wb["Run Sheet"]["C9"] = D3, "Kim"  # a row typed in, without pages
    wb.save(path)
    opts = RunSheetOpts([Row(D4, "Dana", 2, 320)])
    add_takes(job.case, opts, s, [])
    rows = sheet_rows(path)
    assert [r[2] for r in rows] == ["Pat", "Dana", "Pat", "Dana", "Kim", "Dana"]
    assert rows[5][6] == "=H8+1"  # page 320 follows Dana's 319 on row 8, not the empty row 9


def test_what_was_typed_on_our_run_sheet_stays(tmp_path, s):
    job, path = _add_day_1(tmp_path, s)
    wb = load_workbook(path)
    ws = wb["Run Sheet"]
    ws["F6"] = "=1+2"                                # Dana's 3 pages, typed as a sum
    ws["A9"], ws["C9"], ws["F9"] = "TBD", "Kim", 4   # a row with no date yet
    ws["L6"] = "=F6*2"                               # the user's own column
    ws["A8"] = "=see note"                           # a date typed as text that starts with "=" (not a formula)
    ws["A8"].data_type = "s"
    wb.save(path)
    opts = RunSheetOpts([Row(date(2026, 6, 2), "Dana", 10, 300)])  # the day before: goes first, the rest move down
    add_takes(job.case, opts, s, [])
    ws = load_workbook(path)["Run Sheet"]
    assert [ws.cell(r, 3).value for r in range(5, 11)] == ["Dana", "Pat", "Dana", "Pat", "Dana", "Kim"]
    assert ws["F7"].value == "=1+2" and ws["L7"].value == "=F7*2"  # moved with their row
    assert ws["F7"].data_type == ws["L7"].data_type == "f"           # (still formulas)
    assert ws["A10"].value == "TBD" and ws["F10"].value == 4         # the undated row stays where it was
    assert ws["A9"].value == "=see note" and ws["A9"].data_type == "s"  # (text, as typed)


def test_merged_cells_and_notes_beside_the_title(tmp_path, s):
    job, path = _add_day_1(tmp_path, s)
    wb = load_workbook(path)
    ws = wb["Run Sheet"]
    ws.merge_cells("A1:K1")
    ws["M1"] = "Ask Dana about 6/5"
    ws.merge_cells("M1:O1")
    wb.save(path)
    add_takes(job.case, RunSheetOpts([Row(D4, "Dana", 2, 320)]), s, [])
    ws = load_workbook(path)["Run Sheet"]
    assert ws["M1"].value == "Ask Dana about 6/5" and len(sheet_rows(path)) == 5


def test_by_reporter_with_the_tab_renamed_and_names_in_capitals(tmp_path, s):
    job, path = _add_day_1(tmp_path, s)
    wb = load_workbook(path)
    wb["Run Sheet"].title = "Roe's trial"
    wb.save(path)
    add_takes(job.case, RunSheetOpts([Row(D4, "DANA", 2, 320)]), s, [])
    by = load_workbook(path)["By Reporter"]
    assert [by.cell(r, 1).value for r in range(2, 4)] == ["Pat", "Dana"] and by["A4"].value == "Total"
    assert by["C3"].value.startswith("=SUMIF('Roe''s trial'!$C$5:$C$5000,A3,")


def test_nothing_new_leaves_the_run_sheet_alone(tmp_path, s, monkeypatch):
    job, path = _add_day_1(tmp_path, s)
    import minute_filler.runsheet as rs

    def locked(src, dst):
        raise PermissionError(13, "The process cannot access the file")
    monkeypatch.setattr(rs.os, "replace", locked)  # open in Excel: no matter, there is nothing to write
    again = job.runsheet_opts(s)
    add_takes(job.case, again, s, [])
    assert (again.added, again.skipped) == (0, 4)


def test_the_same_transcript_twice_in_one_go(tmp_path, s):
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    rows = job.runsheet_opts(s).rows
    opts = RunSheetOpts(rows + [copy(r) for r in rows])
    add_takes(job.case, opts, s, [])
    assert (opts.added, opts.skipped) == (4, 4) and len(sheet_rows(opts.path)) == 4


def test_a_locked_file_is_tried_again(tmp_path, s, monkeypatch):
    import minute_filler.runsheet as rs
    tries = []
    real = rs.os.replace

    def busy(src, dst):  # Dropbox keeping the file to itself for a moment
        tries.append(1)
        if len(tries) < 3:
            raise PermissionError(13, "The process cannot access the file")
        real(src, dst)
    monkeypatch.setattr(rs.os, "replace", busy)
    monkeypatch.setattr(rs.time, "sleep", lambda t: None)
    job, path = _add_day_1(tmp_path, s)
    assert len(tries) == 3 and len(sheet_rows(path)) == 4


def test_case_names_on_run_sheets():
    assert same_name("Jane Roe v. Sam Poe", "June 2026 712345-2021 Jane Roe v. Sam Poe - Run Sheet")
    assert same_name("Roe v. The City of New York", "June 2026 712345-2021 Roe v. The City of New York - Run Sheet")
    assert same_name("Jane Roe v. Acme Widget Corporation", "Roe v Acme Widget Corp. - Page Sheet")
    assert not same_name("Maria Garcia v. NYC Transit Authority", "Maria Rodriguez v. NYC Transit Authority")
    assert not same_name("Jane Roe v. Sam Poe", "Jane Smith v. Sam Jones")


def _their_sheet(path, headings, rows):
    """A run sheet made elsewhere (not by this app): a "Write here" row, the headings, then rows."""
    path.parent.mkdir(exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(["Write here"] * len(headings))
    ws.append(headings)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


HEADINGS = ["Date", "Weekday", "Reporter", "Day's Take", "Running Take", "Pages written", "Starting Page No.",
            "Ending Page no.", "Witness Start", "Witness End", "Note"]


def test_days_typed_in_ahead_on_a_run_sheet_made_elsewhere(tmp_path, s):
    path = _their_sheet(tmp_path / "sheets" / "June 2026 712345-2021 ROE v Poe - Page Sheet.xlsx", HEADINGS, [
        [date(2026, 6, 2), None, "Pat", 1, 1, 10, 300, "=G3+F3-1"],
        [D3], [D4]])  # the next days, typed in ahead
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    opts = job.runsheet_opts(s)
    add_takes(job.case, opts, s, [])
    ws = load_workbook(path).active
    assert [ws.cell(r, 3).value for r in range(3, 9)] == ["Pat", "Pat", None, "Dana", "Pat", "Dana"]
    assert ws["A4"].value.date() == D3 and ws["A5"].value.date() == D4  # June 4 is still there, for its day
    assert ws["G6"].value == 312  # not "=H5+1": row 5 has no pages yet


def test_headings_of_a_run_sheet_made_elsewhere(tmp_path, s):
    path = _their_sheet(tmp_path / "sheets" / "June 2026 712345-2021 ROE v Poe - Page Sheet.xlsx",
                        ["Trial Date", "Reporter", "No. of\nPages", "Starting Page", "Ending Page"],
                        [[date(2026, 6, 2), "Pat", 10, 300, 309]])
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    add_takes(job.case, job.runsheet_opts(s), s, [])
    ws = load_workbook(path).active
    assert [(ws.cell(r, 1).value.date(), ws.cell(r, 2).value, ws.cell(r, 3).value) for r in (4, 5)] == \
        [(D3, "Pat", 2), (D3, "Dana", 3)]
    bad = _their_sheet(tmp_path / "other" / "June 2026 712345-2021 ROE v Poe - Page Sheet.xlsx",
                       ["Date", "Reporter", "Starting Page", "Ending Page"], [])
    with pytest.raises(ValueError, match="no Pages column"):
        add_takes(job.case, RunSheetOpts(job.runsheet_opts(s).rows, str(bad)), s, [])


def test_what_is_not_a_run_sheet(tmp_path):
    from minute_filler.invoice import TEMPLATE
    assert read_info(TEMPLATE) is None  # the invoice spreadsheet's Page Log
    path = _their_sheet(tmp_path / "June 2026 712345-2021 ROE v Poe - Page Sheet.xlsx", HEADINGS,
                        [[date(2026, 6, 2), None, "Pat", 1, 1, 10, 300]])
    info = read_info(path)
    assert info.index_nos == ["712345/2021"]  # the date 2026-06-02 is not index number 2026/2006
