"""Runs the built program headless on made-up inputs and checks what it made.

After PyInstaller (from anywhere; PYTHONPATH is not needed):
    .venv/Scripts/python.exe .claude/skills/update-the-app/scripts/exe_check.py

It uses a temporary APPDATA and home folder, so the user's own settings, records and run sheets are not
touched (the folder is removed again when all is good); its settings say the user is Pat Reporter (initials
P.R.). The inputs are two days of one trial (a transcript each, June 3 and June 4, written by Pat and Dana
Smith: PR and DS at the foot of the pages) and an e-mail about another case. Expected: the built version, no
errors in the log; for each day a minute agreement, a MOFR, the run sheet and the joint invoice (with its math
saved next to it, marked "YinIt math", as --math asks; the command line puts the files in OUTDIR, the run sheet
too, and makes no case folders), the agreement and the MOFR being one each for both days
(Settings.forms_per_case: listed under each day, as the one joint invoice for both days is) and the invoice
billing Pat's own 10 of the 20 transcript pages, for one party, and has its amounts as fields (and no "Lock
fields" button, removed in 1.3.1); one run sheet with the 4 takes of each day; and --selftest says the
libraries only the window uses are in the build (the Records' Fuzzy and Regex searches, HEIC photos, the look
for a newer version, printing, The math as a PDF, the moving yin-yang, the Help search's word table and typo
fixing, the tick in ticked boxes), and so is the code the courthouse names by "module:name" (each output's maker,
rules and window panel, and the split rules billing firms at different speeds: the spec's hiddenimports). Last, the
build itself holds no records, invoices, settings or logs, and no file of the app's own that git doesn't track.
Exit code 1 if something is off.
"""
import csv
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[4]  # .claude/skills/update-the-app/scripts/ -> the repository
# Before minute_filler is imported (the made-up transcripts are made with it): its settings, records and log go
# to a folder of this check's own, never into the user's real records.db
run = Path(tempfile.mkdtemp(prefix="yinit-exe-check-"))
for _name in ("APPDATA", "USERPROFILE", "HOME"):
    os.environ[_name] = str(run / "import")
(run / "import").mkdir()
sys.path[:0] = [str(ROOT), str(ROOT / "tests")]
import pymupdf  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from minute_filler import __version__  # noqa: E402
from test_runsheet import transcript  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # (the report may hold arrows and ellipses)


def kind(name: str) -> str:
    """'Invoice 2026-0001 - Jane Roe v. Sam Poe.pdf' -> 'Invoice'; the Excel file is the run sheet, '... - the
    math.pdf' the math saved with an invoice (Settings.save_math)."""
    if name.endswith(" - the math.pdf"):
        return "The math"
    return "Run sheet" if name.endswith(".xlsx") else name.split(" - ")[0].split(" 20")[0]


exe = ROOT / "dist" / "YinItAgreementForm" / "YinItAgreementForm.exe"
if not exe.exists():
    shutil.rmtree(run, ignore_errors=True)
    sys.exit(f"PROBLEM: {exe} is not there - build it first (PyInstaller)")
src = run / "in"
src.mkdir()
transcript(src / "Transcript 6-3-2026 Roe v Poe.pdf")  # two reporters, a witness, a word index at the end
transcript(src / "Transcript 6-4-2026 Roe v Poe.pdf", day="June 4, 2026")  # the next day: one invoice for both
(src / "email.txt").write_text("From: someone@examplefirm.com\nSubject: minutes\n\nPlease send the minutes in "
                               "Smith v Jones, Index No. 712222/2024, 5/22/2026, Judge Lopez. About 40 pages.")
# who the user is: the transcripts were written by two reporters, and only the user's pages are billed (without
# the initials, the invoice would be held until Whose pages... says)
settings = run / "appdata" / "YinItAgreementForm" / "settings.json"
settings.parent.mkdir(parents=True)
settings.write_text(json.dumps({"profile": {"name": "Pat Reporter", "initials": "P.R."}}), encoding="utf-8")
env = dict(os.environ, APPDATA=str(run / "appdata"), USERPROFILE=str(run / "home"), HOME=str(run / "home"))
problems = []
try:
    subprocess.run([str(exe), "--batch", str(run / "out"), "--outputs", "agreement,mofr,invoice,runsheet", "--math",
                    str(src)], env=env, timeout=300)  # exit code 1 is expected: the e-mail job has no transcript
except subprocess.TimeoutExpired:
    sys.exit(f"PROBLEM: the program did not finish within 5 minutes (files in {run})")

report_file = run / "out" / "batch.json"
if not report_file.exists():
    sys.exit(f"PROBLEM: the program wrote no batch.json (files in {run})")
report = json.loads(report_file.read_text(encoding="utf-8"))
for j in report["jobs"]:
    print(" ", j["case"], "|", j["forms"], "|", j["error"] or "ok")
days = sorted((j for j in report["jobs"] if "Roe" in (j["case"] or "")), key=lambda j: j["dates"] or "")
kinds = [sorted(kind(f) for f in j["forms"]) for j in days]
if len(days) != 2 or any(j["error"] for j in days) or kinds != [["Invoice", "MOFR", "Minute Agreement",
                                                                  "Run sheet", "The math"]] * 2:
    problems.append("each day of the trial should have an agreement, a MOFR, the run sheet and the joint invoice "
                    "with its math")
maths = sorted({f for j in days for f in j["forms"] if kind(f) == "The math"})
if len(maths) == 1:  # one invoice: its own math only (Settings.save_math "both")
    with pymupdf.open(run / "out" / maths[0]) as doc:
        if doc.metadata.get("creator") != "YinIt math":
            problems.append("the math PDF should be marked YinIt math (so it is never read as an input)")
else:
    problems.append(f"the joint invoice should have one math PDF, not {maths}")
for what in ("Minute Agreement", "MOFR"):  # one for the whole trial, listed under both days
    shared = {f for j in days for f in j["forms"] if kind(f) == what}
    if len(shared) != 1 or "(6-3-2026 to 6-4-2026)" not in next(iter(shared), ""):
        problems.append(f"the two days should share one {what} naming both days, not {sorted(shared)}")
invoices = sorted({f for j in days for f in j["forms"] if kind(f) == "Invoice"})
if len(invoices) != 1:
    problems.append(f"the two days should share one invoice, not {len(invoices)}")
else:
    with pymupdf.open(run / "out" / invoices[0]) as doc:
        widgets = {w.field_name for w in doc[0].widgets()}
    if not {"amount Regular", "amount Expedite"} <= widgets or "DjinnIt lock" in widgets:  # (pdfout.LOCK_BUTTON)
        problems.append("the invoice should have its amounts as fields, and no Lock fields button")
invoices_csv = run / "home" / "Documents" / "YinIt Records" / "invoices.csv"
rows = []
if invoices_csv.exists():
    with open(invoices_csv, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
if len(rows) != 1:
    problems.append(f"one invoice should be recorded, not {len(rows)}")
else:
    row = rows[0]
    print("  invoice:", {k: row.get(k) for k in ("Your pages billed", "Total pages transcribed", "Reporters",
                                                 "Parties")})
    if (row.get("Your pages billed"), row.get("Total pages transcribed"), row.get("Parties")) != ("10", "20", "1") \
            or row.get("Date(s) of proceeding") != "6/3/2026, 6/4/2026":
        problems.append("the invoice should bill Pat's 10 of both days' 20 transcript pages (not the word index), "
                        "for one party")
# the run sheet goes with the case's other files (OUTDIR here), no longer to Documents\YinIt Run Sheets
sheets = list((run / "out").rglob("*.xlsx"))
if list((run / "home" / "Documents" / "YinIt Run Sheets").glob("*.xlsx")):
    problems.append("the run sheet should be made with the case's files, not in Documents\\YinIt Run Sheets")
if sheets:
    ws = load_workbook(sheets[0])["Run Sheet"]
    takes = [(ws.cell(r, 3).value, ws.cell(r, 6).value) for r in range(5, ws.max_row + 1)]
    print("  run sheet:", takes)
    if len(sheets) != 1 or takes != [("Pat", 2), ("Dana", 3), ("Pat", 3), ("Dana", 2)] * 2:
        problems.append("one run sheet should have the takes of both days")
else:
    problems.append("no run sheet was made")
# --selftest: what the window needs that the batch doesn't use (rapidfuzz, regex, pillow-heif, ssl, Qt print, the
# Help's word table, Qt's svg plugin...), and the code the courthouse names by "module:name" (a build bundles it
# only when the spec lists it)
try:
    subprocess.run([str(exe), "--selftest", str(run / "st")], env=env, timeout=120)
except subprocess.TimeoutExpired:
    problems.append("--selftest did not finish within 2 minutes")
st_file = run / "st" / "selftest.json"
st = json.loads(st_file.read_text(encoding="utf-8")) if st_file.exists() else {}
for key, what, lib in (("fuzzy_search", "the Records' Fuzzy search", "rapidfuzz"),
                       ("regex_search", "the Records' Regex search", "regex"),
                       ("heic_photos", "reading iPhone HEIC photos", "pillow-heif"),
                       ("update_check", "the look for a newer version", "ssl"),
                       ("printing", "Print…", "PySide6.QtPrintSupport"),
                       ("math_pdf", "Save as PDF in The math", "pymupdf (Story)"),
                       ("single_instance", "one window at a time", "PySide6.QtNetwork"),
                       ("moving_picture", "the swirling yin-yang", "Qt's imageformats plugin (qwebp)"),
                       ("help_search", "the Help search's word table and typo fixing",
                        "minute_filler/data/help_words.* (the spec's datas) and rapidfuzz.process"),
                       ("tick_marks", "the tick in ticked boxes", "Qt's svg plugin (qsvg)"),
                       ("outputs", "the outputs' code (makers, rules, window panels)",
                        "each module courthouses names (see the spec's hiddenimports)"),
                       ("mixed_speeds", "billing firms that order at different speeds",
                        "the courthouse's split rules (Courthouse.modules, the spec's hiddenimports)")):
    got = st.get(key, "no selftest.json" if not st else "missing from selftest.json")
    print(f"  {key.replace('_', ' ')}:", got)
    if got is not True:
        problems.append(f"{what} does not work in the build ({got}): is {lib} bundled?")
log_file = run / "appdata" / "YinItAgreementForm" / "logs" / "app.log"
log = log_file.read_text(encoding="utf-8") if log_file.exists() else ""
started = [l for l in log.splitlines() if "started: version" in l]
print("  log:", started[-1][20:80] if started else "(no start line)", "| errors:", log.count("ERROR"))
if not started or f"version {__version__} " not in started[-1]:
    problems.append(f"the built program is not version {__version__} (rebuild?)")
if log.count("ERROR"):
    problems.append("the log has errors")
# The build ships no one's records, invoices or settings. The spec copies folders by pattern (rate_sheets/*.csv,
# templates/*.xlsx, forms/*.pdf), so a real invoice or rate sheet saved in one of them would ship: every file of
# the app's own in the build must be one git tracks. Checked after the runs, so a file the program wrote into its
# own folder is caught too.
tracked = set(subprocess.run(["git", "ls-files", "-z", "minute_filler"], cwd=ROOT, capture_output=True, text=True,
                             check=True).stdout.split("\0"))
internal = exe.parent / "_internal"
untracked = sorted(p.relative_to(internal).as_posix() for p in (internal / "minute_filler").rglob("*")
                   if p.is_file() and p.relative_to(internal).as_posix() not in tracked)
private = sorted(p.relative_to(exe.parent).as_posix() for p in exe.parent.rglob("*")
                 if p.name.lower() in ("records.db", "settings.json", "invoices.csv", "activity.csv")
                 or p.suffix.lower() in (".db", ".sqlite", ".log") or "-wal" in p.name or "-shm" in p.name
                 or p.name in ("YinIt Records", "YinIt Run Sheets"))
print("  build's own files:", "all tracked" if not untracked else untracked, "| records or settings:", private or "none")
if untracked:
    problems.append(f"the build has files git doesn't track (someone's own?): {untracked}")
if private:
    problems.append(f"the build has records, settings or logs in it: {private}")
if problems:
    print("\n".join("PROBLEM: " + p for p in problems) + f"\n(files in {run})")
else:
    print("all good")
    shutil.rmtree(run, ignore_errors=True)
sys.exit(1 if problems else 0)
