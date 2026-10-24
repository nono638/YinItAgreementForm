"""Main window: drop zone, job list and paste box on the left, the cards on the right, in the order the job is done
(the user, 2026-10-22: Case, Attorneys, Speed, pages and dates, Invoice, Who ordered what, Who pays what, Who
transcribed what), and the Outputs box at the bottom, one row: a box per output with how many
files Generate makes of it (_show_counts), Options ▾ for the outputs' standing choices (a box under it,
_show_output_options) and the Generate buttons, which say what they make (_generate_labels) and go on a row of
their own when the window is narrow (_place_foot). The outputs are the courthouse's (courthouses.panel_order),
their options laid out by gui/panels.py: the invoice's, this job's own, in a card of the form (OutputSpec.card).
Room for the form on a small screen: the status and the Records, New job and Settings buttons sit at the right of
the menu bar; the divider above the paste box can be dragged; each card folds to a line that sums it up (card,
_fold_card); the drop zone is one slim row once documents are read (DropZone.set_slim), as the whole window takes
them (dropEvent), and the paste box is one line while it is empty (_fit_paste). The status names the first thing
that needs a look, the job's own checks and the warnings in the cards too, and a click goes there (_problems,
_job_status, _go_to_problem); a Before you generate box at the top of the form lists them all, each a link to its
place (_show_todo).
The attorney table shows who; the row selected is spelled out under it (_build_att_details). Each firm of a trial
is marked by a tag and a thin stripe of its colour wherever it is named: the attorney table, Who ordered what, Who
pays what, Excerpts… (_marks, _mark_row). While nothing is loaded, the jobs made lately are listed to open again
(_show_start); fields complete from the records (_fill_completions); New job can be undone for a few seconds
(_undo_clear); the Records button counts the invoices of a status (Settings.records_count: _count_records). View
zooms, picks the theme, folds the cards and shows or hides the job list's Day column. F1 opens the manual at the
topic of the part of the window in use (show_help, gui/help_window.py).
The dates follow the day (workdays.py, Settings -> Dates): a form made on a weekend or a court holiday is dated the
next business day (and one made after the evening hour, when Settings says so, the day after), and the delivery
date is counted in business days; a note under each date says what moved it (_show_date_notes), and both are
worked out again as the day turns (_redate, each minute). Each date has a calendar with the court holidays marked.
The job list is a tree: a row per job (Case / file, Date, Day, Pages, Yours), sorted by date, oldest first, or by
the heading clicked (JobItem, _job_sort_key), and, under a job of several documents, a row per document in the
order read (DocItem); a document with no index number, or with another than its job's, gets a ⚠ there and a note
under the list (_label_docs, _show_index_notes), and Generate asks first about documents whose index numbers
don't match (_ask_index_numbers). A job's tick says whether its day counts: the cards of who ordered, pays and
transcribed what show every day ticked, whichever day is shown (_considered), Generate all makes the days ticked
but not made already (_to_generate, Job.made), and an unticked row is greyed (_dim).
The "Speed, pages and dates" card holds the rate sheet (one missing prices, or with a price lower at a faster
speed, is warned about: _check_rate_sheet; opened from that warning to be filled in, it is read again each time it
is saved: _watch_sheet), the speeds the invoice offers (Speeds offered, each at its price; greyed out, and asked
about before a change, while firms' speeds are chosen in Who ordered what: _speeds_locked) and the one speed and
rate the agreement form names among them (the Settings rule's, Settings.agreement_speed, or the job's own choice;
when an e-mail asks for another speed, the user is asked which). A firm whose speed is set in the "Who ordered
what" card gets its own; firms that ordered the same pages commit to one (Generate asks: _ask_split_speeds).
"Who pays what" sums up the invoices of the days ticked, those to make and those made already, as things change;
it and the Invoice card (the job shown's invoice) link to the whole math. "Who transcribed what" lists each
reporter's takes of every day ticked, with the running heads printed on their pages, to save as an Excel workbook
or a CSV file (_show_takes, _save_takes); a take that came out wrong (a page's initials misread or missing) is
corrected there, and the correction is the transcript's (_fix_take, Doc.fixes). The Invoice card's
Peripherals... (the e-mailed copy and the index, for this job only) and Customize... set what the current job's
invoice shows, Excerpts... which firm ordered which pages of the days ticked of the case (its own window,
gui/excerpts.py), and Whose pages... whose pages of a transcript of several reporters are billed (the user's own,
by default; other reporters' on invoices in their name); with Generate all, the days of one case share one
invoice per attorney (batch.joint_invoice). The files of a case go into a folder of its own inside Save to
(Settings.case_folders), asked about when it has files in it already (_case_folder); its run sheet too, unless
Settings give run sheets a folder of their own. A case's run sheet is also looked for next to its transcripts,
where the records say it is and in the old run sheets folder (_run_sheet_for). The Run sheet box is ticked by how
many reporters wrote a case's transcripts until the user clicks it (_auto_runsheet); New job is greyed out while
there is nothing to clear. File -> Lock finished PDFs saves copies whose fields can no longer be changed.
Generate and Generate all show the files as pictures first (preview.PreviewDialog), with "The math" of the
invoices, unless that is turned off (then the math comes after saving, preview.MathDialog, unless
Settings.show_math is off), and the math is saved as PDFs with the invoices as Settings.save_math says (its
"Save the math" list is under the math too: _math_saves_kept); the box that says what was saved can print the
files. File -> Open recent lists the documents opened lately, and Export / Import settings carry the settings
to another computer. A job once made can be opened again from the Records window, its takes as the user corrected
them (open_past_job, batch.origin_fixes). When the window opens it makes the day's backup of the records and, once
a day, asks whether there is a newer version (update.py: an Update button at the top until it is installed); a
new user is asked a few questions first
(preview.WelcomeDialog), a settings file that couldn't be read is said so instead, and an incomplete rate sheet is
warned about (MainWindow._startup). On the first day the app is used in a new year, the yin-yang's "done" picture
is the New Year one (note_opened).

Documents are read and batches are made on a thread pool (workers.Runner), and the AI is asked on a pool of its
own, one question at a time (MainWindow.ai_pool). Their results are applied on the UI thread; the results of
work started before "New job" are dropped (see MainWindow.gen), and AI questions still waiting are never asked."""
from __future__ import annotations

import html
import json
import re
import tempfile
import textwrap
import time
from calendar import monthrange
from copy import deepcopy
from datetime import MAXYEAR, MINYEAR, date, datetime, timedelta
from collections.abc import Callable
from dataclasses import dataclass, replace
from pathlib import Path

from PySide6.QtCore import QByteArray, QEvent, QPoint, QRect, QSignalBlocker, QSize, QStringListModel, Qt, QThreadPool, QTimer
from PySide6.QtGui import (QActionGroup, QBrush, QColor, QFont, QFontMetrics, QIcon, QKeySequence, QMovie, QPalette,
                           QShortcut)
from PySide6.QtWidgets import (
    QAbstractItemDelegate, QAbstractItemView, QApplication, QBoxLayout, QCheckBox, QComboBox, QCompleter, QFileDialog,
    QFormLayout, QFrame, QHBoxLayout, QHeaderView, QInputDialog, QLabel, QLineEdit, QMainWindow, QMenu, QMessageBox,
    QGridLayout, QPlainTextEdit, QProgressBar, QPushButton, QScrollArea, QSizePolicy, QSplitter,
    QTableWidget,
    QTableWidgetItem, QToolButton, QTreeWidget, QTreeWidgetItem, QVBoxLayout, QWidget,
)

from .. import log as logfile
from ..log import describe, error as log_error, log
from ..batch import (BATCH_EXT, HELD, MISMATCH_LEAD, NOT_INVOICED, Doc, Job, case_dir_for, case_folders, case_reporters,
                     expand_paths,
                     fill_jobs, firm_questions, group, ident, index_number_text, input_folders,
                     billed_split_rule, job_from_origin, job_origin, join_entries, new_invoice_set, origin_fixes,
                     out_dir_for,
                     read_loaders, remerge, same_case, split_forgotten)
from ..extract_llm import OllamaExtractor
from ..extract_regex import find_dates, use_firm_answers
from ..dates import us_date
from ..deliver import backup_folder, backup_records, generate, ledger_for
from ..ingest import ingest_file, ingest_pil, ingest_text
from ..merge import apply_defaults, apply_speed_rule, refresh_dates, refresh_delivery_date, refresh_rate
from ..workdays import date_warning, delivery_day, document_day, quick_day
from ..models import (Attorney, CaseInfo, FIELD_LABELS, FieldState, PROC_TYPES, REQUIRED_KEYS, SRC_AI,
                      SRC_DEFAULT, SRC_DERIVED, SRC_PDF, SRC_RECORDS, SRC_USER)
from ..runsheet import (SheetUnreadable, find_sheets, run_sheet_summary, runsheets_folder, sheet_folders,
                        transcript_pages)
from .. import courthouses
from ..courthouses.base import resolve
from ..settings import LOCAL, OUTPUTS, Settings
from .dialogs import ClarifyDialog, RunSheetDialog, SettingsDialog
from .help_window import help_topic_of, open_help
from .badges import STRIPE_ROLE, TAG_HEIGHT, StripeDelegate, badge_icon, badge_size, tag_color
from .theme import apply_theme
from .widgets import (ASSETS, ClickLabel, ElidedLabel, QuietCombo, open_path, open_url, plural, refill_combo, repolish,
                      round_corners, rounded, set_checks, show_save_error)
from .workers import Runner
from . import zoom as zooming
from .zoom import sized, z

# every kind of file a dropped folder takes (batch.BATCH_EXT), so Browse... offers the same
FILE_FILTER = (f"Documents ({' '.join('*' + e for e in sorted(BATCH_EXT, key=lambda e: (e != '.pdf', e)))});;"
               "All files (*.*)")
ATT_COLS = ["", "Name", "Firm", "Address", "Phone", "Fax", "Email", "Party / role", "Source"]
ATT_FIELDS = [None, "name", "firm", "address", "phone", "fax", "email", "party", "source"]
ORDER_COLS = ["Day", "Attorney", "Pages ordered", "Speed", "Printed page numbers", "Also ordered",
              "Your pages billed"]
ORDER_SPEED = ORDER_COLS.index("Speed")  # (a box to choose it in, see _speed_box)
TAKE_COLS = ["Day", "Reporter", "Printed pages", "Pages", "Headers"]  # the Who transcribed what card
TAKES_NOTE = ("A take wrong (initials misread or missing)? Double-click it to correct where it ends. Every take in "
              "full, kept day after day: tick Run sheet and Generate (the trial's run sheet, in Excel).")
TAKES_TIP = ("The takes, from the initials at the foot of each page.\nDouble-click one (or right-click it) to correct "
             "where it ends, when a page's initials were misread or missing:\nthe pages after it go to the next take, "
             "and your pages (Yours, the invoice) and the run sheet follow.\n✎ = corrected by you.")


def take_cells(row) -> list[str]:
    """A take (runsheet.Row) as the Who transcribed what card shows it: '6/3/2026' ('Roe v. Poe · 6/3/2026' when
    the card shows several cases: Row.case), 'Pat Reporter', 'pp. 241–309', '69', 'PROCEEDINGS, DIRECT - PLAINTIFF -
    JANE ROE'. The last is its pages' running heads as printed, not who took the stand: a witness read from them was
    wrong whenever a head put the examination first (the user, 2026-10-23: "if you try to infer it, it'll be wrong
    and annoying more often")."""
    day = f"{row.day.month}/{row.day.day}/{row.day.year}" if row.day else ""
    day = f"{row.case} · {day}" if getattr(row, "case", "") else day
    pages = (f"pp. {row.start}–{row.end()}" if row.end() is not None and row.end() != row.start else
             f"p. {row.start}" if row.start is not None else "")
    return [day, row.reporter or "(initials not found)", pages, str(row.pages or 0), row.heads]


def reporter_totals(rows) -> str:
    """What each reporter wrote of these takes, most first: 'Pat Reporter 144 pp. · DS 76 pp.'."""
    by: dict[str, int] = {}
    for row in rows:
        who = row.reporter or "(initials not found)"
        by[who] = by.get(who, 0) + (row.pages or 0)
    return " · ".join(f"{who} {n:,} pp." for who, n in sorted(by.items(), key=lambda x: -x[1]))


def save_takes(rows, path: Path, csv_file: bool = False) -> Path:
    """The Who transcribed what table (take_cells, then each reporter's total) saved at path: an Excel workbook,
    or a CSV file (csv_file). A running head that looks like a formula stays text: in the workbook one starting
    with "=", in the CSV file one starting with "=", "+", "-" or "@" (records.csv_text). Returns the path
    written."""
    table = [TAKE_COLS] + [take_cells(r) for r in rows] + [[], ["Each reporter", reporter_totals(rows)]]
    path = Path(path)
    if csv_file:
        import csv
        from ..records import csv_text  # (a head "- CROSS -" is read by Excel as a formula otherwise)
        path = path.with_suffix(".csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerows([[csv_text(x) for x in line] for line in table])
        return path
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Who transcribed what"
    for line in table:
        ws.append([int(x) if x.isdigit() and line is not table[0] else x for x in line])
        for cell in ws[ws.max_row]:  # (a head "== DIRECT ==" is text, not a broken formula)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDE", (12, 24, 18, 8, 60)):
        ws.column_dimensions[col].width = width
    path = path.with_suffix(".xlsx")
    wb.save(path)
    return path
# Speeds offered (under the rate sheet) while firms' speeds are chosen in Who ordered what (_speeds_locked)
SPEEDS_LOCKED_TIP = ("Speeds are chosen for the firms in Who ordered what: these boxes only list speeds for a firm "
                     "that chooses from its invoice.")
SPEEDS_LOCKED_ASK = ("You already chose speeds in Who ordered what.\n\nAre you sure you want to change the speeds "
                     "offered on invoices?")
ORDER_TIPS = {  # the headings' tooltips
    "Pages ordered": "The pages of the day this attorney ordered: every page, or an excerpt (Excerpts…)",
    "Speed": "The speed this attorney ordered its pages of the day at. Firms that ordered the same pages commit "
             "to a speed (Generate asks when one isn't set) and are billed it alone: at different speeds the "
             "original is billed once and divided as Settings → Invoice → Split invoices says. A firm "
             "ordering alone may leave it to choose from its invoice. Different speeds for different pages: "
             "Excerpts…",
    "Printed page numbers": "The page numbers printed on those pages of the transcript",
    "Also ordered": "The other attorneys who also ordered some of the same pages, and which (they split part "
                    "of the cost of those pages, the original among it: see Who pays what)",
    "Your pages billed": "How many of those pages its invoice bills: on a transcript of several reporters only "
                         "yours (or those chosen under Whose pages…)",
}
# the list of jobs on the left: a row per job, and under a job of several documents a row per document
JOB_COLS = ["Case / file", "Date", "Day", "Pages", "Yours"]
JOB_DATE, JOB_DAY, JOB_PAGES, JOB_YOURS = 1, 2, 3, 4
TICKED_ROW = ("Ticked: the cards on the right show this day, and Generate all makes it.\n"
              "Untick it to see the difference (nothing is removed).")
UNTICKED_ROW = "Unticked: left out of the cards on the right and of Generate all. Tick it to bring it back."
MADE_ROW = ("Ticked: the cards on the right show this day. Made already (✓): Generate all leaves it out,\n"
            "but for its invoice while that is still to make (click it, then \"{one}\" makes it all again).")
TICK_TIP = ("Ticked: this day counts. The cards on the right (who ordered, who pays, who transcribed\n"
            "what, Excerpts…) show every ticked day, and Generate all makes them.\n"
            "Unticked (greyed): left out of both, to see the difference it makes. Nothing is removed:\n"
            "tick it again to bring it back.")
YOURS_TIP = ("Yours: the pages with your initials at the foot of the page (your initials: Settings → My info).\n"
             "Pages: every page of the transcript.")
JOB_TIPS = {
    "Case / file": "One job per case and date. Click a job to check or edit it.\n" + TICK_TIP + "\n"
                   "A job read from several documents lists them under it: right-click one to\n"
                   "see its text, move it to a job of its own or remove it.",
    "Date": "The date of the minutes (a document's own date, or what kind of document it is). Click a heading\n"
            "to sort the list by it, again to sort it the other way.",
    "Day": "The day of the week of the job's date (right-click the headings, or View, to hide it)",
    "Pages": "The transcript's own pages (the word index after them left out)",
    "Yours": "The pages you wrote: those with your initials at the foot of the page (every page when you\n"
             "wrote it alone, or when no page has initials).\n"
             "A number typed in Est. number of pages shows as typed.\n"
             "? = can't tell yet: your initials aren't set (Settings → My info), or Whose pages…\n"
             "hasn't said whose the first pages are.",
}
JOB_SORT_ROLE = Qt.UserRole + 20  # what a row sorts by in each column (MainWindow._job_sort_key)
CASE_MIN = 120  # the least the Case / file column keeps, before zoom (_fit_case_column)


class JobItem(QTreeWidgetItem):
    """A job's row in the list, sorted by what its cells mean (their keys in JOB_SORT_ROLE: MainWindow._job_sort_key),
    not by their text: a date as a date (9/3/2026 before 10/14/2026), pages as numbers (87 before 109)."""

    def __lt__(self, other) -> bool:
        tree = self.treeWidget()
        col = tree.sortColumn() if tree is not None else 0
        a, b = self.data(col, JOB_SORT_ROLE), other.data(col, JOB_SORT_ROLE)
        # (without a key yet, by its text: QTreeWidgetItem's own < called from here crashes PySide)
        return self.text(col) < other.text(col) if a is None or b is None else a < b


class DocItem(QTreeWidgetItem):
    """A document's row under its job: in the order the documents were read (its JOB_SORT_ROLE), whichever way the
    list is sorted."""

    def __lt__(self, other) -> bool:
        a, b = self.data(0, JOB_SORT_ROLE), other.data(0, JOB_SORT_ROLE)
        if a is None or b is None:  # (by its text: see JobItem)
            return self.text(0) < other.text(0)
        tree = self.treeWidget()
        # (sorted the other way, Qt asks whether other < self: answered so that the order read stays)
        down = tree is not None and tree.header().sortIndicatorOrder() == Qt.DescendingOrder
        return a > b if down else a < b


def _first_day(text: str) -> date | None:
    """The first date written in text ("6/3/2026, 6/4/2026" -> 6/3/2026); None when there is none."""
    found = [d for _, _, d in find_dates(text or "")]
    try:
        return datetime.strptime(found[0], "%m/%d/%Y").date() if found else None
    except ValueError:
        return None


def _weekday(text: str) -> str:
    """The day of the week of the first date in text ("Wed" for 6/3/2026); "" without one."""
    day = _first_day(text)
    return day.strftime("%a") if day else ""


# the Who ordered what row of a day nobody is ticked on, when other days of its case have attorneys ticked
NOBODY_DAY = "⚠ nobody ticked on this day: no invoice for the case until you tick who ordered it"
# how the Who ordered what card's rows are billed: its ? button's (the line above the table says what Generate all
# bills: _orders_where)
ORDERS_HELP = ("Each attorney is billed for the pages written in its row and nothing else; \"Your pages billed\"\n"
               "counts the pages billed among them: on a transcript of several reporters only yours (or those\n"
               "chosen under Whose pages…), and the Pages field's number when you typed one. Firms that ordered\n"
               "the same pages commit to a speed (Generate asks when one isn't set); a firm ordering alone may\n"
               "choose one from its invoice.")
ATT_DETAIL_COLS = (3, 4, 5, 6, 8)  # the attorney table's columns shown under it for the row selected, not in it
# the fields of the Case card and of the Speed, pages and dates card (a folded card's ⚠: _card_summaries)
CASE_KEYS = ("court", "county", "part", "index_no", "judge", "dates", "case_name", "proc_other")
DETAIL_KEYS = ("delivery", "rate", "copies", "est_pages", "delivery_date", "agreement_date")
# The help topic F1 opens for each part of the window (help_window.help_topic_of: the focused widget's, else the
# nearest parent's): the cards, a few fields, the Outputs box and its Options ▾ panels, the job list and the drop zone.
CARD_TOPICS = {"case": "check-badges", "details": "inv-speeds", "attorneys": "check-attorneys",
               "pays": "inv-who-pays", "orders": "inv-who-ordered", "takes": "run-who-transcribed",
               "invoice": "inv-panel"}
FIELD_TOPICS = {"index_no": "check-index", "est_pages": "check-pages", "copies": "check-copies",
                "delivery_date": "dates-turnaround", "agreement_date": "dates-business-days"}
OUTPUT_TOPICS = {"agreement": "make-agreement", "invoice": "inv-panel", "mofr": "make-mofr",
                 "runsheet": "run-overview"}
# A field's badge: where its value came from, in plain words (the source itself, models.SRC_*, picks its colour)
BADGE_WORDS = {"regex": "found", "derived": "worked out"}
BADGE_TIPS = {"regex": "Found in the document", "AI": "Suggested by the AI model - please check",
              "default": "Your default setting", "derived": "Worked out from the other details",
              "PDF": "Counted from the transcript PDF, without the word index (the line under it says which pages)",
              "records": "From your records: an earlier job with the same index number - please check",
              "you": "Entered by you"}
BADGE_WIDTH = 74  # (px at 100 %: "worked out" fits)
RECORDS_TIP = "Invoices (mark them paid) and everything made so far (Ctrl+R)"
# what the start card says while nothing is loaded (_show_start)
START_TEXT = ("Drop a transcript, an e-mail or a photo anywhere in the window, or paste with Ctrl+V; check the "
              "details, tick what to make, then Generate.")
UNDO_MS = 10000  # how long New job's Undo is offered


def badge_word(source: str) -> str:
    """What a badge says for a source: "regex" -> "found", "derived" -> "worked out", the others as they are."""
    return BADGE_WORDS.get(source, source)


SPEEDS_TIP = ("The speeds the invoice offers, each at its own price (the attorney chooses one). One ticked:\n"
              "the invoice bills that speed alone. They are kept for the next job too.\n"
              "The minute agreement form names one of them (Speed, pages and dates → Speed), so they\n"
              "can be ticked while Invoice is unticked too. A firm whose speed is set in Who ordered what\n"
              "is billed that speed alone.")
EST_PAGES_TIP = ("Counted from the transcript PDF: every page, whoever wrote it (without the word index).\n"
                 "It is read four ways (the pages up to the index, the line numbers, the printed page numbers,\n"
                 "the index from the end); when two of them agree on another count, or the index would take\n"
                 "more than a quarter of the PDF (rounded up; 5 pages of a short one), the field turns amber\n"
                 "and a warning under it says so.\n"
                 "Each attorney's minute agreement shows the pages that attorney ordered (the whole day,\n"
                 "or its excerpt under Excerpts…); the MOFR the pages anyone ordered. Your invoice bills\n"
                 "only your own pages of them (Invoice card: Billed).\n"
                 "A number you type is yours: your invoice bills it, and it is the day's pages on every\n"
                 "agreement (an excerpt counts its pages of it). Typed below the PDF's count, Generate asks\n"
                 "about the index when the two would decide it differently. No transcript (a caption page\n"
                 "only)? Type the pages here and the invoice bills them (a number read from an e-mail isn't\n"
                 "billed).")
AGREEMENT_SPEED_TIP = ("The speed written on the minute agreement forms (and the MOFR), at its rate per page, out\n"
                       "of the speeds offered (ticked above it). Settings → Invoice picks it (first\n"
                       "choice Expedited, else the slowest offered); choose another here for this job (↺ goes\n"
                       "back to the Settings rule). A firm whose speed is set in Who ordered what gets its own\n"
                       "speed on its form, and the MOFR ticks every speed ordered.\n"
                       "When an e-mail asks for another speed, you are asked which.")
# (the tips below say what firms ordering the same pages split, and what each pays for its own, as
# Settings.invoice_index_shared says: format them with _who_splits)
PARTIES_TIP = ("How many parties ordered: they split {shared},\n"
               "and each pays for its own {own} (normally the number of ticked attorneys).\n"
               "No. of copies on the forms follows it.")
WHO_TIP = ("Excerpts: one attorney ordered the whole transcript and another only pages 20 to 40?\n"
           "A table of every day of the case on that invoice: type each run of pages and tick the firms\n"
           "that ordered it, with the prices as you go. Pages ordered together split {shared}.")
WHOSE_TIP = ("Several reporters wrote this transcript (the initials at the foot of the pages change).\n"
             "The invoice bills your own pages. Tick another reporter here to also make invoices in their name\n"
             "for their pages, or bill the whole transcript.")


PAYS_TIP = ("What each firm's invoice comes to, as things stand: its pages, how it ordered them (alone, or\n"
            "shared with other firms: they split {shared}) and what it pays at each speed\n"
            "the invoice offers, with what that comes to a page (its copy, the e-mailed copy and the index\n"
            "counted in). It follows every tick at once; the link spells out the whole math.")


def _who_splits(s) -> dict[str, str]:
    """What firms ordering the same pages split, and what each pays for its own, as Settings say, for the
    tooltips above: {"shared": "the original and the judge's index", "own": "copy, e-mailed copy and index"}
    (no e-mailed copy, or no index, when Settings → Invoice charges none)."""
    from ..invoice import own_words, shared_words
    return {"shared": shared_words(s), "own": own_words(s)}


def _own_words(firms, s) -> str:
    """What each firm pays for its own on these invoices (invoice.firm_invoices), from the lines they charge:
    "copy, e-mailed copy and index", or "copy" for a 30-page day without an index and the e-mailed copy turned
    off for the job. Settings' index on shared pages "split" makes the index one the firms split, not their own."""
    have = {l.charge for f in firms for q in f.quotes for l in q.lines if l.amount > 0}
    own = [w for label, w in (("Copy", "copy"), ("E-mailed copy", "e-mailed copy"), ("Index", "index"))
           if label in have and (label != "Index" or s.invoice_index_shared == "each")]
    return _and(own) or "copy"

WARN_WIDTH = 72  # the longest line of a warning in the Invoice card (characters; see MainWindow._warn)
# the longest line of the Invoice card's Billed and Includes lines (characters): they are broken into lines
# here, not word-wrapped by the label, as a word-wrapped label in a row the form shows later (Billed) was given
# the height of one line less than it needed, its last line cut off
LINE_WIDTH = 40
# how long a rate sheet opened to be filled in is left alone after it changes before it is read again: Excel saves
# in steps (a new file in the old one's place), and a read in between finds it half written or not there
SHEET_SAVED_MS = 1000


def _lines(text: str, width: int = LINE_WIDTH) -> str:
    """Each line of the text broken into lines of at most `width` characters (between words)."""
    return "\n".join(textwrap.fill(line, width) or line for line in text.split("\n"))

def _count_text(key: str, counts: dict | None, everyone: dict | None = None, held: str = "") -> str:
    """The tooltip of the count beside an output's box (_count_short): "Will generate 2 minute agreement forms",
    "Will generate 3 invoices (and 3 detailed copies)", "Will fill in 1 run sheet", "Will generate no invoice" ("No
    invoice until Whose pages… is chosen", "... until the speeds are checked": held, why the job's invoice waits,
    Job.invoice_hold; see _show_counts). With several jobs, what Generate all makes, on a line of its own
    ("Generate all: 6 (+6 detailed)"). Digits, not words: "two minute agreement forms" reads as "two-minute".
    counts, everyone: batch.output_counts of the job shown (None: an empty job, only Generate all's line) and of
    every job ticked. The words are the output's (courthouses.OutputSpec.count_words); a run sheet (group
    "sheet") is filled in."""
    spec = courthouses.output(key)
    one, many = spec.count_words()
    lines = []
    if counts is not None:
        verb = "Will fill in" if spec.group == "sheet" else "Will generate"
        n = counts.get(key, 0)
        text = f"{verb} {n} {one if n == 1 else many}" if n else f"No {one} until {held}" if held else \
            f"{verb} no {one}"
        copies = counts.get("detailed", 0) if key == "invoice" else 0
        if copies:
            text += f" (and {copies} detailed cop{'y' if copies == 1 else 'ies'})"
        lines.append(text)
    if everyone is not None:
        extra = everyone.get("detailed", 0) if key == "invoice" else 0
        lines.append(f"Generate all: {everyone.get(key, 0)}" + (f" (+{extra} detailed)" if extra else ""))
    return "\n".join(lines)


def _count_short(key: str, counts: dict | None, everyone: dict | None = None, held: str = "") -> str:
    """The count beside an output's box in the Outputs row: how many files Generate makes of it for the job shown
    ("(2)"), and with several jobs what Generate all makes ("(2 · all 6)"); a ⚠ first when the job's own is held
    ("⚠ (0)"). Its tooltip says it in words (_count_text). counts, everyone: as _count_text's."""
    parts = [str(counts.get(key, 0))] if counts is not None else []
    if everyone is not None:
        parts.append(f"all {everyone.get(key, 0)}")
    text = f"({' · '.join(parts)})" if parts else ""
    return f"⚠ {text}" if held and counts is not None and not counts.get(key, 0) else text


def _files_words(counts: dict | None) -> str:
    """What a Generate button makes, in words: "3 files: 2 minute agreement forms, 1 MOFR" (the detailed copies and
    the PDFs of the math too, which the counts beside the outputs' boxes leave out); "" for nothing. counts:
    batch.output_counts."""
    parts = []
    for key, n in (counts or {}).items():
        if not n:
            continue
        one, many = ("detailed copy", "detailed copies") if key == "detailed" else \
            ("PDF of the math", "PDFs of the math") if key == "math" else courthouses.output(key).count_words()
        parts.append(f"{n} {one if n == 1 else many}")
    return f"{plural(sum(counts.values()), 'file')}: {', '.join(parts)}" if parts else ""


def _named(keys) -> str:
    """The outputs in a sentence, by their labels: "minute agreement, MOFR, invoice and/or run sheet"."""
    names = [o.label if o.label.isupper() else o.label[:1].lower() + o.label[1:]
             for o in courthouses.outputs() if o.key in set(keys)]
    return ", ".join(names[:-1]) + " and/or " + names[-1] if len(names) > 1 else "".join(names)


def _previewed(outputs) -> list[str]:
    """The outputs a preview shows: those whose files are PDFs the app marks (OutputSpec.mark). The run sheet is
    a spreadsheet, and making it for the preview would add to the real one."""
    return [o for o in outputs if (spec := courthouses.output(o)) is not None and spec.mark]


def _not_previewed(outputs, several: bool = False) -> str:
    """The preview's note on the outputs it doesn't show: "The run sheet is not shown here; it is saved with the
    rest." ("" when it shows them all); several: of Generate all ("The run sheets are ..., they are ...")."""
    left = [spec.count_words()[1 if several else 0] for o in outputs
            if (spec := courthouses.output(o)) is not None and not spec.mark]
    if not left:
        return ""
    many = several or len(left) > 1
    return f"The {' and the '.join(left)} {'are' if many else 'is'} not shown here; " \
           f"{'they are' if many else 'it is'} saved with the rest."


def _price_lines(firms, form_speed: str = "", asked=frozenset()) -> tuple[str, bool]:
    """The Invoice card's prices, two speeds a line so the card stays narrow enough for a small window:
    "Regular $63.00 · Expedite $76.00 each" when every attorney pays the same, else a line per invoice, named
    by Attorney.label ("Alex B. Counsel: Regular $815.50 · Expedite $977.00"), and whether it is so. With
    several speeds, the one the minute agreement form names (form_speed) is marked: "Expedite (form) $76.00".
    A firm of a split order whose speed Generate still asks (asked: Attorney.key()s, batch.speeds_to_ask) is
    priced at the form's speed, with a question mark: "Expedite? $98.00", as the Excerpts window shows it.
    firms: invoice.firm_invoices."""
    from ..invoice_calc import fmt
    from ..rates import speed_key

    def pairs(quotes, ask=False) -> list[str]:
        mark = len(quotes) > 1 and form_speed
        each = [f"{q.speed}{'?' if ask else ''}"
                f"{' (form)' if mark and speed_key(q.speed) == speed_key(form_speed) else ''} "
                f"{fmt(q.per_party)}" for q in quotes]
        return [" · ".join(each[i:i + 2]) for i in range(0, len(each), 2)]

    def asks(f) -> bool:
        return f.atty is not None and f.atty.key() in asked

    if not firms:
        return "", False
    if len({tuple(q.per_party for q in f.quotes) for f in firms}) == 1 and len({asks(f) for f in firms}) == 1:
        shared = len(firms) > 1 or any(q.parties > 1 for q in firms[0].quotes)
        return "\n".join(pairs(firms[0].quotes, asks(firms[0]))) + (" each" if shared else ""), False
    lines, seen = [], set()
    for f in firms:
        name = f.atty.label() if f.atty else "(no attorney)"
        if (name, id(f.quotes)) in seen:
            continue
        seen.add((name, id(f.quotes)))
        name = name if len(name) <= 24 else name[:23] + "…"
        rows = pairs(f.quotes, asks(f))
        lines.append(f"{name}: {rows[0]}" if rows else name)
        lines += ["      " + r for r in rows[1:]]
    return "\n".join(lines), True


def _and(words: list[str]) -> str:
    """['Regular'] -> 'Regular'; ['Regular', 'Expedite', 'Daily'] -> 'Regular, Expedite and Daily'."""
    return " and ".join(words) if len(words) < 3 else ", ".join(words[:-1]) + " and " + words[-1]


def _or(words: list[str]) -> str:
    """['Email'] -> 'Email'; ['Copy', 'Email', 'Index'] -> 'Copy, Email or Index'."""
    return " or ".join(words) if len(words) < 3 else ", ".join(words[:-1]) + " or " + words[-1]


def _file_mark(path: Path) -> tuple | None:
    """A file's size and last change, to tell whether it was saved since (None when it isn't there)."""
    try:
        st = path.stat()
    except OSError:
        return None
    return st.st_size, st.st_mtime_ns


def _pay_html(rows, form_speed: str, form_rate: str, days: int = 1, rates: str = "", own: str = "copy",
              firm_speeds: bool = False, marks: dict | None = None, link: bool = True) -> str:
    """The "Who pays what" card as rich text: the minute agreement form's one speed and rate, then a row per
    invoice: the firm, its pages, how it ordered them and what it pays at each speed the invoice offers
    ("$408.50 ($3.40/pg)"; the agreement form's speed marked "(form)" when there are several), what a page of
    the original costs alone and shared (rates: invoice_math.page_rates) and what each firm pays for its own
    (own: "copy, e-mailed copy and index"), and a link to the whole math. When every invoice bills one speed
    (the speed its firm ordered: a split order commits to one), a Speed and an Amount column instead.
    rows: invoice_math.who_pays; days: how many days the invoices cover (Generate all); firm_speeds: some firms
    have a speed of their own (Job.speeds), which their agreements name; marks: each firm's mark by its key
    (batch.trial_marks), shown as its tag (badges.badge_html) and a stripe of its colour at the row's left edge;
    link: the link to the math at the end (False: the card puts one under several invoices)."""
    from .badges import badge_html
    import html
    from ..invoice_calc import fmt
    from ..rates import speed_key
    single = bool(rows) and all(len(r.prices) == 1 for r in rows)
    speeds = list(dict.fromkeys(sp for r in rows for sp, _, _ in r.prices)) if rows else []
    form = speed_key(form_speed)
    rate = f" at ${form_rate.lstrip('$')} a page" if form_rate and form_speed != "Other" else ""
    offer = (f"Each firm is billed the speed it ordered" if single and len(speeds) > 1 else
             f"The invoices offer {html.escape(_and(speeds))}, each at its own price" if len(speeds) > 1 else
             f"The invoices bill {html.escape(speeds[0])}" if speeds else "")
    span = f" for these {days} days" if days > 1 else ""
    forms = (f"each firm's own speed (<b>{html.escape(form_speed or '(no speed)')}{rate}</b> for a firm without "
             "one)" if firm_speeds else
             f"<b>{html.escape(form_speed or '(no speed)')}{rate}</b> (one speed, one rate)")
    out = [f"Minute agreement forms: {forms}. {offer}{span}:"]
    heads = ["Firm", "Pages", "How it ordered them"] + (["Speed", "Amount"] if single else [
        html.escape(sp) + (" (form)" if len(speeds) > 1 and speed_key(sp) == form else "") for sp in speeds])
    table = ["<table cellspacing='0' cellpadding='3'>",
             "<tr><th width='4'></th>" + "".join(f"<th align='left'>{h}&nbsp;&nbsp;</th>" for h in heads) + "</tr>"]
    for r in rows:
        mark = (marks or {}).get(r.key)
        stripe = f" bgcolor='{tag_color(mark).name()}'" if mark is not None else ""
        cells = [badge_html(mark, r.name), f"{r.pages:,}", html.escape(r.how)]
        if single:
            sp, pays, each = r.prices[0]
            cells += [html.escape(sp), f"{fmt(pays)} ({fmt(each)}/pg)"]
        else:
            got = {sp: (pays, each) for sp, pays, each in r.prices}
            cells += [f"{fmt(got[sp][0])} ({fmt(got[sp][1])}/pg)" if sp in got else "—" for sp in speeds]
        table.append(f"<tr><td width='4'{stripe}></td>" + "".join(f"<td>{c}&nbsp;&nbsp;</td>" for c in cells) + "</tr>")
    table.append("</table>")
    out.append("".join(table))
    if rates:
        out.append(f"{html.escape(rates.rstrip('.'))}. Each firm also pays for its own {html.escape(own)}; /pg is everything "
                   "on its invoice divided by its pages.")
    if link:
        out.append(MATH_LINK)
    return "<br>".join(out)


MATH_LINK = "<a href='math'>How these amounts are worked out…</a>"


def _pages_changed(job, copy) -> bool:
    """The job was given other pages to bill while Generate all worked on its copy (made when the batch
    started): other documents, another page count, other Excerpts... rows or other pages that are the user's
    (Whose pages..., or initials corrected under My info: as many pages, but not the same ones)."""
    return ([id(d) for d in job.docs] != [id(d) for d in copy.docs] or job.portions != copy.portions
            or job.invoice_pages() != copy.invoice_pages() or job.page_basis != copy.page_basis
            or job.front_owner != copy.front_owner or job.own != copy.own)


def _billing_changed(job, copy) -> bool:
    """The job was given what changes its invoice while Generate all worked on its copy: other pages
    (_pages_changed), other attorneys ticked, another Parties number, another speed for the form or for a firm
    (Job.speeds). What the batch billed is then not all there is to bill (an attorney ticked meanwhile was not
    invoiced)."""
    return (_pages_changed(job, copy) or job.ticked_keys() != copy.ticked_keys() or job.parties != copy.parties
            or job.case.get("delivery") != copy.case.get("delivery") or job.speeds != copy.speeds)


def _billed_by_copy(copy) -> list[str]:
    """The attorneys Generate all invoiced for a day (its copy, billed in full): each one ticked on it, as
    InvoiceOpts.skip_key names them ("key", and "key@ds" for another reporter's invoices)."""
    keys = copy.ticked_keys()
    return keys + [f"{k}@{r}" for r in copy.bill_reporters() if r != "me" for k in keys]


def _entry_text(a: Attorney) -> str:
    """An attorney row in a few words, for a question: "Smith Law Group (Dana Smith)", "Mr. Smith"."""
    return f"{a.firm} ({a.name})" if a.firm and a.name else a.firm or a.name


@dataclass(frozen=True)
class Problem:
    """One thing the job on screen needs looked at before Generate (MainWindow._problems): kind "missing" (a
    required field blank), "check" (the job's own checks, as the job list's tooltip says them, and the warnings in
    the cards) or "review" (the guessed fields, in one line); text, what it is; short, a few words of it for the
    status pill; go, what a click on it does (it goes to the place and lights it up; None: nowhere to go)."""
    kind: str
    text: str
    short: str = ""
    go: Callable[[], None] | None = None


def _short(text: str, width: int = 40) -> str:
    """A problem in a few words for the status pill: its first line, without a "⚠", cut at a word."""
    text = text.strip().lstrip("⚠").strip().split("\n")[0]
    text = text[:1].upper() + text[1:]
    return text if len(text) <= width else text[:width].rsplit(" ", 1)[0].rstrip(",:;") + "…"


def _kind_word(doc: Doc) -> str:
    """What a document of a job is, for its row in the list (where a transcript shows its day): "e-mail",
    "invoice", "photo", "PDF", "text" ("transcript" for a photo of one)."""
    kind = doc.regex.doc_kind
    if kind in ("transcript", "invoice"):
        return kind
    if kind == "email" or doc.ing.kind == "email":
        return "e-mail"
    return {"image": "photo", "pdf": "PDF"}.get(doc.ing.kind, "text")


def _cases(jobs: list[Job]) -> list[list[Job]]:
    """These jobs by case (batch.same_case: the same index number or case name), in their order."""
    cases: list[list[Job]] = []
    for j in jobs:
        same = next((g for g in cases if same_case(ident(g[0].case), ident(j.case))), None)
        if same is None:
            cases.append([j])
        else:
            same.append(j)
    return cases


def _held_cases(jobs) -> dict[str, str]:
    """The cases Generate all held (batch.HELD: a file of one of their days couldn't be saved), each with why:
    {"Jane Roe v. Sam Poe": "June 2026 ... Run Sheet.xlsx is open in another program - close it in Excel and try
    again"}, in full (the reason is what the user must act on)."""
    out: dict[str, str] = {}
    for j in jobs:
        for part in j.error.split("; "):
            if part.startswith(HELD) and j.title() not in out:
                out[j.title()] = part[len(HELD):].strip().split("couldn't be saved: ", 1)[-1]
    return out


def _reason(error: str, width: int = 160) -> str:
    """An error for the batch's message: without "PermissionError: " in front, and long enough to say why."""
    error = re.sub(r"^(?:PermissionError|OSError|ValueError): ", "", error)
    return error if len(error) <= width else error[:width - 1] + "…"


# ------------------------------------------------------------------ widgets

class FieldRow(QWidget):
    """Editor + suggestions menu + source badge for one form field (and, for a date, a calendar to pick it from).

    Typing makes the value the user's own (source "you") and calls on_edit(key). Values set by the program
    (set_text, set_state) do not call on_edit. The badge says where the value came from in plain words
    (BADGE_WORDS: "found" for a value read from the document); show_missing False (nothing loaded yet) leaves a
    required field that is blank unmarked.
    """

    def __init__(self, key: str, on_edit, multiline: bool = False, calendar: bool = False):
        super().__init__()
        self.key, self.on_edit, self.multiline = key, on_edit, multiline
        self.state = FieldState()
        self.show_missing = True
        lay = QHBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(6)
        if multiline:
            self.edit = QPlainTextEdit()
            sized(self.edit, "setFixedHeight", 58)
            self.edit.textChanged.connect(self._changed)
        else:
            self.edit = QLineEdit()
            self.edit.textEdited.connect(self._changed)
        # QPlainTextEdit has no textEdited (user-only) signal, so set_text raises this flag instead.
        self._loading = False
        lay.addWidget(self.edit, 1)
        self.calendar = None
        self.marks = None  # a function (year, month) -> {date: name}: the days its calendar marks (the court holidays)
        if calendar:
            self._add_calendar(lay)
        # the other readings of the field: "2 more ▾" (shown only when there are some)
        self.alts = QToolButton()
        self.alts.setObjectName("alts")
        self.alts.setText("▾")
        self.alts.setPopupMode(QToolButton.InstantPopup)
        self.alts.setMenu(QMenu(self.alts))
        sized(self.alts, "setMinimumWidth", 28)
        lay.addWidget(self.alts)
        self.badge = QLabel("")
        self.badge.setObjectName("badge")
        sized(self.badge, "setFixedSize", BADGE_WIDTH, 20)
        self.badge.setAlignment(Qt.AlignCenter)
        lay.addWidget(self.badge, 0, Qt.AlignTop if multiline else Qt.AlignVCenter)

    def _add_calendar(self, lay: QHBoxLayout) -> None:
        """A 📅 button beside the field: a month to click a day in (the field's own date selected), which
        becomes the user's value, as typed ("10/20/2026": a weekend or a court holiday isn't moved, unlike the quick
        buttons; the note under the dates warns of it). The days `marks` gives for the month shown (the court
        holidays) are in bold red, their name as the tooltip."""
        from PySide6.QtGui import QTextCharFormat
        from PySide6.QtWidgets import QCalendarWidget, QWidgetAction
        b = self.calendar = QToolButton()
        b.setText("📅")
        b.setToolTip("Pick the day from a calendar")
        b.setPopupMode(QToolButton.InstantPopup)
        menu = QMenu(b)
        cal = self.month = QCalendarWidget()
        cal.setGridVisible(False)
        cal.setVerticalHeaderFormat(QCalendarWidget.NoVerticalHeader)
        action = QWidgetAction(menu)
        action.setDefaultWidget(cal)
        menu.addAction(action)
        b.setMenu(menu)

        def mark(year: int, month: int) -> None:
            """The marked days of the month shown and the ones beside it (the grid shows a few of each)."""
            from PySide6.QtCore import QDate
            cal.setDateTextFormat(QDate(), QTextCharFormat())  # (a null date: every format cleared)
            if self.marks is None:
                return
            for y, m in ((year - (month == 1), month - 1 or 12), (year, month),
                         (year + (month == 12), month % 12 + 1)):
                for d, name in self.marks(y, m).items():
                    f = QTextCharFormat()
                    f.setForeground(QColor("#dc2626"))
                    f.setFontWeight(700)
                    f.setToolTip(name)
                    cal.setDateTextFormat(QDate(d.year, d.month, d.day), f)

        def shown() -> None:
            """The month of the field's date (today's when it has none), that day selected."""
            from PySide6.QtCore import QDate
            from datetime import datetime
            try:
                d = datetime.strptime(self.text(), "%m/%d/%Y").date()
            except ValueError:
                d = date.today()
            cal.setSelectedDate(QDate(d.year, d.month, d.day))
            mark(d.year, d.month)

        def picked(qd) -> None:
            """A day clicked: the field's value, and the calendar closes."""
            self.choose(f"{qd.month()}/{qd.day()}/{qd.year()}")
            menu.close()

        cal.currentPageChanged.connect(mark)
        menu.aboutToShow.connect(shown)
        cal.clicked.connect(picked)
        cal.activated.connect(picked)
        lay.addWidget(b)

    def text(self) -> str:
        """What the box shows, without leading or trailing spaces."""
        return self.edit.toPlainText().strip() if self.multiline else self.edit.text().strip()

    def set_text(self, t: str) -> None:
        """Shows t without treating it as typed by the user."""
        shown = self.edit.toPlainText() if self.multiline else self.edit.text()
        if shown == t or (self.edit.hasFocus() and shown.strip() == t):
            return  # unchanged: leave the cursor (and a space just typed) where they are
        self._loading = True
        if self.multiline:
            self.edit.setPlainText(t)
        else:
            self.edit.setText(t)
        self._loading = False

    def set_state(self, st: FieldState) -> None:
        """Shows a field's value, source badge and other suggestions."""
        self.state = st
        self.set_text(st.value)
        self._refresh()

    def _refresh(self) -> None:
        """Updates the badge, the review/missing highlight and the suggestions menu from self.state.
        A value not typed by the user is marked for review when the AI suggested it or its confidence is
        below 0.6."""
        st = self.state
        self.badge.setText(badge_word(st.source) if st.value else "")
        self.badge.setProperty("src", st.source if st.value else "")  # (the source itself: the badge's colour)
        self.badge.setToolTip(BADGE_TIPS.get(st.source, ""))
        review = st.value and st.source != SRC_USER and (st.confidence < 0.6 or st.source in (SRC_AI, SRC_RECORDS))
        self.edit.setProperty("review", bool(review))
        # a pill only for a value to look at (the AI's, the records', a guess): "found" on every field drowned them
        self.badge.setProperty("loud", bool(review))
        self.edit.setProperty("missing", self.key in REQUIRED_KEYS and not st.value and self.show_missing)
        others = [a for a in st.alternatives if a != st.value]
        menu = self.alts.menu()
        menu.clear()
        for alt in st.alternatives:
            act = menu.addAction(alt if len(alt) < 110 else alt[:107] + "...")
            act.triggered.connect(lambda _=False, v=alt: self.pick(v))
        self.alts.setVisible(bool(others))
        self.alts.setText(f"{len(others)} more ▾")
        self.alts.setToolTip(f"{plural(len(others), 'other reading')} of this field: click to pick one")
        for w in (self.badge, self.edit):
            repolish(w)

    def pick(self, value: str) -> None:
        """A suggestion picked from the menu: the user's own value (choose), unless it is the value already
        shown, which changes nothing and keeps its source (the transcript's page count picked again stays
        "PDF": as "you" the invoice would bill "the Pages field" instead of the pages counted)."""
        if value == self.state.value and self.text() == value:
            return
        self.choose(value)

    def choose(self, value: str) -> None:
        """Takes a value (a quick delivery date, a day picked from the calendar), as if the user had typed it."""
        self.set_text(value)
        self._changed()

    def _changed(self) -> None:
        """The user changed the text: it becomes their value (source "you", full confidence)."""
        if self._loading:
            return
        self.state = FieldState(self.text(), SRC_USER, 1.0, self.state.alternatives)
        self._refresh()
        self.on_edit(self.key)


# The status while nothing is loaded: plain text, not a pill (it looked like a button: the user, 2026-10-13)
IDLE_STATUS = "Drop a document to begin…"
COMPACT_BELOW = 960  # before the window is shown: window height (px at 100 % zoom) under which the drop zone is small
DROP_TEXT = "Drop documents here"
# the paste box's grey text: an example while it is open, a line to click once a job is loaded (_fit_paste)
PASTE_HINT = ("e.g. \"Please send the minutes for Smith v. Jones, Index 712345/2024, before Justice Lopez on 9/14/2026, "
              "expedited…\"")
PASTE_SLIM = "Paste an e-mail or notes…"
DROP_MORE_TEXT = "Drop more here, or"  # (the slim drop zone, Browse... beside it: the column is narrow)


SWIRLS = ("idle", "working")  # the drop zone's moods that show the swirling yin-yang
LINGER_MS = 10000  # how long the swirl goes on at least, once begun: one full turn of the loop


class Picture(QLabel):
    """The drop zone's yin-yang: a click on the picture itself (not the space beside it) calls on_click."""

    def __init__(self, on_click):
        super().__init__()
        self.on_click = on_click

    def mousePressEvent(self, e) -> None:
        """A left click on the picture calls on_click."""
        pix = self.pixmap()
        if e.button() == Qt.LeftButton and pix is not None and not pix.isNull():
            size = pix.deviceIndependentSize().toSize()
            shown = QRect(QPoint(0, 0), size)
            shown.moveCenter(self.rect().center())
            if shown.contains(e.position().toPoint()):
                self.on_click()
        super().mousePressEvent(e)


class DropZone(QFrame):
    """The big drop target. Dropped files go to on_files(paths), text to on_text(str), a picture to
    on_image(QImage); the Browse button calls on_browse(). It also shows the yin-yang pictures. Three layouts: big,
    compact (a smaller picture, set_compact) and slim (one row, once documents are loaded: set_slim), when the
    whole window takes drops anyway (MainWindow.dropEvent)."""

    def __init__(self, on_files, on_text, on_image, on_browse):
        super().__init__()
        self.setObjectName("drop")
        self.setAcceptDrops(True)
        self.on_files, self.on_text, self.on_image = on_files, on_text, on_image
        self.setMinimumHeight(z(215))
        lay = QVBoxLayout(self)
        lay.setAlignment(Qt.AlignCenter)
        self.icon = icon = QLabel("⭳")
        icon.setObjectName("dropIcon")
        icon.setAlignment(Qt.AlignCenter)
        self.yin = Picture(self.encore)
        self.yin.setAlignment(Qt.AlignCenter)
        self.yin.setVisible(False)
        lay.addWidget(self.yin)
        self.mood = None     # (mood, width) of the picture loaded now, so it is not reloaded for nothing
        self.wanted = None   # the mood asked for last: shown again by set_compact and rezoom, and after a swirl
        self.compact = False
        self.slim = False
        # the first day of the year the app was used ("2027-01-07", see note_opened): "done" is the New Year
        # yin-yang on that day only (checked each time, as the app may be left open overnight)
        self.new_year_day = ""
        # (mood, width) -> picture, made once: the window's _size_drop switches between the sizes as it measures
        self._pixmaps: dict = {}
        # "idle" and "working" swirl: the moving yin-yang (a loop of the video), played only while it is shown
        self.movie = QMovie(str(ASSETS / "yin_working.webp"), QByteArray(), self)  # (goes with the window)
        # A document is usually read in a blink: the swirl then goes on for one turn (linger_ms from when it
        # began) before the "ready" picture. A problem ("stumped") shows at once.
        self.linger_ms = LINGER_MS
        self._swirl_began = 0.0
        self._settle = QTimer(self)
        self._settle.setSingleShot(True)
        self._settle.timeout.connect(lambda: self.set_mood(self.wanted))
        self.movie.frameChanged.connect(self._frame)
        self.text = t = QLabel(DROP_TEXT)
        t.setObjectName("dropText")
        t.setAlignment(Qt.AlignCenter)
        sub = QLabel("PDF, photo or text file - or many at once, or a whole folder: "
                     "documents about the same case and date become one form.")
        sub.setObjectName("muted")
        sub.setAlignment(Qt.AlignCenter)
        sub.setWordWrap(True)
        self.sub = sub
        self.browse_btn = b = QPushButton("Browse...")
        b.clicked.connect(on_browse)
        self._fit_browse()
        for w in (icon, t, sub):
            lay.addWidget(w)
        lay.addSpacing(6)
        lay.addWidget(b, 0, Qt.AlignCenter)

    def set_mood(self, mood: str | None) -> None:
        """Shows the yin-yang: swirling while the app waits for documents ('idle') and reads them ('working'),
        then 'done' (after the swirl has gone on for linger_ms) or 'stumped' (None hides it). On the year's first
        day of use, 'done' is the New Year yin-yang ('newyear')."""
        self.wanted = mood
        self._settle.stop()
        if mood == "done" and self.new_year_day == date.today().isoformat():
            mood = "newyear"
        if mood in ("done", "newyear") and self.mood and self.mood[0] == "working" \
                and self.movie.state() == QMovie.Running:
            left = self.linger_ms - (time.monotonic() - self._swirl_began) * 1000
            if left > 20:
                self._swirl_size()  # (the layout may have changed meanwhile: set_compact)
                self._settle.start(int(left))  # (then set_mood(self.wanted) again)
                return
        width, height = self._size()
        self.sub.setVisible(not self.compact and not self.slim)
        if mood not in SWIRLS:
            self.movie.stop()
        if mood is None:
            self.mood = None  # (else the same mood asked for again is taken as shown, and doesn't swirl)
            self.yin.setVisible(False)
            self.icon.setVisible(not self.compact and not self.slim)
            self.setMinimumHeight(z(48) if self.slim else z(110) if self.compact else z(215))
            return
        self.icon.setVisible(False)
        self.yin.setVisible(True)
        self.setMinimumHeight(height)
        if mood in SWIRLS and self.movie.isValid() and self.movie.state() != QMovie.Running:
            self.movie.start()
        if mood == "working" and (not self.mood or self.mood[0] != "working"):
            self._swirl_began = time.monotonic()
        if (mood, width) == self.mood and self.yin.pixmap() and not self.yin.pixmap().isNull():
            return
        self.mood = (mood, width)
        if (mood, width) not in self._pixmaps:  # (idle swirls too: its first picture is the working one)
            picture = "working" if mood == "idle" else mood
            self._pixmaps[(mood, width)] = rounded(ASSETS / f"yin_{picture}.jpg", width, self._radius())
        self.yin.setPixmap(self._pixmaps[(mood, width)])
        self.yin.setToolTip({"idle": IDLE_STATUS, "working": "Working on it…",
                             "done": "Ready to fill!", "stumped": "Something needs your attention",
                             "newyear": "Happy New Year!"}.get(mood, ""))

    def encore(self) -> None:
        """A click on a still picture (ready, stumped, New Year): the yin-yang swirls for one turn, then the
        picture it showed comes back (unless the app asked for another one meanwhile)."""
        if self.mood is None or self.mood[0] in SWIRLS or not self.movie.isValid():
            return
        wanted = self.wanted
        self.set_mood("working")
        self.wanted = wanted
        self._settle.start(max(1, self.linger_ms))  # (then set_mood(self.wanted): the picture it was)

    def _frame(self, _n: int) -> None:
        """Shows the moving picture's next frame, at the size of the still one (not while the window is
        minimised: nobody sees it)."""
        if self.mood and self.mood[0] in SWIRLS and self.yin.isVisible() and not self.window().isMinimized():
            self.yin.setPixmap(round_corners(self.movie.currentPixmap(), self.mood[1], self._radius()))

    def _size(self) -> tuple[int, int]:
        """The picture's width and the drop zone's least height, in the layout it has now (set_compact,
        set_slim)."""
        if self.slim:
            return z(56), z(48)
        return (z(150), z(170)) if self.compact else (z(300), z(330))

    def _radius(self) -> int:
        """How round the picture's corners are (less on the slim bar's small one)."""
        return z(6) if self.slim else z(12)

    def _swirl_size(self) -> None:
        """While the swirl goes on a while yet (after a read, or a click: _settle), the layout of now: the swirl
        at its width, the line under it, the least height. (Before, a smaller window or another zoom then kept
        the big layout, and the window measured the wrong one.)"""
        width, height = self._size()
        self.sub.setVisible(not self.compact and not self.slim)
        self.setMinimumHeight(height)
        self.mood = ("working", width)

    def rezoom(self) -> None:
        """Shows the picture again at the new zoom's size (the swirl that goes on a while yet goes on)."""
        self._fit_browse()
        if self._settle.isActive():
            self._swirl_size()
            return
        self.mood = None
        self.set_mood(self.wanted)

    def set_compact(self, on: bool) -> None:
        """Switches to the small layout (on) or back: a smaller picture leaves room for the list of jobs, or
        for the inputs in a short window."""
        if on != self.compact:
            self.compact = on
            self.set_mood(self.wanted)

    def set_slim(self, on: bool) -> None:
        """Switches to one row (on: documents are loaded, and the list of jobs and the form need the room) or
        back: a small picture, "Drop more here, or" (DROP_MORE_TEXT) and Browse... side by side."""
        if on == self.slim:
            return
        self.slim = on
        self.layout().setDirection(QBoxLayout.LeftToRight if on else QBoxLayout.TopToBottom)
        self.text.setText(DROP_MORE_TEXT if on else DROP_TEXT)
        self.text.setWordWrap(on)  # (a narrow column: two lines rather than the row cut off)
        self._fit_browse()
        self.set_mood(self.wanted)

    def _fit_browse(self) -> None:
        """Browse... is narrower in the slim row (the left column's width must hold the picture, the words and it)."""
        self.browse_btn.setFixedWidth(z(84) if self.slim else z(120))

    def _hover(self, on: bool):
        """Highlights the drop zone while something is dragged over it (the style sheet's [hover="true"])."""
        self.setProperty("hover", on)
        repolish(self)

    def dragEnterEvent(self, e):
        """Takes a drag of files, text or a picture, and lights the zone up while it is over it."""
        md = e.mimeData()
        if md.hasUrls() or md.hasText() or md.hasImage():
            e.acceptProposedAction()
            self._hover(True)

    def dragLeaveEvent(self, e):
        """The drag left the zone without a drop: its highlight goes."""
        self._hover(False)

    def dropEvent(self, e):
        """Passes what was dropped to its handler (handle_mime). Text from one of the window's own boxes is copied,
        not moved: the box keeps it."""
        self._hover(False)
        handle_mime(e.mimeData(), self.on_files, self.on_text, self.on_image)
        if e.source() is not None:
            e.setDropAction(Qt.CopyAction)
            e.accept()
        else:
            e.acceptProposedAction()


def note_opened(s: Settings, today: date | None = None) -> bool:
    """Notes the app was opened today, and says whether today is its first day in a new year (the New Year
    yin-yang shows all that day, however often the app is opened again). The very first opening (no year noted
    yet: a new user, or one who just updated) is not one, nor is a year that went back (a clock set wrong). Saves
    the settings when they change (once a year), unless the settings file couldn't be read (Settings.unreadable:
    the defaults must not be written over it)."""
    today = today or date.today()
    year = str(today.year)
    if s.opened_year != year:
        # (none yet, or typed oddly in the file: "26", "2026.0" - not a year gone by)
        newer = len(s.opened_year) == 4 and s.opened_year.isdigit() and int(s.opened_year) < today.year
        if newer:
            s.new_year_day = today.isoformat()
        s.opened_year = year
        if not s.unreadable:
            try:
                s.save()
            except OSError as e:
                logfile.error("could not save the settings", e)
    return s.new_year_day == today.isoformat()


def handle_mime(md, on_files, on_text, on_image) -> bool:
    """Passes dropped or pasted data to one handler: local files first, then a picture, then text.
    Returns False when there was nothing it could use."""
    files = [u.toLocalFile() for u in md.urls() if u.isLocalFile()] if md.hasUrls() else []
    if files:
        on_files(files)
        return True
    if md.hasImage():
        img = md.imageData()
        if img is not None and not img.isNull():
            on_image(img)
            return True
    if md.hasText() and md.text().strip():
        on_text(md.text())
        return True
    return False


def _path_key(path: str) -> str:
    """One spelling per file, to tell whether it is loaded already."""
    return str(Path(path).resolve()).lower()


def card(title: str | None = None, key: str = "") -> tuple[QFrame, QVBoxLayout]:
    """A section card (with a heading when title is given) and the layout to put its contents in. The frame keeps
    its heading label as .heading (None without a title). With a key, the card folds to its heading and a line
    that sums it up: a ▾ button before the heading (or a click on the heading) folds it, and its contents are in
    .body (hidden when folded; the card itself stays shown), with .fold (the button), .summary (the line) and .key
    (Settings.folded_cards); MainWindow._fold_card does the folding."""
    fr = QFrame()
    fr.setObjectName("card")
    lay = QVBoxLayout(fr)
    lay.setContentsMargins(16, 14, 16, 16)
    lay.setSpacing(10)
    fr.heading = fr.fold = fr.summary = fr.body = None
    fr.key = key
    if title:
        lbl = ClickLabel(title) if key else QLabel(title)
        lbl.setObjectName("section")
        fr.heading = lbl
        if key:
            top = QHBoxLayout()
            top.setSpacing(6)
            fr.fold = QToolButton()
            fr.fold.setObjectName("fold")
            fr.fold.setArrowType(Qt.DownArrow)  # (drawn, not a character the fonts must be searched for)
            fr.fold.setToolTip("Fold this card to one line (click again to unfold it)")
            top.addWidget(fr.fold)
            top.addWidget(lbl)
            lbl.set_clickable(True)
            fr.summary = ElidedLabel()
            fr.summary.setObjectName("summary")
            fr.summary.setVisible(False)
            top.addWidget(fr.summary, 1)
            lay.addLayout(top)
        else:
            lay.addWidget(lbl)
    if key:
        fr.body = QWidget()
        body = QVBoxLayout(fr.body)
        body.setContentsMargins(0, 0, 0, 0)
        body.setSpacing(10)
        lay.addWidget(fr.body)
        return fr, body
    return fr, lay


# -------------------------------------------------------------- main window

class MainWindow(QMainWindow):
    """The app's window. It holds a list of jobs (one per case and date; more than one is a batch) and shows
    one of them, self.cur, in the editor on the right.

    Background work and the state that guards it:
      gen       bumped by "New job"; a result whose gen is old is dropped.
      work      reads and batches still running (the progress bar shows while it is above 0).
      ai_pending  AI questions not answered yet (waiting their turn on ai_pool, one at a time, or being asked).
      filling   "Generate all" is making files on another thread. It works on copies of the jobs and the
                settings. Until it is done, Generate, Generate all and New job only say "one moment".
      _loading  files being read now, so dropping the same file again does not load it twice.
    """

    def __init__(self, settings: Settings, app: QApplication):
        super().__init__()
        self.s, self.app = settings, app
        # the user's answers about rows that may be one firm: kept ones (Settings.firm_answers) and those for now
        # only (Remember unticked: until the app is closed)
        self._session_answers: list[list] = []
        self._use_answers()
        self.runner = Runner()
        # The AI's questions go on a pool of their own, one at a time: Ollama answers one at a time anyway, and
        # on the global pool each document's question held a thread for up to a minute, so reads and Generate
        # all waited behind them on a laptop with few cores
        self.ai_pool = QThreadPool(self)
        self.ai_pool.setMaxThreadCount(1)
        self.ai_runner = Runner(self.ai_pool)
        self.ai = OllamaExtractor(settings)
        self.ai_ok = False
        self.gen = 0  # bumped by "New job": results of work started before it are dropped
        self.jobs: list[Job] = [Job()]  # more than one = a batch
        self.cur = self.jobs[0]         # the job shown in the editor
        self.ai_pending = 0
        self.work = 0                   # documents being read / batches being made in the background
        self.filling = False            # "Generate all" is running on another thread
        self._loading: set[str] = set()  # files being read right now (so a second drop doesn't add them twice)
        # the incomplete rate sheets warned about this session, by rates.sheet_fingerprint (one edited and read
        # again is warned about again): see _check_rate_sheet
        self._sheets_warned: set[str] = set()
        self._days_box = None  # the question about a rate sheet's old Days while it is open (_check_old_days)
        self._days_later: set[str] = set()  # the sheets answered Not now this session (_check_old_days)
        # the rate sheet changed while Generate all ran: checked when it ends (the sheet in use before then, for Go
        # back; None: nothing waits)
        self._sheet_check_later: str | None = None
        # the incomplete rate sheet opened from that warning to be filled in: read again and checked each time it
        # is saved (_watch_sheet, _sheet_saved). Its path and the file's size and time as last read; the watcher
        # is made when first needed
        self._sheet_watched: tuple[Path, tuple] | None = None
        self._sheet_watch = None
        self._sheet_timer = QTimer(self)
        self._sheet_timer.setSingleShot(True)
        self._sheet_timer.setInterval(SHEET_SAVED_MS)
        self._sheet_timer.timeout.connect(self._sheet_saved)
        # the Run sheet box was ticked or unticked by the user since the last New job: it is no longer ticked by
        # how many reporters wrote the transcripts (see _auto_runsheet)
        self._runsheet_touched = False
        self._runsheet_auto: bool | None = None  # what _auto_runsheet decided; None = as Settings.outputs say
        self._math_previewed = False  # the last preview showed "The math" (see _preview)
        self.excerpts = None  # the Excerpts window, made when first opened (see _who_ordered)
        # what New job cleared, for its Undo (None: nothing to undo; see new_job, _undo_clear)
        self._undo: dict | None = None
        # the Records button's count: bumped by each count asked for, so an older one finishing late is dropped
        # (see _count_records); when the window was last activated (counted again then, not more than every 2 s)
        self._counting = 0
        self._activated = 0.0
        self._count_later = QTimer(self)  # (back within 2 s, from a tick in Records: counted once the 2 s are up)
        self._count_later.setSingleShot(True)
        self._count_later.timeout.connect(self._count_on_activation)
        self._mirrored = self._csv_columns()  # what the CSV copies have columns for (see _records_changed)
        # what the fields complete from (records.past_values): one list each, filled once the records are read
        self._completions: dict[str, QStringListModel] = {}
        self._left_mode = ""  # "single" / "batch": the left column's divider was placed for it (_place_left_split)

        self.setWindowTitle("YinItAgreementForm")
        if zooming.zoom() != settings.zoom:  # (main() sets it before the style sheet is made)
            zooming.set_zoom(settings.zoom)
            apply_theme(app, settings.theme)
        self._fit_minimum()
        self._build()
        if settings.window_geometry:
            self.restoreGeometry(QByteArray.fromBase64(settings.window_geometry.encode()))
        else:
            self.resize(z(1320), z(860))
        self._fit_screen()
        self._wheel = 0  # Ctrl + wheel turned since the last zoom step (see eventFilter)
        # Ctrl + the mouse wheel zooms, the form follows its width, and files dragged over a text box are read
        # (eventFilter)
        app.installEventFilter(self)
        self._show_case()
        self._set_status(IDLE_STATUS, "")
        QTimer.singleShot(50, self._startup)

    # The editor always works on the current job.
    @property
    def case(self) -> CaseInfo:
        """The current job's case (what the editor shows)."""
        return self.cur.case

    @case.setter
    def case(self, value: CaseInfo) -> None:
        self.cur.case = value

    @property
    def proc_touched(self) -> bool:
        """The user changed the current job's proceeding types, so a new merge must keep them."""
        return self.cur.proc_touched

    @proc_touched.setter
    def proc_touched(self, value: bool) -> None:
        self.cur.proc_touched = value

    @property
    def att_touched(self) -> bool:
        """The user changed the current job's attorney table, so a new merge must keep it."""
        return self.cur.att_touched

    @att_touched.setter
    def att_touched(self, value: bool) -> None:
        self.cur.att_touched = value

    # ---------------------------------------------------------- layout
    def _build(self):
        """Builds the window: the status and the Records, New job and Settings buttons at the right of the menu
        bar, the left column (drop zone, job list, paste box), the scrolling form on the right with its cards, the
        Outputs box (a row, with the Generate buttons) under them, the menu and the shortcuts."""
        central = QWidget()
        central.setObjectName("central")
        # On a small screen (or a large zoom) the whole window scrolls, rather than the Outputs box and the
        # Generate buttons being pushed off the bottom of the screen
        self.page_scroll = outer = QScrollArea()
        outer.setWidgetResizable(True)
        outer.setFrameShape(QFrame.NoFrame)
        outer.setWidget(central)
        self.setCentralWidget(outer)
        root = QVBoxLayout(central)
        root.setContentsMargins(18, 8, 18, 14)
        root.setSpacing(10)
        self.setAcceptDrops(True)  # documents dropped anywhere in the window are read (dropEvent, eventFilter)

        # the status and the buttons, at the right of the menu bar (the window's title names the app: a header
        # row of its own took room the form needs on a laptop's screen)
        # (kept on the window: PySide doesn't hand a corner widget over to the menu bar, and would delete it once the
        # local name is gone; _fit_corner sets it again)
        self.head_box = head_box = QWidget()
        head_box.setObjectName("menuCorner")
        head = QHBoxLayout(head_box)
        head.setContentsMargins(0, 3, 12, 3)
        head.setSpacing(8)
        self.status = ClickLabel("")
        self.status.setObjectName("status")
        self.status.clicked.connect(self._go_to_problem)
        self.status.setAttribute(Qt.WA_Hover)  # (its :hover look says it can be clicked)
        head.addWidget(self.status)
        # a newer version is out (_show_update): up here, where it is seen, until it is installed
        self.update_btn = QPushButton()
        self.update_btn.setObjectName("update")
        self.update_btn.setProperty("help_topic", "trouble-install")
        self.update_btn.clicked.connect(lambda: open_url(self.s.update_page))
        self.update_btn.setVisible(False)
        head.addWidget(self.update_btn)
        self.new_btn = new_btn = QPushButton("New job")
        new_btn.setToolTip("Clear everything and start over (Ctrl+N); Undo at the bottom of the window brings it back\n"
                           "for a few seconds (not while documents are still being read)")
        new_btn.clicked.connect(self.new_job)
        set_btn = QPushButton("⚙  Settings")
        set_btn.clicked.connect(self.open_settings)
        self.rec_btn = rec_btn = QPushButton("Records")
        rec_btn.setObjectName("records")  # a hue of its own: the records are apart from making the forms
        rec_btn.setToolTip(RECORDS_TIP)
        rec_btn.clicked.connect(self.open_records)
        head.addWidget(rec_btn)
        head.addWidget(new_btn)
        head.addWidget(set_btn)
        self.menuBar().setCornerWidget(head_box, Qt.TopRightCorner)
        self.busy = QProgressBar()
        self.busy.setRange(0, 0)
        self.busy.setTextVisible(False)
        self.busy.setVisible(False)
        root.addWidget(self.busy)

        # the left column and the form, side by side above the Outputs row
        split = QSplitter(Qt.Horizontal)
        split.setChildrenCollapsible(False)
        root.addWidget(split, 1)

        # left column
        left = QWidget()
        ll = QVBoxLayout(left)
        ll.setContentsMargins(0, 0, 8, 0)
        ll.setSpacing(10)
        self.drop = DropZone(self.add_files, self.add_text, self.add_qimage, self.browse)
        self.drop.setProperty("help_topic", "input-documents")
        note_opened(self.s)
        self.drop.new_year_day = self.s.new_year_day
        ll.addWidget(self.drop)
        self.jobs_label = QLabel("Jobs")
        self.jobs_label.setObjectName("fieldLabel")
        self.jobs_label.setToolTip("⚠ = something to check before Generate (a field missing, no attorney ticked,\n"
                                   "Whose pages… or Excerpts… to choose, Parties counting a party no firm\n"
                                   "named, a speed to choose, documents whose\n"
                                   "index numbers don't match, more than two speeds on the same pages, no pages\n"
                                   "to bill...). Hover over a job to see what. ✓ = saved.\n\n" + TICK_TIP + "\n\n"
                                   + YOURS_TIP)
        ll.addWidget(self.jobs_label)
        # the jobs (a batch) or the job on screen, a row each, and under a job of several documents a row per
        # document (their own pages, the right-click to move one out); see _refresh_jobs. A job's tick says
        # whether the cards on the right and Generate all count its day (_job_ticked)
        self.job_list = QTreeWidget()
        self.job_list.setProperty("help_topic", "batch-job-list")
        self.job_list.setColumnCount(len(JOB_COLS))
        self.job_list.setHeaderLabels(JOB_COLS)
        for c, name in enumerate(JOB_COLS):
            self.job_list.headerItem().setToolTip(c, JOB_TIPS[name])
        hh = self.job_list.header()
        # a click on a heading sorts the list by it, again the other way (by date, oldest first, at first:
        # _job_sort; sorted in _refresh_job_labels); a right-click shows or hides the Day column
        hh.setSectionsClickable(True)
        hh.setSortIndicatorShown(True)
        hh.setSortIndicator(*self._job_sort())
        hh.sortIndicatorChanged.connect(self._job_sort_changed)
        hh.setContextMenuPolicy(Qt.CustomContextMenu)
        hh.customContextMenuRequested.connect(self._job_header_menu)
        hh.setStretchLastSection(False)
        hh.setSectionResizeMode(QHeaderView.Fixed)  # (as wide as their text: _fit_job_columns)
        hh.setSectionResizeMode(0, QHeaderView.Stretch)
        self.job_list.setUniformRowHeights(True)
        self.job_list.setAllColumnsShowFocus(True)
        sized(self.job_list, "setIndentation", 14)
        self.job_list.currentItemChanged.connect(lambda item, _: self._job_selected(self._job_of(item)))
        self.job_list.itemChanged.connect(self._job_ticked)
        self.job_list.itemCollapsed.connect(lambda item: self._job_folded(item, True))
        self.job_list.itemExpanded.connect(lambda item: self._job_folded(item, False))
        self.job_list.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.job_list.setTextElideMode(Qt.ElideRight)
        self.job_list.setContextMenuPolicy(Qt.CustomContextMenu)
        self.job_list.customContextMenuRequested.connect(self._job_menu)
        self._collapsed: set[Job] = set()  # the jobs whose documents the user folded away (kept on a rebuild)
        # the job list above the paste box, a divider between them to drag (placed for one job or a batch:
        # _place_left_split)
        self.left_split = QSplitter(Qt.Vertical)
        self.left_split.setChildrenCollapsible(False)
        top = QWidget()
        tl = QVBoxLayout(top)
        tl.setContentsMargins(0, 0, 0, 4)
        tl.setSpacing(8)
        tl.addWidget(self.job_list, 1)
        # what the documents say about the index number that needs a look (Job.index_number_notes): one with
        # none, or documents whose numbers don't match, with a link to keep them together
        self.index_note = QLabel("")
        self.index_note.setObjectName("pagesWarn")
        self.index_note.setWordWrap(True)
        self.index_note.setTextFormat(Qt.RichText)
        self.index_note.linkActivated.connect(lambda _: self._keep_index_numbers())
        self.index_note.setVisible(False)
        tl.addWidget(self.index_note)
        self.left_split.addWidget(top)
        # (one line once a job is loaded, while it is empty and not typed in: _fit_paste)
        self.paste_pane = bottom = QWidget()
        bl = QVBoxLayout(bottom)
        bl.setContentsMargins(0, 4, 0, 0)
        bl.setSpacing(8)
        self.paste_label = lbl = QLabel("…or paste an e-mail or notes")
        lbl.setObjectName("fieldLabel")
        bl.addWidget(lbl)
        self.paste = QPlainTextEdit()
        self.paste.setPlaceholderText(PASTE_HINT)
        sized(self.paste, "setMinimumHeight", 52)
        self._paste_slim = False
        self._split_unslim = None  # (the divider where it was before the box went to one line: _save_splits)
        # it takes the room left over, but doesn't ask for any: otherwise the left column's scroll area counts
        # its preferred 190 px and shows a scroll bar in a window with room to spare
        self.paste.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Ignored)
        self.paste.textChanged.connect(self._update_header_buttons)
        self.paste.textChanged.connect(self._fit_paste)
        self.paste.setProperty("help_topic", "input-email")
        bl.addWidget(self.paste, 1)
        self.extract_btn = pb = QPushButton("Extract from text")
        pb.clicked.connect(self._extract_paste)
        bl.addWidget(pb)
        self.left_split.addWidget(bottom)
        self.left_split.setStretchFactor(0, 1)
        self.left_split.setStretchFactor(1, 1)
        # it takes the room left over, asking only for its least (as the paste box did on its own): the column's
        # height asked would otherwise leave no room for the big picture (see _size_drop)
        self.left_split.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Ignored)
        ll.addWidget(self.left_split, 1)
        self.ai_label = QLabel("Checking AI…")
        self.ai_label.setObjectName("muted")
        self.ai_label.setWordWrap(True)
        self.ai_label.setTextFormat(Qt.RichText)
        self.ai_label.linkActivated.connect(self._ai_help)
        ll.addWidget(self.ai_label)
        # In a short window the column scrolls, rather than its parts being drawn over each other (the column needs
        # about 620 px in a batch).
        self.left_scroll = left_scroll = QScrollArea()
        left_scroll.setWidgetResizable(True)
        left_scroll.setFrameShape(QFrame.NoFrame)
        left_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        left_scroll.setWidget(left)
        left_scroll.setMinimumWidth(left.minimumSizeHint().width() + 12)  # (room for the scroll bar)
        split.addWidget(left_scroll)

        # right column (scrollable form)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        inner = QWidget()
        rl = QVBoxLayout(inner)
        rl.setContentsMargins(8, 0, 8, 0)
        rl.setSpacing(12)
        scroll.setWidget(inner)
        split.addWidget(scroll)
        self.form_scroll = scroll
        self._pairs: list[tuple[QGridLayout, QWidget, QWidget]] = []  # the cards' two columns (see two_cols)
        split.setStretchFactor(0, 0)
        split.setStretchFactor(1, 1)
        split.setSizes([360, 940])

        self.rows: dict[str, FieldRow] = {}
        self.cards: dict[str, QFrame] = {}  # the cards that fold, by key (see card, _fold_card)

        def row(form: QFormLayout, key: str, multiline=False, calendar=False):
            """The FieldRow of the case field `key`, with its label, added to `form` and kept in self.rows."""
            r = FieldRow(key, self._field_edited, multiline, calendar)
            if key in FIELD_TOPICS:
                r.setProperty("help_topic", FIELD_TOPICS[key])
            self.rows[key] = r
            lab = QLabel(FIELD_LABELS[key])
            lab.setObjectName("fieldLabel")
            form.addRow(lab, r)
            return r

        def folding(title: str, key: str) -> tuple[QFrame, QVBoxLayout]:
            """A card that folds to a line (card), kept in self.cards."""
            fr, lay = card(title, key)
            self.cards[key] = fr
            fr.setProperty("help_topic", CARD_TOPICS.get(key, ""))
            fr.fold.clicked.connect(lambda _=False, f=fr: self._fold_card(f, not f.body.isHidden()))
            fr.heading.clicked.connect(lambda f=fr: self._fold_card(f, not f.body.isHidden()))
            return fr, lay

        def form_in(lay: QVBoxLayout) -> QFormLayout:
            """A form (labels right-aligned beside their fields) added to `lay`."""
            f = QFormLayout()
            f.setLabelAlignment(Qt.AlignRight | Qt.AlignVCenter)
            f.setHorizontalSpacing(12)
            f.setVerticalSpacing(8)
            lay.addLayout(f)
            return f

        def two_cols(lay: QVBoxLayout) -> tuple[QVBoxLayout, QVBoxLayout]:
            """Two columns side by side in a card, one above the other when the form is too narrow for both
            (a small screen, Windows scaling or a zoom: see _place_pairs)."""
            holder = QWidget()
            g = QGridLayout(holder)
            g.setContentsMargins(0, 0, 0, 0)
            g.setHorizontalSpacing(18)
            g.setVerticalSpacing(8)
            wa, wb = QWidget(), QWidget()
            cols = []
            for w in (wa, wb):
                v = QVBoxLayout(w)
                v.setContentsMargins(0, 0, 0, 0)
                cols.append(v)
            g.addWidget(wa, 0, 0)  # side by side until _place_pairs measures the room
            g.addWidget(wb, 0, 1)
            g.setColumnStretch(0, 1)
            g.setColumnStretch(1, 1)
            lay.addWidget(holder)
            self._pairs.append((g, wa, wb))
            return cols[0], cols[1]

        # Before you generate: what the job still needs, each a link to where it is fixed (_show_todo; after
        # GOV.UK's error summary: one list at the top of the page, each item going to its place)
        self.todo = QFrame()
        self.todo.setObjectName("todo")
        self.todo.setProperty("help_topic", "check-status")
        todo_lay = QVBoxLayout(self.todo)
        todo_lay.setContentsMargins(14, 8, 14, 8)
        self.todo_text = QLabel("")
        self.todo_text.setTextFormat(Qt.RichText)
        self.todo_text.setWordWrap(True)
        self.todo_text.linkActivated.connect(self._todo_clicked)
        todo_lay.addWidget(self.todo_text)
        self.todo.setVisible(False)
        self._todo: list[Problem] = []  # (the problems listed, in the links' order)
        rl.addWidget(self.todo)

        # While nothing is loaded: the jobs made lately, to open one again, and how to begin (_show_start)
        self.start_card, sl = card("Recent jobs")
        self.start_list = QVBoxLayout()
        self.start_list.setSpacing(2)
        sl.addLayout(self.start_list)
        foot_row = QHBoxLayout()
        self.start_text = QLabel(START_TEXT)
        self.start_text.setObjectName("muted")
        self.start_text.setWordWrap(True)
        foot_row.addWidget(self.start_text, 1)
        recent_docs = QToolButton()
        recent_docs.setText("Recent documents ▾")
        recent_docs.setObjectName("quick")
        recent_docs.setPopupMode(QToolButton.InstantPopup)
        self.recent_docs = recent_docs  # (its menu is File -> Open recent's, set by _build_menu)
        foot_row.addWidget(recent_docs, 0, Qt.AlignTop)
        sl.addLayout(foot_row)
        # records.recent_jobs, read at startup and when the records change (_read_records_side)
        self._recent_jobs: list[tuple[str, str, str, str]] = []
        rl.addWidget(self.start_card)

        # Case card: two columns, then the case's name and the type of proceeding
        c, cl = folding("Case", "case")
        colA, colB = two_cols(cl)
        fa, fb = form_in(colA), form_in(colB)
        for k in ("court", "county", "part"):
            row(fa, k)
        for k in ("index_no", "judge", "dates"):
            row(fb, k)
        fc = form_in(cl)
        row(fc, "case_name", multiline=True)
        # the types in one row, or two when the form is narrow (_place_pairs: at 140 % the six were wider than the
        # form of a 1320-pixel window)
        self.proc_grid = pr = QGridLayout()
        pr.setHorizontalSpacing(12)
        pr.setVerticalSpacing(4)
        self.proc_boxes: dict[str, QCheckBox] = {}
        for i, p in enumerate(PROC_TYPES):
            cb = QCheckBox(p)
            cb.toggled.connect(self._proc_toggled)
            self.proc_boxes[p] = cb
            pr.addWidget(cb, 0, i)
        pr.setColumnStretch(len(PROC_TYPES), 1)
        self._proc_per_row = len(PROC_TYPES)
        lab = QLabel("Proceeding")
        lab.setObjectName("fieldLabel")
        fc.addRow(lab, pr)
        self.proc_form = fc  # (its labels' width: the room the types have beside them, _place_pairs)
        row(fc, "proc_other")
        rl.addWidget(c)

        # Attorneys card: the table shows who (tick, name, firm, role); the row selected is spelled out under it
        c, cl = folding("Attorneys: tick who ordered", "attorneys")
        c.heading.setToolTip(
            "Each checked attorney gets a minute agreement (and an invoice) of its own: the agreement\n"
            "shows the pages that attorney ordered, whoever wrote them; the invoice bills your pages of them.")
        self.att = QTableWidget(0, len(ATT_COLS))
        self.att.setHorizontalHeaderLabels(ATT_COLS)
        self.att.verticalHeader().setVisible(False)
        self.att.setAlternatingRowColors(True)
        self.att.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.att.setWordWrap(True)
        hh = self.att.horizontalHeader()
        hh.setSectionResizeMode(QHeaderView.Interactive)
        hh.setStretchLastSection(False)
        for i, w in enumerate([30, 140, 190, 230, 115, 105, 190, 170, 60]):
            sized(self.att, "setColumnWidth", i, w, keep=1)
        for i in ATT_DETAIL_COLS:  # (in the details under the table; the columns stay, as the rows' data)
            self.att.setColumnHidden(i, True)
        hh.setSectionResizeMode(1, QHeaderView.Stretch)
        hh.setSectionResizeMode(2, QHeaderView.Stretch)
        self.att.setIconSize(QSize(*badge_size(z(TAG_HEIGHT))))  # (each firm's tag and stripe: _mark_row)
        self.att.setItemDelegateForColumn(1, StripeDelegate(self.att))
        sized(self.att, "setMinimumHeight", 150)
        self.att.itemChanged.connect(self._att_changed)
        self.att.currentCellChanged.connect(lambda *_: self._show_att_details())
        cl.addWidget(self.att)
        self._build_att_details(cl, two_cols, form_in)
        # two rows that may be one firm or attorney ("Smith Law" and "Smith Law Group"): the user says (Same firm,
        # Not the same), and the answer is kept for later documents too (batch.firm_questions, Settings.firm_answers),
        # or with Remember unticked only until the app is closed
        self.firm_ask_box = QWidget()
        fq = QHBoxLayout(self.firm_ask_box)
        fq.setContentsMargins(0, 0, 0, 0)
        fq.setSpacing(8)
        self.firm_ask = QLabel("")
        self.firm_ask.setObjectName("speedAsk")
        self.firm_ask.setWordWrap(True)
        fq.addWidget(self.firm_ask, 1)
        self.firm_same = QPushButton("Same firm")
        self.firm_same.setToolTip("They are one firm (or one attorney): one row, one invoice, one agreement")
        self.firm_same.clicked.connect(lambda: self._answer_firm(True))
        self.firm_apart = QPushButton("Not the same")
        self.firm_apart.setToolTip("They are two firms (or two attorneys): two rows, each billed on its own")
        self.firm_apart.clicked.connect(lambda: self._answer_firm(False))
        fq.addWidget(self.firm_same, 0, Qt.AlignTop)
        fq.addWidget(self.firm_apart, 0, Qt.AlignTop)
        self.firm_remember = QCheckBox("Remember")
        self.firm_remember.setChecked(True)
        self.firm_remember.setToolTip("Keep the answer for later documents naming these two (Settings → Invoice →\n"
                                      "Firms you answered lists them, to forget one). Unticked: for now only.")
        fq.addWidget(self.firm_remember, 0, Qt.AlignTop)
        self.firm_ask_box.setVisible(False)
        self._firm_q = None  # (the days asked about, (row, row)) of the question shown
        cl.addWidget(self.firm_ask_box)
        br = QHBoxLayout()
        add = QPushButton("+ Add attorney")
        add.clicked.connect(self._add_att_row)
        rem = QPushButton("Remove selected")
        rem.clicked.connect(self._remove_att_rows)
        hint = QLabel("Each ticked attorney gets a form of its own. Click a row to see and change all its details "
                      "below.")
        hint.setObjectName("muted")
        hint.setWordWrap(True)
        br.addWidget(add)
        br.addWidget(rem)
        br.addSpacing(12)
        br.addWidget(hint, 1)
        cl.addLayout(br)
        rl.addWidget(c)

        # Speed, pages and dates: the rate sheet, the one speed and one rate the form names (among the
        # speeds the invoice offers, ticked under the rate sheet), No. of copies, the pages and the dates
        c, cl = folding("Speed, pages and dates", "details")
        colA, colB = two_cols(cl)
        fa, fb = form_in(colA), form_in(colB)

        # Rate sheet, with its folder and reload buttons
        sheet_row = QHBoxLayout()
        self.sheet_box = QuietCombo()
        self.sheet_box.setMinimumContentsLength(18)
        self.sheet_box.setSizeAdjustPolicy(QComboBox.AdjustToMinimumContentsLengthWithIcon)
        self.sheet_box.currentIndexChanged.connect(self._sheet_changed)
        folder = QToolButton()
        folder.setText("📂")
        folder.setToolTip("Open the rate sheets folder (add or edit CSV files, then click ⟳)")
        folder.clicked.connect(self._open_sheets_folder)
        reload_btn = QToolButton()
        reload_btn.setText("⟳")
        reload_btn.setToolTip("Reload rate sheets")
        reload_btn.clicked.connect(self._reload_sheets)
        sheet_row.addWidget(self.sheet_box, 1)
        sheet_row.addWidget(folder)
        sheet_row.addWidget(reload_btn)
        lab = QLabel("Rate sheet")
        lab.setObjectName("fieldLabel")
        fa.addRow(lab, sheet_row)
        # the speeds the invoice offers, each at its price (_fill_invoice_speeds): here, by the rate sheet they come
        # from and the speed the agreement names (one of them), not in the Outputs box (the user, 2026-10-22:
        # speeds were chosen in places apart); still kept for every job (Settings.invoice_speeds)
        speeds = QWidget()
        speeds.setToolTip(SPEEDS_TIP)
        self.inv_speeds_grid = QGridLayout(speeds)
        self.inv_speeds_grid.setContentsMargins(0, 0, 0, 0)
        self.inv_speeds_grid.setHorizontalSpacing(12)
        self.inv_speeds_grid.setVerticalSpacing(2)
        self.inv_speed_boxes: dict[str, QCheckBox] = {}
        lab = QLabel("Speeds offered")
        lab.setObjectName("fieldLabel")
        lab.setToolTip(SPEEDS_TIP)
        fa.addRow(lab, speeds)
        # the one speed the minute agreement form names, at its one rate per page
        speed_row = QHBoxLayout()
        speed_row.setSpacing(6)
        self.delivery = QuietCombo()
        # (its items are long, "Expedite · $5.40/pg · copy $1.10 · 5 days": sized to them, the card's two
        # columns no longer fitted side by side in a window of an ordinary width)
        self.delivery.setMinimumContentsLength(20)
        self.delivery.setSizeAdjustPolicy(QComboBox.AdjustToMinimumContentsLengthWithIcon)
        self.delivery.setToolTip(AGREEMENT_SPEED_TIP)
        self.delivery.currentIndexChanged.connect(self._delivery_changed)
        speed_row.addWidget(self.delivery, 1)
        self.speed_reset = QToolButton()
        self.speed_reset.setText("↺")
        self.speed_reset.clicked.connect(self._speed_to_rule)
        speed_row.addWidget(self.speed_reset)
        self.speed_badge = QLabel("")
        self.speed_badge.setObjectName("badge")
        sized(self.speed_badge, "setFixedSize", BADGE_WIDTH, 20)
        self.speed_badge.setAlignment(Qt.AlignCenter)
        speed_row.addWidget(self.speed_badge)
        lab = QLabel("Speed")
        lab.setObjectName("fieldLabel")
        lab.setToolTip(AGREEMENT_SPEED_TIP)
        fa.addRow(lab, speed_row)
        # an e-mail asks for another speed than the rule's: the user is asked to choose (Job.speed_question)
        self.speed_ask_row = QHBoxLayout()
        self.speed_ask_row.setSpacing(8)
        self.speed_ask = QLabel("")
        self.speed_ask.setObjectName("speedAsk")
        self.speed_ask.setWordWrap(True)
        self.speed_ask_row.addWidget(self.speed_ask, 1)
        self.speed_keep = QPushButton("Keep")
        self.speed_keep.setToolTip("Keep the speed shown for this job (the e-mail's request is then answered)")
        self.speed_keep.clicked.connect(self._keep_speed)
        self.speed_ask_row.addWidget(self.speed_keep, 0, Qt.AlignTop)
        fa.addRow("", self.speed_ask_row)
        fa.setRowVisible(self.speed_ask_row, False)
        self.speed_form = fa
        row(fa, "rate")
        self.rate_info = QLabel("")
        self.rate_info.setObjectName("muted")
        self.rate_info.setWordWrap(True)
        fa.addRow("", self.rate_info)
        for k in ("copies", "est_pages"):
            row(fb, k)
        fb.labelForField(self.rows["est_pages"]).setToolTip(EST_PAGES_TIP)
        self.rows["est_pages"].edit.setToolTip(EST_PAGES_TIP)
        # what the number counts ("Transcript pages 378–460 · excludes the 13 word-index pages after them"), and
        # a warning when the PDF's readings of its pages disagree (Job.pages_help); hidden without a transcript
        help_box = QVBoxLayout()
        help_box.setSpacing(4)
        self.pages_help = QLabel("")
        self.pages_help.setObjectName("muted")
        self.pages_help.setWordWrap(True)
        self.pages_help.setToolTip(EST_PAGES_TIP)
        help_box.addWidget(self.pages_help)
        self.pages_warn = QLabel("")
        self.pages_warn.setObjectName("pagesWarn")
        self.pages_warn.setWordWrap(True)
        help_box.addWidget(self.pages_warn)
        fb.addRow("", help_box)
        self.pages_help_row = help_box
        self.order_form = fb
        row(fb, "delivery_date", calendar=True)
        # three buttons a row: six in one row would make the card too wide for a small screen
        quick = QGridLayout()
        quick.setHorizontalSpacing(4)
        quick.setVerticalSpacing(4)
        self.quick_buttons = []  # (button, label, days, months): their tooltips name the day (_quick_tips)
        for n, (label, days, months) in enumerate((("Today", 0, 0), ("Tomorrow", 1, 0), ("1 week", 7, 0),
                                                   ("2 weeks", 14, 0), ("3 weeks", 21, 0), ("1 month", 0, 1))):
            b = QToolButton()
            b.setText(label)
            b.setObjectName("quick")
            b.clicked.connect(lambda _=False, d=days, m=months: self.rows["delivery_date"].choose(
                quick_day(self.s, d, m)))
            self.quick_buttons.append((b, label, days, months))
            quick.addWidget(b, n // 3, n % 3)
        quick.setColumnStretch(3, 1)
        fb.addRow("", quick)
        # what moved the dates (a weekend, a court holiday: Settings -> Dates), or a day typed that is one
        self.date_notes = {}
        for key in ("delivery_date", "agreement_date"):
            if key == "agreement_date":
                row(fb, "agreement_date", calendar=True)
            note = self.date_notes[key] = QLabel("")
            note.setObjectName("muted")
            note.setWordWrap(True)
            fb.addRow("", note)
            fb.setRowVisible(note, False)  # (a hidden row takes no room: _show_date_notes shows it)
        self.date_form = fb
        for key in self.date_notes:
            self.rows[key].marks = self._holidays_of
        self._dated = None  # the day the jobs' dates were last worked out for (_redate)
        self._quick_tips()
        rl.addWidget(c)
        self._fill_sheet_box()

        # The outputs whose options are this job's own choices each get a card here (OutputSpec.card: the
        # invoice's), after the speed and pages they bill and above who ordered what and who pays what, which they
        # change (the user, 2026-10-22: in the Outputs box they were apart from those cards). Hidden while the
        # output is unticked (_refresh_outputs).
        self.output_cards: dict[str, QFrame] = {}
        self.output_opts: dict[str, QWidget] = {}  # each output's options, greyed out while it is unticked
        for spec in courthouses.panel_order():
            if spec.card and spec.panel:
                c, cl = folding(spec.card, spec.key)
                body = QWidget()
                bl = QVBoxLayout(body)
                bl.setContentsMargins(0, 0, 0, 0)
                resolve(spec.panel)(self, form_in(bl))
                cl.addWidget(body)
                self.output_cards[spec.key] = c
                self.output_opts[spec.key] = body
                c.setVisible(False)
                rl.addWidget(c)

        # Who ordered what: every attorney's order of every day ticked in the job list, spelled out (shown while
        # Invoice is ticked and a job has a transcript or pages to bill: _show_orders)
        self.orders_card, cl = folding("Who ordered what", "orders")
        top = QHBoxLayout()
        # (what Generate all bills of those days, in a sentence: _orders_where; how the rows are billed is the ?
        # button's, ORDERS_HELP)
        self.orders_info = QLabel("")
        self.orders_info.setObjectName("muted")
        self.orders_info.setWordWrap(True)
        self.orders_info.setToolTip(ORDERS_HELP)
        top.addWidget(self.orders_info, 1)
        self.orders_help = QToolButton()
        self.orders_help.setText("?")
        self.orders_help.setObjectName("quick")
        self.orders_help.setToolTip(ORDERS_HELP + "\n\nClick for more in the help.")
        self.orders_help.clicked.connect(lambda: self.show_help("inv-who-ordered"))
        top.addWidget(self.orders_help, 0, Qt.AlignTop)
        self.orders_edit = QPushButton("Excerpts…")
        self.orders_edit.clicked.connect(self._edit_orders)
        top.addWidget(self.orders_edit)
        cl.addLayout(top)
        self.orders = QTableWidget(0, len(ORDER_COLS))
        self.orders.setHorizontalHeaderLabels(ORDER_COLS)
        for c, name in enumerate(ORDER_COLS):
            if name in ORDER_TIPS:
                self.orders.horizontalHeaderItem(c).setToolTip(ORDER_TIPS[name])
        self.orders.verticalHeader().setVisible(False)
        self.orders.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.orders.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.orders.setWordWrap(True)
        self.orders.setIconSize(QSize(*badge_size(z(TAG_HEIGHT))))  # (each firm's tag and stripe: _mark_row)
        self.orders.setItemDelegateForColumn(1, StripeDelegate(self.orders))
        # (not the last section stretched: squeezed into what was left, "Your pages billed" lost its heading; it
        # is widened to fill the card instead, when there is room: _show_orders)
        self.orders.horizontalHeader().setStretchLastSection(False)
        self.orders.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)  # as tall as its rows (see _show_orders)
        self.orders.setToolTip("Double-click a row to open Excerpts… at that day (who ordered which pages of the "
                               "case)")
        self.orders.itemDoubleClicked.connect(self._order_row_clicked)
        # the scroll bar comes and goes with the window's width: the table's height makes room for it
        self.orders.horizontalScrollBar().rangeChanged.connect(lambda *_: self._fit_orders_height())
        cl.addWidget(self.orders)
        self.orders_card.setVisible(False)
        rl.addWidget(self.orders_card)

        # Who pays what: each invoice of the days ticked in short, right under the ticks that change it, kept up to
        # date as attorneys, speeds, Parties, Peripherals or Excerpts change (shown while Invoice is ticked and a day
        # ticked has pages to bill: _show_who_pays)
        self.pays_card, cl = folding("Who pays what", "pays")
        self.pays_card.setProperty("result", True)  # (a result to read, not a form to fill: a look of its own)
        self.pays_title = self.pays_card.heading
        self.pays_title.setToolTip(PAYS_TIP.format(**_who_splits(self.s)))
        self.pays = QLabel("")
        self.pays.setObjectName("payTable")
        self.pays.setTextFormat(Qt.RichText)
        self.pays.setWordWrap(True)
        self.pays.setTextInteractionFlags(Qt.TextBrowserInteraction)
        self.pays.setOpenExternalLinks(False)
        self.pays.linkActivated.connect(self._show_live_math)
        self.pays.setToolTip(PAYS_TIP.format(**_who_splits(self.s)))
        cl.addWidget(self.pays)
        # the invoices the card shows (invoice.firm_invoices), for its math and its folded line
        self._pay_firms: list = []
        self._inv_firms: list = []  # the job shown's invoices, priced in the Invoice card (for its link to the math)
        self.pays_card.setVisible(False)
        rl.addWidget(self.pays_card)

        # Who transcribed what: each reporter's takes of every day ticked, from the initials on its pages (the
        # user, 2026-10-21: to review in the window, and save; the run sheet has every take in full)
        self.takes_card, cl = folding("Who transcribed what", "takes")
        self.takes_card.setProperty("result", True)
        self.takes_info = QLabel("")
        self.takes_info.setObjectName("muted")
        self.takes_info.setWordWrap(True)
        cl.addWidget(self.takes_info)
        self.takes = QTableWidget(0, len(TAKE_COLS))
        self.takes.setHorizontalHeaderLabels(TAKE_COLS)
        self.takes.verticalHeader().setVisible(False)
        self.takes.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.takes.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.takes.setWordWrap(True)
        self.takes.horizontalHeader().setStretchLastSection(True)
        self.takes.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)  # as tall as its rows (_show_takes)
        self.takes.setToolTip(TAKES_TIP)
        # a take come out wrong (a page's initials misread): the user moves where it ends (the user, 2026-10-23)
        self.takes.cellDoubleClicked.connect(
            lambda r, _c: self._fix_take(self._take_rows[r]) if 0 <= r < len(self._take_rows) else None)
        self.takes.setContextMenuPolicy(Qt.CustomContextMenu)
        self.takes.customContextMenuRequested.connect(self._takes_menu)
        cl.addWidget(self.takes)
        foot_row = QHBoxLayout()
        self.takes_note = QLabel(TAKES_NOTE)
        self.takes_note.setObjectName("muted")
        self.takes_note.setWordWrap(True)
        foot_row.addWidget(self.takes_note, 1)
        self.takes_save = QPushButton("Save a copy…")
        self.takes_save.setToolTip("This table as an Excel workbook or a CSV file, where you choose")
        self.takes_save.clicked.connect(self._save_takes)
        foot_row.addWidget(self.takes_save, 0, Qt.AlignTop)
        cl.addLayout(foot_row)
        self._take_rows: list = []  # the takes shown (batch.takes_of), for Save a copy... and _fix_take
        self.takes_card.setVisible(False)
        rl.addWidget(self.takes_card)
        rl.addStretch(1)

        # footer: the Outputs box, one row (the user, 2026-10-22: the form needs the room it took) - a box per output
        # to tick, how many files Generate makes of it after it (_show_counts), Options ▾ for the outputs' standing
        # choices (the agreement form and signing, the MOFR's Civil or Criminal, a run sheet already there: a box
        # that opens under it, _show_output_options), then the Generate buttons, at the right, where a task ends.
        # The invoice's own choices are this job's: its card in the form (OutputSpec.card, built above).
        self.foot = foot = QFrame()
        foot.setObjectName("card")
        foot.setProperty("help_topic", "make-outputs")
        fv = QVBoxLayout(foot)
        fv.setContentsMargins(16, 8, 16, 8)
        fv.setSpacing(6)
        self.foot_row = fl = QHBoxLayout()  # (the Generate buttons end it, or a second row when it's narrow)
        fl.setSpacing(14)
        fv.addLayout(fl)
        self.foot_row2 = QHBoxLayout()
        self.foot_row2.addStretch(1)
        fv.addLayout(self.foot_row2)
        title = QLabel("Outputs")  # (the Outputs box, as the help and the messages call it)
        title.setObjectName("section")
        fl.addWidget(title)
        self.output_boxes: dict[str, QCheckBox] = {}
        self.output_counts: dict[str, QLabel] = {}  # how many files Generate makes of each (_show_counts)
        self.option_panels: dict[str, QFrame] = {}  # the panels of Options ▾, by output (gui/panels.py adds to them)
        self.options_pop = QFrame(self, Qt.Popup)
        self.options_pop.setObjectName("optionsPopup")
        # (open, the box is a window of its own: the main window's F1 doesn't reach it; each panel names its topic,
        # the box the Outputs box's)
        self.options_pop.setProperty("help_topic", "make-outputs")
        QShortcut(QKeySequence("F1"), self.options_pop, activated=lambda: self.show_help(
            help_topic_of(self.options_pop.focusWidget() or self.options_pop)))
        pop = QHBoxLayout(self.options_pop)
        pop.setContentsMargins(10, 10, 10, 10)
        pop.setSpacing(10)

        def options_panel(spec) -> QFormLayout:
            """A panel of Options ▾ headed by the output's name, with its options under a line (label: value rows,
            which gui/panels.py fills), greyed out while the output is unticked (output_opts)."""
            holder = QFrame()
            holder.setObjectName("outputPanel")
            holder.setProperty("help_topic", OUTPUT_TOPICS.get(spec.key, "make-outputs"))
            col = QVBoxLayout(holder)
            col.setContentsMargins(12, 8, 12, 10)
            col.setSpacing(6)
            head = QLabel(spec.label)
            head.setObjectName("outputTitle")
            col.addWidget(head)
            rule = QFrame()
            rule.setObjectName("outputRule")
            rule.setFixedHeight(1)
            col.addWidget(rule)
            body = QWidget()
            form = QFormLayout(body)
            form.setContentsMargins(0, 0, 0, 0)
            form.setHorizontalSpacing(8)
            form.setVerticalSpacing(6)
            form.setLabelAlignment(Qt.AlignRight | Qt.AlignVCenter)
            form.setFieldGrowthPolicy(QFormLayout.FieldsStayAtSizeHint)
            col.addWidget(body)
            col.addStretch(1)
            self.output_opts[spec.key] = body
            self.option_panels[spec.key] = holder
            pop.addWidget(holder)
            return form

        for spec in courthouses.panel_order():
            key = spec.key
            cb = QCheckBox(spec.label)
            cb.setObjectName("outputHead")
            cb.setChecked(key in self.s.outputs)
            cb.setToolTip(spec.tip)
            cb.setProperty("help_topic", OUTPUT_TOPICS.get(key, "make-outputs"))  # (F1: that output's topic)
            cb.toggled.connect(lambda _=False, k=key: self._outputs_changed(k))
            self.output_boxes[key] = cb
            count = QLabel("")
            count.setObjectName("outputCount")
            count.setVisible(False)
            self.output_counts[key] = count
            pair = QHBoxLayout()
            pair.setSpacing(4)
            pair.addWidget(cb)
            pair.addWidget(count)
            fl.addLayout(pair)
            if spec.panel and not spec.card:  # (one with a card has its options there)
                resolve(spec.panel)(self, options_panel(spec))
        self.options_btn = QToolButton()
        self.options_btn.setText("Options ▾")
        self.options_btn.setObjectName("quick")
        self.options_btn.setToolTip("Each output's settings, kept for every job: which agreement form and how it is\n"
                                    "signed, the MOFR's Civil or Criminal box, what to do when the case has a run\n"
                                    "sheet already")
        self.options_btn.clicked.connect(self._show_output_options)
        self.options_btn.setVisible(bool(self.option_panels))
        fl.addWidget(self.options_btn)
        fl.addStretch(1)
        self.gen_box = QWidget()  # the Generate buttons, at the end of the row or under it (_place_foot)
        gb = QHBoxLayout(self.gen_box)
        gb.setContentsMargins(0, 0, 0, 0)
        gb.setSpacing(10)
        self.fill_btn = QPushButton("Generate")
        self.fill_btn.setObjectName("generate")  # (the theme makes it stand out: the one thing to click at the end)
        self.fill_btn.clicked.connect(self.fill)
        gb.addWidget(self.fill_btn)
        self.fill_all_btn = QPushButton("Generate all")
        self.fill_all_btn.setObjectName("generate")
        self.fill_all_btn.clicked.connect(self.fill_all_jobs)
        gb.addWidget(self.fill_all_btn)
        fl.addWidget(self.gen_box)
        self._outputs_built = True  # (_refresh_outputs waits for every output's widgets)
        self._fill_invoice_speeds()
        root.addWidget(foot)
        self._restore_splits()
        self._refresh_jobs()

        self._build_menu()
        self.recent_docs.setMenu(self.recent_menu)
        # Undo for New job, at the bottom of the window for a few seconds (_offer_undo)
        self.undo_box = QWidget()
        ub = QHBoxLayout(self.undo_box)
        ub.setContentsMargins(4, 0, 4, 0)
        ub.setSpacing(8)
        self.undo_text = QLabel("")
        ub.addWidget(self.undo_text)
        self.undo_btn = QPushButton("Undo")
        self.undo_btn.setToolTip("Bring back what New job cleared")
        self.undo_btn.clicked.connect(self._undo_clear)
        ub.addWidget(self.undo_btn)
        self.statusBar().addWidget(self.undo_box)
        self.undo_box.setVisible(False)
        self._undo_timer = QTimer(self)
        self._undo_timer.setSingleShot(True)
        self._undo_timer.timeout.connect(self._drop_undo)
        # the dates follow the day: past midnight, or the evening hour of Settings -> Dates (_redate)
        self._day_timer = QTimer(self)
        self._day_timer.setInterval(60_000)
        self._day_timer.timeout.connect(self._redate)
        self._day_timer.start()
        for key, fr in self.cards.items():
            if key in self.s.folded_cards:
                self._fold_card(fr, True, save=False)
        QShortcut(QKeySequence.Paste, self, activated=self._paste_shortcut)
        QShortcut(QKeySequence("Ctrl+N"), self, activated=self.new_job)
        QShortcut(QKeySequence("Ctrl+O"), self, activated=self.browse)
        QShortcut(QKeySequence("Ctrl+Return"), self, activated=self.fill)
        QShortcut(QKeySequence("Ctrl+Shift+Return"), self, activated=self.fill_all_jobs)
        QShortcut(QKeySequence("Ctrl+Enter"), self, activated=self.fill)  # (the Enter of the number keys)
        QShortcut(QKeySequence("Ctrl+Shift+Enter"), self, activated=self.fill_all_jobs)
        QShortcut(QKeySequence("Ctrl+R"), self, activated=self.open_records)
        QShortcut(QKeySequence("F1"), self, activated=lambda: self.show_help(help_topic_of(self.focusWidget())))
        for keys, step in (("Ctrl++", 1), ("Ctrl+=", 1), ("Ctrl+-", -1), ("Ctrl+0", 0)):
            QShortcut(QKeySequence(keys), self, activated=lambda s=step: self.zoom_step(s))

    def _build_menu(self):
        """The File, View and Help menus. View's zoom items only name their keys, after a tab (Ctrl + / Ctrl - /
        Ctrl 0 are the window's shortcuts, see _build, and Ctrl + wheel its eventFilter's): as shortcuts of the menu
        too, they would be taken twice."""
        from .dialogs import AboutDialog, FEEDBACK_URL, WEBSITE_URL
        mb = self.menuBar()
        m = mb.addMenu("&File")
        self.new_action = m.addAction("&New job", self.new_job)  # Ctrl+N handled by the window shortcut
        m.addAction("&Open documents…", self.browse)
        m.addAction("Open a &folder of documents (batch)…", self.browse_folder)
        self.recent_menu = m.addMenu("Open r&ecent")
        self.recent_menu.aboutToShow.connect(self._fill_recent)
        m.addSeparator()
        m.addAction("Open &rate sheets folder", self._open_sheets_folder)
        m.addAction("&Settings…", self.open_settings)
        m.addAction("Ex&port settings (for another computer)…", self.export_settings)
        m.addAction("&Import settings…", self.import_settings)
        m.addSeparator()
        m.addAction("&Records (invoices and history)…", self.open_records)
        m.addAction("Save the invoice &spreadsheet template…", self._save_invoice_template)
        m.addAction("&Lock finished PDFs (no more changes)…", self.lock_pdfs)
        m.addSeparator()
        m.addAction("E&xit", self.close)
        m = self.view_menu = mb.addMenu("&View")
        m.addAction("Zoom &in\tCtrl++", lambda: self.zoom_step(1))
        m.addAction("Zoom &out\tCtrl+-", lambda: self.zoom_step(-1))
        m.addAction("&Actual size\tCtrl+0", lambda: self.zoom_step(0))
        theme = m.addMenu("&Theme")
        self.theme_actions = QActionGroup(self)
        for key, label in (("system", "&As Windows is"), ("light", "&Light"), ("dark", "&Dark")):
            act = theme.addAction(label, lambda k=key: self._set_theme(k))
            act.setCheckable(True)
            act.setData(key)
            act.setChecked(key == (self.s.theme if self.s.theme in ("light", "dark") else "system"))
            self.theme_actions.addAction(act)
        m.addSeparator()
        m.addAction("Fold all &cards", lambda: self._fold_all(True))
        m.addAction("&Unfold all cards", lambda: self._fold_all(False))
        self.weekday_action = m.addAction("Show the &day of the week in the job list",
                                          lambda: self._show_weekday(self.weekday_action.isChecked()))
        self.weekday_action.setCheckable(True)
        self.weekday_action.setChecked(self.s.job_weekday)
        m.addSeparator()
        self.yin_action = m.addAction("Show the &yin-yang", lambda: self._show_yin(self.yin_action.isChecked()))
        self.yin_action.setCheckable(True)
        self.yin_action.setChecked(self.s.show_yin)
        m = mb.addMenu("&Help")
        m.addAction("&How to use YinItAgreementForm…\tF1", lambda: self.show_help())
        m.addAction("The &website (pictures and the latest version)", lambda: open_url(WEBSITE_URL))
        m.addAction("Check for a &newer version", lambda: self._check_update(asked=True))
        m.addSeparator()
        m.addAction("Set up the &AI helper (Ollama)…", self._ai_help)
        m.addAction("Send &feedback…", lambda: open_url(FEEDBACK_URL))
        m.addAction("Open the &log folder", self._open_log_folder)
        m.addAction("&Copy details for a problem report", self._copy_diagnostics)
        m.addSeparator()
        m.addAction("&About YinItAgreementForm", lambda: AboutDialog(self).exec())

    # --------------------------------------------------------- folding, dividers, drops
    def _restore_splits(self) -> None:
        """The divider between the job list and the paste box where the user left it (Settings.window_splits
        "left"), kept for the first job too (_place_left_split places it only when one job becomes a batch or back).
        The "outer" one an older version saved, for a divider above the Outputs box, is not read: that divider is
        gone, and _save_splits drops it."""
        state = self.s.window_splits.get("left")
        if state and self.left_split.restoreState(QByteArray.fromBase64(state.encode())):
            self._left_mode = "single"

    def _save_splits(self) -> None:
        """Remembers where the divider between the job list and the paste box is (closeEvent): where it was before
        the paste box went down to one line (_fit_paste), when it is, so the next start doesn't open with it
        squeezed."""
        state = self._split_unslim if self._paste_slim and self._split_unslim is not None \
            else self.left_split.saveState()
        self.s.window_splits = {"left": bytes(state.toBase64()).decode()}

    def _place_left_split(self, multi: bool) -> None:
        """The divider between the job list and the paste box, when one job becomes a batch or back: one job
        leaves the list a few rows (its documents) and the paste box the rest; a batch gives the list the room,
        leaving the paste box a few lines. Dragged, it stays until the next change."""
        mode = "batch" if multi else "single"
        # (a taller window gives its room to the paste box for one job, to the list for a batch)
        self.left_split.setStretchFactor(0, 1 if multi else 0)
        self.left_split.setStretchFactor(1, 0 if multi else 1)
        if mode == self._left_mode:
            return
        sizes = self.left_split.sizes()
        total = sum(sizes)
        if total <= 0 or not self.left_split.isVisible():
            return  # (not laid out yet: placed when the window shows, see showEvent)
        self._left_mode = mode
        if multi:
            bottom = self.left_split.widget(1).minimumSizeHint().height() + z(16)
            self.left_split.setSizes([max(total - bottom, z(140)), bottom])
        else:
            top = z(110) + (self.index_note.sizeHint().height() if not self.index_note.isHidden() else 0)
            self.left_split.setSizes([top, max(total - top, z(120))])

    def _show_output_options(self) -> None:
        """Options ▾ of the Outputs row: the outputs' standing choices (the agreement form and signing, the MOFR's
        Civil or Criminal box, a run sheet already there), in a box under the button (above it when the screen has
        no room below: the row is at the bottom of the window), closed by a click anywhere else."""
        pop, btn = self.options_pop, self.options_btn
        pop.adjustSize()
        size = pop.sizeHint()
        below = btn.mapToGlobal(QPoint(0, btn.height()))
        screen = (btn.screen() or QApplication.primaryScreen()).availableGeometry()
        y = below.y() if below.y() + size.height() <= screen.bottom() else \
            btn.mapToGlobal(QPoint(0, 0)).y() - size.height()
        x = min(below.x(), screen.right() - size.width())
        pop.move(max(screen.left(), x), max(screen.top(), y))
        pop.show()

    def _fold_card(self, fr: QFrame, folded: bool, save: bool = True) -> None:
        """Folds a card to its heading and a line that sums it up (_card_summaries), or unfolds it. The card stays
        shown (or hidden) as it was. Kept for the next time (Settings.folded_cards)."""
        fr.body.setVisible(not folded)
        fr.summary.setVisible(folded)
        fr.fold.setArrowType(Qt.RightArrow if folded else Qt.DownArrow)
        fr.fold.setToolTip("Unfold this card" if folded else "Fold this card to one line (click again to unfold it)")
        if folded:
            self._card_summaries()
        keys = [k for k in self.s.folded_cards if k != fr.key] + ([fr.key] if folded else [])
        if keys != self.s.folded_cards:
            self.s.folded_cards = keys
            if save:
                self._save_settings()

    def _fold_all(self, folded: bool) -> None:
        """View -> Fold all cards / Unfold all cards."""
        for fr in self.cards.values():
            self._fold_card(fr, folded, save=False)
        self._save_settings()

    def _card_summaries(self) -> None:
        """The line each folded card shows: what it holds in short, "⚠ " first when something in it needs a look
        (a required field blank, nobody ticked, a speed to choose, a warning in the Invoice card or Who ordered
        what, nothing to bill until something is fixed)."""
        if not hasattr(self, "cards"):
            return
        folded = {k for k, fr in self.cards.items() if fr.body.isHidden()}
        if not folded:
            return
        case, job = self.case, self.cur
        missing = set(case.missing_required()) if not job.is_empty() else set()
        lines = {}
        if "case" in folded:
            judge = case.get("judge")
            bits = [case.get("case_name"), case.get("index_no"), case.get("dates"),
                    f"Justice {judge}" if judge else "", f"Part {case.get('part')}" if case.get("part") else "",
                    ", ".join(sorted(case.proc_types))]
            lines["case"] = (bool(missing & set(CASE_KEYS)), bits)
        if "details" in folded:
            pages, copies = case.get("est_pages"), case.get("copies")
            rate = case.get("rate")
            bits = [self.sheet_box.currentText(),
                    f"{case.get('delivery')} ${rate}/pg" if case.get("delivery") and rate else case.get("delivery"),
                    f"{pages} pages" if pages else "", f"{copies} {'copy' if copies == '1' else 'copies'}" if copies
                    else "", f"delivery {case.get('delivery_date')}" if case.get("delivery_date") else ""]
            lines["details"] = (bool(missing & set(DETAIL_KEYS)) or bool(job.speed_question()), bits)
        if "attorneys" in folded:
            atts = [a for a in case.attorneys if (a.name or a.firm).strip()]
            ticked = [a.name or a.firm for a in atts if a.checked]
            text = f"{len(ticked)} of {len(atts)} ticked" + (f": {', '.join(ticked)}" if ticked else "")
            lines["attorneys"] = (bool(atts) and not ticked, [text if atts else "No attorneys yet"])
        if "pays" in folded:
            # (each firm once: it has an invoice for each set of days, made and to make)
            names = list(dict.fromkeys((f.atty.firm or f.atty.name) if f.atty else "(no one ticked)"
                                       for f in self._pay_firms))
            # (nothing to bill because of something to fix, as the card says: its ⚠ on the folded line too)
            lines["pays"] = (not names and "⚠" in self.pays.text(),
                             [f"{plural(len(self._pay_firms), 'invoice')}: {', '.join(names)}" if names else
                              "Nothing to bill yet"])
        if "takes" in folded:
            lines["takes"] = (False, [self.takes_info.text() or "No takes yet"])
        if "invoice" in folded and hasattr(self, "inv_info"):  # (the Invoice card: OutputSpec.card)
            billed = self.inv_pages_info.text() if self.inv_form.isRowVisible(self.inv_pages_row) else ""
            warn = any(w.text().startswith("⚠") for w in (self.inv_info, self.inv_who, self.inv_whose))
            lines["invoice"] = (warn, [f"Parties {self.inv_parties.value()}", " ".join(billed.split()),
                                       self.inv_info.text().split("\n")[0].lstrip("⚠ ")])
        if "orders" in folded:
            t = self.orders
            rows = [" ".join(t.item(r, c).text() for c in (0,) if t.item(r, c)) + ": " +
                    ", ".join(t.item(r, c).text() for c in (1, 2) if t.item(r, c) and t.item(r, c).text())
                    for r in range(t.rowCount())]
            lines["orders"] = (self.orders_info.text().startswith("⚠"), ["; ".join(rows) or "Nobody yet"])
        for key, (warn, bits) in lines.items():
            text = "  ·  ".join(b for b in bits if b and b.strip())
            self.cards[key].summary.setText(("⚠ " if warn else "") + text)

    # (a file dragged over the window: read wherever it is dropped, as on the drop zone; a text box would take
    # its path as typed text)
    @staticmethod
    def _files_or_picture(md) -> bool:
        """Whether a drag holds files of this computer or a picture (what any part of the window takes)."""
        return (md.hasUrls() and any(u.isLocalFile() for u in md.urls())) or md.hasImage()

    def dragEnterEvent(self, e) -> None:
        """Files, a picture or text dragged over the window (outside the drop zone, which takes its own): the drop
        zone lights up. Text dragged from one of the window's own boxes is not taken: let go over a card by mistake,
        it would be cut from its box and read as a new e-mail."""
        md = e.mimeData()
        if e.source() is not None and not self._files_or_picture(md):
            return
        if md.hasUrls() or md.hasText() or md.hasImage():
            e.acceptProposedAction()
            self.drop._hover(True)

    def dragMoveEvent(self, e) -> None:
        """Keeps taking the drag while it moves over the window."""
        e.acceptProposedAction()

    def dragLeaveEvent(self, e) -> None:
        """The drag left the window: the drop zone's light goes off."""
        self.drop._hover(False)

    def dropEvent(self, e) -> None:
        """Dropped anywhere in the window: read as on the drop zone (handle_mime)."""
        self.drop._hover(False)
        handle_mime(e.mimeData(), self.add_files, self.add_text, self.add_qimage)
        e.acceptProposedAction()

    def _drag_filter(self, obj, e) -> bool:
        """(eventFilter) Files or a picture dragged over any part of this window but the drop zone (which takes its
        own): taken by the window, so a text box or table doesn't take a path as typed text. True when the event was
        handled."""
        w = obj if isinstance(obj, QWidget) else None
        if w is None or w.window() is not self or w is self.drop or self.drop.isAncestorOf(w):
            return False
        md = e.mimeData()
        if not self._files_or_picture(md):
            return False
        if e.type() == QEvent.Drop:
            self.dropEvent(e)
        else:
            e.acceptProposedAction()
            self.drop._hover(True)
        return True

    # --------------------------------------------------------- the attorney table's details
    def _build_att_details(self, cl: QVBoxLayout, two_cols, form_in) -> None:
        """Under the attorney table: every detail of the row selected (the table shows who, firm and role),
        typed into here (_att_detail_edited) and written into the table's row, as a cell typed into is."""
        self.att_details = QWidget()
        lay = QVBoxLayout(self.att_details)
        lay.setContentsMargins(0, 0, 0, 0)
        self.att_which = QLabel("")  # whose details these are ("Details of the row selected: Dana Smith")
        self.att_which.setObjectName("muted")
        lay.addWidget(self.att_which)
        colA, colB = two_cols(lay)
        fa, fb = form_in(colA), form_in(colB)
        self.att_edits: dict[int, QWidget] = {}
        # the boxes typed in (or filled from their list of past values) since they last showed or wrote their row:
        # only those are kept through a refresh and written back (_att_typed); one only clicked into holds nothing
        # of the user's, and what the AI or a read puts in the row meanwhile wins
        self._att_dirty: set[int] = set()
        self._att_filling = False  # (set while _show_att_details fills the boxes: that isn't typing)
        for col, form in ((1, fa), (2, fa), (3, fa), (4, fb), (5, fb), (6, fb), (7, fb)):
            if col == 3:
                w = QPlainTextEdit()
                sized(w, "setFixedHeight", 58)
                w.setPlaceholderText("One line per line of the address")
                w.installEventFilter(self)  # (written into the table when it loses the focus: eventFilter)
                w.textChanged.connect(lambda c=col: self._att_box_changed(c))
            else:
                w = QLineEdit()
                w.editingFinished.connect(lambda c=col: self._att_detail_edited(c))
                w.textChanged.connect(lambda _text, c=col: self._att_box_changed(c))
            self.att_edits[col] = w
            lab = QLabel(ATT_COLS[col])
            lab.setObjectName("fieldLabel")
            form.addRow(lab, w)
        self.att_source = QLabel("")
        self.att_source.setObjectName("muted")
        lab = QLabel("Source")
        lab.setObjectName("fieldLabel")
        fb.addRow(lab, self.att_source)
        cl.addWidget(self.att_details)
        self._att_row = -1  # the row the details show
        self._att_job = None  # and of which job
        self._show_att_details()

    def _show_att_details(self) -> None:
        """The details of the row selected (the first when none is); hidden while the table has no row."""
        if not hasattr(self, "att_edits"):
            return
        r = self.att.currentRow()
        if r < 0 and self.att.rowCount():
            r = 0
        # (the same row of the same job: what is being typed stays; another row, or another job shown meanwhile, as
        # when a read finishes, has details of its own, which replace it)
        same = r == self._att_row and self.cur is self._att_job
        self._att_row, self._att_job = r, self.cur
        on = 0 <= r < self.att.rowCount()
        self.att_details.setVisible(on)  # (no rows: nothing to spell out)
        for col, w in self.att_edits.items():
            if same and w.hasFocus() and self._att_typed(col):
                continue  # (being typed in: a refresh meanwhile mustn't take the typing back)
            item = self.att.item(r, col) if on else None
            text = item.text() if item else ""
            self._att_filling = True
            try:
                if col == 3:
                    lines = "\n".join(p.strip() for p in re.split(r"\s+/\s*|\s*/\s+", text) if p.strip())
                    if w.toPlainText() != lines:
                        w.setPlainText(lines)
                elif w.text() != text:
                    w.setText(text)
            finally:
                self._att_filling = False
            self._att_dirty.discard(col)  # (what the row holds: nothing to write back)
        source = self.att.item(r, 8).text() if on and self.att.item(r, 8) else ""
        self.att_source.setText(badge_word(source))
        self.att_source.setToolTip(BADGE_TIPS.get(source, ""))
        who = " — ".join(self.att.item(r, c).text() for c in (1, 2) if on and self.att.item(r, c)
                         and self.att.item(r, c).text().strip())
        self.att_which.setText(f"Details of row {r + 1}: {who}" if on else "")

    def _att_detail_edited(self, col: int) -> None:
        """A detail typed (Enter, or the box left): written into the row's cell, which updates the job as typing
        in the cell does (_att_changed). An address's lines are joined with " / ", as the table writes them."""
        r = self._att_row
        if not 0 <= r < self.att.rowCount():
            return
        w = self.att_edits[col]
        text = (" / ".join(line.strip() for line in w.toPlainText().splitlines() if line.strip()) if col == 3
                else w.text().strip())
        item = self.att.item(r, col)
        if item is None:
            item = QTableWidgetItem("")
            self.att.setItem(r, col, item)
        self._att_dirty.discard(col)
        if item.text() != text:
            item.setText(text)  # (fires itemChanged: _att_changed)

    def _att_box_changed(self, col: int) -> None:
        """A detail box's text changed: by the user (typed, pasted, picked from its list), unless the window is
        filling it from the row."""
        if not self._att_filling:
            self._att_dirty.add(col)

    def _att_typed(self, col: int) -> bool:
        """Whether the detail box of that column was typed in since it last showed or wrote the row (a box only
        clicked into holds nothing of the user's: what the AI or a read put in the row meanwhile wins)."""
        return col in self._att_dirty

    def _commit_att_details(self) -> None:
        """A detail still being typed is written into the table first (Ctrl+Enter generates without leaving the
        box: see _commit_edit)."""
        if not hasattr(self, "att_edits"):
            return
        fw = QApplication.focusWidget()
        for col, w in self.att_edits.items():
            if w is fw and self._att_typed(col):
                self._att_detail_edited(col)

    # --------------------------------------------------------- what the records hold (read off the window's thread)
    def _read_records_side(self) -> None:
        """Reads what the window shows from the records: the jobs made lately (the start card), the values the
        fields complete from, and the Records button's count. At startup and after the records change."""
        from ..records import past_values, recent_jobs

        def jobs_read(found) -> None:
            self._recent_jobs = list(found or [])
            self._show_start()

        self.runner.start(recent_jobs, on_done=jobs_read, on_error=lambda _m: None)
        self.runner.start(past_values, on_done=self._fill_completions, on_error=lambda _m: None)
        self._count_records()

    def _fill_completions(self, values: dict) -> None:
        """The court, part and judge fields and the attorney details' name and firm complete what is typed from
        the records' values (records.past_values): a list under the box, any part of a value matching."""
        boxes = {"court": [self.rows["court"].edit], "part": [self.rows["part"].edit],
                 "judge": [self.rows["judge"].edit], "firm": [self.att_edits[2]], "name": [self.att_edits[1]]}
        for key, edits in boxes.items():
            model = self._completions.get(key)
            if model is None:
                model = self._completions[key] = QStringListModel(self)
                for edit in edits:
                    comp = QCompleter(model, edit)
                    comp.setCaseSensitivity(Qt.CaseInsensitive)
                    comp.setFilterMode(Qt.MatchContains)
                    comp.setCompletionMode(QCompleter.PopupCompletion)
                    edit.setCompleter(comp)
            model.setStringList(list((values or {}).get(key, [])))

    def _show_start(self) -> None:
        """The start card, shown while nothing is loaded: a link per job made lately (records.recent_jobs) that
        opens it again, Recent documents, and how to begin. Hidden as soon as anything is loaded, typed or pasted,
        or being read (_loaded)."""
        if not hasattr(self, "start_card"):
            return
        empty = not self._loaded()
        self.start_card.setVisible(empty)
        if not empty or getattr(self, "_start_shown", None) == self._recent_jobs:
            return
        self._start_shown = list(self._recent_jobs)  # (the links are made again only when the list changed)
        while self.start_list.count():
            item = self.start_list.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
        for case, dates, ts, origin in self._recent_jobs:
            made = f"made {us_date(ts[:10])}" if ts else ""
            text = "  ·  ".join(b for b in (case or "(no case name)", dates, made) if b)
            link = ClickLabel(text)
            link.setTextFormat(Qt.PlainText)
            link.setObjectName("startLink")
            link.set_clickable(True)
            link.setToolTip("Open this job again, as it was made")
            link.clicked.connect(lambda o=origin: self.open_past_job(o))
            self.start_list.addWidget(link)
        self.start_card.heading.setText("Recent jobs" if self._recent_jobs else "Start")

    # --------------------------------------------------------- the Records button's count
    def _count_records(self) -> None:
        """Counts the invoices of the status Settings -> Records names (records_count: "Ready to send", "Not paid",
        or "" for none) off the window's thread, for the Records button (_show_count). A colour no row can take with
        the boxes in use counts Not paid instead (tracking.counted: Ready to send, the default, with Paid alone). A
        count asked for meanwhile replaces this one."""
        from ..records import live_invoices
        from ..tracking import count_of, counted
        self._counting += 1
        boxes, rules = self.s.boxes(), self.s.rules()
        n, name = self._counting, counted(self.s.records_count, boxes, rules)
        if not name:
            self._show_count(name, 0)
            return

        def count() -> int | None:
            try:
                return count_of(name, live_invoices(), boxes, rules)
            except OSError:  # (the records can't be read just now, a restore: the count stays as it was; not an
                return None  # error each time the window comes back, which would fill the log)

        def done(found: int | None) -> None:
            if n == self._counting and found is not None:  # (a newer count is on its way: this one is out of date)
                self._show_count(name, found)

        self.runner.start(count, on_done=done, on_error=lambda _m: None)

    def _show_count(self, name: str, n: int) -> None:
        """The Records button with its count ("Records · Ready to send: 3"), or plain at 0 or with none chosen."""
        if n and name:
            self.rec_btn.setText(f"Records · {name}: {n}")
            self.rec_btn.setToolTip(f"{plural(n, 'invoice')} {'is' if n == 1 else 'are'} {name} (Settings → Records → "
                                    f"Count on the Records button changes or turns off this count).\n{RECORDS_TIP}")
        else:
            self.rec_btn.setText("Records")
            self.rec_btn.setToolTip(RECORDS_TIP)
        self._fit_corner()

    def _fit_corner(self) -> None:
        """The menu bar makes room for its corner (the status and the buttons) again when their text got longer or
        shorter: it measures the corner only when it is set, not as it changes."""
        box = getattr(self, "head_box", None)
        if box is None:
            return
        hint = box.sizeHint()
        if hint != getattr(self, "_corner_hint", None):
            self._corner_hint = hint
            self.menuBar().setCornerWidget(box, Qt.TopRightCorner)

    def changeEvent(self, e) -> None:
        """Back to the window (from the Records window, where invoices are ticked, or another program): the
        Records button counts again, at most every 2 s, and the dates follow the day if it changed meanwhile
        (_redate: the computer woke up the next morning). Another theme: the unticked jobs are greyed in its
        colour (_redim)."""
        super().changeEvent(e)
        if e.type() == QEvent.ActivationChange and self.isActiveWindow() and hasattr(self, "rec_btn"):
            self._redate()
            self._count_on_activation()
        elif e.type() in (QEvent.PaletteChange, QEvent.ApplicationPaletteChange):
            self._redim()

    def _count_on_activation(self) -> None:
        """The Records button counts again, unless it did less than 2 s ago: then once those are up (a tick in the
        Records window and straight back would otherwise leave the count as it was)."""
        wait = 2 - (time.monotonic() - self._activated)
        if wait > 0:
            if not self._count_later.isActive():
                self._count_later.start(int(wait * 1000) + 1)
            return
        self._activated = time.monotonic()
        self._count_records()

    # --------------------------------------------------------- startup
    def _startup(self):
        """Runs once the window is on screen: the welcome questions on the first run (no name yet; asked once),
        then the AI check, the day's backup of the records, what the window shows from the records (the jobs made
        lately, the completions, the Records button's count), the Update button of a newer version found before and,
        once a day, the look for a newer version, a warning when the rate sheet in use is incomplete and the
        user hasn't said to use it as it is (_check_rate_sheet), and the question whether a sheet still with
        the Days of the template before 2.8.0 should say 5 and 15 (_check_old_days). A settings file that couldn't
        be read is said so instead of the welcome: the app runs on the defaults, and nothing is written over the
        file until the user presses Save in Settings or imports settings (Settings.save(force=True))."""
        if self.s.unreadable:
            QMessageBox.warning(self, "Your settings couldn't be read",
                                "Your settings file couldn't be read (it may be damaged, or held by another "
                                "program), so the app starts with its defaults for now. Nothing is saved over it "
                                "until you press Save in Settings, which replaces it with the settings you see "
                                f"there; to keep what it holds, close the app and look at the file first:\n\n"
                                f"{self.s.path}")
        elif not self.s.profile.name and not self.s.welcomed:
            self.s.welcomed = True
            self.welcome()
        self._check_ai()
        self.runner.start(backup_records, backup_folder(self.s))  # off the UI thread: a big database takes a moment
        self._read_records_side()  # (the jobs made lately, what fields complete from, the Records button's count)
        self._show_known_update()  # (found on an earlier day, or earlier today: shown every time until installed)
        if self.s.check_updates and self.s.update_checked != date.today().isoformat():
            self._check_update()
        self._check_rate_sheet()  # (edited since the app last ran, or incomplete and never answered)
        self._check_old_days()  # (a sheet filled in before 2.8.0: 7 and 21 were calendar days)

    def welcome(self) -> None:
        """The first-run questions (name, address, rate sheet, where files go). Skipped: asked no more;
        Settings has all of it."""
        from .preview import WelcomeDialog
        if WelcomeDialog(self.s, self).exec():
            self._settings_changed()
        else:
            self._save_settings()

    def _save_settings(self) -> None:
        """Saves the settings; a problem is logged, not shown (they are saved again when the window closes)."""
        try:
            self.s.save()
        except OSError as e:
            logfile.error("could not save the settings", e)

    def _check_update(self, asked: bool = False) -> None:
        """Asks GitHub, in the background, whether there is a newer version (see update.py) and shows the Update
        button at the top when there is (_show_update; remembered, so it shows every time until installed).
        asked: Help -> Check for a newer version, which also says so when there is none, or when GitHub can't be
        reached; the daily look says nothing then."""
        from .. import update

        def done(found):
            self.s.update_checked = date.today().isoformat()
            if found:
                self._show_update(*found)
                if asked:
                    QMessageBox.information(self, "A newer version", f"Version {found[0]} is available. The "
                                            "Update button at the top of the window opens its download page.")
            else:  # (none newer: one remembered was installed, or taken back)
                self.s.update_version = self.s.update_page = ""
                self._show_known_update()
                if asked:
                    QMessageBox.information(self, "Up to date",
                                            f"You have the latest version ({update.__version__}).")
            self._save_settings()

        def failed(msg):
            if asked:
                QMessageBox.information(self, "Could not check", "Could not reach GitHub to look for a newer "
                                        f"version (are you online?).\n\n{msg}")

        self.runner.start(update.newer, on_done=done, on_error=failed)

    def _show_update(self, version: str, page: str) -> None:
        """A newer version is out: remembered (Settings.update_version, update_page), and a button at the top of
        the window, beside Records, that opens its download page. (A line in the status bar, as before 2.8.0, went
        unseen at the bottom of a laptop's screen.)"""
        from .. import update
        self.s.update_version, self.s.update_page = version, page
        self.update_btn.setText(f"⬆  Update to {version}")
        self.update_btn.setToolTip(f"Version {version} is out (you have {update.__version__}). Click to open its "
                                   "download page in your browser; nothing is installed by itself.\n"
                                   "Settings → Options turns the daily look for a newer version off.")
        self.update_btn.setVisible(True)
        self._fit_corner()

    def _show_known_update(self) -> None:
        """The newer version found earlier, shown again (each time the app starts, until this copy is that version
        or newer: then it is forgotten). Hidden while the daily look is turned off in Settings."""
        from .. import update
        version = self.s.update_version
        if version and update.version_key(version) <= update.version_key(update.__version__):
            self.s.update_version = self.s.update_page = version = ""  # (installed)
        if version and self.s.check_updates:
            # (read back from settings.json, which anyone could edit: the button opens only a page of this project
            # in the browser, as update.latest gives it)
            page = self.s.update_page
            self._show_update(version, page if page.startswith(update.RELEASES_PAGE + "/")
                              else update.RELEASES_PAGE + "/latest")
        else:
            self.update_btn.setVisible(False)
            self._fit_corner()

    # ------------------------------------------------------ recent files, settings to and from a file
    def _fill_recent(self) -> None:
        """File -> Open recent, filled as it opens: the documents and folders opened lately, newest first
        (one that is gone is greyed out)."""
        m = self.recent_menu
        m.clear()
        for p in self.s.recent_files:
            path = Path(p)
            act = m.addAction(f"{path.name}   ·   {path.parent}".replace("&", "&&"),
                              lambda p=p: self.add_files([p], batch=Path(p).is_dir()))
            act.setEnabled(path.exists())
        if not self.s.recent_files:
            m.addAction("(nothing opened yet)").setEnabled(False)
        m.addSeparator()
        m.addAction("Clear this list", lambda: self._set_opt("recent_files", [])).setEnabled(bool(self.s.recent_files))

    def export_settings(self) -> None:
        """File -> Export settings: saves the settings and the rate sheets as one file (Settings.export_to),
        after asking whether the user's own details go into it (yes for another computer of theirs, no for a
        colleague)."""
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Question)
        box.setWindowTitle("Export settings")
        box.setText("Include your personal details in the file?\n\n"
                    "They are: your name, title, address, phone, fax, e-mail, website and initials (My info), the "
                    "payment text of your invoices, the names of other reporters (Run sheet) and your folders.\n\n"
                    "•  Include them to move to another computer of your own.\n"
                    "•  Leave them out to give the file to a colleague: they keep their own details, and get "
                    "your options, invoice wording and rate sheets.\n\n"
                    "Your signature picture and your records are never in the file.")
        with_mine = box.addButton("Include my details", QMessageBox.YesRole)
        without = box.addButton("Leave them out", QMessageBox.NoRole)
        box.addButton(QMessageBox.Cancel)
        box.setDefaultButton(without)  # (the careful choice, should Enter be pressed without reading)
        box.exec()
        if box.clickedButton() not in (with_mine, without):
            return
        personal = box.clickedButton() is with_mine
        name = "YinIt settings.json" if personal else "YinIt settings (no personal details).json"
        dest, _ = QFileDialog.getSaveFileName(self, "Export settings", str(Path.home() / "Documents" / name),
                                              "Settings file (*.json)")
        if not dest:
            return
        try:
            self.s.export_to(dest, personal=personal)
        except OSError as e:
            show_save_error(self, e, "Could not export the settings")
            return
        holds = ("your details (name, address, contact, payment text), your options and your rate sheets"
                 if personal else "your options, invoice wording and rate sheets, and none of your personal details")
        QMessageBox.information(self, "Settings exported", f"Saved as\n{dest}\n\nOn the other computer: File → "
                                f"Import settings.\n\nThe file holds {holds}. Your signature picture and your "
                                "records are not in it.")

    def import_settings(self) -> None:
        """File -> Import settings: replaces the settings with those of a file made by Export settings, after
        asking; the settings as they were are kept next to the settings file as "settings before import.json"."""
        if self._batch_running():
            return
        src, _ = QFileDialog.getOpenFileName(self, "Import settings", str(Path.home() / "Documents"),
                                             "Settings file (*.json)")
        if not src:
            return
        try:
            new = Settings.import_from(src, self.s)
        except (ValueError, OSError) as e:
            QMessageBox.warning(self, "Could not import", str(e))
            return
        who = (f"reporter: {new.profile.name or '(no name)'}" if new.imported_personal else
               "it has no personal details: your name, contact details, payment text and folders stay as they are")
        if QMessageBox.question(self, "Import settings", f"Replace your settings with those of this file "
                                f"({who})?\n\nYour settings as they are now are kept as "
                                "\"settings before import.json\" in the app's settings folder. "
                                "Your records are not changed.") != QMessageBox.Yes:
            return
        sheet = self.s.sheet().name  # (another sheet imported, and incomplete: its warning can go back to this one)
        try:
            new.write_imported_sheets()  # (only now: answering No leaves the rate sheet folder as it was)
            if self.s.path.exists():
                self.s.path.with_name("settings before import.json").write_bytes(self.s.path.read_bytes())
            # the Settings object is shared (the Records window, the AI helper): filled in place. This computer's
            # own (LOCAL) are left as they are now: `new` copied them before the question, and the day's look for
            # a newer version, or a divider moved, may have changed them while it waited
            for name in Settings.__dataclass_fields__:
                if name not in LOCAL:
                    setattr(self.s, name, getattr(new, name))
            self.s.save(force=True)  # (the user's own choice: written even over a settings file that couldn't be read)
        except OSError as e:
            show_save_error(self, e, "Could not import the settings")
            return
        self._settings_changed(sheet=sheet)
        self._toast("Settings imported.")

    def show_help(self, topic: str = ""):
        """Help -> How to use, and F1 (the topic of the part of the window that has the focus): the Help window, at
        that topic (or where it was), left open beside the window (gui/help_window)."""
        return open_help(topic, self)

    def _ai_help(self, _link=""):
        """Shows the Ollama setup help, then checks the AI again (a model may have been installed)."""
        from .dialogs import OllamaHelpDialog
        OllamaHelpDialog(self.s.ollama_model, self.s.ollama_host, self).exec()
        self._check_ai()

    def _ai_text(self, text: str):
        """Shows the AI's state under the paste box. A longer or shorter line changes the left column's height,
        so the drop zone is sized again."""
        self.ai_label.setText(text)
        self._size_drop_later()

    def _check_ai(self):
        """Asks Ollama in the background whether the model is ready and shows the answer under the paste box.
        Documents of a single job dropped while it was checking are then sent to the AI."""
        if not self.s.use_ai:
            self.ai_ok = False
            self._ai_text("AI is off (Settings → AI). Using rules only.")
            return
        self._ai_text("Checking AI…")
        self.ai = ai = OllamaExtractor(self.s)

        def done(res):
            if ai is not self.ai or not self.s.use_ai:  # settings changed since (another model, or AI off)
                return
            self.ai_ok, msg = res
            link = "" if self.ai_ok else '  <a href="setup">How to set up</a>'
            self._ai_text(("● " if self.ai_ok else "○ ") + html.escape(msg, quote=False) + link)  # (the label is rich text)
            # inputs dropped while the check was still running were not sent to the AI
            waiting = [d for d in self.cur.docs if d.ai is None]
            if self.ai_ok and waiting and self.ai_pending == 0 and len(self.jobs) == 1:
                self._maybe_ai(self.cur, waiting)

        self.runner.start(self.ai.status, on_done=done, on_error=lambda m: done((False, m)))

    # ------------------------------------------------------ inputs
    def browse(self):
        """Asks for one or more documents and reads them (Ctrl+O)."""
        files, _ = QFileDialog.getOpenFileNames(self, "Choose document(s)", "", FILE_FILTER)
        if files:
            self.add_files(files)

    def browse_folder(self):
        """Asks for a folder and reads every document in it as a batch."""
        folder =QFileDialog.getExistingDirectory(self, "Choose a folder of documents")
        if folder:
            self.add_files([folder], batch=True)

    def _paste_shortcut(self):
        """Ctrl+V. In a text box it pastes as usual; anywhere else the clipboard's files, picture or text
        are read as a new input."""
        if self.focusWidget() in (self.paste,) or isinstance(self.focusWidget(), (QLineEdit, QPlainTextEdit)):
            fw = self.focusWidget()
            fw.paste()
            return
        md = QApplication.clipboard().mimeData()
        if not handle_mime(md, self.add_files, self.add_text, self.add_qimage):
            self._toast("Nothing to paste.")

    def _extract_paste(self):
        """Extract from text: reads what is in the paste box as an input."""
        text = self.paste.toPlainText().strip()
        if text:
            self.add_text(text)

    def add_text(self, text: str):
        """Reads pasted or dropped text as an input named "Pasted text 1", "Pasted text 2"… and shows it in the
        paste box."""
        n =sum(1 for j in self.jobs for d in j.docs if d.ing.name.startswith("Pasted text")) + 1
        self._ingest([lambda: ingest_text(text, f"Pasted text {n}")], [f"Pasted text {n}"])
        if self.paste.toPlainText().strip() != text.strip():
            self.paste.setPlainText(text)

    def add_qimage(self, qimg):
        """Reads a pasted or dropped picture (a QImage), e.g. a screenshot of an e-mail, with OCR."""
        from PIL import Image
        from PySide6.QtCore import QBuffer, QIODevice
        buf = QBuffer()
        buf.open(QIODevice.WriteOnly)
        qimg.save(buf, "PNG")
        import io
        pil = Image.open(io.BytesIO(bytes(buf.data())))
        self._ingest([lambda: ingest_pil(pil, "Pasted image")], ["Pasted image"])

    def add_files(self, paths: list[str], batch: bool = False, same_job: bool = False, split_rule: str = "",
                  fixes: dict | None = None):
        """Reads files, and the documents in folders. Files already loaded or still being read are skipped.
        Several documents about different cases (or batch=True) start a batch; same_job: they all go to the job
        on screen whatever they say (a past job opened again: that is how it was put together). split_rule: the
        jobs made of them were billed by that split rule (an invoice of several days opened again: Job.split_rule).
        fixes: the takes corrected in them, by Doc.key() (a past job opened again: batch.origin_fixes; Doc.fixes).
        What was opened is kept for File -> Open recent."""
        gen = self.gen
        for p in reversed(paths):
            self.s.remember_file(p)
        self._save_settings()

        def find():  # off the UI thread: a big folder takes a while to go through
            return [(p, _path_key(p)) for p in expand_paths(paths)]

        def found(everything):
            if not self._work_done(gen):
                return
            # judged now, not when dropped: the same file dropped twice in a row is still being read
            loaded = {_path_key(d.path) for j in self.jobs for d in j.docs if d.path} | self._loading
            fresh = [(p, key) for p, key in everything if key not in loaded]
            if len(fresh) < len(everything):
                self._toast(f"Skipped {len(everything) - len(fresh)} document(s) that are already loaded.")
            if fresh:
                files = [p for p, _ in fresh]
                self._ingest([(lambda p=p: ingest_file(p)) for p in files], [Path(p).name for p in files],
                             files, batch, {key for _, key in fresh}, same_job, split_rule, fixes)
            else:
                self._update_status()
                if not everything:
                    self._toast("No documents found there.")

        def failed(msg):
            if not self._work_done(gen):
                return
            self._update_status()
            QMessageBox.warning(self, "Could not open that", msg)

        self.work += 1
        if any(Path(p).is_dir() for p in paths):
            self._set_status("Looking for documents…", "busy")
            self.busy.setVisible(True)
        self.runner.start(find, on_done=found, on_error=failed)

    def _work_done(self, gen: int) -> bool:
        """A background read or batch has ended. False when "New job" was clicked since it started:
        its result is then dropped (and the counters were already reset)."""
        if gen != self.gen:
            return False
        self.work = max(0, self.work - 1)
        self._idle()
        return True

    def _ingest(self, loaders, names, paths=None, batch=False, keys=frozenset(), same_job=False, split_rule="",
                fixes=None):
        """Reads inputs in the background and adds them to the jobs.

        loaders: one function per input that returns the Ingested document; names: what to call each in
        messages; paths: the files (None for pasted text or pictures); keys: the files' _path_key, held in
        self._loading until the read ends. With a single job and documents all about one case, they are
        added to that job (with same_job, whatever they are about); otherwise (or with batch=True) they are
        sorted into jobs by case and date, the new ones billed by split_rule when one is given (the days of an
        invoice opened again: Job.split_rule). fixes: the takes corrected in the documents, by Doc.key() (a past
        job opened again: batch.origin_fixes); each document read whose key is there gets them (Doc.fixes). A cell
        of the attorney table being typed in when the read ends is taken as typed, and opened again afterwards
        while the same job is on screen (_resume_edit).
        """
        # the reading thread gets its own copy of the settings: Settings may change them meanwhile
        gen, target, s = self.gen, self.cur, deepcopy(self.s)
        self.work += 1
        self._loading |= keys
        self._set_status(f"Reading {names[0]}…" if len(names) == 1 else f"Reading {len(names)} documents…", "busy")
        self.busy.setVisible(True)

        def work(progress):
            return read_loaders(loaders, names, s, paths, progress)

        def step(i, n, name):
            if gen == self.gen and n > 1:
                self.busy.setRange(0, n)
                self.busy.setValue(i)
                self._set_status(f"Reading {i + 1} of {n}:  {name[:40]}", "busy")

        def done(result):
            if not self._work_done(gen):
                return
            self._loading -= keys
            docs, errors = result
            for d in docs:
                for w in d.ing.warnings:
                    self._toast(w)
                if fixes and fixes.get(d.key()):  # (a past job opened again: its takes as they were corrected)
                    d.fixes = dict(fixes[d.key()])
            if docs:
                typing = self._sync_from_ui()
                shown = self.cur
                one_job = len(self.jobs) == 1 and target in self.jobs and not batch
                if one_job and (same_job or len(docs) == 1 or len(group(docs, self.s)) == 1):
                    # the usual way: everything dropped is about the job on screen
                    target.docs += docs
                    target.unbill()  # new documents may mean new pages to bill
                    remerge(target, self.s)
                    self._auto_runsheet()
                    self._show_job()
                    self._maybe_ai(target, docs)
                else:
                    before = [j for j in self.jobs if not j.is_empty()]
                    count = len(before)
                    self.jobs = group(docs, self.s, before) or [Job()]
                    for j in self.jobs:  # (the days of an invoice opened again: billed as it was, also when a
                        # file dropped meanwhile was read first and took them in)
                        if split_rule and not j.split_rule and any(d is e for d in docs for e in j.docs):
                            j.split_rule = split_rule
                    if self.cur not in self.jobs:  # (the first of the list: the earliest day, at first)
                        self.cur = self._jobs_in_list_order()[0]
                    self._auto_runsheet()
                    self._refresh_jobs()
                    self._show_job()
                    new = len(self.jobs) - count
                    self._toast(f"{len(docs)} document(s) read:  {new} new job(s), "
                                f"{len(self.jobs)} in total.")
                if self.cur is shown:  # (still the job the user was typing in)
                    self._resume_edit(typing)
            else:
                self._update_status()
            log.info("read %d of %d document(s)%s", len(docs), len(loaders), " (batch)" if batch else "")
            if errors:
                self._unreadable(errors, len(loaders))

        def failed(msg):
            if not self._work_done(gen):
                return
            self._loading -= keys
            self._set_status("Could not read that input", "warn")
            QMessageBox.warning(self, "Could not read input", msg)

        self.runner.start(work, on_done=done, on_error=failed, on_progress=step)

    def _unreadable(self, errors: list[str], total: int):
        """Says which inputs could not be read: the error itself for a single input, else a list (first 8)."""
        if total == 1:
            self._set_status("Could not read that input", "warn")
            QMessageBox.warning(self, "Could not read input", errors[0])
            return
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Warning)
        box.setWindowTitle("Some documents could not be read")
        box.setText(f"{len(errors)} of {total} documents could not be read and were left out:\n\n"
                    + "\n".join(e[:110] for e in errors[:8]) + ("\n…" if len(errors) > 8 else ""))
        box.setDetailedText("\n".join(errors))
        box.exec()

    def _maybe_ai(self, job: Job, docs: list):
        """Asks the model about documents added to a single job (batches use the rules only): pictures always,
        e-mails and text when Settings say so, PDFs only while a required field is still blank (or the case name
        only a suggestion from the records). The questions
        are asked one at a time (ai_runner, on ai_pool). Each answer merges the job the document is in by then
        again (a later read may have merged its job into another), unless "New job" was clicked, the document
        was removed or the AI turned off meanwhile."""
        if not (self.s.use_ai and self.ai_ok):
            return
        todo = []
        for d in docs:
            if d.ing.kind == "image":
                todo.append(d)
            elif d.ing.kind in ("email", "text") and self.s.ai_for_text:
                todo.append(d)
            elif d.ing.kind == "pdf" and (job.case.missing_required()
                                          or job.case.fields["case_name"].source == SRC_RECORDS):
                todo.append(d)
        if not todo:
            return
        gen = self.gen
        self.ai_pending += len(todo)
        self.busy.setVisible(True)
        self._update_status()
        for d in todo:
            def done(ex, d=d):
                if gen != self.gen:
                    return
                self.ai_pending -= 1
                holder = next((j for j in self.jobs if any(x is d for x in j.docs)), None)
                if holder is not None and self.s.use_ai:
                    d.ai = ex
                    self._remerge(holder)
                self._update_status()

            def failed(msg, d=d):
                if gen != self.gen:
                    return
                self.ai_pending -= 1
                self._ai_text(f"○ AI failed on {html.escape(d.ing.name, quote=False)}: {html.escape(msg[:120], quote=False)}")
                self._update_status()

            self.ai_runner.start(self.ai.extract, d.ing, on_done=done, on_error=failed)

    # ------------------------------------------------------------ jobs
    # The list on the left: a row per job of a batch (or the job on screen alone), and under a job of several
    # documents a row per document. The rows keep their job (and document) in these roles.
    _JOB, _DOC, _DOCS = Qt.UserRole, Qt.UserRole + 1, Qt.UserRole + 2

    def _shown_jobs(self) -> list[Job]:
        """The jobs the list shows, a row each: every job of a batch; the job on screen alone once it has
        documents (nothing dropped yet: no row)."""
        if len(self.jobs) > 1:
            return self.jobs
        return [self.cur] if self.cur.docs else []

    def _job_sort(self) -> tuple[int, Qt.SortOrder]:
        """The column and the way the job list is sorted (Settings.job_sort): by date, oldest first, until the user
        clicks a heading (the user, 2026-10-22: a trial's days came newest first, as the documents were read)."""
        name, way = (list(self.s.job_sort) + ["", ""])[:2]
        col = JOB_COLS.index(name) if name in JOB_COLS else JOB_DATE
        return col, Qt.DescendingOrder if way == "descending" else Qt.AscendingOrder

    def _job_sort_changed(self, col: int, order: Qt.SortOrder) -> None:
        """A heading clicked: the list sorted by it, kept for next time (Settings.job_sort)."""
        if 0 <= col < len(JOB_COLS):
            self._set_opt("job_sort", [JOB_COLS[col], "descending" if order == Qt.DescendingOrder else "ascending"])
            self._fit_job_columns()  # (the arrow is in another column now)

    def _job_sort_key(self, job: Job, col: int) -> tuple:
        """What a job sorts by in the list's column col: its first date (Date and Day; a job without one after
        those with one), its pages (Pages, Yours; blank after numbers) or its case (Case / file); then its date,
        its case and the order it was read in, so two jobs alike keep that order."""
        day = _first_day(job.case.get("dates"))
        when = (0, day.toordinal()) if day else (1, 0)
        case = (job.case.get("case_name") or job.title() or "").lower()
        order = self.jobs.index(job) if job in self.jobs else len(self.jobs)
        if col in (JOB_PAGES, JOB_YOURS):
            has = bool(job.transcripts()) or job.pages_typed()
            n = (job.transcript_pages() if job.transcripts() else None) if col == JOB_PAGES else \
                (job.your_pages() if has else None)
            return ((0, n) if n is not None else (1, 0)) + when + (case, order)
        if col == 0:
            return (case,) + when + (order,)
        return when + (case, order)

    def _jobs_in_list_order(self) -> list[Job]:
        """The jobs in the order the list shows them (_job_sort): the first is the one shown after a batch is read
        or a job is taken out of it."""
        col, order = self._job_sort()
        return sorted(self.jobs, key=lambda j: self._job_sort_key(j, col), reverse=order == Qt.DescendingOrder)

    def _fit_job_columns(self) -> None:
        """The job list's Date, Day, Pages and Yours as wide as their text (the widest cell, or the heading, with
        its arrow for the column sorted by), the case taking the rest: sized by Qt, every column kept room for the
        sort arrow, and the case's name was cut to two letters in a column 360 pixels wide."""
        lst = self.job_list
        hh = lst.header()
        bold = QFont(hh.font())
        bold.setBold(True)  # (the headings are bold: the style sheet's QHeaderView::section)
        heads, cells = QFontMetrics(bold), lst.fontMetrics()
        rows = []
        for i in range(lst.topLevelItemCount()):
            top = lst.topLevelItem(i)
            rows += [top] + [top.child(k) for k in range(top.childCount())]
        for c in (JOB_DATE, JOB_DAY, JOB_PAGES, JOB_YOURS):
            arrow = z(16) if c == hh.sortIndicatorSection() else 0
            need = max([cells.horizontalAdvance(r.text(c)) for r in rows] +
                       [heads.horizontalAdvance(JOB_COLS[c]) + arrow])
            lst.setColumnWidth(c, need + z(16))
        self._fit_case_column()

    def _fit_case_column(self) -> None:
        """The Case / file column takes the room the others leave, but never less than CASE_MIN (zoomed in, in a
        narrow list, it held the tick box and two letters): with less, it keeps that much and the list scrolls
        sideways."""
        lst = self.job_list
        hh = lst.header()
        room = lst.viewport().width() - sum(lst.columnWidth(c) for c in range(1, len(JOB_COLS))
                                            if not lst.isColumnHidden(c))
        narrow = room < z(CASE_MIN)
        hh.setSectionResizeMode(0, QHeaderView.Interactive if narrow else QHeaderView.Stretch)
        if narrow:
            lst.setColumnWidth(0, z(CASE_MIN))
        lst.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded if narrow else Qt.ScrollBarAlwaysOff)

    def _show_weekday(self, on: bool) -> None:
        """The job list's Day column shown or hidden (right-click on its headings, or View), kept for next time
        (Settings.job_weekday)."""
        self.job_list.setColumnHidden(JOB_DAY, not on)
        self._fit_case_column()
        if hasattr(self, "weekday_action"):
            self.weekday_action.setChecked(on)
        if self.s.job_weekday != on:
            self._set_opt("job_weekday", on)

    def _job_header_menu(self, pos) -> None:
        """Right-click on the job list's headings: Day of the week, shown or not."""
        menu = QMenu(self)
        act = menu.addAction("Day of the week")
        act.setCheckable(True)
        act.setChecked(self.s.job_weekday)
        act.toggled.connect(self._show_weekday)
        menu.exec(self.job_list.header().mapToGlobal(pos))

    def _job_of(self, item: QTreeWidgetItem | None) -> Job | None:
        """The job of a row of the list: its own, or the job of a document's row; None when it is gone."""
        if item is None:
            return None
        job = (item.parent() or item).data(0, self._JOB)
        return job if job in self.jobs else None

    def _job_item(self, job: Job) -> QTreeWidgetItem | None:
        """The row of a job in the list (None when it has none)."""
        lst = self.job_list
        return next((lst.topLevelItem(i) for i in range(lst.topLevelItemCount())
                     if lst.topLevelItem(i).data(0, self._JOB) is job), None)

    def _select_job(self, job: Job) -> None:
        """Selects a job's row, which shows the job (_job_selected)."""
        item = self._job_item(job)
        if item is not None:
            self.job_list.setCurrentItem(item)

    def _job_text(self, job: Job) -> str:
        """A job's name in the list (its first cell), e.g. "✓ Jane Roe v. Sam Poe" (the column cuts it short).
        ✓ = saved (or billed on another day's invoice), ⚠ = it failed or needs a look: something missing for
        the outputs ticked, or an invoice held until a choice is made (its tooltip says which, see _job_tip)."""
        done = (job.saved or job.invoiced) and not job.error
        mark = "✓ " if done else "⚠ " if job.to_check(self._outputs()) else ""
        return mark + job.title()

    @staticmethod
    def _job_date(job: Job) -> str:
        """A job's date in the list: the first date of its Date(s) of Minutes ("6/3/2026", "6/3/2026 +1" with
        another), or the field as it is when no date is read in it."""
        text = job.case.get("dates")
        found = list(dict.fromkeys(d for _, _, d in find_dates(text)))
        return found[0] + (f" +{len(found) - 1}" if len(found) > 1 else "") if found else text

    def _job_tip(self, job: Job) -> str:
        """A job's tooltip: the case, its dates and index number in full, its documents, the pages typed (when
        Yours shows them), what to check (the reasons for its ⚠, see Job.issues, the index numbers that don't
        match among them), the documents with no index number (Job.index_number_notes), who the form is for and
        what was saved."""
        head = [job.case.get("case_name") or job.title(), job.case.get("dates"), job.case.get("index_no")]
        tip = ["  ·  ".join(dict.fromkeys(v for v in head if v))]
        names = [d.ing.name for d in job.docs]
        tip += ([f"{len(names)} documents:"] if len(names) > 1 else []) + (names or ["(no documents)"])
        if job.pages_typed():
            tip.append(f"Yours: the {job.invoice_pages()} pages typed in Est. number of pages")
        issues = job.issues(self._outputs())
        if issues:
            tip.append("To check before Generate:")
            tip += ["⚠ " + p for p in issues]
        tip += [n for n in job.index_number_notes() if not n.startswith(MISMATCH_LEAD)]  # (a mismatch: in issues)
        who = [a.label() for a in job.case.attorneys if a.checked]
        tip.append("Form for: " + ("; ".join(who) if who else "(blank attorney block)"))
        tip += [f"Saved: {p.name}" for p in job.saved]
        if job.invoiced:
            tip.append("Invoiced: Generate all won't bill this day again")
        return "\n".join(tip)

    def _refresh_jobs(self):
        """Rebuilds the list on the left: a row per job (_shown_jobs), its documents' rows under it (_sync_job_row),
        then their text (_refresh_job_labels), the job on screen selected. A batch gets a tick per job, Generate
        all beside the Generate for the job shown (their words: _generate_labels), and most of the column (the
        divider above the paste box: _place_left_split); a single job a short list. The jobs whose documents the
        user folded away stay folded."""
        multi = len(self.jobs) > 1
        self.fill_all_btn.setVisible(multi)
        # (set here, not with sized(): a zoom change must not give a single job the batch's height)
        self.job_list.setMinimumHeight(z(140) if multi else z(56))
        # (the divider's least height, its panes' least: asked of the column although it asks for no more)
        self.left_split.setMinimumHeight(self.left_split.minimumSizeHint().height())
        self._place_left_split(multi)
        self._size_drop()
        self._generate_labels()
        self.fill_btn.setObjectName("" if multi else "generate")  # (Generate all is the one to click then)
        repolish(self.fill_btn)
        self._collapsed &= set(self.jobs)
        lst = self.job_list
        blocked = lst.blockSignals(True)
        lst.setSortingEnabled(False)  # (sorted once the rows are filled in: _refresh_job_labels)
        lst.clear()
        for job in self._shown_jobs():
            item = JobItem(lst, [""] * len(JOB_COLS))
            item.setData(0, self._JOB, job)
            for c in (JOB_PAGES, JOB_YOURS):
                item.setTextAlignment(c, Qt.AlignRight | Qt.AlignVCenter)
            self._sync_job_row(job, item)
        cur = self._job_item(self.cur)
        if cur is not None:
            lst.setCurrentItem(cur)
        lst.blockSignals(blocked)
        self._refresh_job_labels()
        self._size_drop_later()  # measured again once the column's new contents are laid out

    def _sync_job_row(self, job: Job, item: QTreeWidgetItem | None = None) -> None:
        """Puts a row per document under a job of several (none under a job of one), rebuilt only when its
        documents changed: a click that shows a job never deletes the row clicked. When a document's row was the
        current one, the job's row becomes it. item: the job's row (_refresh_jobs has it); found here when
        not given, and the whole list is rebuilt when the job's row is missing or shouldn't be there (a single
        job's row comes with its first document and goes with its last)."""
        if item is None:
            item = self._job_item(job)
            if (item is None) == (job in self._shown_jobs()):
                self._refresh_jobs()
                return
            if item is None:
                return
        docs = job.docs if len(job.docs) > 1 else []
        keys = " ".join(str(id(d)) for d in docs)  # (text: Qt would hand a tuple back as a list)
        if item.data(0, self._DOCS) == keys and item.childCount() == len(docs):
            return
        lst = self.job_list
        blocked = lst.blockSignals(True)
        was_child = lst.currentItem() is not None and lst.currentItem().parent() is item
        item.takeChildren()
        for n, d in enumerate(docs):
            child = DocItem(item, [""] * len(JOB_COLS))
            child.setData(0, self._DOC, d)
            child.setData(0, JOB_SORT_ROLE, n)  # (in the order read, however the list is sorted)
            for c in (JOB_PAGES, JOB_YOURS):
                child.setTextAlignment(c, Qt.AlignRight | Qt.AlignVCenter)
        item.setData(0, self._DOCS, keys)
        item.setExpanded(job not in self._collapsed)
        if was_child:
            lst.setCurrentItem(item)
        # the strip for the arrows only when a job has documents to fold: room for the case's name otherwise
        lst.setRootIsDecorated(any(lst.topLevelItem(i).childCount() for i in range(lst.topLevelItemCount())))
        lst.blockSignals(blocked)
        self._label_docs(job, item)

    def _job_folded(self, item: QTreeWidgetItem, folded: bool) -> None:
        """A job's documents were folded away (or shown again) by the user: kept so when the list is rebuilt."""
        job = self._job_of(item)
        if job is not None:
            (self._collapsed.add if folded else self._collapsed.discard)(job)

    def _refresh_job_labels(self, counts: dict | None = None):
        """Updates each row of the list (the job's name, date and day of the week, its transcripts' pages and the
        user's, its tooltip (on the tick, what ticking does first; on Yours, what it counts), what it sorts by in
        each column, its tick, greyed while unticked (_dim), and its documents' rows: _label_docs), sorts
        the list (_job_sort) and fits its columns (_fit_job_columns); then the heading above it and the Generate
        buttons' words (_generate_labels: their numbers count every file, the math PDFs too). Pages and Yours are
        hidden when no job has a transcript or pages typed in, Day while Settings.job_weekday is off. A single job
        has no tick, and its heading says "This job".
        counts: batch.output_counts of the jobs Generate all makes, when just counted (_show_counts)."""
        multi = len(self.jobs) > 1
        lst = self.job_list
        blocked = lst.blockSignals(True)  # setting the ticks must not look like the user clicking them
        lst.setSortingEnabled(False)  # (a row moving while the rows are set would be missed: sorted at the end)
        pages_shown = False
        for item in [lst.topLevelItem(i) for i in range(lst.topLevelItemCount())]:
            job = self._job_of(item)
            if job is None:
                continue
            has = bool(job.transcripts()) or job.pages_typed()
            pages_shown |= has
            yours = job.your_pages()
            cells = {0: self._job_text(job), JOB_DATE: self._job_date(job), JOB_DAY: _weekday(job.case.get("dates")),
                     JOB_PAGES: str(job.transcript_pages()) if job.transcripts() else "",
                     JOB_YOURS: ("?" if yours is None else str(yours)) if has else ""}
            tip = self._job_tip(job)
            for c, text in cells.items():
                item.setText(c, text)
                item.setToolTip(c, tip)
            if multi:  # (what the tick does, where the tick is; what Yours counts, on its number)
                head = (UNTICKED_ROW if not job.include else
                        MADE_ROW.format(one=self._one_words(job)) if job.made else TICKED_ROW)
                item.setToolTip(0, f"{head}\n\n{tip}")
            item.setToolTip(JOB_YOURS, f"{YOURS_TIP}\n\n{tip}")
            for c in range(len(JOB_COLS)):
                item.setData(c, JOB_SORT_ROLE, self._job_sort_key(job, c))
            if multi:
                item.setCheckState(0, Qt.Checked if job.include else Qt.Unchecked)
            self._dim(item, multi and not job.include)
            self._label_docs(job, item)
        lst.setColumnHidden(JOB_PAGES, not pages_shown)
        lst.setColumnHidden(JOB_YOURS, not pages_shown)
        lst.setColumnHidden(JOB_DAY, not self.s.job_weekday)
        lst.setSortingEnabled(True)  # (sorted by the heading chosen: _job_sort, _job_sort_key)
        self._fit_job_columns()
        lst.blockSignals(blocked)
        if not multi:
            self.jobs_label.setText("This job")
            return
        chosen = self._to_generate()
        outputs = self._outputs()
        need = sum(1 for j in self.jobs if j.to_check(outputs))
        saved = sum(1 for j in self.jobs if j.saved or j.invoiced)
        text = f"Jobs: {len(self.jobs)}"
        if need:
            text += f"  ·  {need} to check (⚠)"
        if saved:
            text += f"  ·  {saved} saved (✓)"
        self.jobs_label.setText(text)
        # days of a case share a run sheet, invoices too
        from ..batch import output_counts  # (looked up now, as _show_counts does: a test counts the calls)
        self._generate_labels(counts if counts is not None else output_counts(chosen, outputs, self.s))
        self.fill_all_btn.setEnabled(bool(chosen) and bool(outputs) and not self.filling)

    def _label_docs(self, job: Job, item: QTreeWidgetItem) -> None:
        """The rows of a job's documents (under it, for a job of several): the name, a transcript's own day and its
        day of the week or what kind of document it is (_kind_word), a transcript's pages and the user's of them
        (Job.your_pages_of).
        A ⚠ before a document with no index number (as Job.index_number_notes lists them), and, while Generate
        would ask about them (Job.index_number_asks), before one whose index number isn't the job's: the job's are
        those of the group of documents (Job.index_number_mismatch) with the index number the case shows, else
        the first group's. Its tooltip says why above the document's text."""
        if not item.childCount():
            return
        from ..batch import index_key
        mismatch = job.index_number_mismatch() if job.index_number_asks() else []
        own = index_key(job.case.get("index_no"))
        ref = next((g for g in mismatch if any(own in d.index_numbers() for d in g)), mismatch[0]) \
            if mismatch else []
        odd = {id(d) for g in mismatch if g is not ref for d in g}
        theirs = ", ".join(sorted({n for d in ref for n in d.index_numbers()}))
        transcripts = {id(d) for d in job.transcripts()}
        lst = self.job_list
        blocked = lst.blockSignals(True)
        for i in range(item.childCount()):
            child = item.child(i)
            d = child.data(0, self._DOC)
            if d is None:
                continue
            why = ("No index number found in this document" if not d.index_numbers() else
                   f"Its index number ({', '.join(sorted(d.index_numbers()))}) isn't this job's ({theirs})"
                   if id(d) in odd else "")
            if id(d) in transcripts:
                yours = job.your_pages_of(d)
                cells = [job.day_text(d) or "transcript", _weekday(job.day_text(d)), str(transcript_pages(d.ing)),
                         "?" if yours is None else str(yours)]
            else:
                cells = [_kind_word(d), "", "", ""]
            child.setText(0, ("⚠ " if why else "") + d.ing.name)
            for c, text in zip((JOB_DATE, JOB_DAY, JOB_PAGES, JOB_YOURS), cells):
                child.setText(c, text)
            text = d.ing.text[:1500] or "(no text)"
            if d.ing.ocr_used:
                text = "(read with OCR)\n" + text
            tip = f"⚠ {why}\n\n{text}" if why else text
            for c in range(len(JOB_COLS)):
                child.setToolTip(c, tip)
            if id(d) in transcripts:
                child.setToolTip(JOB_YOURS, f"{YOURS_TIP}\n\n{tip}")
            self._dim(child, len(self.jobs) > 1 and not job.include)
        lst.blockSignals(blocked)

    def _dim(self, item: QTreeWidgetItem, on: bool) -> None:
        """A row of the list greyed while its job is unticked (the theme's muted text: the user, 2026-10-23, "subtle,
        but something"), or as the others again. The caller blocks the list's signals (a colour set is an
        itemChanged: _job_ticked)."""
        brush = QBrush(QApplication.palette().color(QPalette.PlaceholderText)) if on else None
        for c in range(len(JOB_COLS)):
            item.setData(c, Qt.ForegroundRole, brush)

    def _redim(self) -> None:
        """The unticked rows greyed again in the new theme's colour: _dim sets a fixed colour on each row, which
        doesn't follow a theme chosen afterwards by itself."""
        lst = getattr(self, "job_list", None)
        if lst is None:
            return
        multi = len(self.jobs) > 1
        with QSignalBlocker(lst):
            for i in range(lst.topLevelItemCount()):
                item = lst.topLevelItem(i)
                job = self._job_of(item)
                if job is not None:
                    for row in [item] + [item.child(k) for k in range(item.childCount())]:
                        self._dim(row, multi and not job.include)

    def _considered(self) -> list[Job]:
        """The jobs the cards on the right show: every one ticked in the list, not just the one shown (the user,
        2026-10-23: the cards followed the day shown, Who transcribed what every day, and "the checkboxes ...
        should dictate what files are being considered"); the job shown when it is the only one. Empty jobs
        are left out."""
        if len(self.jobs) < 2:
            return [] if self.cur.is_empty() else [self.cur]
        return [j for j in self.jobs if j.include and not j.is_empty()]

    def _to_generate(self) -> list[Job]:
        """The jobs Generate all makes: those ticked, but not those made already (Job.made: their ✓; Generate for
        the job shown makes one again), unless a made day's invoice is still to make (made while Invoice was
        unticked, say): Generate all then makes its invoice alone, with its case's other days (batch.fill_jobs),
        as the cards say it bills it. A held one too (Whose pages... to choose, say): Generate all asks about it,
        or says why its case gets no invoice, as for a day not made yet, rather than billing the others alone."""
        bill = self._ticked("invoice")
        return [j for j in self.jobs if j.include and not j.is_empty() and (
            not j.made or bill and not j.invoiced and j.invoice_pages())]

    def _show_job(self):
        """Puts the current job in the editor, and its documents in the list (_sync_job_row)."""
        self._sync_job_row(self.cur)
        self._show_case()
        self._update_status()

    def _show_index_notes(self) -> None:
        """What the current job's documents say about the index number that needs a look (Job.index_number_notes),
        under the list: a document with none, documents whose numbers don't match (with "Keep them together");
        and a ⚠ on each such document's row (_label_docs), its tooltip saying why. Kept up to date with the case
        (_show_case), as the AI's answer can bring a number."""
        job = self.cur
        notes = job.index_number_notes()
        mismatch = job.index_number_mismatch() if job.index_number_asks() else []
        item = self._job_item(job)
        if item is not None:
            self._label_docs(job, item)
        lines = ["⚠ " + html.escape(n) for n in notes]
        if mismatch:
            lines.append("<a href='keep'>Keep them together</a> (or right-click a document → Move to a job of its "
                         "own)")
        self.index_note.setText("<br>".join(lines))
        self.index_note.setVisible(bool(lines))

    def _keep_index_numbers(self) -> None:
        """"Keep them together": the job's documents are about one case though their index numbers don't match
        (Job.index_numbers_ok); Generate won't ask, until another document comes."""
        self.cur.index_numbers_ok = [d.key() for d in self.cur.docs]
        self._show_index_notes()
        self._refresh_job_labels()

    def _job_selected(self, job: Job | None):
        """A row was clicked in the list (a job, or one of its documents): keep the edits made to the job on
        screen, then show that job."""
        if job is None or job is self.cur:
            return
        self._sync_from_ui()
        self.cur = job
        self._show_job()

    def _job_ticked(self, item: QTreeWidgetItem, column: int = 0):
        """A job's tick changed: the cards on the right show it or leave it out (_considered: the invoices and their
        math with them), and so does Generate all; its row is greyed while unticked (_refresh_job_labels)."""
        job = self._job_of(item)
        if column != 0 or item.parent() is not None or job is None or len(self.jobs) < 2:
            return
        job.include = item.checkState(0) == Qt.Checked
        self._update_status()  # (the list of what to check too: a day unticked isn't billed)

    def _job_menu(self, pos):
        """Right-click on the list. A document's row: show its text, move it to a job of its own, remove it. A
        job's row: show the text of its document, for a job of one; Remove this job in a batch (a single job:
        Remove from job, for its one document); Tick all and Untick all in a batch. The jobs can change while the
        menu is open (a read that ends merges them): what is chosen acts on the job and document it was opened
        on, if they are still there (_remove_doc, _remove_job)."""
        item = self.job_list.itemAt(pos)
        job = self._job_of(item)
        doc = item.data(0, self._DOC) if item is not None and item.parent() is not None else None
        multi = len(self.jobs) > 1
        if job is not None and doc is None and len(job.docs) == 1:
            doc = job.docs[0]  # (a job of one document: its row is the document's too)
        m = QMenu(self)
        view = m.addAction("Show extracted text") if doc is not None else None
        split = m.addAction("Move to a job of its own") if doc is not None and len(job.docs) > 1 else None
        rem_doc = m.addAction("Remove from job") if doc is not None and (item.parent() is not None or not multi) \
            else None
        rem = m.addAction("Remove this job") if job is not None and multi and item.parent() is None else None
        tick = untick = None
        if multi:
            m.addSeparator()
            tick = m.addAction("Tick all")
            untick = m.addAction("Untick all")
        if m.isEmpty():
            return
        act = m.exec(self.job_list.viewport().mapToGlobal(pos))
        if act is None:
            return
        if act in (tick, untick):
            for j in self.jobs:
                j.include = act == tick
            self._update_status()  # the days Generate all bills together, and what to check
        elif act == view:
            self._doc_text(doc)
        elif act in (split, rem_doc):
            self._remove_doc(job, doc, own_job=act == split)
        elif act == rem:
            self._remove_job(job)

    def _doc_text(self, doc: Doc) -> None:
        """Show extracted text: the text read from a document, in a box."""
        box = QMessageBox(self)
        box.setWindowTitle(doc.ing.name)
        box.setText("Text read from this document:")
        box.setDetailedText(doc.ing.text or "(no text)")
        box.exec()

    def _remove_doc(self, job: Job, doc: Doc, own_job: bool = False) -> None:
        """Remove from job, or (own_job) Move to a job of its own, right after it: a document taken out of its
        job, which is merged again without it. The job may be another than the one on screen. A job of a batch
        left with no document goes, as with Remove this job; a single job stays, empty."""
        if job not in self.jobs or doc not in job.docs:  # (the list may have changed while the menu was open)
            return
        self._sync_from_ui()
        job.docs.remove(doc)
        remerge(job, self.s)
        if own_job:
            new = Job(docs=[doc], batch=True)
            remerge(new, self.s)
            self.jobs.insert(self.jobs.index(job) + 1, new)
        elif not job.docs and len(self.jobs) > 1:
            # its last document, in a batch (a menu opened on the job alone, then a read made it a batch): the
            # job goes, as with Remove this job, rather than stay in the list empty
            self.jobs.remove(job)
            if self.cur not in self.jobs:
                self.cur = self._jobs_in_list_order()[0]
        self._auto_runsheet()
        self._refresh_jobs()
        self._show_job()

    def _remove_job(self, job: Job) -> None:
        """Remove this job: a job of a batch taken out of the list (and of Generate all); the last one leaves a
        blank job."""
        if job not in self.jobs:  # (the list may have changed while the menu was open)
            return
        self._sync_from_ui()
        self.jobs.remove(job)
        if not self.jobs:
            self.jobs = [Job()]
        if self.cur not in self.jobs:
            self.cur = self._jobs_in_list_order()[0]
        self._auto_runsheet()
        self._refresh_jobs()
        self._show_job()

    # ----------------------------------------------------- merge/show
    def _remerge(self, job: Job | None = None):
        """Merges a job's documents again (default: the current job). For the job on screen, the edits in
        the editor are taken first and the editor shows the result, with a cell of the attorney table that was
        being typed in opened again (_resume_edit: an AI answer arrives while the user types)."""
        job = job or self.cur
        typing = self._sync_from_ui() if job is self.cur else None
        remerge(job, self.s)
        if job is self.cur:
            self._show_case()
            self._resume_edit(typing)

    def _show_case(self):
        """Shows the current job's case in the editor: fields, speed, proceeding boxes, attorneys, invoice line.
        No. of copies is worked out again first (the ticks may have changed while another job was shown), and
        the agreement form's speed by the Settings rule unless chosen (a new job gets the rule's speed and rate,
        not those the window showed last)."""
        self.cur.refresh_copies(self.s)
        apply_speed_rule(self.case, self.s)
        loaded = self._loaded()
        for key, r in self.rows.items():
            r.show_missing = loaded  # (an empty window has nothing missing yet)
            r.set_state(self.case.fields[key])
        self._fill_speeds()
        for p, cb in self.proc_boxes.items():
            cb.blockSignals(True)
            cb.setChecked(p in self.case.proc_types)
            cb.blockSignals(False)
        self._show_attorneys()
        self._refresh_outputs()
        self._show_index_notes()
        self._show_start()
        self._show_date_notes()

    def _loaded(self) -> bool:
        """Whether anything is loaded or typed: a job with documents or details, pasted text, or a read going
        on. Until then the window shows the start card and no field is marked missing."""
        return not all(j.is_empty() for j in self.jobs) or bool(self.paste.toPlainText().strip()) or self.work > 0

    def _mark_missing(self) -> None:
        """The required fields left blank are marked once something is loaded (FieldRow.show_missing), and no
        longer when the window is emptied."""
        loaded = self._loaded()
        for r in self.rows.values():
            if r.show_missing != loaded:
                r.show_missing = loaded
                r._refresh()

    def _show_attorneys(self):
        """Fills the attorney table from the current case (without counting as an edit by the user), and the
        details under it with the row that was selected (by its place), or the first."""
        row = self.att.currentRow()
        self.att.blockSignals(True)
        self.att.setRowCount(0)
        marks = self._marks(self.cur)
        for a in self.case.attorneys:
            self._append_att(a)
            self._mark_row(self.att, self.att.rowCount() - 1, marks.get(a.key()))
        self.att.resizeRowsToContents()
        if row >= 0 and self.att.rowCount():  # (none selected before: none now, so Remove selected removes nothing)
            self.att.setCurrentCell(min(row, self.att.rowCount() - 1), 1)
        self.att.blockSignals(False)
        self._show_att_details()

    def _marks(self, job: Job | None) -> dict:
        """The marks of the firms of job's trial, by Attorney.key() (batch.trial_marks: plaintiffs 1, 2, defense
        A, B, others 001): the same on every day of it, so a firm is told apart at a glance wherever it is named
        (the user, 2026-10-21)."""
        from ..batch import trial_marks
        if job is None:
            return {}
        try:
            return trial_marks(job, self.jobs)
        except Exception as e:  # (a mark is a help: never a reason for the window to fail)
            log_error("could not mark the firms", e)
            return {}

    def _mark_row(self, table: QTableWidget, row: int, mark) -> None:
        """A firm's mark on its row of a table (the attorneys', Who ordered what's): its tag beside its name
        (column 1) and a stripe of its colour at the name's left edge (badges.StripeDelegate); nothing without a
        mark. No shade behind the row: it read as the amber of a field to check (the user, 2026-10-22)."""
        item = table.item(row, 1)
        if item is None:
            return
        item.setIcon(badge_icon(mark, z(TAG_HEIGHT)) if mark is not None else QIcon())
        item.setData(STRIPE_ROLE, tag_color(mark).name() if mark is not None else None)

    def _mark_attorneys(self) -> None:
        """Each row of the attorney table marked as its firm is now (_mark_row): an edit can change a firm's side,
        its name or the rows before it, and a zoom the tags' size. The rows are those of self.case.attorneys."""
        marks = self._marks(self.cur)
        with QSignalBlocker(self.att):  # (an icon set is an itemChanged: _att_changed would take it as an edit)
            for row, a in enumerate(self.case.attorneys[:self.att.rowCount()]):
                self._mark_row(self.att, row, marks.get(a.key()))

    def _append_att(self, a: Attorney):
        """Adds a table row for one attorney. Address lines show joined by " / "; a placeholder row is greyed."""
        r =self.att.rowCount()
        self.att.insertRow(r)
        chk = QTableWidgetItem()
        chk.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled)
        chk.setCheckState(Qt.Checked if a.checked else Qt.Unchecked)
        self.att.setItem(r, 0, chk)
        for c, f in enumerate(ATT_FIELDS[1:], 1):
            val = getattr(a, f) or ""
            if f == "address":
                val = " / ".join(l for l in val.splitlines() if l.strip())
            it = QTableWidgetItem(val)
            if f == "source":
                it.setFlags(Qt.ItemIsEnabled)
            if a.is_placeholder():
                it.setForeground(self.palette().placeholderText())
            self.att.setItem(r, c, it)

    def _read_attorneys(self) -> list[Attorney]:
        """The attorney table as Attorney objects, with " / " in the address turned back into line breaks."""
        out = []
        for r in range(self.att.rowCount()):
            vals = {}
            for c, f in enumerate(ATT_FIELDS[1:], 1):
                it = self.att.item(r, c)
                vals[f] = it.text().strip() if it else ""
            # " / " separates lines; "c/o" and "12-1/2" are left alone
            vals["address"] = "\n".join(p.strip() for p in re.split(r"\s+/\s*|\s*/\s+", vals["address"])
                                        if p.strip())
            a = Attorney(**vals)
            a.checked = self.att.item(r, 0).checkState() == Qt.Checked if self.att.item(r, 0) else False
            out.append(a)
        return out

    def _sync_from_ui(self, commit: bool = True) -> tuple[int, str, int] | None:
        """Copies the user's edits in the editor into the current job's case: typed fields, the proceeding
        boxes and the attorney table. Call it before the case is merged again, switched or filled.

        commit: a cell still being typed in is taken as typed (an AI answer or a second document arriving
        must not throw it away), and its editor closes; returns what _commit_edit returns, for _resume_edit.
        With commit=False (the status pill, which only reads the fields) an open editor is left alone."""
        typing = self._commit_edit() if commit else None
        for key, r in self.rows.items():
            if r.state.source == SRC_USER:
                self.case.fields[key] = FieldState(r.text(), SRC_USER, 1.0, r.state.alternatives)
        # (the speed has no FieldRow: _delivery_changed, Keep and ↺ put it in the case as it is chosen)
        self.case.proc_types = {p for p, cb in self.proc_boxes.items() if cb.isChecked()}
        self.case.attorneys = self._read_attorneys()
        return typing

    # ---------------------------------------------------- UI events
    def _field_edited(self, key: str):
        """The user typed in a field (or picked a suggestion): it goes into the case straight away. The
        transcripts' own page count typed in Est. number of pages stays the count (badge PDF): as the user's
        number it would bill the other reporters' pages too (see Job.pages_typed). (Only the whole count: this
        runs at each key, and a field emptied or a number half typed must stay as it is. A blank field, or the
        user's own pages typed, are no number of the user's either: Job.is_the_count.)"""
        st = self.rows[key].state
        whole = self.cur.transcript_pages() if key == "est_pages" else 0
        if whole and st.source == SRC_USER and st.value.strip() == str(whole):
            st = FieldState(str(whole), SRC_PDF, self.cur.count_confidence(), st.alternatives)
            self.rows[key].set_state(st)
        self.case.fields[key] = st
        if key in self.date_notes:  # (a day typed or picked: a weekend or a court holiday is said)
            self._show_date_notes()
        self._update_status()

    def _proc_toggled(self, _):
        """A proceeding box was clicked: the tick goes into the case straight away (an answer about the firms
        shows the case again from it), and later merges keep the user's choice."""
        self.proc_touched = True
        self.case.proc_types = {p for p, cb in self.proc_boxes.items() if cb.isChecked()}

    def _att_changed(self, _item):
        """A cell or tick of the attorney table changed: later merges keep the user's table, and the details under
        it show the change."""
        self.att_touched = True
        self._attorneys_edited()
        self._show_att_details()

    def _attorneys_edited(self):
        """A tick or an edit in the attorney table: No. of copies (the ordering parties, unless typed), the
        invoice's parties and the file counts follow, and so do the Excerpts... rows when an attorney's name
        (or firm) was changed."""
        old, new = self.case.attorneys, self._read_attorneys()
        if len(old) == len(new):  # the same rows (none added or removed): the attorney of each row is the same
            keys = {a.key() for a in new}
            for before, after in zip(old, new):
                if before.key() != after.key() and before.key() not in keys:
                    self.cur.rename_in_portions(before.key(), after.key())
        self.case.attorneys = new
        self._show_copies()
        self._update_status()
        self._mark_attorneys()  # (its own rows too, as Who ordered what's: a party typed can change a tag)

    def _show_copies(self, jobs: list[Job] | None = None):
        """No. of copies follows the ordering parties (Job.refresh_copies: the Parties number when set, else the
        attorneys ticked) of these jobs (default: the current one), unless typed; the editor shows the current
        job's."""
        for job in jobs or [self.cur]:
            if job is self.cur and self.rows["copies"].state.source == SRC_USER:
                continue  # typed in the editor (and perhaps not yet in the case)
            job.refresh_copies(self.s)
            if job is self.cur:
                self.rows["copies"].set_state(self.case.fields["copies"])

    def _add_att_row(self):
        """Adds a ticked, empty attorney row and starts editing its Name cell."""
        self.att.blockSignals(True)
        self._append_att(Attorney(source=SRC_USER, checked=True))
        self.att.blockSignals(False)
        self.att_touched = True
        self._attorneys_edited()
        last = self.att.rowCount() - 1
        self.att.setCurrentCell(last, 1)  # (so the details under the table are the new row's)
        self.att.editItem(self.att.item(last, 1))

    def _remove_att_rows(self):
        """Removes the selected attorney rows."""
        rows = sorted({i.row() for i in self.att.selectedIndexes()}, reverse=True)
        for r in rows:
            self.att.removeRow(r)
        if rows:
            self.att_touched = True
            self._attorneys_edited()
            self._show_att_details()  # (the row they showed may be gone: the current cell moved before it went)

    # ------------------------------------------------ rate sheets / speed
    def _fill_sheet_box(self):
        """Lists the rate sheets (each tooltip shows its speeds, the headings read as another price's, the prices
        it is missing and those lower at a faster speed), says which files were skipped, then the speeds."""
        from ..rates import cheaper_faster, missing_prices
        sheets, problems = self.s.sheets()
        current = self.s.sheet()
        shown = sheets or [current]

        def missing(sh) -> str:
            """'\n⚠ Missing: Regular: Email, Index; Realtime: Original', then any price lower at a faster speed ('' for
            a complete sheet in order)."""
            gaps, cheaper = missing_prices(sh), cheaper_faster(sh)
            out = ("\n⚠ Missing: " + "; ".join(f"{name}: {', '.join(what)}" for name, what in gaps)) if gaps else ""
            return out + (("\n⚠ Lower at a faster speed: " + "; ".join(cheaper)) if cheaper else "")

        days = "business days" if self.s.turnaround_business_days else "days"
        tips = ["\n".join(sp.label() + (f"  ·  {sp.days} {days}" if sp.days is not None else "") for sp in sh.speeds)
                + (f"\nRates last updated {sh.updated}" if sh.updated else "")
                + "".join(f'\nThe "{heading}" column is read as the {label} price'
                          for heading, label in sh.read_as.items())
                + missing(sh) for sh in shown]
        refill_combo(self.sheet_box, [(sh.name, sh.name) for sh in shown], current.name, tips)
        for p in problems:
            self._toast(f"Rate sheet skipped - {p}")
        self._fill_speeds()

    def _fill_speeds(self):
        """Fills the agreement form's Speed box: the speeds the invoice offers, cheapest first as the boxes above
        (with their turnaround days; every speed of the sheet when none is ticked), plus "Other", and shows the
        job's speed (else the Settings rule's: never the box's last choice, which was another job's)."""
        keep = self.case.get("delivery") or self.s.agreement_speed()
        items = []
        for sp in self.s.offered_speeds():
            days = self.s.days_for(sp.name)
            items.append((sp.label() + (f"  ·  {plural(days, 'day')}" if days is not None else ""), sp.name))
        refill_combo(self.delivery, items + [("Other (type the rate yourself)", "Other")])
        self._select_speed(keep)

    def _select_speed(self, name: str):
        """Selects `name` (matched loosely, e.g. 'Expedited' -> 'Expedite') without firing change events. A speed
        the box doesn't list shows as the Settings rule's ("Other" when the sheet doesn't have it at all)."""
        sp = self.s.sheet().find(name)
        target = sp.name if sp else ("Other" if name else self.s.agreement_speed())
        if self.delivery.findData(target) < 0 and target != "Other":
            target = self.s.agreement_speed()
        with QSignalBlocker(self.delivery):
            self.delivery.setCurrentIndex(max(0, self.delivery.findData(target)))
        self._show_rate_info()
        self._show_speed_state()

    def _show_speed_state(self):
        """The agreement form speed's badge ("default": the Settings rule's; "you": chosen for this job), its ↺
        (back to the rule, shown for a speed chosen) and the question when an e-mail asks for another speed
        than the one shown (Job.speed_question): the box is marked and Keep answers it."""
        f = self.case.fields["delivery"]
        src = f.source if f.value else ""
        self.speed_badge.setText(badge_word(src))
        self.speed_badge.setProperty("src", src)
        self.speed_badge.setProperty("loud", src == SRC_AI)
        self.speed_badge.setToolTip({SRC_USER: "Chosen by you for this job",
                                     SRC_DEFAULT: "Picked by the rule in Settings → Invoice"}.get(src, ""))
        repolish(self.speed_badge)
        rule = self.s.delivery_name(self.s.agreement_speed())
        self.speed_reset.setVisible(src == SRC_USER)
        self.speed_reset.setToolTip(f"Back to the speed Settings → Invoice picks ({rule})")
        ask = self.cur.speed_question()
        if ask:
            asked, _, _ = self.cur.asked_speed()
            name = self.s.delivery_name(asked)
            text = f"⚠ {ask[0].upper()}{ask[1:]}. Please choose a speed for this job:"
            if not self.s.speed_offered(asked):
                text += f" ({name} isn't offered on the invoice: tick it under Speeds offered to choose it)"
            self.speed_ask.setText(text)
            self.speed_keep.setText(f"Keep {self.delivery.currentData() or rule}")
        self.speed_form.setRowVisible(self.speed_ask_row, bool(ask))
        self.delivery.setProperty("ask", bool(ask))
        repolish(self.delivery)

    def _ask_speeds(self, jobs: list[Job]) -> bool:
        """Before agreement forms or MOFRs are made: for the jobs whose e-mail asks for another speed than the
        one their form names (Job.speed_question), asks which speed (SpeedDialog: the e-mail's preselected when
        it is offered). Each answer is the user's choice for the job, its rate and delivery date following.
        False when Go back was clicked, or Generate all started meanwhile (nothing is made then)."""
        from .dialogs import SpeedDialog
        ask = [j for j in jobs if j.speed_question()]
        if not ask:
            return True
        speeds = [(sp.label(), sp.name) for sp in self.s.offered_speeds()]
        rows = []
        for j in ask:
            asked, _, _ = j.asked_speed()
            pre = self.s.delivery_name(asked) if self.s.speed_offered(asked) else j.case.get("delivery")
            title = ", ".join(x for x in (j.title(), j.case.get("dates")) if x)
            rows.append((f"{title}: {j.speed_question()}", speeds, pre))
        dlg = SpeedDialog(rows, self)
        if dlg.exec() != SpeedDialog.Accepted or self._batch_running():
            return False
        for j, name in zip(ask, dlg.chosen()):
            if name and j in self.jobs:  # (a read that ended meanwhile may have merged it into another job)
                j.case.fields["delivery"] = FieldState(name, SRC_USER, 1.0, [name])
                refresh_rate(j.case, self.s)
                if self.s.fill_delivery_date:
                    refresh_delivery_date(j.case, self.s)
        if any(j is self.cur for j in ask):
            self._fill_speeds()
            self._apply_speed()
        else:
            self._update_status()
        return True

    def _ask_parties(self, jobs: list[Job]) -> bool:
        """Before anything is made: a day split under Excerpts... whose Parties number, set by hand, isn't the
        number of firms that ordered (Job.parties_mismatch) may have a firm left unticked by mistake. The
        invoices and No. of copies count the firms that ordered; Go back leaves it to be checked (False), Go on
        makes the files so (True)."""
        odd = [j for j in jobs if j.parties_mismatch()]
        if not odd:
            return True
        lines = "\n".join(f"•  {', '.join(x for x in (j.title(), j.case.get('dates')) if x)}: {j.parties_mismatch()}"
                           for j in odd[:8]) + ("\n…" if len(odd) > 8 else "")
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Warning)
        box.setWindowTitle("Parties and Excerpts disagree")
        box.setText(f"{lines}\n\nThe invoices and No. of copies count the firms that ordered pages under "
                    "Excerpts…, not the Parties number. If a firm that ordered is missing, go back and tick it "
                    "under Excerpts….")
        back = box.addButton("Go back", QMessageBox.RejectRole)
        go_on = box.addButton("Go on (count the firms ticked)", QMessageBox.AcceptRole)
        box.setDefaultButton(back)
        box.setEscapeButton(back)
        box.exec()
        return box.clickedButton() is go_on and not self._batch_running()

    def _ask_split_speeds(self, jobs: list[Job], button: str = "Generate") -> bool:
        """Before the files are made: the firms of a split order whose speed isn't set (batch.speeds_to_ask:
        firms that ordered the same pages commit to a speed, as the courthouse says) are asked about at once
        (dialogs.SplitSpeedsDialog, each set to the firm's speed on its other pages of the case, else the agreement
        form's). Each answer is the firm's speed for all its pages of the day that have none of their own
        (Job.speeds; one set meanwhile is kept). Generate asks again after the Clarify box, for the firms ticked
        there. False when Go back was clicked, or Generate all started meanwhile (nothing is made then)."""
        from ..batch import speeds_to_ask
        from .dialogs import SplitSpeedsDialog
        asks = speeds_to_ask(jobs, self.s)
        if not asks:
            return True
        marks = {id(job): self._marks(job) for job in dict.fromkeys(a.job for a in asks)}  # (by trial: a firm's
        # mark in one can be another's in the next)
        dlg = SplitSpeedsDialog(asks, [(sp.label(), sp.name) for sp in self.s.sheet().speeds], button, self,
                                marks)
        if dlg.exec() != SplitSpeedsDialog.Accepted or self._batch_running():
            return False
        from ..extract_regex import same_entry
        for a, name in zip(asks, dlg.chosen()):
            if name and a.job in self.jobs:  # (a read that ended meanwhile may have merged it into another job)
                rows = a.job.case.attorneys
                key = a.key
                if a.attorney is not None and not any(x.key() == key for x in rows):
                    # (a document read meanwhile renamed the firm: "Dana Smith" of "Smith Law" now)
                    key = next((x.key() for x in rows if same_entry(a.attorney, x)), key)
                a.job.speeds.setdefault(key, name)
        self._update_status()
        return True

    def _ask_index(self, groups: list[list[Job]]) -> bool:
        """Before invoices are made: for each set of days billed together (groups) where a Pages number typed
        below the transcript's own count changes whether there is an index (batch.index_question), asks whether
        to include one: the user made the count their own, so the app doesn't decide. The answer is kept on the
        typed days (Job.index_on_typed: No judges them on the number typed, Yes on the transcript), so it isn't
        asked again; Peripherals... shows it. When the rate sheet has no index price for the invoice's speeds
        (IndexQuestion.priced), no index is charged either way: a warning says so instead, and nothing is kept, so
        Generate asks once a price is added. False when Go back was clicked (nothing is made), or Generate all
        started meanwhile. Answers given before a Go back are kept, as the box says."""
        from ..batch import index_question
        asked = False
        try:
            for group in groups:
                group = [j for j in group if j in self.jobs]  # (a read that ended meanwhile may have merged one)
                q = index_question(group, self.s) if group else None
                if q is None:
                    continue
                dates = [d for j in group if (d := j.case.get("dates"))]
                title = ", ".join(x for x in (group[0].title(), ", ".join(dates)) if x)
                several = len(group) > 1
                # under the rule "each" day on its own, the answer is about the days typed only
                whom = ("those days" if len(q.jobs) > 1 else "that day") \
                    if several and self.s.invoice_index_rule == "each" else "the invoice"
                which = f" for {', '.join(d for j in q.jobs if (d := j.case.get('dates')))}" if several else ""
                has = "the transcripts have" if q.transcripts > 1 else "the transcript has"
                gets, it = ("get", "they get") if whom == "those days" else ("gets", "it gets")
                box = QMessageBox(self)
                if not q.priced:  # no index can be charged either way: nothing to choose, but say why there's none
                    box.setIcon(QMessageBox.Warning)
                    box.setWindowTitle("No index price")
                    box.setText(f"{title}\n\nYou typed {q.typed:,} pages in Est. number of pages{which}; {has} "
                                f"{q.counted:,}.\n\nJudged on the transcript, {whom} {gets} an index "
                                f"({q.pdf_why}), but your rate sheet has no index price for the speeds this invoice "
                                "offers, so no index is charged.")
                    box.setInformativeText("To charge one, add a price in the Index column of the rate sheet (File → "
                                           "Open rate sheets folder); Generate then asks about the index.")
                    go = box.addButton("Go on without an index", QMessageBox.AcceptRole)
                    back = box.addButton("Go back", QMessageBox.RejectRole)
                    box.setDefaultButton(go)
                    box.setEscapeButton(back)
                    box.exec()
                    if box.clickedButton() is not go or self._batch_running():  # (closed: as Go back)
                        return False
                    continue
                box.setIcon(QMessageBox.Question)
                box.setWindowTitle("Index for this invoice?")
                box.setText(f"{title}\n\nYou typed {q.typed:,} pages in Est. number of pages{which}; {has} "
                            f"{q.counted:,}.\n\nJudged on the transcript, {whom} {gets} an index ({q.pdf_why}); "
                            f"judged on the pages you typed, {it} none ({q.typed_why}).\n\nInclude an index (and the "
                            f"judge's) {'on this invoice' if whom == 'the invoice' else 'for ' + whom}?")
                box.setInformativeText("Your answer is kept for this job: Peripherals… in the Invoice card "
                                       "changes it.")
                yes = box.addButton("Yes, an index", QMessageBox.YesRole)
                no = box.addButton("No index", QMessageBox.NoRole)
                back = box.addButton("Go back", QMessageBox.RejectRole)
                box.setDefaultButton(back)
                box.setEscapeButton(back)
                box.exec()
                if box.clickedButton() not in (yes, no) or self._batch_running():  # (closed: as Go back)
                    return False
                for j in q.jobs:
                    if j in self.jobs:
                        j.index_on_typed = box.clickedButton() is no
                asked = True
            return True
        finally:
            if asked:  # (the Includes line and the prices follow the answers, a later Go back too)
                self._update_status()

    def _speed_to_rule(self):
        """↺: the job's speed goes back to the one Settings → Invoice picks; the rate and delivery date follow."""
        self.case.fields["delivery"] = FieldState()
        self._apply_speed_rule()

    def _keep_speed(self):
        """Keep: the speed shown is the user's choice for this job (answers an e-mail's other speed)."""
        name = self.delivery.currentData() or ""
        if name:
            self.case.fields["delivery"] = FieldState(name, SRC_USER, 1.0, [name])
        self._apply_speed()

    def _show_rate_info(self):
        """The line under the Rate field (Speed, pages and dates): the copy rate, the speed's other prices
        and when the rates changed."""
        sheet = self.s.sheet()
        sp = sheet.find(self.delivery.currentData() or "")
        bits = []
        if sp and sp.copy:
            bits.append(f"Copies ${sp.copy}/pg")
        if sp:
            bits += [f"{k} ${v.lstrip('$')}" for k, v in sp.extras.items()]
        if sheet.updated:
            bits.append(f"rates updated {sheet.updated}")
        self.rate_info.setText("  ·  ".join(bits))

    def _sheet_changed(self, _idx):
        """Another rate sheet was picked: it becomes the default, and every job is priced again from it (the
        agreement form's speed by the Settings rule again, unless chosen and still offered: _sheet_in_use). An
        incomplete one is warned about first (_check_rate_sheet), before anything is priced, with the one in use
        before to go back to: the jobs are then as they were (a speed chosen that the new sheet lacks isn't
        reset)."""
        name = self.sheet_box.currentData()
        if not name or name == self.s.rate_sheet:
            return
        # (the one in use: the sheet Settings names may be gone, another being used in its place)
        previous = self.s.sheet().name
        picked = next((sh for sh in self.s.sheets()[0] if sh.name == name), None)
        # checked before Settings names it: a read or an AI answer that ends while the box is open is priced from
        # the sheet in use, which Go back keeps
        if picked is None or self._check_rate_sheet(previous, again=True, priced=False, sheet=picked):
            self.s.rate_sheet = name
            self._sheet_in_use()
            if picked is None:
                self._check_rate_sheet(previous, again=True)
            self._check_old_days()

    def _sheet_in_use(self) -> None:
        """The rate sheet in use changed: saved, and every job priced again from it."""
        self._drop_undo()  # (the jobs New job cleared were priced from the sheet before)
        self._save_settings()
        self._fill_invoice_speeds()
        self._apply_speed_rule()
        self._remerge_others()

    def _check_rate_sheet(self, previous: str = "", again: bool = False, priced: bool = True,
                          sheet=None) -> bool:
        """Warns when the rate sheet in use leaves out prices (rates.missing_prices): a speed left off the sheet is
        one not offered, which is fine, but every speed on it needs its Original, Copy, Email and Index price, or
        invoices come out wrong (the user, 2026-10-09); or when a price is lower at a faster speed
        (rates.cheaper_faster), which can't be right, though the prices are the user's to get right (2026-10-20).
        Called whenever the sheet in use may have changed: at start
        (it may have been edited meanwhile), when another is picked or the sheets are read again (⟳), and after
        Settings or an import. Use it anyway remembers the sheet as it was read (Settings.rate_sheets_ok, by
        rates.sheet_fingerprint, the last RATE_SHEETS_OK_KEPT): not asked again until its file is changed and read
        again. Go back to <previous> (the sheet in use before; offered when it is another, still in the folder)
        picks that one again, which is checked in turn; Open the sheet to fill it in leaves this one in use and
        opens its file, read again and checked each time it is saved (_watch_sheet); Not now (offered, and Esc,
        when there is no Go back: at start, ⟳) leaves it in use as it is. Asked once a session for the same file as
        read unless `again` (picked, ⟳, another chosen in Settings). priced: the jobs were priced from the sheet
        already (they are priced from `previous` again on Go back). sheet: the one to check, when it isn't in use
        yet (picked in the dropdown: Settings names it only once it is kept). False = Go back was clicked (the
        sheet before is in use again)."""
        from ..rates import cheaper_faster, missing_prices, sheet_fingerprint
        from ..settings import RATE_SHEETS_OK_KEPT
        if self.filling:  # (Generate all runs, on its own copy of the settings: asked when it ends, _batch_ended)
            if self._sheet_check_later is None:  # (the first sheet before: Go back goes there)
                self._sheet_check_later = previous
            return True
        sheet = sheet or self.s.sheet()
        gaps, cheaper = missing_prices(sheet), cheaper_faster(sheet)
        mark = sheet_fingerprint(sheet)
        if (not (gaps or cheaper) or not mark or mark in self.s.rate_sheets_ok
                or (mark in self._sheets_warned and not again)):
            return True
        self._sheets_warned.add(mark)
        lines = "\n".join(f"•  {name}: no {_or(what)} price"
                          + (" (so it isn't offered at all)" if what == ["Original"] else "") for name, what in gaps[:8])
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Warning)
        box.setWindowTitle("Your rate sheet is incomplete" if gaps else "Your rate sheet's prices look wrong")
        said = [f'The rate sheet "{sheet.name}" is missing some prices:\n\n{lines}'
                + ("\n…" if len(gaps) > 8 else "") + "\n\nLeaving off a speed you don't offer is fine, but every "
                "speed on the sheet needs all four prices: Original, Copy, Email and Index. Without them, "
                "invoices can come out wrong (no index charged, or copies at $0.00)."
                + (' The Email and Index prices are read from the columns headed so, or close to it ("E-mail", '
                   '"Index Price") when only one column is.' if any({"Email", "Index"} & set(what)
                                                                    for _, what in gaps) else "")] if gaps else []
        if cheaper:
            said.append(("It also has" if gaps else f'The rate sheet "{sheet.name}" has')
                        + " prices that are lower at a faster speed:\n\n" + "\n".join(f"•  {c}" for c in cheaper[:8])
                        + ("\n…" if len(cheaper) > 8 else "") + "\n\nA faster speed shouldn't cost less: please "
                        "check these against your own rates.")
        box.setText("\n\n".join(said) + "\n\nUse this rate sheet anyway?")
        use = box.addButton("Use it anyway", QMessageBox.AcceptRole)
        names = [sh.name for sh in self.s.sheets()[0]]
        back = box.addButton(f"Go back to {previous}", QMessageBox.RejectRole) \
            if previous and previous != sheet.name and previous in names else None
        fix = box.addButton("Open the sheet to fill it in" if gaps else "Open the sheet to fix it",
                            QMessageBox.ActionRole)
        fix.setToolTip(f'Opens "{sheet.path.name}" in Excel (or the program that opens CSV files on this '
                       "computer).\nEach time you save it, it is read again and checked.")
        later = box.addButton("Not now", QMessageBox.RejectRole) if back is None else None
        box.setDefaultButton(back or fix)
        box.setEscapeButton(back or later)
        box.exec()
        if box.clickedButton() is use:
            self.s.rate_sheets_ok = [m for m in self.s.rate_sheets_ok if m != mark][-(RATE_SHEETS_OK_KEPT - 1):] + [mark]
            self._save_settings()
            if self._sheet_watched is not None and self._sheet_watched[0] == sheet.path:
                self._unwatch_sheet()  # (opened to be fixed, then kept as it is: no note at each save)
        elif back is not None and box.clickedButton() is back:
            self.s.rate_sheet = previous
            self._fill_sheet_box()
            if priced:
                self._sheet_in_use()
            self._check_rate_sheet()  # (that one may be incomplete too)
            return False
        elif box.clickedButton() is fix:
            self._watch_sheet(sheet.path)
            open_path(sheet.path)
        return True

    def _check_old_days(self) -> None:
        """Asks once whether the rate sheet in use, which still has the Days of the template before 2.8.0
        (Expedite 7, Regular 21: calendar days then; rates.old_days), should say 5 and 15: turnaround days are
        business days now, and 7 and 21 of them come out about 9 and 30 calendar days. Change writes the numbers
        into the sheet (rates.set_new_days) and prices every job again; Keep is remembered for that sheet
        (Settings.sheet_days_kept); Not now (and Esc, and the box's X) asks again at the next start. Called where
        _check_rate_sheet is, after it, and once a sheet opened to be filled in has all its prices (_sheet_saved):
        not while it misses some (one question at a time, and the sheet waits for its prices first; a price lower
        at a faster speed doesn't hold it, as the user may mean it). Not asked while
        Generate all runs, nor with business days off (Settings -> Dates: 7 and 21 are calendar days again then).
        Not modal to the app (open, not exec): nothing waits for the answer."""
        from ..rates import OLD_DAYS, cheaper_faster, missing_prices, old_days
        from ..settings import RATE_SHEETS_OK_KEPT
        sheet = self.s.sheet()
        watched = (self._sheet_watched is not None and self._sheet_watched[0] == sheet.path
                   and (bool(missing_prices(sheet)) or not cheaper_faster(sheet)))
        if (self._days_box is not None or self.filling or watched or not self.s.turnaround_business_days
                or not old_days(sheet) or sheet.name in self.s.sheet_days_kept or sheet.name in self._days_later):
            return
        # each speed once, Expedite first (a sheet can have two rows of one speed: Expedite and Rush)
        found = {k: sp.name for k in OLD_DAYS for sp in reversed(sheet.speeds) if sp.key == k and sp.days is not None}

        def numbers(i: int) -> str:
            return " and ".join(str(OLD_DAYS[k][i]) for k in found)
        says = " and ".join(f"{name} {OLD_DAYS[k][0]} days" for k, name in found.items())
        about = " and ".join(f"{OLD_DAYS[k][0]} business days come out about {OLD_DAYS[k][2]} calendar days"
                             for k in found)
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Question)
        box.setWindowTitle("Your rate sheet's turnaround days")
        box.setText(f'The rate sheet "{sheet.name}" gives {says}. Before version 2.8.0 those were calendar days. '
                    f"The app now counts business days (Settings → Dates), so {about}, and delivery dates come "
                    f"later than before.\n\nChange the sheet to {numbers(1)} business days, about the same as "
                    "before? Only its Days column changes.")
        change = box.addButton(f"Change to {numbers(1)}", QMessageBox.AcceptRole)
        keep = box.addButton(f"Keep {numbers(0)}", QMessageBox.ActionRole)
        keep.setToolTip("Not asked about this sheet again")
        later = box.addButton("Not now", QMessageBox.RejectRole)
        later.setToolTip("Asked again the next time the app starts")
        box.setDefaultButton(change)
        box.setEscapeButton(later)

        def answered(_result) -> None:
            self._days_box = None
            box.deleteLater()
            if self.s.sheet().path != sheet.path or not self.s.turnaround_business_days:
                return  # (another sheet in use since, or calendar days: asked again when it is so again)
            if box.clickedButton() is change:
                self._set_new_days(sheet)
            elif box.clickedButton() is keep:
                kept = [n for n in self.s.sheet_days_kept if n != sheet.name]
                self.s.sheet_days_kept = kept[-(RATE_SHEETS_OK_KEPT - 1):] + [sheet.name]
                self._save_settings()
            else:
                self._days_later.add(sheet.name)

        box.finished.connect(answered)
        self._days_box = box
        box.open()

    def _set_new_days(self, sheet) -> None:
        """Change on the question about the old Days (_check_old_days): the sheet says 5 and 15, it is read again
        and every job priced again. An incomplete sheet used anyway (Settings.rate_sheets_ok), or warned about
        this session (_sheets_warned), stays so as changed: only its Days changed, not its prices. A sheet that
        can't be written (open in Excel), or still has old Days after it (edited meanwhile), is said so, and asked
        about again with ⟳ or the next start."""
        from ..rates import OLD_DAYS, old_days, set_new_days, sheet_fingerprint
        before = sheet_fingerprint(sheet)
        try:
            set_new_days(sheet.path)
        except (OSError, ValueError) as e:
            logfile.error("could not change a rate sheet's Days", e)
            return self._days_not_changed(sheet, "it may be open in Excel. Close it there")
        self._read_sheets_again()
        now = next((sh for sh in self.s.sheets()[0] if sh.path == sheet.path), None)
        if now is None or old_days(now):  # (no longer read as a sheet, or edited in between)
            return self._days_not_changed(sheet, "its Days column may have been changed meanwhile. Look at it")
        mark = sheet_fingerprint(now)
        if before in self._sheets_warned:
            self._sheets_warned.add(mark)
        if before in self.s.rate_sheets_ok:
            self.s.rate_sheets_ok = [mark if m == before else m for m in self.s.rate_sheets_ok]
            self._save_settings()
        says = ", ".join(f"{sp.name} {sp.days}" for k in OLD_DAYS for sp in now.speeds
                         if sp.key == k and sp.days is not None)
        self._toast(f'"{now.name}" says {says} business days now.')

    def _days_not_changed(self, sheet, why: str) -> None:
        """Says (without waiting for an answer) that the old Days of a rate sheet weren't changed, and why."""
        box = QMessageBox(QMessageBox.Warning, "The rate sheet wasn't changed",
                          f'"{sheet.path.name}" couldn\'t be changed: {why}, then press ⟳ next to the Rate sheet '
                          "dropdown to be asked again.", QMessageBox.Ok, self)
        box.setAttribute(Qt.WA_DeleteOnClose)
        box.open()

    def _watch_sheet(self, path: Path) -> None:
        """The rate sheet at `path` was opened to be filled in: each time it is saved, it is read again and
        checked (_sheet_saved). Its folder is watched too: Excel saves a new file in place of the old one, which
        a watch on the file alone loses."""
        from PySide6.QtCore import QFileSystemWatcher
        if self._sheet_watch is None:
            self._sheet_watch = QFileSystemWatcher(self)
            self._sheet_watch.fileChanged.connect(lambda _p: self._sheet_timer.start())
            self._sheet_watch.directoryChanged.connect(lambda _p: self._sheet_timer.start())
        self._unwatch_sheet()
        self._sheet_watched = (path, _file_mark(path))
        self._sheet_watch.addPaths([str(path), str(path.parent)])

    def _batch_ended(self) -> None:
        """Generate all (or its preview) ended: the rate sheet check held back while it ran (_check_rate_sheet: a
        sheet picked or read again meanwhile) is made soon (_held_sheet_check)."""
        if self._sheet_check_later is not None:
            QTimer.singleShot(0, self, self._held_sheet_check)

    def _held_sheet_check(self) -> None:
        """The rate sheet check held back while Generate all ran, then its old Days question (both wait for the
        batch). Not under the batch's own windows (Batch finished, The math, the preview's Save or Go back): again
        in half a second while one is open; a new batch checks when it ends. Go back goes to the sheet in use
        before the batch's picks."""
        if self._sheet_check_later is None or self.filling:
            return
        if QApplication.activeModalWidget() is not None:
            QTimer.singleShot(500, self, self._held_sheet_check)
            return
        previous, self._sheet_check_later = self._sheet_check_later, None
        self._check_rate_sheet(previous, again=True)
        self._check_old_days()

    def _unwatch_sheet(self) -> None:
        """No rate sheet is watched any more (_watch_sheet)."""
        self._sheet_watched = None
        self._sheet_timer.stop()
        if self._sheet_watch is not None:
            paths = self._sheet_watch.files() + self._sheet_watch.directories()
            if paths:
                self._sheet_watch.removePaths(paths)

    def _sheet_saved(self) -> None:
        """The rate sheet opened to be filled in, or its folder, changed (a second ago: SHEET_SAVED_MS). When the
        file itself was saved and can be read: the sheets are read again and every job priced again, as ⟳ does,
        and the note at the bottom says whether the sheet has all its prices now or what it still misses (no box:
        the user may save as they go), and which prices are still lower at a faster speed (rates.cheaper_faster).
        Watched until it has them all and none is lower, until Use it anyway keeps it as it is (_check_rate_sheet), or
        until it isn't the sheet in use any more; the old Days question waits only for missing prices. Waits
        while a question is on screen or Generate all runs: what they show isn't priced again under them."""
        from ..rates import cheaper_faster, load_sheet, missing_prices, sheet_fingerprint
        if self._sheet_watched is None:
            return
        if QApplication.activeModalWidget() is not None or self.filling:
            self._sheet_timer.start()  # (again in a second)
            return
        path, mark = self._sheet_watched
        if path.exists() and all(Path(p) != path for p in self._sheet_watch.files()):
            self._sheet_watch.addPath(str(path))  # (Excel's new file in place of the old one)
        now = _file_mark(path)
        if now is None or now == mark:
            return  # (gone for a moment while Excel saves, or another file in the folder: Excel's "~$" one)
        if self.s.sheet().path != path:
            self._unwatch_sheet()  # (another sheet was picked meanwhile)
            return
        try:
            load_sheet(path)
        except (OSError, ValueError, UnicodeDecodeError):
            return  # (half written: read when it is saved again)
        self._sheet_watched = (path, now)
        self._read_sheets_again()
        sheet = self.s.sheet()
        if sheet.path != path:  # (no longer read as a sheet: _fill_sheet_box said why)
            self._unwatch_sheet()
            return
        gaps, cheaper = missing_prices(sheet), cheaper_faster(sheet)
        self._sheets_warned.add(sheet_fingerprint(sheet))  # (told below: no box about it again this session)
        if gaps or cheaper:
            still = ["Still missing: " + "; ".join(f"{name}: {_or(what)}" for name, what in gaps)] if gaps else []
            still += ["Still lower at a faster speed: " + "; ".join(cheaper)] if cheaper else []
            self._toast(f'Read "{sheet.name}" again. ' + ". ".join(still))
        else:
            self._unwatch_sheet()
            self._toast(f'Read "{sheet.name}" again: its prices are all there and in order now.')
        if not gaps:
            self._check_old_days()  # (its Days, waiting for the prices)

    def _remerge_others(self):
        """New rates or settings also apply to the batch's other jobs."""
        for job in self.jobs:
            if job is not self.cur and job.docs:
                remerge(job, self.s)

    def _delivery_changed(self, _idx):
        """The user picked a speed: it counts as their choice (it also answers an e-mail's other speed), and the
        rate and delivery date follow it."""
        name = self.delivery.currentData() or ""
        self.case.fields["delivery"] = FieldState(name, SRC_USER, 1.0, [name])
        self._apply_speed()

    def _apply_speed(self):
        """Re-derives rate and delivery date from the job's speed (unless typed by the user), and shows them
        and what follows (the prices, the job list's ⚠ of a speed to choose). The case's speed is what counts:
        the box only shows it (it was once copied from the box, and so a stale choice of another job's, or the
        box's first item, became the job's speed)."""
        if self.case.get("delivery") == "Other" and self.rows["rate"].state.source != SRC_USER:
            self.case.fields["rate"] = FieldState()  # (no sheet rate for a speed of the user's own)
        refresh_rate(self.case, self.s)
        if self.s.fill_delivery_date:
            refresh_delivery_date(self.case, self.s)
        for k in ("rate", "delivery_date"):
            if self.rows[k].state.source != SRC_USER:
                self.rows[k].set_state(self.case.fields[k])
        self._show_date_notes()
        self._show_rate_info()
        self._show_speed_state()
        self._update_status()  # (a job's ⚠ of a speed to choose comes or goes, and so does a held invoice)

    def _apply_speed_rule(self):
        """What decides the agreement form's speed changed (the speeds ticked, the rate sheet, the settings, ↺):
        every job's speed is the Settings rule's again unless the user chose one still offered, its rate and
        delivery date follow (merge.apply_speed_rule), and the window shows the current job's."""
        for job in self.jobs:
            apply_speed_rule(job.case, self.s)
        self._fill_speeds()
        self._apply_speed()

    # --------------------------------------------------------- the dates (Settings -> Dates, workdays.py)
    def _holidays_of(self, year: int, month: int) -> dict:
        """The court holidays of a month, {date: name}: what the date fields' calendars mark ({} for a month of a
        year a date can't have: the calendar asks for the months beside the one shown, 12/9999's too)."""
        if not MINYEAR <= year <= MAXYEAR:
            return {}
        first = date(year, month, 1)
        last = first.replace(day=monthrange(year, month)[1])
        return dict(self.s.holiday_calendar().between(first, last))

    def _quick_tips(self) -> None:
        """The quick buttons' tooltips, with the day each gives as of now ("Estimated delivery 10/20/2026 - a
        week from the form's date")."""
        since = "the form's date" if self.s.turnaround_from == "form" else "today"
        words = {"Tomorrow": "a day", "1 week": "a week", "1 month": "a month"}
        for b, label, days, months in self.quick_buttons:
            when = quick_day(self.s, days, months)
            how = f": {since}" if label == "Today" else f" - {words.get(label, label)} from {since}"
            b.setToolTip(f"Estimated delivery {when}{how}\n(a weekend or a court holiday becomes the next business "
                         "day)")

    def _show_date_notes(self) -> None:
        """Under the two dates, what moved them: the form dated the next business day ("Dated Tue 10/13/2026, the
        next business day: today is a Sunday and Mon 10/12 is Columbus Day."), the delivery date counted in
        business days ("5 business days from Tue 10/13/2026") or moved off a weekend; or, for a day typed, picked
        or read from a document, a warning when it is a weekend or a court holiday ("⚠ 11/11/2026 is Veterans Day,
        a court holiday."). A note with nothing to say takes no room."""
        if not hasattr(self, "date_notes"):  # (the window is still being built)
            return
        f = self.case.fields
        dating = document_day(self.s)
        for key, note in self.date_notes.items():
            st, text = f[key], ""
            if key == "agreement_date" and st.source == SRC_DEFAULT and st.value == us_date(dating.day):
                text = dating.note()
            elif key == "delivery_date" and st.source == SRC_DERIVED:
                n = self.s.days_for(self.case.get("delivery"))
                if n is not None:
                    d, why = delivery_day(self.s, n, dating.day)
                    text = why if us_date(d) == st.value else ""
                    sp = self.s.sheet().find(self.case.get("delivery"))
                    if text and sp is not None and sp.days is not None:
                        # (where the number comes from: a sheet filled in before 2.8.0 still says 7 and 21, which
                        # were calendar days then; Settings -> Dates' 5 and 15 are used only without a Days column)
                        text += " (the rate sheet's Days column)"
            elif st.value:
                text = date_warning(st.value, self.s)
            if note.text() != text:
                note.setText(text)
            look = "dateWarn" if text.startswith("⚠") else "muted"  # (a day to look at stands out)
            if note.objectName() != look:
                note.setObjectName(look)
                repolish(note)
            self.date_form.setRowVisible(note, bool(text))

    def _redate(self, force: bool = False) -> None:
        """The day a form made now is dated (workdays.document_day) changed: midnight, the evening hour of
        Settings -> Dates, or its settings (force). Every job's date of agreement filled in by default and
        delivery date not typed become that day's (merge.refresh_dates), and the window shows them. Runs each
        minute (_day_timer), when the window is active again and after Settings."""
        if not hasattr(self, "date_notes"):
            return
        day = document_day(self.s).day
        # (and the real day it is: with the turnaround counted from the day you make the form, Settings.turnaround_from
        # "today", the delivery date moves at midnight even when the form's date doesn't, after the evening hour or
        # over a weekend)
        dated = (day, day.made_on)
        if dated != self._dated or force:
            self._dated = dated
            for job in self.jobs:
                refresh_dates(job.case, self.s, day)
            for key in self.date_notes:
                if self.rows[key].state.source != SRC_USER:
                    self.rows[key].set_state(self.case.fields[key])
            self._quick_tips()
            self._card_summaries()  # (Details folded: its line names the delivery date)
        self._show_date_notes()

    def _open_log_folder(self):
        open_path(logfile.log_dir())

    def _copy_diagnostics(self):
        """Copies the app and system details for a problem report (no case details) to the clipboard."""
        QApplication.clipboard().setText(logfile.diagnostics(self.s))
        self._toast("Copied: paste it into your message. It has no case details.")

    def _open_sheets_folder(self):
        from ..rates import sheets_dir
        open_path(sheets_dir(self.s.rate_sheets_dir))

    def _reload_sheets(self):
        """⟳: reads the rate sheet files again (after they were edited) and prices every job again. The sheet in
        use, still incomplete or newly so, is warned about again, unless the user said to use it as it is
        (_check_rate_sheet)."""
        self._read_sheets_again()
        self._toast(f"Loaded {len(self.s.sheets()[0])} rate sheet(s).")
        self._check_rate_sheet(again=True)  # (edited in Excel: still incomplete, or newly so)
        self._check_old_days()

    def _read_sheets_again(self) -> None:
        """The rate sheet files are read again (⟳, or the one opened to be filled in was saved: _sheet_saved), and
        every job priced again from the sheet in use."""
        self._drop_undo()  # (the jobs New job cleared were priced from the sheets as they were)
        self.s.reload_rates()
        self._fill_sheet_box()
        self._fill_invoice_speeds()  # the sheet's speeds may have changed
        self._apply_speed_rule()
        self._remerge_others()

    def _set_opt(self, name, value):
        """Sets a setting and saves it: the outputs ticked, a standing choice under Options ▾ (gui/panels.py), the
        speeds offered, the job list's sort and Day column, the zoom, the list of recent files cleared. A file that
        can't be written is logged, not shown (the window must follow the option all the same; the settings are
        saved again when it closes)."""
        setattr(self.s, name, value)
        self._save_settings()

    # ------------------------------------------------------- status
    def _set_status(self, text: str, state: str, mood: str | None = None):
        """Sets the status at the top. state: "" (plain text, as IDLE_STATUS is), or "busy", "ok" or "warn" (a
        pill of its colour). When the yin-yang is on (Settings → Options, Settings.show_yin), state also picks its
        picture unless mood names one: it swirls while busy ("working") and while the app waits (plain: "idle"); ok
        shows "done" once the swirl's turn is over, warn "stumped" at once (DropZone.set_mood)."""
        self.status.setText(text)
        self.status.setProperty("state", state)
        repolish(self.status)
        self._fit_corner()
        if self.s.show_yin:
            self.drop.set_mood(mood or {"busy": "working", "ok": "done", "warn": "stumped"}.get(state, "idle"))
        else:
            self.drop.set_mood(None)
        self._update_header_buttons()  # (the status changes whenever work starts or ends, or jobs come or go)

    def _update_status(self):
        """Brings the Outputs box and the cards that follow from the ticks (the Invoice card, Who ordered what, Who
        pays what, Who transcribed what: _refresh_outputs, which also brings the Excerpts window along), the
        question about rows that may be one firm, the status pill and the Before you generate box (_job_status),
        the list of jobs on the left and the folded cards' lines up to date."""
        self._all_counts = None
        self._refresh_outputs()  # (it counts what Generate all makes: _all_counts)
        self._show_firm_question()
        self._job_status()
        self._refresh_job_labels(self._all_counts)
        self._card_summaries()

    def _show_firm_question(self) -> None:
        """The Attorneys card's question when two rows of the job, or of its case's other days, may be one firm or
        attorney (batch.firm_questions: "Smith Law" and "Smith Law Group"): Same firm or Not the same answers it
        (_answer_firm); hidden when there is none."""
        group = self._invoice_group() if self.cur in self.jobs else [self.cur]
        qs = firm_questions(group)
        self._firm_q = (group, qs[0]) if qs else None
        if qs:
            a, b = qs[0]
            more = f"  ({len(qs)} to answer)" if len(qs) > 1 else ""
            self.firm_ask.setText(f"⚠ Are these one firm?{more}\n•  {_entry_text(a)}\n•  {_entry_text(b)}")
        self.firm_ask_box.setVisible(bool(qs))

    def _answer_firm(self, same: bool) -> None:
        """Same firm / Not the same, for the question shown (see _show_firm_question): kept for later documents
        while the card's Remember is ticked, else until the app is closed."""
        if self._firm_q is None or self._batch_running():
            return
        _, (a, b) = self._firm_q
        self._commit_edit()  # (a cell still being typed in: the table is shown again from the case)
        self._record_firm_answer(a, b, same, self.firm_remember.isChecked())
        self._show_case()
        self._update_status()

    def _use_answers(self) -> None:
        """The answers about rows that may be one firm, as same_entry reads them: those kept in Settings and those
        given for now only (extract_regex.use_firm_answers)."""
        use_firm_answers(self.s.firm_answers + self._session_answers)

    def _record_firm_answer(self, a: Attorney, b: Attorney, same: bool, remember: bool = True) -> None:
        """The user's answer about rows a and b: kept (remember: Settings.firm_answers, saved, with the names as
        written, so later documents naming them are read so too) or for now only (until the app is closed).
        When they are one, they become one row on each day of every job loaded (batch.join_entries, a case at a
        time); No. of copies follows."""
        pair = {a.key(), b.key()}
        row = [a.key(), b.key(), same, _entry_text(a), _entry_text(b)]
        self._session_answers = [r for r in self._session_answers if {r[0], r[1]} != pair]
        if remember:
            self.s.firm_answers = [r for r in self.s.firm_answers if {r[0], r[1]} != pair] + [row]
            self._save_settings()
        else:
            self._session_answers.append(row)
        self._use_answers()
        if same:  # on every job loaded, not only the case asked about: it won't be asked about again
            for case in _cases(self.jobs):
                join_entries(case, a, b)
        for j in self.jobs:
            j.refresh_copies(self.s)

    def _ask_firms(self, groups: list[list[Job]]) -> bool:
        """Before anything is made: rows of these days (each list the days of a case) that may be one firm or
        attorney and haven't been answered are asked about, one pair at a time (Same firm, Not the same). The
        Remember box starts as the card's and carries from one question to the next, and back to the card. False
        when Go back was clicked (nothing is made), or Generate all started meanwhile."""
        asked = False
        ticked = self.firm_remember.isChecked()
        for group in groups:
            # the questions are worked out again after each answer: an answer joins rows, and a read that ends
            # while the box is open may merge the job again (new rows); an answered pair isn't asked again
            while qs := firm_questions([j for j in group if j in self.jobs]):
                a, b = qs[0]
                box = QMessageBox(self)
                box.setIcon(QMessageBox.Question)
                box.setWindowTitle("One firm or two?")
                case = ", ".join(x for x in (group[0].title(), group[0].case.get("dates")) if x)
                box.setText(f"{case}: are these one firm (or one attorney)?\n\n•  {_entry_text(a)}\n"
                            f"•  {_entry_text(b)}\n\nOne firm gets one invoice and one minute agreement.")
                remember = QCheckBox("Remember my answer for later documents")
                remember.setChecked(ticked)
                remember.setToolTip("Unticked: the answer holds until you close the app. Settings → Invoice →\n"
                                    "Firms you answered lists your answers (those for now too), to forget one.")
                box.setCheckBox(remember)
                same = box.addButton("Same firm", QMessageBox.YesRole)
                apart = box.addButton("Not the same", QMessageBox.NoRole)
                back = box.addButton("Go back", QMessageBox.RejectRole)
                box.setEscapeButton(back)
                box.exec()
                clicked = box.clickedButton()
                ticked = remember.isChecked()
                self.firm_remember.setChecked(ticked)
                if clicked not in (same, apart) or self._batch_running():
                    return False
                self._record_firm_answer(a, b, clicked is same, ticked)
                asked = True
        if asked:
            self._show_case()
            self._update_status()
        return True

    def _follow_excerpts(self) -> None:
        """The Excerpts window, when open, shows the current job's days as they are now (it is called from
        _refresh_outputs, which every change reaches: a speed, the parties, a job ticked, the delivery). When
        the job has no days to show (a new job, one without a transcript, a job removed), the window closes
        (ExcerptsWindow.clear): else it would keep showing, and changing, days no longer among the jobs."""
        w = self.excerpts
        if w is None or not w.isVisible():
            return
        if self.cur not in self.jobs or not self.cur.invoice_days():
            w.clear()
            return
        w.load(self._excerpt_group())

    def _job_status(self):
        """The status pill for the job on screen: reading (documents still being read, whatever the job on
        screen says: a single file read while a folder is still coming in), asking the AI, what needs a look
        ("⚠  1 missing, 2 to check: Index Number is missing ›": required fields blank, the job's own checks as the
        job list's ⚠ has them and the warnings in the cards, guessed fields: _problems, the first named), or
        ready. Its tooltip lists them; a click goes to the first (_go_to_problem)."""
        busy = self.ai_pending > 0
        self.busy.setVisible(busy or self.work > 0)
        self.status.set_clickable(False)
        if self.work > 0:
            if self.status.property("state") != "busy":  # ("Reading 2 of 5: ..." is kept as it is)
                self._set_status("Reading…", "busy")
            self._show_todo([])
            return
        if not self.cur.docs:
            self._set_status(IDLE_STATUS, "")
            self.status.setToolTip("")
            self._show_todo([])
            return
        self._sync_from_ui(commit=False)  # (a name being typed in the attorney table stays open)
        problems = self._problems()
        self._show_todo(problems)
        missing = [FIELD_LABELS[k] for k in self.case.missing_required()]
        checks = [p.text for p in problems if p.kind == "check"]
        review = [FIELD_LABELS.get(k, k) for k, r in self.rows.items() if r.edit.property("review")]
        if busy:
            self._set_status(f"Asking {self.s.ollama_model}…  (you can keep editing)", "busy")
        elif problems:
            parts = []
            if missing:
                parts.append(f"{len(missing)} missing")
            if checks:
                parts.append(f"{len(checks)} to check")
            if review:
                parts.append(f"{len(review)} to review")
            lines = ([f"Missing: {', '.join(missing)}"] if missing else []) + [f"⚠ {c}" for c in checks] + (
                [f"Guesses to look at (highlighted): {', '.join(review)}"] if review else [])
            lead = "⚠  " if missing or checks else "Ready to fill  ·  "
            # (the first one named, and an arrow: the pill is to be clicked, the user, 2026-10-22)
            text = f"{lead}{', '.join(parts)}: {problems[0].short} ›"
            # the yin-yang only cracks when something required is missing
            self._set_status(text, "warn", "stumped" if missing else "done")
            self.status.setToolTip("\n".join(lines) + "\n\nClick to go to the first one.")
            self.status.set_clickable(True)
        else:
            self._set_status("✓  Ready to fill", "ok")
            self.status.setToolTip("")

    def _job_checks(self) -> list[str]:
        """What the job on screen needs looked at besides its blank required fields: the job list's ⚠ lines
        (Job.issues: nobody ticked, a speed to choose, Whose pages…, Excerpts… or Parties, index numbers that
        don't match, no pages to bill...), without the "… is missing" ones the pill counts already."""
        missing = {f"{FIELD_LABELS[k]} is missing" for k in self.case.missing_required()}
        return [i for i in self.cur.issues(self._outputs()) if i not in missing]

    def _problems(self) -> list[Problem]:
        """Everything the job on screen needs looked at before Generate, each with where it is fixed, in this
        order (not Generate's, which asks about the index numbers first): the required fields blank; nobody
        ticked, two rows that may be one firm, a speed to choose, index numbers that don't match, Whose pages… or
        Excerpts… or Parties, another day of the case nobody is ticked on (no invoice for the case: it goes to that
        day), no pages to bill, no transcript for the run sheet (the job's checks, _job_checks, with the firm
        question and the case's other days, which those don't list); the page count to check and a date on a
        weekend or a court holiday (the warnings under those fields); any other check (nowhere to go); then the
        guessed fields, in one line. The status pill names the first and goes to it; the Before you generate box
        shows them all (_show_todo)."""
        job, case = self.cur, self.case
        if job.is_empty():
            return []
        from ..invoice import NO_INVOICE
        from ..runsheet import NO_RUNSHEET
        outputs = self._outputs()

        def field(key: str) -> Callable[[], None]:
            return lambda: self._go_to(self.rows[key].edit, "case" if key in CASE_KEYS else "details")

        def attorneys() -> None:
            if self.att.rowCount():
                self.att.setCurrentCell(0, 1)
            self._go_to(self.att, "attorneys")

        hold = job.invoice_hold() if "invoice" in outputs else ""

        def held() -> None:
            kind = job.hold_kind()  # (the one the hold names: invoice_hold)
            self._go_to(self.inv_whose if kind == "ownership" else
                        self.inv_parties if kind == "unnamed" else self.inv_who, "invoice")

        def runsheet() -> None:
            self._flash(self.output_boxes["runsheet"])

        ask = job.speed_question()
        found: list[tuple[int, Problem]] = []  # (its place in the order, the problem)
        for text in self._job_checks():
            if text == "no attorney is ticked":
                rank, go = 0, attorneys
            elif ask and text.startswith(f"{ask}: choose the speed"):
                rank, go = 2, lambda: self._go_to(self.delivery, "details")
            elif text.startswith("the documents' index numbers don't match"):
                rank, go = 3, lambda: self._flash(self.index_note)
            elif hold and text == hold:
                rank, go = 4, held
            elif text == NO_INVOICE:
                rank, go = 5, field("est_pages")
            elif text == NO_RUNSHEET and "runsheet" in self.output_boxes:
                rank, go = 6, runsheet
            else:
                rank, go = 9, None
            found.append((rank, Problem("check", text, _short(text), go)))
        if self._firm_q is not None:
            _, (a, b) = self._firm_q
            text = f"Are these one firm: {_entry_text(a)} and {_entry_text(b)}?"
            found.append((1, Problem("check", text, "Are these one firm?",
                                     lambda: self._go_to(self.firm_same, "attorneys"))))
        if "invoice" in outputs and self.cur in self.jobs:  # another day of the case nobody is ticked on
            from ..batch import group_problem
            group = self._invoice_group()
            # (not when nobody is ticked on this day: its own check says so; nor for a day invoiced already:
            # Generate all doesn't bill it again)
            stuck = group_problem(group) if job.ticked_keys() and not job.invoiced else ""
            empty = next((j for j in group if not j.ticked_keys()), None)
            if stuck and empty is not None:
                def other_day(day: Job = empty) -> None:
                    self._select_job(day)
                    attorneys()
                found.append((4, Problem("check", f"No invoice for the case yet: {stuck}",
                                         "No invoice for the case yet", other_day)))
        if self.pages_warn.text().strip():
            found.append((7, Problem("check", self.pages_warn.text().strip(), "Check the page count",
                                     field("est_pages"))))
        for key, note in self.date_notes.items():
            if note.text().startswith("⚠"):
                found.append((8, Problem("check", note.text(), _short(note.text()), field(key))))
        out = [Problem("missing", f"{FIELD_LABELS[k]} is missing", f"{FIELD_LABELS[k]} is missing", field(k))
               for k in case.missing_required()]
        out += [p for _, p in sorted(found, key=lambda x: x[0])]  # (stable: the job's own order within a rank)
        review = [r for r in self.rows.values() if r.edit.property("review")]
        if review:
            names = [FIELD_LABELS.get(r.key, r.key) for r in review]
            text = (f"{names[0]} is a guess: check it" if len(names) == 1 else
                    f"{len(names)} guesses to look at: {', '.join(names)}")
            out.append(Problem("review", text, names[0], field(review[0].key)))
        return out

    def _go_to_problem(self) -> None:
        """A click on the status pill: to the first thing that needs a look (_problems, in its order). Its card is
        unfolded and scrolled to; a field gets the cursor, and it, or the note or button to look at, lights up for a
        moment (_flash)."""
        if self.cur.is_empty() or self.work > 0:
            return
        self._commit_att_details()  # (the pill takes no focus: an attorney detail being typed is kept first)
        first = next((p for p in self._problems() if p.go is not None), None)
        if first is not None:
            first.go()

    def _show_todo(self, problems: list[Problem]) -> None:
        """The Before you generate box at the top of the form: each thing to look at (_problems) as a link to where
        it is fixed (a click goes there and lights it up, as the status pill's does for the first; one with nowhere
        to go is plain text). Hidden when there is nothing, while documents are read and before any is (the user,
        2026-10-22, after GOV.UK's error summary: one list at the top, each item going to its place)."""
        self._todo = problems
        self.todo.setVisible(bool(problems))
        if not problems:
            return
        lines = []
        for i, p in enumerate(problems):
            text = html.escape(" ".join(p.text.split()).lstrip("⚠ "))
            lines.append(f"•&nbsp;&nbsp;<a href='{i}'>{text}</a>" if p.go is not None else f"•&nbsp;&nbsp;{text}")
        self.todo_text.setText("<b>⚠ Before you generate</b><br>" + "<br>".join(lines))

    def _todo_clicked(self, href: str) -> None:
        """A link of the Before you generate box: to that problem's place (_show_todo)."""
        try:
            p = self._todo[int(href)]
        except (ValueError, IndexError):
            return
        if p.go is not None and not self.cur.is_empty() and self.work == 0:
            self._commit_att_details()  # (as the status pill's click: an attorney detail being typed is kept)
            p.go()

    def _go_to(self, w: QWidget, card_key: str = "") -> None:
        """Unfolds the card (by key) holding w, scrolls the form to it and gives it the cursor."""
        fr = self.cards.get(card_key)
        if fr is not None and fr.body.isHidden():
            self._fold_card(fr, False)
        self.form_scroll.ensureWidgetVisible(w, 0, z(60))
        w.setFocus(Qt.OtherFocusReason)
        self._flash(w)

    def _flash(self, w: QWidget, ms: int = 1500) -> None:
        """Lights a widget up for a moment (the style sheet's [flash="true"]), to show where to look."""
        if w.isHidden():
            return
        if self.form_scroll.widget().isAncestorOf(w):
            self.form_scroll.ensureWidgetVisible(w, 0, z(60))
        w.setProperty("flash", True)
        repolish(w)

        def off() -> None:
            w.setProperty("flash", False)
            repolish(w)

        QTimer.singleShot(ms, w, off)

    def _idle(self):
        """A piece of background work is done: the progress bar goes back to 'busy' style, and stays only
        while the AI or another read is still going."""
        self.busy.setRange(0, 0)
        self.busy.setVisible(self.ai_pending > 0 or self.work > 0)

    def _saved_box(self, title: str, text: str, folders: list, details: str = "", warn: bool = False,
                   files: list | None = None):
        """The 'files saved' message, with a button that opens the folder(s) and, when `files` are given,
        one that prints them (see preview.print_files)."""
        box = QMessageBox(self)
        box.setWindowTitle(title)
        box.setIcon(QMessageBox.Warning if warn else QMessageBox.Information)
        box.setText(text)
        if details:
            box.setDetailedText(details)
        print_btn = box.addButton("Print…", QMessageBox.ActionRole) if files else None
        open_folder = box.addButton("Open folder", QMessageBox.ActionRole) if folders else None
        box.setDefaultButton(box.addButton(QMessageBox.Ok))
        box.exec()
        if open_folder is not None and box.clickedButton() == open_folder:
            for f in folders[:3]:
                open_path(f)
        elif print_btn is not None and box.clickedButton() == print_btn:
            from .preview import print_files
            print_files(self, files)

    def _toast(self, msg: str):
        """A short note at the bottom of the window that goes away after 8 seconds."""
        self.statusBar().showMessage(msg, 8000)

    # --------------------------------------------------------- fill
    def _batch_running(self) -> bool:
        """True (with a note to the user) while Generate all is still making files."""
        if self.filling:
            self._toast("Still making the files of the batch - one moment.")
        return self.filling

    def _commit_edit(self) -> tuple[int, str, int] | None:
        """A cell of the attorney table still being typed in (Ctrl+Enter pressed in it, a merge, an answer about
        the firms) is taken as typed. The editor is found in the table itself, not by the keyboard focus: a
        background answer arrives while the focus may be anywhere. Returns (column, text, cursor position) when
        the user was typing in it (it had the focus), for _resume_edit, else None. A detail being typed under the
        table is written into it first (_commit_att_details)."""
        self._commit_att_details()
        if self.att.state() != QAbstractItemView.EditingState or getattr(self, "_committing", False):
            return None  # (committing fires itemChanged, whose handlers sync the editor again: once is enough)
        editor = QApplication.focusWidget()
        if editor is None or editor is self.att or not self.att.isAncestorOf(editor):
            viewport = self.att.viewport()
            # (the editor is the viewport's own child; one closed before is hidden until Qt deletes it)
            editor = next((w for w in viewport.findChildren(QWidget) if w.parent() == viewport
                           and w.isVisibleTo(viewport)), None)
        if editor is None:
            return None
        typing = None  # (the window's focus, kept while another program is in front; elsewhere = typing elsewhere)
        if isinstance(editor, QLineEdit) and self.focusWidget() is editor:
            typing = (self.att.currentIndex().column(), editor.text().strip(), editor.cursorPosition())
        self._committing = True
        try:
            self.att.commitData(editor)
            self.att.closeEditor(editor, QAbstractItemDelegate.NoHint)
        finally:
            self._committing = False
        return typing

    def _resume_edit(self, typing: tuple[int, str, int] | None) -> None:
        """Opens the cell editor again after a merge from the background (a document read, an AI answer) closed
        it, on the row that now holds the text typed (or starts with it: the merge may have added the firm's other
        names), with the cursor where it was: the next key the user types goes on with "Dana Smi" instead of
        replacing it. Nothing happens when that row can't be told."""
        if not typing or not typing[1]:
            return
        col, text, cursor = typing
        cells = [(r, it.text()) for r in range(self.att.rowCount()) if (it := self.att.item(r, col))]
        rows = [r for r, t in cells if t == text] or [r for r, t in cells if t.startswith(text)]
        if len(rows) != 1:
            return
        index = self.att.model().index(rows[0], col)
        self.att.setCurrentIndex(index)
        self.att.edit(index)
        viewport = self.att.viewport()
        editor = next((w for w in viewport.findChildren(QLineEdit) if w.parent() == viewport
                       and w.isVisibleTo(viewport)), None)
        if editor is not None:
            editor.setFocus()
            editor.deselect()
            editor.setCursorPosition(min(cursor, len(editor.text())))

    def fill(self):
        """Generate: makes the ticked outputs for the job on screen, on the UI thread (Ctrl+Enter).

        Asks first, as needed, in this order: whether documents whose index numbers don't match belong together
        (_ask_index_numbers); whether to go on without the invoice (no pages to bill: no transcript and none
        typed) or the run sheet (no transcript); whose pages to bill, or whether to go on without an invoice
        that is held (_without_unchecked_invoice); whether rows that may be one firm are one (_ask_firms); which
        speed, when an e-mail asks for another one (_ask_speeds); a Parties number that isn't the firms that
        ordered (_ask_parties); the speed of each firm of a split order that has none (_ask_split_speeds), then
        whether to go on without the invoice when the speeds hold it (more on the same pages than are billed for
        now); whether there is an index, when a page count typed below the transcript's changes it (_ask_index);
        about blank required fields, competing values and a case name only the records suggest, and who ordered
        when no attorney is ticked (ClarifyDialog), then the split-order speeds of the firms ticked there; which
        run sheet the takes go on, and to close it first when it is open in Excel (nothing is made while it
        can't be read or written); the index numbers again, for a document read meanwhile; and, when
        the case's own folder is there with files in it, whether to add to it (_case_folder). Then the preview
        (_preview, when Settings.preview_before_saving), and the files are made.
        Background work can finish while a question is on screen, so after each one the job is checked to still
        exist and is taken as it is now (an invoice held meanwhile is asked about again before the files are
        made). Go back at the preview leaves the case folder to be asked about next time.
        """
        if self._batch_running():
            return
        self._commit_edit()
        self._sync_from_ui()
        job = self.cur
        case = job.case
        outputs = self._outputs()
        if not outputs:
            QMessageBox.information(self, "Nothing to make", f"Tick what to make first: {_named(OUTPUTS)} "
                                    "(in the Outputs box).")
            return
        # documents whose index numbers don't match may be about different cases: asked before anything else
        if not self._ask_index_numbers([job]) or job not in self.jobs:
            return
        missing = [o for o in outputs if o not in job.makeable(outputs) and not job.unneeded(o)]
        if missing:
            left_out = " or ".join(OUTPUTS[o].lower() for o in missing)
            reasons = {"invoice": "An invoice needs pages to bill, and the Pages field says 0."
                       if job.transcript_pages() or job.pages_typed() else
                       "An invoice needs pages to bill, and this job has no transcript PDF: type the pages to bill "
                       "in Est. number of pages (Speed, pages and dates) to make one.",
                       "runsheet": "A run sheet needs a transcript PDF (the initials on its pages), and this job "
                                   "has none."}
            why = "\n".join(reasons[o] for o in missing if o in reasons)
            if QMessageBox.question(
                    self, f"No {left_out}", f"{why}\n\nMake the other outputs without the {left_out}?"
            ) != QMessageBox.Yes:
                return
            if job not in self.jobs or self._batch_running():
                return
            case = job.case  # an AI answer that came in meanwhile merged the job again
        outputs = job.makeable(outputs)  # (a run sheet one reporter's case doesn't need goes without a word)
        if not outputs:
            return
        outputs = self._without_unchecked_invoice(job, outputs)
        if not outputs:
            return
        # rows that may be one firm: asked before the speeds and the parties (one firm is one agreement, one
        # invoice; the run sheet doesn't name them). Each question is asked when an output chosen asks it
        # (courthouses.OutputSpec.asks).
        if courthouses.asks(outputs, "firms") and (not self._ask_firms([[job]]) or job not in self.jobs):
            return
        # the e-mail asks for another speed than the Settings rule's: which one the forms name
        if courthouses.asks(outputs, "speed") and (not self._ask_speeds([job]) or job not in self.jobs):
            return
        if courthouses.asks(outputs, "parties") and (not self._ask_parties([job]) or job not in self.jobs):
            return
        # the firms of a split order commit to a speed: the speeds not set yet (they change the prices)
        if courthouses.asks(outputs, "speeds") and (not self._ask_split_speeds([job]) or job not in self.jobs):
            return
        if "invoice" in outputs and job.invoice_hold():  # (more speeds answered on the same pages than are billed)
            outputs = self._without_unchecked_invoice(job, outputs)
            if not outputs or job not in self.jobs:
                return
        case = job.case  # (Whose pages... or a read that ended meanwhile may have merged the job again)
        # Ask about required fields that are blank, fields with competing values, and a case name only the records
        # suggest (only the case name when neither the agreement nor the MOFR is made)
        questions = []
        for key in REQUIRED_KEYS if courthouses.asks(outputs, "fields") else ["case_name"]:
            fs = case.fields[key]
            if not fs.value or fs.source == SRC_RECORDS or (
                    fs.source != SRC_USER and len([a for a in fs.alternatives if a != fs.value]) > 0
                    and fs.confidence < 0.8):
                questions.append((key, fs.value, fs.alternatives))
        real = [a for a in case.attorneys if not a.is_placeholder() and (a.name or a.firm)]
        # who ordered: only agreements and invoices are addressed to an attorney
        ask_att = real and not any(a.checked for a in case.attorneys) and courthouses.asks(outputs, "attorney")
        if questions or ask_att:
            asked = list(case.attorneys)
            dlg = ClarifyDialog(questions, asked if ask_att else None, self)
            if dlg.exec() != ClarifyDialog.Accepted or job not in self.jobs or self._batch_running():
                return
            # While the question was on screen an AI answer may have come in (the job was merged again) or
            # documents still being read may have put another job in the editor: the answers go to the job
            # they were asked about, as it is now.
            case = job.case
            for key, val in dlg.answers().items():
                case.fields[key] = FieldState(val, SRC_USER, 1.0, case.fields[key].alternatives)
            chosen = dlg.checked_attorneys()
            if chosen is not None:
                # the rows as asked about, never by their place: an AI answer may have put them in another order
                # (or filled in a firm, a new key: then by key, or as the same entry)
                from ..extract_regex import same_entry
                rows = [asked[i] for i in chosen]
                keys = {x.key() for x in rows if x.key()}
                for a in case.attorneys:
                    a.checked = any(a is x for x in rows) or a.key() in keys or any(same_entry(x, a) for x in rows)
                job.att_touched = True  # a later AI answer must not undo the choice
                job.refresh_copies(self.s)  # No. of copies: the parties just ticked
            if job is self.cur:
                self._show_case()
            # the firms just ticked may share pages: their speeds are asked now (no box when there's none to ask)
            if chosen is not None and courthouses.asks(outputs, "speeds") and (
                    not self._ask_split_speeds([job]) or job not in self.jobs):
                return
            if "invoice" in outputs and job.invoice_hold():  # (more speeds answered on the same pages)
                outputs = self._without_unchecked_invoice(job, outputs)
                if not outputs or job not in self.jobs:
                    return
        # a Pages number typed below the transcript's count that changes whether there is an index: which. Asked
        # once the firms who ordered and their speeds are known: whether an index can be charged is their speeds'
        if "invoice" in outputs and (not self._ask_index([[job]]) or job not in self.jobs):
            return

        sheet = None
        if "runsheet" in outputs:
            target = self._run_sheet_for(job)
            if target is None or job not in self.jobs or self._batch_running():
                return
            if not self._sheets_free([target], "Generate"):
                return
            sheet = job.runsheet_opts(self.s)
            sheet.target = target
        # the files are made from the job as it is now: an answer from the AI or a read that ended while a
        # question was on screen merged it again (and what is on screen with it)
        if "invoice" in outputs and job.invoice_hold():  # (an attorney was unticked meanwhile)
            outputs = self._without_unchecked_invoice(job, outputs)
            if not outputs or job not in self.jobs:
                return
        # (a document read while a question was on screen, with an index number of its own: asked about too)
        if not self._ask_index_numbers([job]) or job not in self.jobs:
            return
        # the case's own folder inside Save to (Settings.case_folders): asked about when it has files already
        folder_was = job.folder, job.folder_named  # (put back when nothing is saved: asked about next time)
        if not self._case_folder([job], outputs) or job not in self.jobs:
            return
        case = job.case
        out_dir = out_dir_for(job, self.s)
        self._note_trial_rules([job], fresh=True)  # (billed as an earlier invoice of the trial was, if any)
        opts, origin = job.invoice_sets(), job_origin(job)  # (the user's invoices, and other reporters' ticked)
        ordered = job.ordered_pages()  # each attorney's agreement: the pages it ordered, whoever wrote them
        speeds = job.ordered_speeds()  # and the speed(s) it ordered them at, where set
        docs = list(job.docs)
        previewed = False  # the preview showed the math
        # the day every file is dated (today, or the next business day: Settings -> Dates), the preview's too: one
        # made at 4:59 PM is saved as it was shown, though the evening hour came while it was open
        dating = document_day(self.s)
        if self.s.preview_before_saving:
            # what is saved is what was shown: the case as it is now, whatever comes in while the preview is up
            case = deepcopy(case)
            if not self._preview(case, outputs, opts, ordered, speeds, dating.day):
                job.folder, job.folder_named = folder_was
                return
            previewed = self._math_previewed
            if job not in self.jobs:  # (a read that ended meanwhile put its documents with another job)
                self._toast("The job changed while the preview was open: nothing was saved. Generate again.")
                return
            if self._batch_running():
                return
        keys: list[str] = []  # the attorneys invoiced
        math: list = []  # (FirmInvoice, number) of each invoice made: "The math" shows how its amounts were reached
        detailed: list[Path] = []  # the invoices' detailed copies: kept for later, not opened or printed now
        # billed this way now, and so again (Job.split_rule; the invoices' record keeps it, the other files' not,
        # should the invoice fail): the rule the invoices (and the preview, by Settings') are priced by; put back when
        # no invoice comes of it. The record is still the job as it was shown (origin above), whatever a read that
        # ended during the preview added since.
        rule_was = job.split_rule
        billed_origin = None
        if "invoice" in outputs:
            job.settle_split_rule(self.s)
            billed_origin = {**origin, "job": {**origin["job"], "split_rule": job.split_rule}}
        try:
            paths = generate(case, self.s, out_dir, outputs, opts, runsheet=sheet, folders=input_folders(job),
                             invoiced=keys, origin=origin, math=math, detailed=detailed, ordered=ordered,
                             speeds=speeds, day=dating.day, invoice_origin=billed_origin)
        except Exception as e:
            if not keys:
                job.split_rule = rule_was
            made = list(getattr(e, "made", []))
            job.saved, job.error = made, f"{type(e).__name__}: {e}"
            if keys:
                # some attorneys' invoices were made: trying again doesn't bill them twice. A day invoiced before
                # is billable again for the others only (billed_keys() is empty while it is invoiced).
                job.invoiced_keys = list(dict.fromkeys(job.billed_keys() + keys))
                job.invoiced = False
            show_save_error(self, e, "Could not make the files", "\n\nSaved before the problem:\n"
                            + "\n".join(p.name for p in made) if made else "")
            self._update_status()
            self._records_changed()
            return
        log.info("made %d file(s): %s, form type %s", len(paths), "+".join(outputs), self.s.form_choice)
        job.saved, job.error = paths, ""
        if not keys:  # (no invoice made: no transcript, say)
            job.split_rule = rule_was
        # Generate all won't bill this day again, nor make it (Job.made), unless a document was added while the
        # preview was open: its pages are not on the invoice just made
        if job.docs == docs:
            job.made = True
            if "invoice" in outputs:
                job.invoiced, job.invoice_set = True, new_invoice_set()
        # the detailed copies and the math PDFs are kept for when they're asked for: not opened or printed now
        from ..invoice_math import is_math_file
        kept = [p for p in paths if p in detailed or is_math_file(p)]
        if self.s.open_after:
            for p in paths:
                if p not in kept:
                    open_path(p)
        if math and self.s.show_math and not previewed:  # (the preview showed it, with the files)
            self._show_math(math, self.s.folder_for("invoice", out_dir), [p for p in kept if is_math_file(p)])
        made = [p for p in paths if not (sheet and p == sheet.path)]
        folders = list(dict.fromkeys(p.parent for p in made))  # an output may have a folder of its own
        where = f"\n\nin {folders[0]}" if len(folders) == 1 else "\n\nin " + ", ".join(map(str, folders))
        text = f"Saved {plural(len(made), 'file')}:\n\n" + "\n".join(p.name for p in made) + where if made else ""
        if made and dating.moved():  # ("Dated Tue 10/13/2026, the next business day: today is a Sunday and ...")
            text = f"{dating.note()}\n\n{text}"
        if sheet and sheet.path:
            text += ("\n\n" if text else "") + run_sheet_summary(sheet)
            folders.append(sheet.path.parent)
        # (Print...: the PDFs but those kept; the run sheet is a workbook, printed from Excel)
        self._saved_box("Saved", text, list(dict.fromkeys(folders)),
                        files=[p for p in paths if p.suffix.lower() == ".pdf" and p not in kept])
        self._redate()  # (generate dated the job's case as its files, or the preview's copy: the window catches up)
        self._update_status()
        self._set_status(f"✓  Saved {len(paths)} file(s)", "ok")
        self._records_changed()

    def _show_math(self, made: list, folder: Path, saved: list | None = None) -> None:
        """The math of the invoices just made (preview.MathDialog), to glance at and close; `saved`: the PDFs of
        it Generate saved (Settings.save_math), which it names."""
        from .preview import MathDialog
        try:
            dlg = MathDialog(made, self.s, self, folder, saved=saved)
        except Exception as e:  # the invoices are saved; this is only an aside
            logfile.error("could not spell out the math", e)
            return
        before = self.s.save_math
        dlg.exec()
        self._math_saves_kept(before)

    def _math_saves_kept(self, before: str) -> None:
        """After a window showing the math, whose "Save the math" list may have changed Settings.save_math
        (before: as it was): the counts follow, Generate all's number counting the math PDFs (output_counts)."""
        if self.s.save_math != before:
            self._show_counts()
            self._refresh_job_labels()

    def _preview(self, case: CaseInfo, outputs: list[str], opts, ordered: dict | None = None,
                 speeds: dict | None = None, day: date | None = None) -> bool:
        """Shows the files Generate is about to make, as pictures (preview.PreviewDialog); True = save them.
        They are made in a temporary folder with a records database of its own (Ledger.preview_copy: the
        invoice shows the number it will get, and none is taken). The run sheet is not shown: it is a
        spreadsheet, and making it would add to the real one. True too when there is nothing to show. When the
        preview can't be made, the user is asked whether to save without one (a problem with the files
        themselves is then said as usual). With invoices among them, "The math" is a tab of the preview (unless
        Settings.show_math is off); self._math_previewed then says it was shown, so it isn't shown again after
        saving. ordered, speeds: each attorney's pages and speeds (Job.ordered_pages, Job.ordered_speeds), as
        generate() takes them; day: the day the files are dated, the one the saved files get too."""
        from .preview import PreviewDialog
        self._math_previewed = False
        show = _previewed(outputs)
        if not show:
            return True
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
            s = deepcopy(self.s)
            s.output_dir, s.output_dirs = tmp, {}
            s.save_math = "off"  # (its tab shows the math; Save saves it with the invoices)
            math: list = []
            try:
                files = generate(deepcopy(case), s, Path(tmp), show, opts, math=math, ordered=ordered, speeds=speeds,
                                 ledger=ledger_for(self.s).preview_copy(Path(tmp) / "records"), day=day)
                dlg = PreviewDialog([f for f in files if f.suffix.lower() == ".pdf"], self.s, self,
                                    _not_previewed(outputs), math if self.s.show_math else None)
                self._math_previewed = dlg.math is not None
            except Exception as e:
                logfile.error("could not make the preview", e)
                return QMessageBox.question(self, "No preview", "The preview could not be made "
                                            f"({type(e).__name__}).\n\nSave the files without it?") == QMessageBox.Yes
            before = self.s.save_math
            accepted = dlg.exec() == PreviewDialog.Accepted
            self._math_saves_kept(before)
            return accepted

    def _without_unchecked_invoice(self, job: Job, outputs: list[str]) -> list[str]:
        """The outputs to make, less the invoice when it is held (Job.invoice_hold): whose pages to bill is asked
        first (Whose pages...); Excerpts... rows that need checking, and more speeds on the same pages than are
        billed for now (Job.speed_problems), are not. Asks whether to make the others without it ([] = make
        nothing)."""
        if "invoice" in outputs and job.ownership_problem() and not job.portions_problem():
            self._whose_pages(job, job.ownership_problem())
            if job not in self.jobs or self._batch_running():
                return []
        if "invoice" not in outputs or not job.invoice_hold():
            return outputs
        others = [o for o in outputs if o != "invoice"]
        kind = job.hold_kind()  # (the title and the text about the one the hold names)
        title = {"portions": "Excerpts… needs checking", "ownership": "Whose pages to bill",
                 "unnamed": "Parties needs checking"}.get(kind, "The speeds need checking")
        if kind == "portions":
            why = (f"{job.portions_check()}.\n\nNo invoice is made for this day until you check it: open "
                   "Excerpts… in the Invoice card and tick who ordered which pages, or choose \"Remove this "
                   "day's excerpts\".")
        elif kind == "unnamed":  # parties not named, on pages ordered at different speeds
            why = (f"{job.unnamed_problem()}.\n\nNo invoice is made for this day until every party is named or "
                   "Parties is changed (the Parties box in the Invoice card is lit up).")
        elif kind == "speeds":  # more speeds on the same pages than are billed for now
            why = (f"{job.invoice_hold()}.\n\nNo invoice is made for this day until you change a firm's speed "
                   "(Who ordered what, or Excerpts… for a run of pages).")
        else:
            why = (f"{job.ownership_problem()[0].upper()}{job.ownership_problem()[1:]}.\n\nNo invoice is made "
                   "for this job until you choose whose pages to bill (Whose pages… in the Invoice card).")
        if not others:
            QMessageBox.warning(self, title, why)
            self._show_parties_problem(job)
            return []
        answer = QMessageBox.question(self, title, why + "\n\nMake the other outputs without the invoice?")
        self._show_parties_problem(job)
        if answer != QMessageBox.Yes:
            return []
        if job not in self.jobs or self._batch_running():
            return []
        return others

    def _show_parties_problem(self, job: Job) -> None:
        """After the box saying the job's Parties number counts more parties than the firms named, who ordered at
        different speeds (Job.unnamed_problem): the Parties box lights up for a few seconds, the job shown, so
        the reporter sees what to change (the user, 2026-10-21). Nothing for another problem (one the hold names
        first: Job.hold_kind), or a job not shown."""
        if job is self.cur and job.hold_kind() == "unnamed":
            fr = self.cards.get("invoice")
            if fr is not None and fr.body.isHidden():
                self._fold_card(fr, False)
            self._flash(self.inv_parties, 4000)

    def _sheets_free(self, targets: list[str], button: str) -> bool:
        """Before anything is made: the run sheets the takes go on can be written. One open in another program
        (Excel) is named in a message saying to close it, and False is returned: nothing is made, so that the
        takes of a trial never go on its run sheet with a day missing. targets: the run sheets chosen ("" or
        None: a new one, nothing to check); button: "Generate" or "Generate all", to press again."""
        from ..runsheet import is_locked
        locked = list(dict.fromkeys(t for t in targets if t and is_locked(t)))
        if not locked:
            return True
        names = "\n".join(f"•  {Path(t).name}" for t in locked)
        box = QMessageBox(QMessageBox.Warning, "Close the run sheet first",
                          f"{'This run sheet is' if len(locked) == 1 else 'These run sheets are'} open in another "
                          f"program (probably Excel), so the takes can't be added:\n\n{names}\n\n"
                          f"Close {'it' if len(locked) == 1 else 'them'} (or the other copy of YinItAgreementForm "
                          f"that has {'it' if len(locked) == 1 else 'them'} open), then press {button} again. "
                          "Nothing was made.", QMessageBox.Ok, self)
        box.exec()
        return False

    def _run_sheet_for(self, job: Job, button: str = "Generate") -> str | None:
        """The run sheet to add the job's takes to ("" = a new one), as Settings → Run sheet says; asks when the
        case seems to have one already. It is looked for where a new one would go (the case's folder, before
        Generate has chosen it: batch.case_dir_for; or the run sheets' own folder under Options → Folders), next to
        its transcripts, where the records say the case's run sheets are and in the old run sheets folders
        (runsheet.sheet_folders). None = cancelled, also when a workbook named for the case can't be opened just
        now (runsheet.SheetUnreadable: the user is told which, and to press `button` again; nothing is made)."""
        if self.s.runsheet_existing == "new":
            return ""
        try:
            found = find_sheets(job.case, sheet_folders(self.s, case_dir_for(job, self.s), input_folders(job),
                                                        job.case))
        except SheetUnreadable as e:
            QMessageBox.warning(self, "Run sheet", f"{e}\n\nNothing was made. Press {button} again once it can "
                                                   "be opened.")
            return None
        if not found:
            return ""
        if self.s.runsheet_existing == "add":
            return str(found[0].path)
        dlg = RunSheetDialog(found, job.title(), self)
        return dlg.choice() if dlg.exec() == RunSheetDialog.Accepted else None

    def _ask_index_numbers(self, jobs: list[Job], button: str = "Generate") -> bool:
        """Generate's first question, and again just before the files are made (a document read meanwhile): when
        a job's documents have index numbers that don't match (Job.index_number_asks: they may be about
        different cases), says which document has which, the jobs of Generate all in one box, and asks whether to
        go on ("<button> anyway"). Going on keeps them together (Job.index_numbers_ok: not asked again until
        another document comes). The answer covers the documents the box named: one read while it was open is
        asked about again. True = go on."""
        while True:
            # (after an answer: the jobs still there, a document that came meanwhile asking again)
            asking = [j for j in jobs if j.index_number_asks() and j in self.jobs]
            if not asking:
                return True
            shown = {id(j): [d.key() for d in j.docs] for j in asking}
            if len(asking) == 1:
                text = index_number_text(asking[0].index_number_mismatch()) + "."
            else:
                text = (f"In {len(asking)} jobs the documents' index numbers don't match, so they may not be about "
                        "the same case:\n\n" + "\n".join(f"•  {j.title()}: "
                                                         f"{index_number_text(j.index_number_mismatch(), '')}"
                                                         for j in asking[:8]) + ("\n…" if len(asking) > 8 else ""))
            box = QMessageBox(QMessageBox.Warning, "Different index numbers",
                              text + "\n\nIf a document belongs to another case, go back, right-click it in the "
                              "list on the left (under its job) and choose Move to a job of its own.", parent=self)
            back = box.addButton("Go back", QMessageBox.RejectRole)
            go = box.addButton(f"{button} anyway", QMessageBox.AcceptRole)
            box.setDefaultButton(back)
            box.setEscapeButton(back)
            box.exec()
            if box.clickedButton() is not go or self._batch_running():
                return False
            for j in asking:
                j.index_numbers_ok = shown[id(j)]
            if self.cur in asking:
                self._show_index_notes()
            self._refresh_job_labels()  # (the ⚠ and "don't match" of the list go too)

    def _case_folder(self, jobs: list[Job], outputs: list[str], button: str = "Generate") -> bool:
        """Chooses each job's own case folder (Settings.case_folders: batch.case_folders, Job.folder). One that is
        already there with files in it is added to, given a new one beside it ("712345-2021 (2)"), or asked
        about, as Settings.case_folder_existing says; the question can say not to ask again. A job that chose its
        folder at an earlier Generate keeps it (until its folder's name changes). The caller puts the folders back
        when nothing is saved after all (Go back at the preview). False = Go back (nothing chosen)."""
        folders = case_folders(jobs, self.s, outputs)
        there = [f for f in folders if f.files]
        new = self.s.case_folder_existing == "new"
        if there and self.s.case_folder_existing == "ask":
            answer = self._ask_case_folders(there, button)
            if answer is None or self._batch_running():
                return False
            new = answer
        for f in folders:
            f.use(new=new)
        return True

    def _ask_case_folders(self, there: list, button: str) -> bool | None:
        """The question of _case_folder: these case folders (batch.CaseFolder) are already there with files in
        them; add the new files to them (False), or make new ones beside them (True)? None = Go back. "Don't ask
        again" makes the answer Settings.case_folder_existing."""
        one = len(there) == 1
        if one:
            f = there[0]
            text = (f"A folder for this case is already there, with {plural(f.files, 'file')} in it:\n\n{f.path}\n\n"
                    f"Add the new files to it, or make a new folder beside it (\"{f.new_name()}\")?")
        else:
            text = (f"{len(there)} case folders are already there, with files in them:\n\n"
                    + "\n".join(f"•  {f.path}  ({plural(f.files, 'file')})" for f in there[:8])
                    + ("\n…" if len(there) > 8 else "")
                    + "\n\nAdd the new files to them, or make a new folder beside each (\"… (2)\")?")
        box = QMessageBox(QMessageBox.Question, "The case folder is there" if one else "The case folders are there",
                          text, parent=self)
        add = box.addButton("Add to it" if one else "Add to them", QMessageBox.AcceptRole)
        new = box.addButton(f"New folder \"{there[0].new_name()}\"" if one else "New folders", QMessageBox.ActionRole)
        back = box.addButton("Go back", QMessageBox.RejectRole)
        box.setDefaultButton(add)
        box.setEscapeButton(back)
        never = QCheckBox("Don't ask again (Settings → Options → Folders)")
        box.setCheckBox(never)
        box.exec()
        if box.clickedButton() not in (add, new):
            return None
        if never.isChecked():
            self.s.case_folder_existing = "new" if box.clickedButton() is new else "add"
            self._save_settings()
        return box.clickedButton() is new

    def fill_all_jobs(self):
        """Generate all: makes the outputs of every ticked job not made already (_to_generate: a made day whose
        invoice is still to make gets that invoice alone) on another thread (Ctrl+Shift+Enter). The jobs it makes
        stay ticked, the cards showing them, but are made (Job.made) from then on: pressed again, it makes only
        what is left.

        Asks first, as needed, in this order: the jobs whose documents' index numbers don't match, in one box
        (_ask_index_numbers); whose pages to bill, once per job that needs it; rows of a case's days that may be
        one firm (_ask_firms); the speed an e-mail asks for (_ask_speeds); a Parties number that isn't the firms
        that ordered (_ask_parties); the speeds of the firms of split orders that have none, every case in one
        box (_ask_split_speeds). Then the days whose invoice still waits for a choice (Excerpts..., whose pages,
        the speeds, a case with a day nobody is ticked on) are listed, to go back or go on without them. Blank
        fields are not asked about, but incomplete jobs are listed, to leave out or make anyway; then, for each
        invoice of the days made, whether there is an index when a page count typed below the transcript's
        changes it (_ask_index). Where the takes go is asked once per case when it has a run sheet already
        (Settings → Run sheet; the jobs of one case share a run sheet group, so its days go on one sheet), and
        those open in Excel are to be closed first. Last, the case folders already there with files in them, in
        one box (_case_folder), and the index numbers again, for a document read meanwhile.
        The batch works on copies of the jobs and the settings (shown first in a preview when
        Settings.preview_before_saving: _preview_batch, whose Go back leaves the case folders to be asked again);
        when it ends, each job's saved files, error, whether it was invoiced and with which days (Job.invoice_set)
        are copied back, and the days made without a problem are marked made (_run_batch).
        With a single job this is the same as Generate.
        """
        if len(self.jobs) < 2:
            return self.fill()
        if self._batch_running():
            return
        self._commit_edit()
        self._sync_from_ui()
        chosen = self._to_generate()  # (ticked, and not made already)
        if not chosen:
            return
        outputs = self._outputs()
        if not outputs:
            return
        # jobs whose documents' index numbers don't match (they may be about different cases): asked first
        if not self._ask_index_numbers(chosen, "Generate all"):
            return
        chosen = [j for j in chosen if j in self.jobs]
        if "invoice" in outputs:  # whose pages to bill is asked now, once per job that needs it
            for j in [j for j in chosen if j.ownership_problem() and not j.portions_problem()]:
                if j in self.jobs:
                    self._whose_pages(j, f"{j.title()}: {j.ownership_problem()}")
                if self._batch_running():
                    return
            chosen = [j for j in chosen if j in self.jobs]
        # Rows of a case's days that may be one firm: asked now (one firm is one agreement, one invoice; the run
        # sheet doesn't name them). Each question is asked when an output chosen asks it (OutputSpec.asks).
        if courthouses.asks(outputs, "firms") and not self._ask_firms(_cases(chosen)):
            return
        chosen = [j for j in chosen if j in self.jobs]
        # The speed of the jobs whose e-mail asks for another speed than the Settings rule's: asked now, at once
        if courthouses.asks(outputs, "speed"):
            if not self._ask_speeds(chosen):
                return
            chosen = [j for j in chosen if j in self.jobs]
        if courthouses.asks(outputs, "parties") and not self._ask_parties(chosen):
            return
        chosen = [j for j in chosen if j in self.jobs]
        # The firms of split orders commit to a speed: those not set yet, asked at once for every case
        if courthouses.asks(outputs, "speeds") and not self._ask_split_speeds(chosen, "Generate all"):
            return
        chosen = [j for j in chosen if j in self.jobs]
        # Days whose invoice still waits for a choice (Excerpts..., whose pages to bill, the speeds): say which,
        # before anything is made. Going on bills the other days of their case without them. A case with a day
        # nobody is ticked on gets no invoice at all (batch.group_problem): said too, as fill_jobs groups them.
        held = [j for j in chosen if "invoice" in outputs and not j.invoiced and j.invoice_hold()]
        stuck: list[tuple[list[Job], str]] = []  # (the days of a case, why none of them is invoiced)
        if "invoice" in outputs:
            from ..batch import group_problem, invoice_groups
            billed = [j for j in chosen if not j.invoiced and not j.invoice_hold()]
            stuck = [(g, group_problem(g)) for g in invoice_groups(billed, self.s) if group_problem(g)]
        if held or stuck:
            lines = [f"•  {j.title()} {j.case.get('dates')}: {j.invoice_hold()}" for j in held]
            lines += [f"•  {g[0].title()} ({plural(len(g), 'day')}): {why}; going on makes no invoice for this "
                      "case" for g, why in stuck]
            box = QMessageBox(self)
            box.setIcon(QMessageBox.Warning)
            box.setWindowTitle("Not every day is decided")
            box.setText("Some days still need your choice before they can be invoiced:\n\n"
                        + "\n".join(lines[:8]) + ("\n…" if len(lines) > 8 else "")
                        + "\n\nGo back and make the choices (Whose pages… or Excerpts… in the Invoice card, a "
                        "firm's speed in Who ordered what, or tick who ordered a day in its attorney table), or go "
                        "on anyway: the days listed get no "
                        "invoice now, and the other days of their case are invoiced without them, unless the case "
                        "has a day nobody is ticked on (they can be invoiced later).")
            back = box.addButton("Go back and decide", QMessageBox.RejectRole)
            go_on = box.addButton("Go on anyway", QMessageBox.AcceptRole)
            box.setDefaultButton(back)
            box.setEscapeButton(back)
            box.exec()
            if box.clickedButton() is not go_on or self._batch_running():  # (closed: as Go back)
                # show the first day to decide (for a case, its day with nobody ticked)
                first = held[0] if held else next((j for j in stuck[0][0] if not j.ticked_keys()), None)
                if first is not None and first in self.jobs and first is not self.cur and not self._batch_running():
                    self._select_job(first)
                return
            chosen = [j for j in chosen if j in self.jobs]

        def problems(j: Job) -> list[str]:
            """What is incomplete about a job, less the invoice choices just warned about (a made day: about its
            invoice alone, the one thing made for it)."""
            return [p for p in j.output_problems(["invoice"] if j.made else outputs)
                    if not (j in held and p == j.invoice_hold())]

        gaps = [j for j in chosen if problems(j)]
        if gaps:
            box = QMessageBox(self)
            box.setIcon(QMessageBox.Question)
            box.setWindowTitle("Some jobs are incomplete")
            box.setText(f"{len(gaps)} of {len(chosen)} jobs (marked ⚠) are incomplete:\n\n"
                        + "\n".join(f"•  {j.title()}:  {', '.join(problems(j))}" for j in gaps[:8])
                        + ("\n…" if len(gaps) > 8 else ""))
            ready = box.addButton(f"Do the {len(chosen) - len(gaps)} complete ones", QMessageBox.AcceptRole)
            everything = box.addButton("Do all (blanks left, no invoice without pages, no run sheet without a "
                                       "transcript)", QMessageBox.ActionRole)
            box.addButton(QMessageBox.Cancel)
            ready.setEnabled(len(chosen) > len(gaps))
            box.exec()
            if box.clickedButton() == ready:
                chosen = [j for j in chosen if not problems(j)]
            elif box.clickedButton() != everything:
                return
        if "invoice" in outputs:
            # a typed page count that changes whether there is an index: asked per invoice, of the days that are
            # made (after the incomplete ones were left out: an invoice of fewer days may decide it another way)
            from ..batch import group_problem, invoice_groups
            chosen = [j for j in chosen if j in self.jobs]
            billed = [j for j in chosen if not j.invoiced and not j.invoice_hold()]
            if not self._ask_index([g for g in invoice_groups(billed, self.s) if not group_problem(g)]):
                return
        # Where each case's takes go: asked once per case (the days of a trial share one run sheet)
        sheet_for: dict[int, tuple[str, int]] = {}  # id(job) -> (the run sheet chosen, its case's number)
        asked: list[tuple[Job, str]] = []
        for j in chosen:
            if "runsheet" not in j.makeable(outputs) or j.made:  # (a made day: its invoice alone, no takes)
                continue
            case_no = next((n for n, (k, _) in enumerate(asked) if same_case(ident(k.case), ident(j.case))), None)
            if case_no is None:
                target = self._run_sheet_for(j, "Generate all")
                if target is None or self._batch_running():
                    return
                case_no = len(asked)
                asked.append((j, target))
            sheet_for[id(j)] = (asked[case_no][1], case_no)
        # A read that ended while a question was on screen may have joined two jobs into one
        chosen = [j for j in chosen if j in self.jobs]
        if not chosen:
            return
        sheet_for = {k: v for k, v in sheet_for.items() if k in {id(j) for j in chosen}}
        if not self._sheets_free([t for t, _ in sheet_for.values()], "Generate all"):
            return
        # each case's own folder inside Save to: asked about once for all those with files in them already, before
        # the jobs are copied (the copies carry the folder chosen: Job.folder)
        folder_was = {id(j): (j.folder, j.folder_named) for j in chosen}
        if not self._case_folder(chosen, outputs, "Generate all") or self._batch_running():
            return
        # (a document read while a question was on screen, with an index number of its own: asked about too)
        if not self._ask_index_numbers(chosen, "Generate all"):
            for j in chosen:  # (Go back: nothing saved, the folders asked about next time)
                j.folder, j.folder_named = folder_was.get(id(j), (j.folder, j.folder_named))
            return
        chosen = [j for j in chosen if j in self.jobs]
        if not chosen:
            return
        # The batch is made on another thread while the window stays usable, so it gets its own copy of the
        # jobs and the settings: editing a job, an AI answer or a change in Settings can't reach files half made.
        def copy_jobs() -> dict:
            return {id(j): replace(j, case=deepcopy(j.case), docs=list(j.docs), saved=[], error="",
                                   invoiced_keys=list(j.invoiced_keys), portions=deepcopy(j.portions),
                                   speeds=dict(j.speeds),
                                   page_basis=deepcopy(j.page_basis), front_owner=dict(j.front_owner),
                                   runsheet_to=sheet_for.get(id(j), (None, None))[0],
                                   runsheet_group=sheet_for.get(id(j), (None, None))[1])
                    for j in self.jobs}
        copies = copy_jobs()
        todo, settings = [copies[id(j)] for j in chosen], deepcopy(self.s)
        dating = document_day(self.s)  # (every file's date, the preview's too: today, or the next business day)
        if self.s.preview_before_saving and _previewed(outputs):
            # the files are shown first, with the math, made from copies of their own: what is saved is what was
            # shown (the copies above), and Go back saves nothing
            def save(previewed: bool) -> None:
                # (the preview's math tab may have just changed them: the math is saved as it was last shown)
                settings.save_math, settings.math_layout = self.s.save_math, self.s.math_layout
                for j in chosen:  # (the folders chosen: kept, so the same jobs generating again use them)
                    j.folder, j.folder_named = copies[id(j)].folder, copies[id(j)].folder_named
                self._run_batch(chosen, copies, todo, settings, outputs, previewed, dating)
            shown = copy_jobs()
            for j in chosen:  # (until Save: Go back saves nothing, and the folders are asked about next time)
                j.folder, j.folder_named = folder_was.get(id(j), (j.folder, j.folder_named))
            self._preview_batch(chosen, shown, outputs, save, dating.day)
        else:
            self._run_batch(chosen, copies, todo, settings, outputs, False, dating)

    def _preview_batch(self, chosen: list[Job], copies: dict, outputs: list[str], then,
                       day: date | None = None) -> None:
        """Generate all's preview before saving: the batch made on another thread into a temporary folder, with a
        records database of its own (Ledger.preview_copy: the numbers the invoices will get, none taken), then
        shown in one PreviewDialog with "The math" of every invoice. Save calls then(math shown); Go back saves
        nothing. The run sheets are not made for it (they are spreadsheets, and would add to the real ones). day:
        the day the files are dated, the saved ones too."""
        from .preview import PreviewDialog
        tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        s = deepcopy(self.s)
        s.output_dir, s.output_dirs = tmp.name, {}
        s.save_math = "off"  # (its tab shows the math; Save saves it with the invoices)
        shown = _previewed(outputs)
        math: list = []
        try:  # (before the window is marked busy: a records database that can't be opened left it busy for good)
            ledger = ledger_for(self.s).preview_copy(Path(tmp.name) / "records")
        except Exception as e:
            tmp.cleanup()
            logfile.error("could not make the preview of the batch", e)
            if QMessageBox.question(self, "No preview", f"The preview could not be made ({type(e).__name__})."
                                    "\n\nSave the files without it?") == QMessageBox.Yes \
                    and not self._batch_running():
                then(False)
            return
        self.filling = True
        self.work += 1
        self._update_header_buttons()
        self.fill_all_btn.setEnabled(False)
        self.fill_btn.setEnabled(False)
        self.busy.setVisible(True)
        self._set_status("Making the preview…", "busy")

        def ended():
            self.filling = False
            self.work = max(0, self.work - 1)
            self._update_header_buttons()
            self._idle()
            self.fill_btn.setEnabled(True)
            self._refresh_job_labels()
            self._batch_ended()

        def done(paths):
            ended()
            try:
                try:
                    dlg = PreviewDialog([p for p in paths if p.suffix.lower() == ".pdf"], self.s, self,
                                        _not_previewed(outputs, several=True),
                                        math if self.s.show_math else None)
                except Exception as e:  # (a page that can't be drawn: as for one job, see _preview)
                    logfile.error("could not show the preview of the batch", e)
                    self._job_status()
                    if QMessageBox.question(self, "No preview", f"The preview could not be made ({type(e).__name__})."
                                            "\n\nSave the files without it?") == QMessageBox.Yes \
                            and not self._batch_running():
                        then(False)
                    return
                before = self.s.save_math
                save = dlg.exec() == PreviewDialog.Accepted
                previewed = dlg.math is not None
                self._math_saves_kept(before)
            finally:
                tmp.cleanup()
            self._job_status()
            if save and not self._batch_running():
                then(previewed)

        def failed(msg):
            ended()
            tmp.cleanup()
            self._job_status()
            logfile.error("could not make the preview of the batch", RuntimeError(msg))
            if QMessageBox.question(self, "No preview", "The preview could not be made.\n\nSave the files "
                                    "without it?") == QMessageBox.Yes and not self._batch_running():
                then(False)

        self.runner.start(fill_jobs, [copies[id(j)] for j in chosen], s, batch=list(copies.values()),
                          outputs=shown, ledger=ledger, math=math, day=day, on_done=done, on_error=failed)

    def _run_batch(self, chosen: list[Job], copies: dict, todo: list[Job], settings, outputs: list[str],
                   previewed: bool, dating=None) -> None:
        """Makes Generate all's files on another thread from the copies of the jobs, and copies what happened
        back to the jobs when it ends. previewed: the preview showed the math (else it is shown after, when
        Settings.show_math is on). dating: the day every file is dated (workdays.document_day, as Generate all
        started; default: now's), which the message at the end says when it isn't today."""
        dating = dating or document_day(self.s)
        math: list = []
        self.filling = True
        self.work += 1
        self._update_header_buttons()
        self.fill_all_btn.setEnabled(False)
        self.fill_btn.setEnabled(False)
        self.busy.setVisible(True)

        # No self.gen check here: New job is refused while self.filling, so these jobs are still the window's.
        def ended():
            self.filling = False
            self.work = max(0, self.work - 1)
            self._update_header_buttons()
            self._idle()
            self.fill_btn.setEnabled(True)
            self._batch_ended()

        def step(i, n, name):
            self.busy.setRange(0, n)
            self.busy.setValue(i)
            self._set_status(f"Making {i + 1} of {n}…", "busy")

        def done(paths):
            ended()
            # (ids of) jobs whose invoice changed while the batch ran (_billing_changed): not marked made, so the
            # next Generate all bills them as they are now
            changed = set()
            for j in chosen:
                copy = copies[id(j)]
                j.saved, j.error = copy.saved, copy.error
                j.split_rule = j.split_rule or copy.split_rule  # (the rule its invoices were just billed by)
                # (a read that ended during the run put this day's documents with another day: that job's now)
                k = j if any(j is x for x in self.jobs) else next(
                    (x for x in self.jobs if any(d is e for d in j.docs for e in x.docs)), None)
                if k is not None and not k.split_rule:
                    k.split_rule = copy.split_rule
                if _billing_changed(j, copy):
                    changed.add(id(j))
                    if copy.invoiced and not _pages_changed(j, copy):
                        # only who orders (or the speed, the parties) changed: those just invoiced aren't billed
                        # again by the next run, but an attorney ticked meanwhile is
                        j.invoiced_keys = list(dict.fromkeys(j.billed_keys() + _billed_by_copy(copy)))
                        j.invoiced = False
                else:
                    j.invoiced, j.invoiced_keys = copy.invoiced, copy.invoiced_keys
                    j.invoice_set = copy.invoice_set  # (the days billed together: the cards show them so)
            # days whose only "error" is the invoice held back, as the user chose before it started (see
            # batch.fill_jobs: those messages come last, so an error that starts with one has nothing else)
            held_days = [j for j in chosen if HELD in j.error]  # (a held case's days: listed once, under the case)
            unbilled = [j for j in chosen if j.error.startswith(NOT_INVOICED) and not any(j is h for h in held_days)]
            failed = [j for j in chosen if j.error and not j.saved and j not in unbilled]
            partial = [j for j in chosen if j.error and j.saved and j not in unbilled]
            made = len(chosen) - len(failed) - sum(1 for j in unbilled if not j.saved)
            for j in chosen:  # (still ticked: the cards go on showing them)
                if (j.saved or j.invoiced) and not j.error and id(j) not in changed:
                    j.made = True  # "Generate all" again only does what is left
            changed = [j for j in chosen if id(j) in changed and any(j is k for k in self.jobs)]
            self._update_status()
            folders = list(dict.fromkeys(str(p.parent) for p in paths))
            text = f"Saved {len(paths)} file{'s' if len(paths) != 1 else ''} for {plural(made, 'job')}"
            text += f" in\n{folders[0]}" if len(folders) == 1 else f" in {len(folders)} folders." if folders else "."
            if paths and dating.moved():  # ("Dated Tue 10/13/2026, the next business day: today is a Sunday ...")
                text = f"{dating.note()}\n\n{text}"
            held = _held_cases(held_days)  # a file of a case couldn't be saved: nothing (more) made for its days
            if held:
                # ("more": a trial's run sheet is written first, and may be saved before another of its files fails)
                text += (f"\n\nNothing more was made for {plural(len(held), 'case')} once a file couldn't be saved:\n"
                         + "\n".join(f"•  {case}: {why}" for case, why in held.items())
                         + "\n\nClose the file (in Excel, a PDF viewer or another copy of YinItAgreementForm), then "
                         "press Generate all again: it makes those days again.")
                # (the day it happened on, whose own error doesn't say HELD, is among them: the same case)
                in_held = [ident(h.case) for h in held_days]
                failed = [j for j in failed if not any(same_case(ident(j.case), k) for k in in_held)]
                partial = [j for j in partial if not any(same_case(ident(j.case), k) for k in in_held)]
            if unbilled:
                text += f"\n\n{plural(len(unbilled), 'day')} not invoiced, as you chose:\n" + "\n".join(
                    f"•  {j.title()} {j.case.get('dates')}: {_reason(j.error[len(NOT_INVOICED):].strip())}"
                    for j in unbilled[:6])
            if failed:
                text += f"\n\n{len(failed)} could not be saved (is a file - a PDF or the Excel run sheet - open in " \
                        "another program?):\n" + "\n".join(f"•  {j.title()}: {_reason(j.error)}" for j in failed[:6])
            if partial:
                text += f"\n\n{len(partial)} job(s) are not complete:\n" + \
                        "\n".join(f"•  {j.title()}: {_reason(j.error)}" for j in partial[:6])
            if changed:
                text += (f"\n\n{len(changed)} job(s) changed while the files were made (a document was added, "
                         "the pages, Excerpts…, the attorneys ticked, the parties or the speed changed), so "
                         "Generate all again bills them as they are now:\n" + "\n".join(f"•  {j.title()}" for j in changed[:6]))
            if math and self.s.show_math and not previewed:
                from ..invoice_math import is_math_file
                self._show_math(math, Path(folders[0]) if folders else self.s.folder_for("invoice"),
                                [p for p in paths if is_math_file(p)])
            self._saved_box("Batch finished", text, folders, "\n".join(str(p) for p in paths),
                            warn=bool(failed or held))
            self._set_status(f"✓  Saved {len(paths)} file(s)", "warn" if failed or held else "ok")
            self._records_changed()

        def crashed(msg):
            ended()
            self._update_status()
            QMessageBox.critical(self, "Could not fill the forms", msg)

        self.runner.start(fill_jobs, todo, settings, batch=list(copies.values()), outputs=outputs, math=math,
                          day=dating.day, on_done=done, on_error=crashed, on_progress=step)

    # ------------------------------------------------------- misc
    def new_job(self):
        """Clears every job and starts over (Ctrl+N), without asking: Undo at the bottom of the window brings
        them back for a few seconds (_offer_undo). Work already running in the background is not stopped, but
        its results are dropped (self.gen goes up), and AI questions still waiting their turn are never asked
        (_clear_jobs); with such work, there is no Undo (what it would have added is lost). Nothing happens when
        there is nothing to clear (the button is greyed out then)."""
        if not self._anything_to_clear():
            return
        if self._batch_running():
            return
        self._sync_from_ui()  # (an attorney's name being typed, and the details under the table, kept)
        snapshot = None
        if self.work == 0 and self.ai_pending == 0:
            snapshot = {"jobs": self.jobs, "cur": self.cur, "paste": self.paste.toPlainText(),
                        "collapsed": set(self._collapsed), "touched": self._runsheet_touched,
                        "auto": self._runsheet_auto,
                        "runsheet": self.output_boxes["runsheet"].isChecked() if "runsheet" in self.output_boxes
                        else None}
        n = sum(1 for j in self.jobs if not j.is_empty())
        self._clear_jobs()
        self._offer_undo(snapshot, n)

    def _offer_undo(self, snapshot: dict | None, n: int) -> None:
        """After New job: "Cleared 4 jobs." and Undo at the bottom of the window for UNDO_MS (just "Cleared."
        without a snapshot)."""
        self._undo = snapshot
        self.statusBar().clearMessage()
        what = plural(n, "job") if n else "the pasted text"
        self.undo_text.setText(f"Cleared {what}." if snapshot else "Cleared.")
        self.undo_btn.setVisible(snapshot is not None)
        self.undo_box.setVisible(True)
        self._undo_timer.start(UNDO_MS)

    def _drop_undo(self) -> None:
        """New job's Undo is no longer offered: its time is up, something new was loaded or typed (an undo would
        throw it away), or the settings or rate sheets the cleared jobs were priced with changed."""
        self._undo = None
        if hasattr(self, "undo_box"):
            self._undo_timer.stop()
            self.undo_box.setVisible(False)

    def _undo_clear(self) -> None:
        """Undo: the jobs New job cleared come back as they were, the job that was shown on screen again, with the
        pasted text and the Run sheet box. (The Excerpts window, closed by New job, stays closed.)"""
        snap = self._undo
        self._drop_undo()
        # (a file dropped or a recent one opened since New job is still being read: bringing the old jobs back
        # would drop it without a word)
        if snap is None or self._batch_running() or self.work or self.ai_pending:
            return
        self.gen += 1  # (anything started since is dropped)
        self.ai_runner.drop_queued()
        self.ai_pending = self.work = 0
        self._loading.clear()
        self.jobs, self.cur = snap["jobs"], snap["cur"]
        self._collapsed = set(snap["collapsed"])
        self._runsheet_touched, self._runsheet_auto = snap["touched"], snap["auto"]
        if snap["touched"] and snap["runsheet"] is not None:  # (the user's own tick, as they left it)
            with QSignalBlocker(self.output_boxes["runsheet"]):
                self.output_boxes["runsheet"].setChecked(snap["runsheet"])
        else:
            self._auto_runsheet()
        with QSignalBlocker(self.paste):
            self.paste.setPlainText(snap["paste"])
        self._refresh_jobs()
        self._show_job()
        self._redate(force=True)  # (dated as of now: midnight or the evening hour may have passed meanwhile)
        self._update_status()
        self._toast("Brought back what New job cleared.")

    def _anything_to_clear(self) -> bool:
        """Whether New job would do anything: a job with documents or typed details, pasted text, or documents
        still being read."""
        return not all(j.is_empty() for j in self.jobs) or bool(self.paste.toPlainText().strip()) or self.work > 0

    def _update_header_buttons(self) -> None:
        """New job (button and menu) is greyed out while there is nothing to clear, or Generate all runs. The
        start card shows while nothing is loaded, and blank required fields are marked once something is (as the
        status changes, or text is pasted)."""
        if not hasattr(self, "new_action"):  # still building the window
            return
        on = self._anything_to_clear() and not self.filling
        self.new_btn.setEnabled(on)
        self.new_action.setEnabled(on)
        if self._undo is not None and self._loaded():
            self._drop_undo()  # (something new was loaded or typed since New job: an Undo would throw it away)
        self._show_start()
        self._mark_missing()

    def _clear_jobs(self, job: Job | None = None) -> None:
        """Drops every job and starts with one (a blank one, or `job`); results of work still running are
        dropped (self.gen goes up), and AI questions still waiting their turn are never asked."""
        self.gen += 1
        self.ai_runner.drop_queued()
        if self.excerpts is not None:  # (it shows days of the jobs dropped)
            self.excerpts.clear()
        self.jobs = [job or Job()]
        self.cur = self.jobs[0]
        self._runsheet_touched = False  # the Run sheet box goes back to the saved choice (see _auto_runsheet)
        self._runsheet_auto = None
        self._auto_runsheet()
        self.ai_pending = self.work = 0
        self._loading.clear()
        self.busy.setRange(0, 0)
        self._refresh_jobs()
        self.paste.clear()
        self._show_case()
        self.busy.setVisible(False)
        self._set_status(IDLE_STATUS, "")
        self._show_todo([])  # (the jobs it listed are gone)

    def open_past_job(self, origin: str) -> bool:
        """Opens a job again from its record (Records -> Open this job again), for a corrected invoice or
        another day of the case. origin: Activity.origin (JSON), or what the Records window put together from
        an older record's columns. The case comes back as it was filled in (batch.job_from_origin) and the
        documents it was read from are read again when they are still where they were, their takes as the user
        corrected them (batch.origin_fixes, given to add_files). An invoice of several
        days is read from its documents alone, a job per day. Asks first when there is work on screen; False
        when nothing was opened (during Generate all a box says so, over the Records window the click came from:
        this window's status bar is behind it)."""
        if self.filling:
            QMessageBox.information(getattr(self, "_records_win", None) or self, "Open this job again",
                                    "Still making the files of the batch - one moment, then try again.")
            return False
        try:
            data = json.loads(origin)
        except ValueError:
            data = None
        if not isinstance(data, dict):
            QMessageBox.information(self, "Open this job again", "This record doesn't say what it was made from.")
            return False
        if any(not j.is_empty() for j in self.jobs) and QMessageBox.question(
                self, "Open this job again", "Clear what is in the window now and open this job instead?"
        ) != QMessageBox.Yes:
            return False
        if self._batch_running():
            return False
        sources = [p for p in data.get("sources") or [] if isinstance(p, str)]
        there = [p for p in sources if Path(p).is_file()]
        several_days = data.get("joint") is True and there
        try:  # (billed one way: that way again)
            rule = billed_split_rule(origin, ledger_for(self.s))
        except Exception as e:
            logfile.error("could not look up the records", e)
            rule = ""
        if several_days:
            self._clear_jobs()
        else:
            try:
                job = job_from_origin(data, self.s, rule)
            except Exception as e:  # (a record damaged in a way job_from_origin doesn't expect)
                logfile.error("could not open a past job", e)
                QMessageBox.information(self, "Open this job again", "This record can't be opened again.")
                return False
            if len(there) < len(sources) or not sources:
                # Not every document will be read again: what was read from the missing ones would be lost
                # when the others are merged, so it is kept as if typed (not the defaults: they are worked out).
                # Not a page count of a transcript (SRC_PDF) when every PDF is read again: it is counted again
                # from the transcript. Kept as typed, the invoice would bill it as "the Pages field"
                # (Job.pages_typed) instead of the user's own pages of a transcript of several reporters. (When
                # a PDF is missing, it may be the transcript: an order letter saved as a PDF is no count.)
                # Nor a speed a document named: the agreement form's speed is the Settings rule's unless the user
                # chose one (merge.refresh_speed), and an e-mail's speed was never the user's choice.
                pdfs = [p for p in sources if p.lower().endswith(".pdf")]
                pdf_back = bool(pdfs) and all(p in there for p in pdfs)
                for key, f in job.case.fields.items():
                    if f.value and f.source not in (SRC_DEFAULT, SRC_DERIVED) and not (
                            f.source == SRC_PDF and pdf_back) and key != "delivery":
                        f.source = SRC_USER
            apply_defaults(job.case, self.s)
            self._clear_jobs(job)
            self._show_job()
        if there:
            self.add_files(there, batch=bool(several_days), same_job=not several_days,
                           split_rule=rule if several_days else "", fixes=origin_fixes(data))
        gone = len(sources) - len(there)
        if gone:
            self._toast(f"{plural(gone, 'document')} of this job {'is' if gone == 1 else 'are'} no longer where "
                        f"{'it was' if gone == 1 else 'they were'}: drop {'it' if gone == 1 else 'them'} here to "
                        "make an invoice again.")
        self.raise_()
        self.activateWindow()
        return True

    def open_settings(self, _checked: bool = False, tab: str = ""):
        """Opens Settings (not while Generate all is making files: it reads the settings, and the answers about
        the firms, on its thread), at the tab named (Records' "Tick boxes and colours…": "Records"). After Save,
        the theme, the window's options, the rate sheets and the AI check are refreshed and every job is merged
        and priced again, with the answers about rows that may be one firm as they are now (one forgotten, for now
        or for good, splits its rows again and is asked about again)."""
        if self._batch_running():
            return
        def pairs() -> set[frozenset]:
            return {frozenset(r[:2]) for r in self.s.firm_answers + self._session_answers}
        before, sheet = pairs(), self.s.sheet().name  # (the one in use: the one Settings names may be gone)
        dlg = SettingsDialog(self.s, self, self._session_answers)
        if tab:
            dlg.show_tab(tab)
        ok = dlg.exec()
        dlg.deleteLater()  # (closed is only hidden: each visit's dialog, and a help window it opened, would stay)
        if ok:
            self._session_answers = dlg.session_answers
            self._settings_changed({k for pair in before - pairs() for k in pair}, sheet)

    def _settings_changed(self, forgotten: set[str] | frozenset = frozenset(), sheet: str | None = None) -> None:
        """The settings were changed (Settings saved, the welcome questions answered, settings imported): the
        theme, the window's options, the rate sheets and the AI check are refreshed and every job is merged
        and priced again (with the answers about rows that may be one firm as they are now). forgotten: the
        keys of rows whose answer was forgotten: a job that joined them reads its rows again (split_forgotten).
        sheet: the rate sheet Settings named before (Settings, an import); the one in use now is checked
        (_check_rate_sheet), with Go back to that one, and asked about even if it was this session when it is
        another."""
        self._drop_undo()  # (the jobs New job cleared were merged and priced with the settings before)
        self._use_answers()
        if forgotten:
            self._sync_from_ui()
            for job in self.jobs:
                if split_forgotten(job, forgotten, self.s) and job is self.cur:
                    self._show_attorneys()  # (the merge below starts from the table)
        if self.s.zoom != zooming.zoom():
            self.set_zoom(self.s.zoom)  # (the theme too)
        apply_theme(self.app, self.s.theme)
        self.per_email.setChecked(self.s.per_email)
        self.sign_rep.setChecked(self.s.sign_reporter)
        for combo, value in ((self.form_choice, self.s.form_choice), (self.mofr_division, self.s.mofr_division),
                             (self.rs_existing, self.s.runsheet_existing)):
            with QSignalBlocker(combo):
                combo.setCurrentIndex(max(0, combo.findData(value)))
        set_checks(self.output_boxes, self.s.outputs)
        self._auto_runsheet()
        self.s.reload_rates()
        # an incomplete sheet is warned about before the jobs are priced from it: Go back leaves them as they were
        if not self._check_rate_sheet(sheet or "", again=sheet is not None and sheet != self.s.rate_sheet,
                                      priced=False):
            self._save_settings()  # (Settings, or the import, saved the sheet gone back from)
        self._fill_sheet_box()
        self._fill_invoice_speeds()
        self._check_ai()
        if self.cur.docs:
            self._remerge()
        self._apply_speed_rule()  # (a job without documents too: the rule may be another now)
        self._remerge_others()
        self._redate(force=True)  # (Settings -> Dates: the business days, the evening hour, the holidays)
        self._update_status()
        for act in self.theme_actions.actions():  # (View -> Theme and Show the yin-yang, as Settings has them now)
            act.setChecked(act.data() == (self.s.theme if self.s.theme in ("light", "dark") else "system"))
        self.yin_action.setChecked(self.s.show_yin)
        self._records_changed()  # (an open Records window: its tick boxes and colours, Settings -> Records; the count)
        self._show_known_update()  # (Settings -> Options: the daily look turned off hides the Update button)
        self._check_old_days()  # (another sheet named, or business days turned on again)

    # ------------------------------------------------------- outputs and records
    def _auto_runsheet(self) -> None:
        """Ticks the Run sheet box when two or more reporters wrote the transcripts of a case (the initials at
        the foot of the pages, the names on the title pages: batch.case_reporters, the days of a case together), and
        unticks it when one did. Only for these jobs: Settings.outputs keeps the user's own choice, and once the
        user clicks the box it is left as they set it (until New job). Without a transcript the box is as saved."""
        days: dict[tuple, list[Job]] = {}
        for job in self.jobs:
            if job.transcripts():
                days.setdefault(job.name_key(), []).append(job)
        cases = {k: case_reporters(js) for k, js in days.items()}
        for job in self.jobs:  # ticked for a case of several reporters, the box leaves out the others
            job.runsheet_unneeded = not self._runsheet_touched and len(cases.get(job.name_key(), ())) == 1
        self._runsheet_auto = any(len(who) > 1 for who in cases.values()) if cases else None
        box = getattr(self, "output_boxes", {}).get("runsheet")  # (none while the window is built)
        if self._runsheet_touched or box is None:
            return
        want = "runsheet" in self._outputs()
        if box.isChecked() != want:
            with QSignalBlocker(box):
                box.setChecked(want)
            self._refresh_outputs()
            self._refresh_job_labels()

    def _ticked(self, key: str) -> bool:
        """The output's box in the Outputs box is ticked (False for an output the courthouse doesn't make)."""
        box = self.output_boxes.get(key)
        return box is not None and box.isChecked()

    def _outputs(self) -> list[str]:
        """The outputs to make (keys of OUTPUTS): Settings.outputs, the boxes ticked in the Outputs box, but for
        the run sheet, ticked or not for these jobs by how many reporters wrote them (_auto_runsheet) until the
        user clicks its box."""
        out = [k for k in OUTPUTS if k in self.s.outputs and k != "runsheet"]
        auto = None if self._runsheet_touched else self._runsheet_auto
        if "runsheet" in OUTPUTS and (("runsheet" in self.s.outputs) if auto is None else auto):
            out.append("runsheet")
        return out

    def _outputs_changed(self, key: str = ""):
        """The user ticked or unticked an output: the outputs ticked are saved straight away, as the default for
        next time. A run sheet ticked by _auto_runsheet keeps the saved choice for it until the user clicks it."""
        if key == "runsheet":
            self._runsheet_touched = True
            self._auto_runsheet()  # (every case gets the run sheet the user asked for)
        saved = [k for k in OUTPUTS if self.output_boxes[k].isChecked()]
        if not self._runsheet_touched and self._runsheet_auto is not None:
            saved = [k for k in saved if k != "runsheet"] + (["runsheet"] if "runsheet" in self.s.outputs else [])
        self._set_opt("outputs", [k for k in OUTPUTS if k in saved])
        self._update_status()  # (an output unticked takes its problems off the list of what to check)

    def _refresh_outputs(self):
        """Greys out the options of the outputs not ticked, and hides the card of one (the Invoice card) while it
        is unticked or nothing is loaded. For the current job: the tooltips and warnings of the invoice (it needs
        pages to bill: a transcript, or pages typed) and the run sheet (it needs a transcript), the minute
        agreement's speed and the Invoice card's parties and prices; for every day ticked (_considered), the Who
        pays what, Who ordered what and Who transcribed what cards; the counts beside the outputs' boxes and the
        Generate buttons' words (_show_counts) and, when open, the Excerpts window (_follow_excerpts)."""
        if not getattr(self, "_outputs_built", False):  # still building the window
            return
        for key, body in self.output_opts.items():
            body.setEnabled(self.output_boxes[key].isChecked())
        job = self.cur
        for key, fr in self.output_cards.items():
            fr.setVisible(self.output_boxes[key].isChecked() and not job.is_empty())
        pages = job.invoice_pages()
        if "invoice" in self.output_boxes:
            self.output_boxes["invoice"].setToolTip(
                "An invoice for each ticked attorney, priced from the rate sheet" if pages else
                "Needs pages to bill: a transcript PDF among the inputs, or the pages\n"
                "typed in Est. number of pages (a number read from an e-mail isn't billed).")
        if "runsheet" in self.output_boxes:
            self.output_boxes["runsheet"].setToolTip(
                "Adds the transcript's takes (who wrote which pages, from the initials on each page)\n"
                "to the case's run sheet in Excel, or starts one (Settings → Run sheet)" if job.transcript_pages()
                else "Needs a transcript PDF among the inputs (the initials on its pages).\n"
                     "Jobs without one get no run sheet.")
            self.rs_info.setText("" if job.transcript_pages() or not job.docs else
                                 "⚠ no transcript PDF: no run sheet for this job")
            self.rs_info.setVisible(bool(self.rs_info.text()))
        self._show_pages_help(job)
        self._show_form_speed()
        self._show_who_pays(self._show_invoice_prices()[0])  # (the invoices priced once for both)
        self._show_counts()
        self._show_orders()
        self._show_takes()
        self._update_header_buttons()
        self._follow_excerpts()
        self._card_summaries()  # (a folded card's line: the speed and its price may have changed)

    def _show_invoice_prices(self) -> tuple[list | None, int]:
        """The Invoice card's parties, Includes line (Peripherals…) and prices for the current job. When Generate
        all bills it with other days of its case, or billed it with them already ("Invoiced already ✓"), the first
        line says so and the prices are the joint invoice's (_invoice_group).
        When the attorneys don't all pay the same (they ordered different days or pages), a line per attorney. A
        firm of a split order whose speed Generate still asks is priced at the form's speed, with a "?"
        (_price_lines). No prices while the invoice is held (whose pages, Excerpts... to check, more speeds on the
        same pages than are billed, a day of the case nobody is ticked on, unless the day is invoiced already).
        Returns the invoices priced (invoice.firm_invoices; None when there are none to price, the card saying
        why), for Who pays what and the Invoice card's link to the math, and how many days they cover."""
        from ..batch import joint_invoice
        job = self.cur
        group = self._invoice_group()
        self._note_trial_rules(group)  # (priced as an earlier invoice of the trial was, if there is one)
        case, opts = joint_invoice(group)
        split = job.portions is not None  # Excerpts... decides this day's parties (or needs checking)
        check = job.portions_check()
        with QSignalBlocker(self.inv_parties):
            # never fewer parties than the attorneys ticked on a day: each would pay a whole original
            self.inv_parties.setMinimum(max([1] + [len(j.ticked_keys()) for j in group]))
            self.inv_parties.setValue(opts.parties)
        self.inv_parties.setEnabled(not split)
        self.inv_parties.setToolTip("This day's pages are split between the attorneys under Excerpts…" if split
                                    else PARTIES_TIP.format(**_who_splits(self.s)))
        why = self._who_ordered_why(job)
        self.inv_who.setEnabled(not why or split)  # rows that need checking can always be opened, to fix them
        self.inv_who.setText("⚠ Excerpts…" if check else "✓ Excerpts…" if split else "Excerpts…")
        self.inv_who.setToolTip(check + "\n(click to check it)" if check else why or WHO_TIP.format(**_who_splits(self.s)))
        owned = self._show_whose_pages(job, group)
        self._show_speeds_lock()
        with QSignalBlocker(self.inv_detail):
            self.inv_detail.setChecked(job.invoice_detail)  # this job's own (used when it is generated alone)
        from ..batch import index_question, typed_days
        from ..invoice import index_priced
        on_typed = any(j.index_on_typed for j in typed_days(group))
        self.inv_peripherals_info.setText(_lines(self._peripherals_text(
            opts, index_question(group, self.s), on_typed, len(group) > 1, index_priced(case, self.s, opts))))
        self.inv_info.setToolTip("")
        if not job.invoice_pages() and not owned:
            if not job.docs:
                self.inv_info.setText("")  # an empty job has nothing to warn about yet
            elif job.transcript_pages() or job.pages_typed():
                self._warn("⚠ the Pages field says 0: no invoice for this job")
            else:
                self._warn("⚠ no transcript PDF: type the pages to bill in Est. number of pages")
            return None, 1
        if owned:  # whose pages to bill isn't known: no prices until Whose pages... says
            # (the reason names the file, which can be very long: it goes in the tooltip)
            self.inv_info.setText("⚠ Choose whose pages to bill (Whose pages…)")
            self.inv_info.setToolTip(f"{owned[0].upper()}{owned[1:]}\n\n"
                                     "No invoice is made for this job until you say whose pages to bill.")
            return None, 1
        if check:  # no prices: nothing is billed for this day until the rows are checked
            self._warn(f"⚠ {check}")
            self.inv_info.setToolTip("No invoice is made for this day until Excerpts… is checked:\n"
                                     "tick who ordered which pages, or choose \"Remove this day's excerpts\".")
            return None, 1
        if job.speed_problems():  # more speeds on the same pages than are billed for now: no prices either
            self._warn(f"⚠ {job.invoice_hold()}")
            self.inv_info.setToolTip("No invoice is made for this day until a firm's speed is changed\n"
                                     "(Who ordered what, or Excerpts… for a run of pages).")
            return None, 1
        if job.unnamed_problem():  # parties not named at different speeds: nobody's share can be worked out
            self._warn(f"⚠ {job.unnamed_problem()}")
            self.inv_info.setToolTip("No invoice is made for this day until every party is named\n"
                                     "(ticked in the attorney table, its speed set in Who ordered what) or Parties is changed.")
            return None, 1
        from ..batch import group_problem
        # (invoiced already: its invoice was made, so nothing holds it; priced as things are now, below)
        held = group_problem(group) if not job.invoiced else ""
        if held:  # nobody ticked on a day of the joint invoice: no prices until it says who ordered it
            self._warn(f"⚠ Generate all: no invoice for these {len(group)} days yet:\n{held}")
            self.inv_info.setToolTip("Tick the attorneys who ordered that day's pages (in its attorney table);\n"
                                     "until then Generate all makes the other outputs but no invoice for the case.")
            return None, 1
        try:
            from ..invoice import firm_invoices
            firms = firm_invoices(case, self.s, opts)
        except Exception as e:  # a broken rate sheet must not break the window
            self._warn(f"⚠ {e}")
            return None, 1
        if not firms:  # (every attorney ticked was invoiced already by a run that stopped, or ordered no pages)
            self._warn("No invoice left to make: every attorney ticked is invoiced already or ordered no pages")
            return None, 1
        from ..batch import speeds_to_ask
        asked = {a.key for a in speeds_to_ask(group, self.s)}  # (priced at the form's speed meanwhile)
        prices, per_firm = _price_lines(firms, self.case.get("delivery"), asked)
        ask_tip = ("A speed with \"?\" is not set yet: firms that order the same pages commit to a speed.\n"
                   "Generate asks for it (set it beforehand in Who ordered what)." if asked else "")
        self.inv_info.setToolTip(ask_tip)
        if len(group) > 1 and job.invoiced:  # (the days one invoice billed already: batch.bill_sets)
            self.inv_info.setText(f"Invoiced already ✓: one invoice for {len(group)} days ({opts.pages} pp.)\n"
                                  f"{prices}")
            self.inv_info.setToolTip(
                f"These {len(group)} days were billed on one invoice for each attorney: Generate all won't bill\n"
                f"them again. \"{self._one_words(job)}\" bills the day shown again, alone ({job.invoice_pages()} pp.).")
        elif len(group) > 1:
            self.inv_info.setText(f"Generate all: one invoice for {len(group)} days ({opts.pages} pp.)\n{prices}")
            forms = ("It also makes one minute agreement per attorney (its days and pages) and one MOFR\n"
                     "for these days (Settings → Options).\n" if self.s.forms_per_case else
                     "The minute agreements and MOFRs are made for each day (Settings → Options).\n")
            self.inv_info.setToolTip(
                f"Generate all bills these {len(group)} days of the case on one invoice for each attorney\n"
                f"(the days it is ticked on, and the pages it ordered).\n{forms}"
                f"\"{self._one_words(job)}\" makes the day shown alone ({job.invoice_pages()} pp.)."
                + (f"\n\n{ask_tip}" if ask_tip else ""))
        else:
            self.inv_info.setText(f"{opts.pages} pp." + ("\n" if per_firm else " · ") + prices)
        return firms, len(group)

    def _warn(self, text: str) -> None:
        """A warning in the Invoice card's Prices row, its long lines broken (a file name too): a label is as
        wide as its longest line, and a long one made the card, and the window, wider than the screen."""
        self.inv_info.setText("\n".join(textwrap.fill(line, WARN_WIDTH) for line in text.split("\n")))

    def _form_speed_text(self) -> str:
        """The one speed and rate the minute agreement form names: 'Expedite · $5.65 a page'; with firms that
        have a speed of their own (Who ordered what), 'Each firm's own speed (Expedite · $5.65 a page without
        one)'."""
        speed, rate = self.case.get("delivery"), self.case.get("rate")
        if not speed:
            text = "(none: choose one under Speed, pages and dates)"
        elif not rate:
            text = f"{speed} · no rate (type it under Speed, pages and dates)"
        else:
            text = f"{speed} · ${rate.lstrip('$')} a page"
        return f"Each firm's own speed ({text} without one)" if self.cur.ordered_speeds() else text

    def _show_form_speed(self) -> None:
        """The Speed row of the Minute agreement panel under Options ▾ (see _form_speed_text)."""
        self.ag_speed.setText(self._form_speed_text())

    def _show_who_pays(self, firms: list | None) -> None:
        """The "Who pays what" card: the invoices of every day ticked in the list (_considered; the user,
        2026-10-23: not just the day shown's; batch.bill_sets: those Generate all makes, and those made already),
        each a row per firm (the firm, its pages, how it ordered them, what it pays at each speed and what that
        comes to a page, or its one speed and amount when every invoice bills one; invoice_math.who_pays) under the
        speed and rate the minute agreement forms name (each firm's own, when some have one: _pay_html). With
        several invoices, a heading over each names its days (and case, with several cases), "invoiced already"
        for one made; an invoice that can't be priced yet says why (_price_set), and so does a day whose invoice
        waits for a choice. Shown while Invoice is ticked and a day ticked has pages to bill (or whose pages to
        bill is to be chosen). firms: what _show_invoice_prices priced for the job shown's invoice
        (_invoice_group), for the Invoice card's link to its math, and used again here rather than priced twice
        (this runs at every keystroke)."""
        import html
        from ..batch import bill_sets
        from ..invoice_math import page_rates, who_pays
        considered = self._considered()
        self._inv_firms = list(firms or []) if self._ticked("invoice") else []
        self.inv_form.setRowVisible(self.inv_math, bool(self._inv_firms))  # (the Invoice card's link to the math)
        on = self._ticked("invoice") and any(j.invoice_pages() or j.ownership_problem() for j in considered)
        self.pays_card.setVisible(on)
        self._pay_firms = []
        if not on:
            return
        tip = PAYS_TIP.format(**_who_splits(self.s))  # (as Settings say now)
        self.pays_title.setToolTip(tip)
        self.pays.setToolTip(tip)
        sets, held = bill_sets(considered, self.s)
        own = [id(j) for j in self._invoice_group()]
        several = len(_cases([j for g in sets for j in g])) > 1
        blocks = []
        for g in sets:
            if [id(j) for j in g] == own and firms is not None:
                got, why = list(firms), ""  # (the job shown's: priced already)
            else:
                got, why = self._price_set(g)
            head = ""
            if len(sets) > 1 or several:
                head = ", ".join(j.case.get("dates") or j.title() for j in g)
                if several:
                    head = f"{g[0].case.get('case_name') or g[0].title()}: {head}"
                if g[0].invoiced:
                    head += " · invoiced already ✓"
                head = f"<b>{html.escape(head)}</b><br>"
            elif g[0].invoiced:
                head = "Invoiced already ✓ (Generate all won't bill " + ("it" if len(g) == 1 else "them") + " again)<br>"
            if got:
                case = self.case if any(j is self.cur for j in g) else g[0].case
                blocks.append(head + _pay_html(
                    who_pays(got), case.get("delivery"), case.get("rate"), len(g), page_rates(got),
                    _own_words(got, self.s), any(j.speeds or j.ordered_speeds() for j in g), self._marks(g[0]),
                    link=False))
                self._pay_firms += got
            else:
                why = why or (self.inv_info.text() if [id(j) for j in g] == own else "") or "Nothing to bill yet."
                blocks.append(head + html.escape(why).replace("\n", "<br>"))
        for j in held:  # (no invoice until it is checked: Generate all leaves it out)
            blocks.append(html.escape(f"⚠ No invoice yet for {j.case.get('dates') or j.title()}: {j.invoice_hold()}"))
        if not blocks:
            blocks.append(html.escape(self.inv_info.text() or "Nothing to bill yet.").replace("\n", "<br>"))
        text = "<br><br>".join(blocks) + (f"<br>{MATH_LINK}" if self._pay_firms else "")
        if text != self.pays.text():  # (set only when it changes: it is worked out at every keystroke)
            self.pays.setText(text)

    def _price_set(self, group: list[Job]) -> tuple[list, str]:
        """The invoices of these days, as Who pays what shows an invoice other than the job shown's (batch.
        joint_invoice, invoice.firm_invoices), priced as their trial was before (_note_trial_rules); or none and
        why: a day nobody is ticked on (batch.group_problem; not for days invoiced already), a rate sheet that
        can't price them, or nobody left to bill."""
        from ..batch import group_problem, joint_invoice
        from ..invoice import firm_invoices
        # (days invoiced already are one set, all invoiced or none: batch.bill_sets)
        held = group_problem(group) if not group[0].invoiced else ""
        if held:
            return [], f"⚠ No invoice for these {len(group)} days yet: {held}"
        self._note_trial_rules(group)
        try:
            case, opts = joint_invoice(group)
            firms = firm_invoices(case, self.s, opts)
        except Exception as e:  # (a broken rate sheet must not break the window)
            return [], f"⚠ {e}"
        return firms, "" if firms else ("No invoice left to make: every attorney ticked is invoiced already or "
                                        "ordered no pages")

    def _show_counts(self) -> None:
        """Beside each ticked output's box, how many files Generate makes of it for the job shown and, with several
        jobs, how many Generate all makes of the days it makes (_to_generate) ("(2 · all 6)": _count_short; in
        words in its tooltip: _count_text; batch.output_counts); it follows every tick (attorneys, outputs,
        Excerpts...). Hidden while the output is unticked, and for an empty job when it is the only one. The
        Generate buttons' words follow too (_generate_labels). What Generate all makes is kept in _all_counts for
        _refresh_job_labels, so a keystroke counts every job once, not twice."""
        from ..batch import output_counts
        outputs, job = self._outputs(), self.cur
        shown = not job.is_empty()
        # (again: the job shown, generated alone, bills a day invoiced already once more; Generate all doesn't)
        counts = output_counts([job], outputs, self.s, again=True) if shown else None
        everyone = None
        if len(self.jobs) > 1:
            everyone = output_counts(self._to_generate(), outputs, self.s)
        self._all_counts = everyone
        self._job_counts = counts
        self._generate_labels(everyone)
        held = ""
        if shown and "invoice" in outputs and job.invoice_hold():  # (Generate asks first: see invoice_hold)
            held = {"portions": "Excerpts… is checked", "ownership": "Whose pages… is chosen",
                    "unnamed": "Parties is checked"}.get(job.hold_kind(), "the speeds are checked")
        for key, label in self.output_counts.items():
            on = key in outputs and (shown or everyone is not None)
            why = held if key == "invoice" else ""
            label.setText(_count_short(key, counts, everyone, why) if on else "")
            label.setToolTip(_count_text(key, counts, everyone, why) if on else "")
            label.setVisible(on)

    def _generate_labels(self, everyone: dict | None = None) -> None:
        """The Generate buttons say what they make (the user, 2026-10-22: in a trial, "Generate this job (14 files)"
        beside "Generate all (10 files)" was a puzzle): for the days of one case, the day shown and the days
        ("Generate 6/4/2026 only", "Generate all 3 days"), else the job shown and the jobs ("Generate this job",
        "Generate all 5 jobs"; "the 2 days ticked" when some are unticked, "the 1 day left" when some are made
        already: _to_generate), then how many files, and in the tooltip what they are (_files_words). Generate alone
        for a single job. everyone: batch.output_counts of the jobs Generate all makes (None: its button is left as
        it is)."""
        multi = len(self.jobs) > 1
        counts = getattr(self, "_job_counts", None) if not self.cur.is_empty() else None
        n = sum(counts.values()) if counts else 0
        shown = [j for j in self.jobs if not j.is_empty()]
        one_case = multi and len(_cases(shown)) == 1
        words = self._one_words(self.cur)
        self.fill_btn.setText(f"{words}  ({plural(n, 'file')})" if n else words)
        what = _files_words(counts)
        self.fill_btn.setToolTip((f"Makes {what}" if what else "Makes nothing yet") +
                                 (", for the day shown alone" if one_case else ", for the job shown alone" if multi
                                  else "") + " (Ctrl+Enter)")
        if everyone is None or not multi:
            return
        ticked = [j for j in shown if j.include]
        todo = {id(j) for j in self._to_generate()}
        chosen = [j for j in ticked if id(j) in todo]
        unit = "day" if one_case else "job"
        scope = (f"the {plural(len(chosen), unit)} left" if len(chosen) < len(ticked) else
                 f"all {plural(len(chosen), unit)}" if len(chosen) == len(shown) else
                 f"the {plural(len(chosen), unit)} ticked")
        files = sum(everyone.values())
        self.fill_all_btn.setText(f"Generate {scope}  ({plural(files, 'file')})" if chosen else "Generate all")
        what = _files_words(everyone)
        self.fill_all_btn.setToolTip(
            (f"Makes {what}" if what else "Makes nothing yet") + f", for {scope} (Ctrl+Shift+Enter)" if chosen else
            f"Every {unit} ticked is made already (✓). To make one again, click it, then "
            + ("Generate (its date) only." if one_case else "Generate this job.")
            if ticked else "Nothing is ticked in the list (Ctrl+Shift+Enter)")
        self._place_foot()  # (longer or shorter names: the row may hold them now, or no longer)

    def _one_words(self, job: Job) -> str:
        """What the Generate button (the job shown alone) says for this job, without its count, as the texts that
        name it say it too: "Generate 6/4/2026 only" for a day of the one case loaded, "Generate this job" among
        other jobs, "Generate" for a job alone (_generate_labels)."""
        if len(self.jobs) < 2:
            return "Generate"
        if len(_cases([j for j in self.jobs if not j.is_empty()])) == 1:
            return f"Generate {self._job_date(job) or 'this day'} only"
        return "Generate this job"

    def _show_live_math(self, link: str = "") -> None:
        """The links of the Who pays what card and the Invoice card ("invoice"): the whole math of the invoices
        the card shows (Who pays what: every day ticked; the Invoice card: the job shown's invoice), as things
        stand (preview.MathDialog)."""
        firms = self._inv_firms if link == "invoice" else self._pay_firms
        if not firms:
            return
        from .preview import MathDialog
        try:
            dlg = MathDialog([(f, "") for f in firms], self.s, self, live=True)
        except Exception as e:  # (as _show_math: the math of an odd invoice must not break the window)
            logfile.error("could not show the math", e)
            self._toast(f"The math could not be shown ({type(e).__name__}).")
            return
        before = self.s.save_math
        dlg.exec()
        self._math_saves_kept(before)

    def _show_pages_help(self, job: Job) -> None:
        """The lines under Est. number of pages: what the number counts, and a warning when the transcript PDF's
        readings of its pages disagree (Job.pages_help); the row is hidden for a job without a transcript."""
        line, warning = job.pages_help()
        # (the row first: showing it shows every label in it)
        self.order_form.setRowVisible(self.pages_help_row, bool(line or warning))
        self.pages_help.setText(line)
        self.pages_help.setVisible(bool(line))
        self.pages_warn.setText(warning)
        self.pages_warn.setVisible(bool(warning))

    def _show_whose_pages(self, job: Job, group: list[Job] | None = None) -> str:
        """The Invoice card's Billed row, shown for a transcript of several reporters: "You wrote 65 of 153 total
        pages" (the whole transcripts' own pages, the word index left out; "Billing 153 of 153 total pages (chosen
        under Whose pages…)" when Whose pages... chose others, "The Pages field: 40 of 153 total pages" when typed),
        and the Whose pages... button. group: the days Generate all bills on one invoice with the job
        (_invoice_group): the row is the whole invoice's, "One invoice for 3 days: you wrote 144 of 341 total
        pages", a line per day in its tooltip, so it stays the same whichever of them is shown (the user,
        2026-10-21: it said the day's pages, though the invoice is the days'). Returns why the invoice is held
        until Whose pages... says ("" when it isn't): the job's own."""
        days = list(group) if group and len(group) > 1 else [job]
        shared = any(j.shared_transcripts() for j in days)
        held = job.ownership_problem()
        self.inv_form.setRowVisible(self.inv_pages_row, bool(shared or held))
        if not (shared or held):
            return ""
        whole = sum(j.transcript_pages() for j in days)
        billed = sum(j.invoice_pages() for j in days)
        mine = all(j.basis_of(d)[0] == "me" for j in days for d in j.transcripts())
        if held:
            text = "Whose pages?"
        elif any(j.pages_typed() for j in days):
            text = f"The Pages field: {billed} of {whole} total pages"
        elif mine:
            text = f"You wrote {billed} of {whole} total pages"
        else:
            text = f"Billing {billed} of {whole} total pages\n(chosen under Whose pages…)"
        if len(days) > 1 and not held:
            text = f"One invoice for {len(days)} days: {text[0].lower()}{text[1:]}"
        self.inv_pages_info.setToolTip("\n".join(
            f"{j.case.get('dates') or j.title()}: {j.invoice_pages()} of {j.transcript_pages()} pages"
            for j in days) if len(days) > 1 else "")
        others = list(dict.fromkeys(r for j in days for r in j.bill_reporters() if r != "me"))
        if others and not held:  # invoices in other reporters' names too (Whose pages... ticked them)
            if not any("me" in j.bill_reporters() for j in days):
                text = "Not yours"
            text += "\nalso " + ", ".join(
                f"{r.upper()}'s ({sum(j.billed_by(r).invoice_pages() for j in days if r in j.bill_reporters())})"
                for r in others)
            bare = [r.upper() for r in others if not self.s.reporter(r).has_details()]
            if bare:
                text += f"\n⚠ no invoice details for {', '.join(bare)}: Settings → Run sheet → Invoices for"
        self.inv_pages_info.setText(_lines(text))
        self.inv_whose.setText("⚠ Whose pages…" if held else "Whose pages…")
        self.inv_whose.setToolTip(held or WHOSE_TIP)
        return held

    def _whose_pages(self, job: Job, reason: str = "") -> bool:
        """Whose pages…: which pages of the job's transcripts of several reporters (or of the one held, see
        Job.ownership_problem) are billed. Returns True when OK was clicked and the job is still there. The job
        is merged again, so the Pages field shows the pages billed."""
        from .dialogs import PagesOwnerDialog
        docs = [d for d in job.transcripts() if len(d.reporters()) > 1 or (
            d.reporters() and "*" not in job.basis_of(d) and not job.billed_pages_of(d))]
        if not docs:
            QMessageBox.information(self, "Whose pages", "The transcripts of this job have one reporter's pages "
                                    "each (or no initials at all): they are billed in full.")
            return False
        dlg = PagesOwnerDialog([(d.key(), d.ing.name, d.owners()) for d in docs], set(job.own), job.page_basis,
                               job.front_owner, reason, self, self.s)
        if not dlg.exec() or not any(job is j for j in self.jobs):
            return False
        basis, front = dlg.values()
        job.page_basis = {**job.page_basis, **basis}
        job.front_owner = {**job.front_owner, **front}
        self._remerge(job)  # the Pages field shows the pages billed
        self._update_status()
        return True

    def _show_orders(self):
        """The "Who ordered what" card: for every day ticked in the list (_considered; the user, 2026-10-23: not
        just the day shown), a row per ordering attorney (one per Attorney.key: a firm's row naming several
        attorneys is one) saying which pages it ordered (the whole day, or which stretch), the speed it ordered
        them at (a box to choose it in: _speed_box), who else on its invoice also ordered them and how many of your
        pages that bills, with its firm's mark (_mark_row). The days come earliest first; with days of several
        cases, by case, its name before the date. A row's invoice is that of its day's set (batch.bill_sets: the
        days Generate all bills together, or those an invoice made already billed, whose dates have a ✓); its
        "Also ordered" and marks are of those days. Nothing is left to guess: an
        attorney not ticked on a day says it orders nothing, a day waiting for a choice says so, a day nobody is
        ticked on holds its case's invoice back (batch.group_problem), and a row that bills none of the pages says
        it gets no invoice for them. The line above says what Generate all bills and leaves out (_orders_where).
        Shown while Invoice is ticked and some job has a transcript or pages to bill; never an empty table (a row
        says why a day has nothing to bill, and with no day ticked the table is hidden)."""
        from ..batch import OrderLine, bill_sets, group_attorneys, group_problem
        from ..pdfout import short_caption
        job = self.cur
        on = self._ticked("invoice") and any(j.transcript_pages() or j.invoice_pages() for j in self.jobs)
        self.orders_card.setVisible(on)
        if not on:
            return
        days = [j for j in self._considered() if j.transcript_pages() or j.invoice_pages()]
        several = len(_cases(days)) > 1  # (by case, then: its name before each date)
        days.sort(key=lambda j: self._job_sort_key(j, 0 if several else JOB_DATE))
        sets, held = bill_sets(days, self.s)
        set_of = {id(j): g for g in sets for j in g}
        # by id() of a set's list of days: its attorneys, its problem (group_problem), its firms' marks
        everyone: dict[int, list] = {}
        stuck_of: dict[int, str] = {}
        marks_of: dict[int, dict] = {}
        lines: list[tuple[Job, list[Job], OrderLine]] = []
        for j in days:
            g = set_of.get(id(j), [j])  # (a day held, or with nothing to bill: alone)
            if id(g) not in stuck_of:
                everyone[id(g)] = group_attorneys(g) if len(g) > 1 else []
                # a day nobody is ticked on: no invoice for its case (a set made already is made)
                stuck_of[id(g)] = group_problem(g) if not g[0].invoiced else ""
                marks_of[id(g)] = self._marks(g[0])
            if stuck_of[id(g)] and not j.ticked_keys():
                got = [OrderLine(j.case.get("dates"), "", "(nobody ticked)", note=NOBODY_DAY)]
            else:
                got = j.order_lines(everyone[id(g)])
            if not got:  # say why the day has nothing to bill rather than leave it out
                why = (j.invoice_hold() or ("the Pages field says 0: no invoice for this job"
                                            if not j.invoice_pages() else
                                            "tick the attorneys who ordered in the Attorneys table"))
                got = [OrderLine(j.case.get("dates"), "", "—" if j.ticked_keys() else "(nobody ticked)",
                                 note="⚠ " + why)]
            lines += [(j, g, line) for line in got]
        mine = all(j.basis_of(d)[0] == "me" for j in days for d in j.transcripts())
        none = "none of these pages are yours" if mine else "none of these pages are the ones billed"
        t = self.orders
        t.setRowCount(0)
        t.setVisible(bool(lines))
        for j, g, line in lines:
            r = t.rowCount()
            t.insertRow(r)
            stuck = stuck_of[id(g)]
            if line.note:
                ordered, printed, shared, billed = line.note, "", "", ""
            else:
                part = ", ".join(f"{a}–{b}" if a != b else str(a) for a, b in line.spans)
                if line.typed:  # the Pages field's number is billed, not the transcript's pages: say so
                    ordered = (f"every page (Pages field typed: {line.pages} billed)" if line.whole else
                               f"excerpt: {part} of the {line.pages} pages typed in the Pages field")
                else:
                    ordered = f"every page, 1–{line.pages}" if line.whole else f"excerpt: {part} of {line.pages}"
                printed = ", ".join(p for p in line.printed if p)
                shared = "; ".join(f"{self._name_of(g, o)} ({', '.join(f'{a}–{b}' for a, b in spans)})"
                                   for o, spans in line.shared.items()) or "nobody (alone)"
                if line.parties:
                    shared = f"split {line.parties} ways (the Parties number)"
                billed = str(line.billed)
                if not line.billed:  # an attorney with no pages billed gets no invoice for them (see firm_pages)
                    elsewhere = any(o.key == line.key and o.billed and not o.note for _, h, o in lines if h is g)
                    billed = f"0 ({none})" if elsewhere else f"0 (no invoice: {none})"
                elif stuck:
                    billed += " (held: no invoice yet, see ⚠)"
            day = line.day + (" ✓" if j.invoiced else "")
            if several:
                day = f"{short_caption(j.case.get('case_name') or j.title(), 24)} · {day}"
            for c, text in enumerate((day, line.name, ordered, "", printed, shared, billed)):
                it = QTableWidgetItem(text)
                it.setData(Qt.UserRole, id(j))
                if (c == 2 and line.note.startswith("⚠")) or (c == 6 and "no invoice" in text):
                    it.setForeground(QColor("#d97706"))  # (the amber of the warnings, light and dark)
                if c == 0 and j.invoiced:
                    it.setToolTip("✓ Invoiced already: Generate all won't bill this day again")
                t.setItem(r, c, it)
            if not line.note and line.key:
                t.setCellWidget(r, ORDER_SPEED, self._speed_box(j, line.key))
            self._mark_row(t, r, marks_of[id(g)].get(line.key))
        t.resizeColumnsToContents()
        for c in range(2, len(ORDER_COLS)):  # long notes wrap instead of making the table wider than the card
            t.setColumnWidth(c, min(t.columnWidth(c), z(320)))
        t.setColumnWidth(0, min(t.columnWidth(0), z(260)))  # (a case's name before the date)
        last = len(ORDER_COLS) - 1  # (as wide as its heading at least, and to the card's edge when there's room)
        room = t.viewport().width() - sum(t.columnWidth(c) for c in range(last))
        t.setColumnWidth(last, max(t.columnWidth(last), t.horizontalHeader().sectionSizeHint(last), room))
        t.resizeRowsToContents()
        self._fit_orders_height()
        whose = bool(job.ownership_problem()) and not job.portions_problem()
        self.orders_edit.setText("⚠ Whose pages…" if whose else "Excerpts…")
        self.orders_edit.setToolTip(
            "Choose whose pages of the day shown are billed (the same as Whose pages… in the Invoice card)"
            if whose else "Who ordered which pages of each day of the case, as one table (the same as Excerpts… "
                          "in the Invoice card)")
        # (how the rows are billed: the ? beside it, ORDERS_HELP)
        self.orders_info.setText(self._orders_where(days, sets, held, stuck_of))

    def _orders_where(self, days: list[Job], sets: list[list[Job]], held: list[Job], stuck: dict) -> str:
        """The line above Who ordered what's table: what Generate all bills of the days it shows ("Generate all
        bills these 2 days of the case on one invoice per attorney."), which are invoiced already ("Invoiced
        already (✓): 6/3/2026. Generate all won't bill it again."), whose invoice waits for a choice, and which
        days are unticked, so left out ("Unticked, so left out: 6/4/2026."). A case with a day nobody is ticked on
        comes first, with ⚠. sets, held: batch.bill_sets of the days; stuck: each set's batch.group_problem, by
        id() of its list of days."""
        def dates(jobs: list[Job]) -> str:
            return ", ".join(j.case.get("dates") or j.title() for j in jobs)

        parts = [f"⚠ No invoice for these {len(g)} days of the case yet: {stuck[id(g)]}."
                 for g in sets if stuck.get(id(g))]
        billing = [g for g in sets if not g[0].invoiced and not stuck.get(id(g))]
        joint = [g for g in billing if len(g) > 1]
        if len(joint) == 1:
            parts.append(f"Generate all bills these {len(joint[0])} days of the case on one invoice per attorney" + (
                ", with one minute agreement per attorney and one MOFR for them." if self.s.forms_per_case else
                "."))
        elif joint:
            parts.append(f"Generate all bills each case's days on one invoice per attorney ({len(joint)} cases).")
        billed = [j for g in sets if g[0].invoiced for j in g]
        if billed:
            again = f" ({self._one_words(self.cur)} does)" if any(j is self.cur for j in billed) else ""
            parts.append(f"Invoiced already (✓): {dates(billed)}. Generate all won't bill "
                         f"{'it' if len(billed) == 1 else 'them'} again{again}.")
        if held:
            parts.append(f"⚠ No invoice yet for {dates(held)}: {held[0].invoice_hold()}.")
        left_out = [j for j in self.jobs if not j.include and (j.transcript_pages() or j.invoice_pages())]
        if left_out and len(self.jobs) > 1:
            parts.append(f"Unticked, so left out: {dates(left_out)}.")
        if not days:
            parts.append("Nothing ticked in the job list: tick the days to see them here.")
        return "  ".join(parts)

    def _speed_box(self, job: Job, key: str) -> QuietCombo:
        """The "Who ordered what" card's Speed box of one firm on one day: the speed it ordered (Job.speeds, or the
        one set on its only run in Excerpts...), the rate sheet's speeds, cheapest first, and "Chooses on the
        invoice" while it orders alone (a firm sharing pages commits to a speed: until one is set it reads "Speed?
        (Generate asks)", in amber). "Several (set in Excerpts…)" when its runs of the day are at different speeds
        (or some have none); a speed the rate sheet doesn't have is kept, and says so. Choosing one sets it for
        all the firm's pages of the day (Job.set_speed)."""
        box = QuietCombo()
        split = key in job.split_keys()
        runs = {sp for _, keys, speeds in job.runs_ordered() for k, sp in zip(keys, speeds) if k == key}
        # (the speed in effect on its runs: one set on its only run in Excerpts... is the firm's too)
        now = next(iter(runs)) if len(runs) == 1 else job.speeds.get(key, "")
        if len(runs) > 1:
            box.addItem("Several (set in Excerpts…)", "*")
        elif not now and split:
            box.addItem("Speed? (Generate asks)", "?")
            box.setStyleSheet("QComboBox { color: #d97706; }")
        if not split:
            box.addItem("Chooses on the invoice", "")
        for sp in self.s.sheet().speeds:
            box.addItem(sp.name, sp.name)
        sheet = self.s.sheet()
        if now and sheet.find(now) is None:  # (a speed the rate sheet doesn't have: kept, and said)
            box.addItem(f"{now} (not on the rate sheet)", now)
        at = box.findData("*") if len(runs) > 1 else box.findData(sheet.find(now).name if sheet.find(now) else now)
        box.setCurrentIndex(max(0, at))
        box.setToolTip(ORDER_TIPS["Speed"])
        box.activated.connect(lambda _i, j=job, k=key, b=box: self._speed_chosen(j, k, b.currentData()))
        return box

    def _speed_chosen(self, job: Job, key: str, speed) -> None:
        """A speed chosen in the "Who ordered what" card: the firm's speed for all its pages of the day (or none,
        "Chooses on the invoice"); the card, the prices and the Excerpts window follow (after the box's own
        signal: the card is filled again, the box with it)."""
        if speed in ("*", "?", None) or job not in self.jobs:
            return
        job.set_speed(key, speed)
        QTimer.singleShot(0, self._update_status)

    def _show_takes(self) -> None:
        """The Who transcribed what card: a row per take of every day ticked in the list (_considered;
        batch.takes_of: the day, the reporter, the printed pages, how many, the running heads as printed;
        take_cells), a take the user corrected marked ✎ (Row.corrected, _fix_take), and what each reporter wrote
        in all. Shown while a day ticked has a transcript."""
        from ..batch import takes_of
        try:
            rows = takes_of(self._considered(), self.s)
        except Exception as e:  # (a transcript that can't be read again must not break the window)
            log_error("could not list who transcribed what", e)
            rows = []
        self._take_rows = rows
        self.takes_card.setVisible(bool(rows))
        if not rows:
            return
        t = self.takes
        t.setRowCount(len(rows))
        for r, row in enumerate(rows):
            for c, text in enumerate(take_cells(row)):
                item = QTableWidgetItem(text + (" ✎" if row.corrected and c == TAKE_COLS.index("Printed pages") else ""))
                if row.corrected:
                    item.setToolTip("Corrected by you (right-click: undo the corrections of this transcript)")
                t.setItem(r, c, item)
        t.resizeColumnsToContents()
        t.resizeRowsToContents()
        t.setFixedHeight(t.horizontalHeader().height() + sum(t.rowHeight(r) for r in range(t.rowCount())) + 4)
        days = len({row.day for row in rows})
        self.takes_info.setText(f"{'This day' if days == 1 else f'These {days} days'}: " + reporter_totals(rows))

    def _take_doc(self, row) -> tuple[Job, Doc] | None:
        """The job and transcript a take of the card (runsheet.Row) was read from, or None (gone meanwhile, or not
        read from a transcript's pages)."""
        if not row.doc or row.at is None:
            return None
        for job in self._considered():
            for d in job.transcripts():
                if d.key() == row.doc:
                    return job, d
        return None

    def _takes_menu(self, pos) -> None:
        """Right-click on a take of Who transcribed what: correct where it ends, or undo the corrections of its
        transcript."""
        r = self.takes.rowAt(pos.y())
        if not 0 <= r < len(self._take_rows):
            return
        # (the take itself, not its row: a read that ends while the menu is open draws the card again)
        row = self._take_rows[r]
        found = self._take_doc(row)
        menu = QMenu(self)
        fix = menu.addAction("Correct where this take ends…")
        fix.triggered.connect(lambda: self._fix_take(row))
        if found and found[1].fixes:
            undo = menu.addAction("Undo my corrections to this transcript's takes")
            undo.triggered.connect(lambda: self._undo_take_fixes(found[0], found[1]))
        menu.exec(self.takes.viewport().mapToGlobal(pos))

    def _ask_take_change(self, job: Job, doc: Doc, fixes: dict, what: str) -> bool:
        """The "Are you sure?" box before the takes of a transcript change (fixes: Doc.fixes as they would be):
        what changes (`what`), and the user's pages that day before and after; files made already aren't changed.
        True = go ahead."""
        before = job.your_pages()
        kept = doc.fixes
        doc.fixes = fixes
        try:
            after = job.your_pages()
        finally:
            doc.fixes = kept
        text = what
        if before != after and None not in (before, after):
            text += f"\n\nYour pages that day: {before} → {after}."
        if job.made or job.invoiced:
            text += ("\n\nFiles made already aren't changed: Generate this day again for a corrected invoice. A run "
                     "sheet made already keeps its takes as they were (correct it in Excel).")
        box = QMessageBox(QMessageBox.Question, "Correct the takes", text + "\n\nAre you sure?", parent=self)
        yes = box.addButton("Yes, correct it", QMessageBox.AcceptRole)
        no = box.addButton(QMessageBox.Cancel)
        box.setDefaultButton(no)
        box.setEscapeButton(no)
        box.exec()
        return box.clickedButton() is yes and not self._batch_running()

    def _holder(self, doc: Doc) -> Job | None:
        """The job a document is in now (a read that ended while a box was open may have moved it), or None when
        it was removed."""
        return next((j for j in self.jobs if any(d is doc for d in j.docs)), None)

    def _fix_take(self, row) -> None:
        """Correct where this take ends… (Who transcribed what, a take double-clicked or right-clicked: row, its
        runsheet.Row): a page's initials misread or missing put pages in the wrong take. Asks for the printed page
        it really ends on, then, in an "Are you sure?" box, which pages move, where the take ends then (the same
        reporter's next take joins it when it reaches it) and how the user's pages change. Ended later, the pages
        up to there become the take's; ended sooner, its pages after there become the next take's
        (takes.moved_end). The correction is the transcript's (Doc.fixes): the takes, Yours, the pages billed and
        the run sheet follow, and the records of the files made from then on keep it (batch.page_fixes). Refused,
        with a box that says why, for the transcript's last take, a take whose initials aren't known yet (Whose
        pages… says) and pages without printed numbers; nothing is done while Generate all runs."""
        from ..runsheet import title_page_total
        from ..takes import find_takes, fix_marks, moved_end, page_numbers, page_owners
        if self._batch_running():  # (Generate all reads the same pages)
            return
        found = self._take_doc(row)
        title = "Correct where this take ends"
        if found is None:
            QMessageBox.information(self, title, "This take wasn't read from the initials on a transcript's pages "
                                    "(a scan, say), so there is nothing to correct here.")
            return
        job, doc = found
        marks, front, titles = doc.marks(), doc.front(job.front_owner.get(doc.key(), "")), title_page_total(doc.ing)
        takes = find_takes(marks, titles, front)
        take = next((t for t in takes if t.at == row.at), None)
        numbers = page_numbers(marks)
        if take is None:
            return  # (the card is behind: it is shown again with the next change)
        last = take.at + take.pages - 1
        if take is takes[-1]:
            QMessageBox.information(self, title, "This is the transcript's last take: it ends on its last page. To "
                                    "move where it starts, correct where the take before it ends.")
            return
        if not take.initials:
            QMessageBox.information(self, title, "Whose these pages are isn't known yet: Whose pages… (on the "
                                    "Invoice card) says.")
            return
        # (a cover page without a printed number: the take's first numbered page)
        first_no = next((numbers[i] for i in range(take.at, last + 1) if numbers[i] is not None), None)
        if None in (first_no, numbers[-1]):
            QMessageBox.information(self, title, "These pages have no printed page numbers to go by.")
            return
        who, day = row.reporter or take.initials.upper(), take_cells(row)[0]
        now = numbers[last]
        end_no, ok = QInputDialog.getInt(
            self, title, f"{who}'s take{' on ' + day if day else ''} starts on page {first_no} and ends on "
                         f"page {now}.\n\nOn which page does it really end?", now, first_no, numbers[-1])
        if not ok or end_no == now:
            return
        end = next((i for i in range(take.at, len(marks)) if numbers[i] == end_no), None)
        if end is None:
            QMessageBox.information(self, title, f"There is no page {end_no} after the take's first page.")
            return
        try:
            fixes = moved_end(marks, take.at, end, titles, front)
        except ValueError as e:
            QMessageBox.information(self, title, f"The take can't end there: {e}.")
            return
        names = {x.initials: x.reporter for x in self._take_rows if x.doc == doc.key() and x.initials}
        name = lambda ini: names.get(ini) or ini.upper()
        owners = page_owners(marks)
        # (leaving out a page moved_end only pins to whose it is already: it doesn't move)
        pages = sorted(i for i in fixes if owners[i] != fixes[i])
        span = (f"page {numbers[pages[0]]}" if len(pages) == 1 else
                f"pages {numbers[pages[0]]}–{numbers[pages[-1]]}")
        was = sorted({name(owners[i]) for i in pages if owners[i]})
        # where it ends then: the same reporter's next take, when the pages reach it, joins it
        after = next((t for t in find_takes(fix_marks(marks, fixes), titles, front) if t.at == take.at), take)
        ends = numbers[after.at + after.pages - 1]
        what = (f"{who}'s take{' on ' + day if day else ''} will end on page {ends} instead of {now}: {span} "
                f"{'becomes' if len(pages) == 1 else 'become'} {name(fixes[pages[0]])}'s"
                + (f" (now {' and '.join(was)}'s)" if was else "")
                + (f", and the next take of {who}'s, up to page {ends}, joins it." if ends != end_no else "."))
        if not self._ask_take_change(job, doc, {**doc.fixes, **fixes}, what):
            return
        holder = self._holder(doc)  # (a read that ended while the box was open may have moved the transcript)
        if holder is None:
            self._toast("The transcript was removed meanwhile: nothing was corrected.")
            return
        doc.correct(fixes)
        self._takes_changed(holder)

    def _undo_take_fixes(self, job: Job, doc: Doc) -> None:
        """Undo my corrections to this transcript's takes: its pages counted by the initials on them again
        (Doc.fixes emptied), after an "Are you sure?" box. job: where the transcript was when the menu opened; the
        job holding it now is used (_holder: a read that ended meanwhile may have moved it)."""
        job = self._holder(doc)
        if job is None or not doc.fixes or self._batch_running():
            return
        if not self._ask_take_change(job, doc, {}, f"The takes of {doc.ing.name} go back to what the initials on "
                                                   "its pages say."):
            return
        holder = self._holder(doc)
        if holder is not None:
            doc.fixes = {}
            self._takes_changed(holder)

    def _takes_changed(self, job: Job) -> None:
        """After a transcript's takes changed (corrected, or the corrections undone): the job merged again (its
        pages are worked out from the takes as they are now) and every card, the job list's Yours among them,
        shown again."""
        self._remerge(job)
        self._update_status()

    def _save_takes(self) -> None:
        """Save a copy...: the Who transcribed what table as an Excel workbook or a CSV file, where the user
        chooses (in the records folder at first). The rows are those shown when it was clicked (a read that
        ends while the save box is open changes the card)."""
        rows = list(self._take_rows)
        if not rows:
            return
        from ..pdfout import safe_filename
        name = safe_filename(f"Who transcribed what - {self.cur.title() or 'trial'}") + ".xlsx"
        path, chosen = QFileDialog.getSaveFileName(self, "Save who transcribed what",
                                                   str(self.s.records_folder() / name),
                                                   "Excel workbook (*.xlsx);;CSV file (*.csv)")
        if not path:
            return
        try:
            out = save_takes(rows, Path(path), csv_file="csv" in chosen.lower() or
                             path.lower().endswith(".csv"))
        except OSError as e:
            QMessageBox.warning(self, "Not saved", f"The file could not be saved ({e}).\n\nIs it open in Excel?")
            return
        except Exception as e:  # (a header with a character a workbook can't hold, ...)
            log_error("could not save who transcribed what", e)
            QMessageBox.warning(self, "Not saved", f"The file could not be saved ({describe(e)}).")
            return
        self._toast(f"Saved {out.name}")

    def _fit_orders_height(self) -> None:
        """The card's table as tall as its rows (it doesn't scroll up and down), with room for the horizontal
        scroll bar when the columns are wider than the card: else it covers the last row."""
        t = self.orders
        head = t.horizontalHeader()
        bar = t.horizontalScrollBar()
        height = (max(head.height(), head.sizeHint().height()) + sum(t.rowHeight(r) for r in range(t.rowCount()))
                  + 2 * t.frameWidth() + (bar.sizeHint().height() if bar.maximum() > 0 else 0))
        t.setFixedHeight(max(height, z(60)))

    def _edit_orders(self) -> None:
        """The card's button (and a row double-clicked): Excerpts… for the day shown, or Whose pages… when that
        is what holds its invoice back (Excerpts… has no pages to split until it is chosen)."""
        job = self.cur
        if job.ownership_problem() and not job.portions_problem():
            self._whose_pages(job, job.ownership_problem())
        else:
            self._who_ordered()

    @staticmethod
    def _name_of(group: list[Job], key: str) -> str:
        """An attorney's Attorney.label (the attorney, or the firm of a row naming several) by Attorney.key(),
        from the days of an invoice; the key itself when none of them has it."""
        for j in group:
            for a in j.case.attorneys:
                if a.key() == key:
                    return a.label()
        return key

    def _order_row_clicked(self, item) -> None:
        """A row of "Who ordered what" double-clicked: shows its day and opens its Excerpts window (or Whose
        pages…, see _edit_orders)."""
        job = next((j for j in self.jobs if id(j) == item.data(Qt.UserRole)), None)
        if job is None:
            return
        if job is not self.cur:
            self._select_job(job)
        if job is self.cur:
            self._edit_orders()

    def _who_ordered_why(self, job: Job) -> str:
        """Why Excerpts… can't be used for this job ("" when it can): it needs pages to bill (see
        Job.portions_unavailable; a job of several days is ordered whole in it)."""
        return job.portions_unavailable()

    def _who_ordered(self, job: Job | None = None):
        """Excerpts…: who ordered which pages of the days billed with the job, as one table (gui/excerpts.py; the
        days ticked of its case, as a rule: _excerpt_group). It stays open beside the window: what is changed there
        is kept at once, and what is changed here shows there."""
        from .excerpts import ExcerptsWindow
        job = job or self.cur
        why = self._who_ordered_why(job)
        if why:
            QMessageBox.information(self, "Excerpts", why)
            return
        if self.excerpts is None:
            self.excerpts = ExcerptsWindow(self.s, self._excerpts_changed, self, marks=self._marks)
        self.excerpts.load(self._excerpt_group(job))
        rows = [i for i, r in enumerate(self.excerpts.runs) if r.day.job is job]
        if rows:
            self.excerpts.table.selectRow(rows[0])
        self.excerpts.show()
        self.excerpts.raise_()
        self.excerpts.activateWindow()

    def _excerpt_group(self, job: Job | None = None) -> list[Job]:
        """The days Excerpts… shows with a job: those ticked that are billed on one invoice with it (as
        _invoice_group: every day ticked of its case, as a rule), its days waiting for a choice included
        (Excerpts… is where they are fixed); for a day invoiced already, the days that invoice billed. Just the
        job itself when it is unticked."""
        from ..batch import invoice_groups
        cur = job or self.cur
        if cur.invoiced:
            return self._invoice_group(cur)
        days = [j for j in self._considered() if not j.invoiced and j.invoice_days()]
        if cur in days:
            for g in invoice_groups(days, self.s):
                if any(j is cur for j in g):
                    return g
        return [cur]

    def _excerpts_changed(self) -> None:
        """A change kept in the Excerpts window: the attorney table, No. of copies (each day's ordering
        parties, see excerpts.store), the prices and the card follow it."""
        if self.excerpts is not None and any(r.day.job is self.cur for r in self.excerpts.runs):
            self._show_attorneys()
            self._show_copies()
        self._update_status()

    def _invoice_group(self, job: Job | None = None) -> list[Job]:
        """The days billed on one invoice with a job (default: the current one), as the cards show them
        (batch.bill_sets of the days ticked, _considered): those of its case Generate all bills together, as
        Settings.invoice_joint says (the ticked days not invoiced yet whose invoice isn't held, Job.invoice_hold:
        Excerpts... rows to check, or whose pages to bill), or, for a day invoiced already, the days that invoice
        billed. Just the job itself when it is unticked, its invoice is held or it has no pages to bill."""
        from ..batch import bill_sets
        cur = job or self.cur
        if not cur.invoice_pages() or not any(j is cur for j in self._considered()):
            return [cur]
        sets, _ = bill_sets(self._considered(), self.s)
        return next((g for g in sets if any(j is cur for j in g)), [cur])

    def _peripherals_text(self, opts, question=None, on_typed: bool = False, several: bool = False,
                          priced: bool = True) -> str:
        """What the invoice includes besides the original and the copies, and whether there is an index and why
        (invoice.index_reason, on every page of the transcripts), why on a line of its own: 'E-mailed copy,
        index\n(83 total pages, 50 or more)', 'E-mailed copy, no index\n(40 total pages, under 50)', 'No e-mailed
        copy (this job), index\n(turned on for this job)'. With no pages yet, the rule: 'E-mailed copy, index from
        50 pages'. question: a typed page count that changes whether there is an index (batch.index_question):
        'E-mailed copy, index? Generate asks\n(40 typed, the transcript has 83)' ('Generate all asks' for an
        invoice of several days: several). on_typed: a day of it is judged on its typed count, as answered
        (Job.index_on_typed): '(on the pages you typed: 40 total pages, under 50)'. priced False: the rate sheet has
        no index price for the invoice's speeds (those offered and those firms ordered at: invoice.index_priced),
        so days that get an index are charged none: 'E-mailed copy, no index charged\n(the rate sheet has no index
        price)'."""
        from ..invoice import index_mode, index_reason, indexed_days
        s = self.s
        email = s.invoice_include_email if opts.email is None else opts.email
        text = ("E-mailed copy" if email else "No e-mailed copy") + (" (this job)" if opts.email is not None else "")
        if not priced and opts.pages and any(indexed_days(opts, s)):
            return f"{text}, no index charged\n(the rate sheet has no index price)"
        if question is not None:
            has = "the transcripts have" if question.transcripts > 1 else "the transcript has"
            return (f"{text}, index? Generate{' all' if several else ''} asks\n"
                    f"({question.typed:,} typed, {has} {question.counted:,})")
        if index_mode(opts, s) == "auto" and not opts.pages:
            return f"{text}, index from {s.invoice_index_threshold} pages"
        mine = "on the pages you typed: " if on_typed and index_mode(opts, s) == "auto" else ""
        why = f"({mine}{index_reason(opts, s)})"  # (on a line of its own)
        if len(why) > LINE_WIDTH and ", " in why:  # broken after its comma, not with "more)" left on a line alone
            head, _, tail = why.rpartition(", ")
            why = f"{head},\n{tail}"
        return f"{text}, {'index' if any(indexed_days(opts, s)) else 'no index'}\n{why}"

    def _invoice_peripherals(self):
        """Peripherals…: the e-mailed copy and the index for this job's invoice, and the other days on it; and,
        for its days typed below their transcripts' count (batch.typed_days), the answer to Generate's question
        (Job.index_on_typed)."""
        from ..batch import joint_invoice, typed_days
        from .dialogs import InvoicePeripheralsDialog
        job = self.cur
        group = self._invoice_group(job)
        opts = joint_invoice(group)[1]
        typed = typed_days(group)
        answers = {j.index_on_typed for j in typed}
        dlg = InvoicePeripheralsDialog(
            opts.email, opts.index, self.s, len(group), self,
            typed=(sum(j.invoice_pages() for j in typed), sum(j.transcript_pages() for j in typed)) if typed else None,
            on_typed=answers.pop() if len(answers) == 1 else None)
        if dlg.exec():
            email, index = dlg.values()
            group = self._invoice_group(job)  # as it is now: a read may have ended meanwhile
            for day in group:
                day.invoice_email, day.invoice_index = email, index
            if typed:
                for day in typed_days(group):
                    day.index_on_typed = dlg.on_typed()
            self._update_status()

    def _invoice_show(self):
        """Customize…: what granular detail shows on this job's invoice, and the other days on it."""
        from ..batch import joint_invoice
        from .dialogs import InvoiceShowDialog
        job = self.cur
        dlg = InvoiceShowDialog(joint_invoice(self._invoice_group(job))[1].show, self.s, self)
        if dlg.exec():
            items = dlg.values()
            for day in self._invoice_group(job):  # as it is now: a read may have ended meanwhile
                day.invoice_show = items

    def _place_pairs(self) -> None:
        """Puts the two columns of the Case and Speed, pages and dates cards side by side when the form has
        room for both, else one above the other, and the types of proceeding in one row or two, so the form never
        needs scrolling sideways; then the Generate buttons in the Outputs row or under it (_place_foot)."""
        scroll = getattr(self, "form_scroll", None)
        if scroll is None:
            return
        room = scroll.viewport().width() - z(16) - z(32)  # the form's and the card's margins
        for g, wa, wb in self._pairs:
            side = wa.minimumSizeHint().width() + wb.minimumSizeHint().width() + 18 <= room
            if (g.itemAtPosition(0, 1) is not None) == side and g.count() == 2:
                continue
            g.removeWidget(wa)
            g.removeWidget(wb)
            g.addWidget(wa, 0, 0)
            g.addWidget(wb, 0, 1) if side else g.addWidget(wb, 1, 0)
            g.setColumnStretch(0, 1)
            g.setColumnStretch(1, 1 if side else 0)
        # the types of proceeding: in one row when they fit beside their label, else in two rows
        boxes = list(self.proc_boxes.values())
        need = sum(b.sizeHint().width() for b in boxes) + 12 * (len(boxes) - 1)
        labels = [self.proc_form.itemAt(i, QFormLayout.LabelRole) for i in range(self.proc_form.rowCount())]
        label_w = max((x.widget().sizeHint().width() for x in labels if x is not None and x.widget()), default=0)
        per_row = len(boxes) if need <= room - label_w - z(12) else (len(boxes) + 1) // 2
        if per_row != self._proc_per_row:
            self._proc_per_row = per_row
            for b in boxes:
                self.proc_grid.removeWidget(b)
            for i, b in enumerate(boxes):
                self.proc_grid.addWidget(b, i // per_row, i % per_row)
            for col in range(len(boxes) + 1):
                self.proc_grid.setColumnStretch(col, 1 if col == per_row else 0)
        self._place_foot()

    def _place_foot(self) -> None:
        """The Generate buttons at the end of the Outputs row when it has room for them, else on a row of their own
        under it, at the right (a narrow window, a zoom, or the long names of "Generate 6/4/2026 only" and
        "Generate all 3 days" with every output ticked): the page never scrolls sideways for the row."""
        if not getattr(self, "_outputs_built", False):
            return
        row, box = self.foot_row, self.gen_box
        in_row = row.indexOf(box) >= 0
        mine = box.minimumSizeHint().width() + row.spacing()
        rest = row.minimumSize().width() - (mine if in_row else 0)
        room = self.page_scroll.viewport().width() - 36 - 32  # (the window's margins and the row's)
        fits = rest + mine <= room
        if fits and not in_row:
            self.foot_row2.removeWidget(box)
            row.addWidget(box)
        elif not fits and in_row:
            row.removeWidget(box)
            self.foot_row2.addWidget(box)

    def _size_drop(self):
        """The drop zone is one slim row once documents are loaded (the whole window takes more of them), and
        otherwise shows the big picture when the left column has room for it above the list and the paste box,
        the small one when not. It is tried and measured, since the column's height depends on fonts and wrapped
        text."""
        slim = len(self.jobs) > 1 or any(j.docs for j in self.jobs)
        self.drop.set_slim(slim)
        self._fit_paste()
        if slim:
            return
        scroll = getattr(self, "left_scroll", None)
        if scroll is None or not scroll.isVisible():
            self.drop.set_compact(self.height() < z(COMPACT_BELOW))
            return
        self.drop.set_compact(False)
        left = scroll.widget()
        left.layout().activate()
        width = scroll.viewport().width()
        # the height the column asks for at that width (wrapped text included)
        need = left.heightForWidth(width) if left.hasHeightForWidth() else left.minimumSizeHint().height()
        if max(need, left.minimumSizeHint().height()) > scroll.viewport().height():
            self.drop.set_compact(True)

    def _fit_paste(self) -> None:
        """The paste box is one line once a job is loaded, while it is empty and not typed in ("Paste an e-mail or
        notes…", without its label and Extract from text: the job list gets the room; the user, 2026-10-22), as the
        drop zone is one slim row then; a click in it, or text pasted, opens it again, and New job too."""
        if not hasattr(self, "paste_pane"):
            return
        # (the window's focus widget, not hasFocus: switching to Outlook to copy the e-mail leaves the box open)
        slim = (any(j.docs for j in self.jobs) and not self.paste.toPlainText().strip()
                and self.focusWidget() is not self.paste)
        was = self._paste_slim
        if slim == was:
            return
        if slim and was is False:  # (where the divider was, kept for next time: _save_splits; not after a zoom)
            self._split_unslim = self.left_split.saveState()
        self._paste_slim = slim
        self.paste_label.setVisible(not slim)
        self.extract_btn.setVisible(not slim)
        self.paste.setPlaceholderText(PASTE_SLIM if slim else PASTE_HINT)
        if slim:
            line = self.paste.fontMetrics().lineSpacing() + z(14)
            self.paste.setMinimumHeight(line)
            self.paste.setMaximumHeight(line)
            self.paste.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)  # (one line: no bar beside it)
            self.paste_pane.setMaximumHeight(line + z(6))
            total = sum(self.left_split.sizes())
            if total > 0:  # (the divider keeps the pane's old height otherwise: the room goes to the list)
                self.left_split.setSizes([max(total - line - z(6), 1), line + z(6)])
        else:
            self.paste.setMinimumHeight(z(52))
            self.paste.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
            self.paste.setMaximumHeight(16777215)  # (QWIDGETSIZE_MAX)
            self.paste_pane.setMaximumHeight(16777215)
            if was:  # (opened again: the divider placed again for the box open, _place_left_split)
                self._left_mode = None
                QTimer.singleShot(0, self, lambda: self._place_left_split(len(self.jobs) > 1))

    def _size_drop_later(self):
        """_size_drop once Qt has laid the window out (sizes measured before that are not the final ones). The
        window is the timer's context: a window closed meanwhile is not measured."""
        QTimer.singleShot(0, self, self._size_drop)

    def resizeEvent(self, e):
        """The cards' two columns and the Generate buttons' place follow the window's width (_place_pairs), and the
        drop zone its height."""
        super().resizeEvent(e)
        self._place_pairs()
        if hasattr(self, "drop"):
            self._size_drop()

    def showEvent(self, e):
        """When the window first shows, the drop zone is sized again once the window is laid out: measured
        earlier, the left column looks shorter than it is."""
        super().showEvent(e)
        self._place_pairs()
        self._size_drop_later()
        # the divider once the window is laid out (measured earlier, it would be placed in proportion to sizes
        # that aren't the window's)
        QTimer.singleShot(0, self, lambda: self._place_left_split(len(self.jobs) > 1))

    def _fill_invoice_speeds(self):
        """Speeds offered, under the rate sheet: a box per speed of the current rate sheet with its price
        ("Regular  $4.30/pg"), cheapest first as on the invoice, two a row, ticked for the speeds invoices offer
        (Settings.invoice_speeds)."""
        from ..rates import speed_key
        while self.inv_speeds_grid.count():
            item = self.inv_speeds_grid.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
        self.inv_speed_boxes = {}
        offered = {speed_key(x) for x in self.s.invoice_speeds}
        try:
            speeds = list(self.s.sheet().speeds)
        except Exception:  # a broken rate sheet must not break the window
            speeds = []
        names = [sp.name for sp in speeds]
        per_row = 2  # (more in a row would make the card's left column too wide for its two columns side by side)
        for i, sp in enumerate(speeds):  # cheapest first: the order of the Speed box and of the invoice
            name = sp.name
            cb = QCheckBox(f"{name}  ${sp.original}/pg" if sp.original else name)
            days = self.s.days_for(name)
            cb.setToolTip(sp.label() + (f"  ·  {plural(days, 'day')}" if days is not None else ""))
            cb.setChecked(speed_key(name) in offered)
            cb.toggled.connect(self._invoice_speeds_changed)
            self.inv_speed_boxes[name] = cb
            self.inv_speeds_grid.addWidget(cb, i // per_row, i % per_row)
        if not names:
            self.inv_speeds_grid.addWidget(QLabel("(none on the rate sheet)"), 0, 0)
        self._show_speeds_lock()

    def _speeds_locked(self) -> bool:
        """The Speeds offered boxes are soft locked: a firm of the invoice shown (its days: _invoice_group) has its
        speed chosen in Who ordered what (or Excerpts…), and the user hasn't said yes to changing them anyway for
        this job (_speeds_unlocked). The user, 2026-10-21: the speed is chosen in one place; the boxes look greyed
        out, and a click asks before changing them."""
        job = getattr(self, "cur", None)
        if job is None or job.is_empty() or getattr(self, "_speeds_unlocked", None) is job:
            return False
        return any(sp for j in self._invoice_group(job) for _, _, speeds in j.runs_ordered() for sp in speeds)

    def _show_speeds_lock(self) -> None:
        """The Speeds offered boxes greyed out (still clickable: _invoice_speeds_changed asks) while soft locked."""
        locked = self._speeds_locked()
        for name, cb in getattr(self, "inv_speed_boxes", {}).items():
            if cb.property("softlock") == locked:
                continue
            cb.setProperty("softlock", locked)
            cb.setStyleSheet("QCheckBox { color: gray; }" if locked else "")
            tip = cb.toolTip().split("\n\n")[0]
            cb.setToolTip(tip + ("\n\n" + SPEEDS_LOCKED_TIP if locked else ""))

    def _invoice_speeds_changed(self, _=None):
        """The speeds ticked are saved straight away. Speeds of other rate sheets stay as they were. With one
        speed ticked the invoice bills that speed alone; with none, the agreement form's speed. Every job's
        agreement form speed is the Settings rule's among them again, unless the user chose one still offered
        (Expedite unticked and ticked again: back to Expedite, not left on Regular). While the boxes are soft
        locked (_speeds_locked: speeds chosen in Who ordered what), a click asks first; No puts the box back."""
        from ..rates import speed_key
        box = self.sender()
        if isinstance(box, QCheckBox) and self._speeds_locked():
            if QMessageBox.question(self, "Speeds chosen in Who ordered what", SPEEDS_LOCKED_ASK) != QMessageBox.Yes:
                with QSignalBlocker(box):
                    box.setChecked(not box.isChecked())
                return
            self._speeds_unlocked = self.cur  # (asked once: this job's boxes are the user's to change now)
            self._show_speeds_lock()
        here = {speed_key(n) for n in self.inv_speed_boxes}
        keep = [x for x in self.s.invoice_speeds if speed_key(x) not in here]
        self._set_opt("invoice_speeds", keep + [n for n, cb in self.inv_speed_boxes.items() if cb.isChecked()])
        self._apply_speed_rule()

    def _open_run_sheets_folder(self):
        """Opens the folder of the shown job's run sheet: where the case's run sheet is (runsheet.find_sheets in
        runsheet.sheet_folders: it may be next to the transcripts, where the records say, or in the old run sheets
        folder), else where it will go (the case's folder, or the folder Settings → Options → Folders gives run
        sheets, made first if needed). A case folder not made yet is not made for this: the nearest folder of it
        that is there is opened (for a job that names no case yet, its Save to folder)."""
        from ..batch import save_to_for
        job = self.cur
        own = self.s.output_dirs.get("runsheet")
        named = not job.is_empty() and bool(job.case.get("index_no") or job.case.get("case_name"))
        where = case_dir_for(job, self.s) if named else save_to_for(job, self.s)
        found = []
        if named:
            try:
                found = find_sheets(job.case, sheet_folders(self.s, where, input_folders(job), job.case))
            except SheetUnreadable:
                pass  # (held by another program just now: the folder a new one would go in is opened instead)
        if found:
            folder = found[0].path.parent
        elif own:
            folder = runsheets_folder(self.s, where)
            try:
                folder.mkdir(parents=True, exist_ok=True)
            except OSError as e:
                QMessageBox.warning(self, "Could not open the folder", str(e))
                return
        else:
            folder = where
            while not folder.is_dir() and folder.parent != folder:
                folder = folder.parent
        open_path(folder)

    def lock_pdfs(self):
        """File → Lock finished PDFs: a copy of each PDF picked, with its fields flattened so the values can no
        longer be changed ('<name> (locked).pdf' next to it). Only PDFs this app made that still have fields
        are locked."""
        from ..pdfout import has_fields, is_generated, lock_pdf
        start = self.s.output_dirs.get("agreement") or self.s.output_dir or str(Path.home() / "Documents")
        files, _ = QFileDialog.getOpenFileNames(self, "Lock finished PDFs", start, "PDF files (*.pdf)")
        if not files:
            return
        made, skipped, no_fields, failed = [], [], [], []
        for f in map(Path, files):
            if not is_generated(f):
                skipped.append(f.name)
                continue
            try:
                if not has_fields(f):  # flattened when saved, or a locked copy: nothing to lock
                    no_fields.append(f.name)
                    continue
                made.append(lock_pdf(f))
            except Exception as e:
                failed.append(f"{f.name}: {e}")
        copies = "1 locked copy" if len(made) == 1 else f"{len(made)} locked copies"
        lines = [f"Saved {copies} next to the originals."] if made else []
        if skipped:
            lines.append("Not made by this app, so left alone:\n  " + "\n  ".join(skipped))
        if no_fields:
            lines.append("Already locked or flattened (no fields), so left alone:\n  " + "\n  ".join(no_fields))
        if failed:
            lines.append("Could not be locked (is it open in a PDF viewer?):\n  " + "\n  ".join(failed))
        QMessageBox.information(self, "Lock finished PDFs", "\n\n".join(lines))

    def _invoice_detail_changed(self, on: bool):
        """"Show granular detail" was clicked: it applies to this job's invoice (and the other days on it) only;
        a new job starts with it off."""
        for job in self._invoice_group():
            job.invoice_detail = on
        self._refresh_outputs()

    def _invoice_parties_changed(self, n: int):
        """A number typed here stays; set back to the number of ticked attorneys, it follows them again. It
        applies to the other days on the same invoice too, but only those several attorneys ordered: a day one
        attorney ordered alone is that attorney's to pay. No. of copies follows it (unless typed)."""
        from ..batch import group_attorneys
        group = self._invoice_group()
        ticked = max(1, len(group_attorneys(group)) if len(group) > 1 else len(self.cur.ticked_keys()))
        for job in group:
            alone = len(group) > 1 and len(job.ticked_keys()) < 2
            job.parties = 0 if n == ticked or alone else n
        self._show_copies(group)  # No. of copies is the number of ordering parties
        self._update_status()  # (Parties set right takes its hold off the list of what to check)

    def open_records(self):
        """Shows the Records window (Ctrl+R), made once and reloaded each time it is opened. Opened afresh, it
        lists the user's own invoices: those made in a colleague's name only after its Show them (again)."""
        from .records_window import RecordsWindow
        win = getattr(self, "_records_win", None)
        if win is None:
            win = self._records_win = RecordsWindow(self.s, self)
            win.on_reopen = self.open_past_job
            win.on_settings = lambda tab: self.open_settings(tab=tab)
            win.on_changed = self._records_edited
        if not win.isVisible():
            win.theirs = False
        win.reload()
        if win.isMinimized():  # (show() alone leaves it minimized; a maximized one stays maximized)
            win.setWindowState(win.windowState() & ~Qt.WindowMinimized)
        win.show()
        win.raise_()
        win.activateWindow()
        self._recap(win)

    def _recap(self, win, today: date | None = None) -> None:
        """The first time Records is opened in a month: what last month came to ("Last month (September 2026)
        you made $870.00 with 243 pages"), and the first time in a year, last year too (records.recap). The box
        has OK and a tick to show no more recaps (Settings -> Options turns them back on). A month without
        invoices says nothing. On April 1 the recap of last month is first a joke (records.april_fools)."""
        if not self.s.recaps:
            return
        from ..records import april_fools, recap
        try:
            lines, month, year = recap(win.ledger.invoices(), today, self.s.recap_month, self.s.recap_year)
        except Exception as e:  # the records still open
            logfile.error("could not sum up last month", e)
            return
        if (month, year) != (self.s.recap_month, self.s.recap_year):
            self.s.recap_month, self.s.recap_year = month, year
            self._save_settings()
        if not lines:
            return
        joke = april_fools(today)
        if joke and lines[0].startswith("Last month"):
            fool = QMessageBox(win)
            fool.setIcon(QMessageBox.Information)
            fool.setWindowTitle("Recap")
            fool.setText(joke)
            fool.addButton("Really?!", QMessageBox.AcceptRole)
            fool.exec()
            lines = ["April Fools! Here's what really happened:"] + lines
        box = QMessageBox(win)
        box.setIcon(QMessageBox.Information)
        box.setWindowTitle("Recap")
        box.setText("\n\n".join(lines))
        stop = QCheckBox("Don't show me monthly or annual recaps anymore")
        stop.setToolTip("Settings → Options turns them back on.")
        box.setCheckBox(stop)
        box.addButton(QMessageBox.Ok)
        box.exec()
        if stop.isChecked():
            self.s.recaps = False
            self._save_settings()

    def _records_changed(self):
        """Files were made, or Settings saved: an open Records window shows them, and the window reads again what
        it shows from the records (the jobs made lately, the fields' completions, the Records button's count). A
        tick box switched on, off or renamed in Settings -> Records gets its column in the CSV copies at once, as
        the Records window does when it is open."""
        win = getattr(self, "_records_win", None)
        open_ = win is not None and win.isVisible()
        if open_:
            win.reload()
        columns = self._csv_columns()
        if columns != self._mirrored and not open_:
            from ..records import Ledger
            folder, boxes = self.s.records_folder(), list(columns[1])  # (read here: the thread has its own copy)
            self.runner.start(lambda: Ledger(mirror_dir=folder, boxes=boxes).mirror(), on_done=lambda _r: None,
                              on_error=lambda _m: None)
        self._mirrored = columns
        self._trial_rules = None  # (an invoice made or deleted: read again when next needed)
        self._read_records_side()

    def _records_edited(self) -> None:
        """The Records window changed the records (a delete, a restore, an undo, a backup put back): the way each
        trial was first billed is read again (an invoice in the trash no longer counts), and the cards priced by
        it follow."""
        self._trial_rules = None
        self._update_status()

    def _note_trial_rules(self, jobs: list, fresh: bool = False) -> None:
        """Gives these jobs the split rule their trial was first billed by (Job.trial_rule, batch.trial_split_rules),
        read from the records once until they change (_records_changed); fresh: read now (Generate, so what it
        saves is what its preview shows). Records that can't be read leave each as Settings say; none yet, nothing
        is read (an empty database isn't made just to look)."""
        from ..batch import note_trial_rules, trial_split_rules
        from ..records import default_db
        rules = getattr(self, "_trial_rules", None)
        if rules is None or fresh:
            try:
                rules = trial_split_rules(ledger_for(self.s)) if default_db().exists() else {}
            except Exception as e:
                log_error("could not read the records for the trials' split rules", e)
                rules = {}
            self._trial_rules = rules
        note_trial_rules(jobs, rules)

    def _csv_columns(self) -> tuple:
        """What the CSV copies of the records have columns for (deliver.ledger_for): the records folder and the
        tick boxes in use besides Paid."""
        return self.s.records_dir, tuple((b.key, b.name) for b in self.s.boxes() if b.on and b.kind == "tick")

    def _save_invoice_template(self):
        """Saves a copy of the invoice spreadsheet template that ships with the app where the user picks."""
        from ..invoice import TEMPLATE
        if not TEMPLATE.is_file():
            QMessageBox.warning(self, "Not found", "The invoice spreadsheet template is missing from this install.")
            return
        dest, _ = QFileDialog.getSaveFileName(self, "Save the invoice spreadsheet template",
                                              str(Path.home() / "Documents" / TEMPLATE.name),
                                              "Excel workbook (*.xlsx)")
        if dest:
            import shutil
            try:
                shutil.copyfile(TEMPLATE, dest)
            except OSError as e:
                QMessageBox.warning(self, "Could not save", str(e))
                return
            open_path(Path(dest).parent)

    # ------------------------------------------------------------- zoom and screen
    def zoom_step(self, step: int) -> None:
        """Zoom in (1), out (-1) or back to 100 % (0): Ctrl + / Ctrl - / Ctrl 0, Ctrl + wheel."""
        now = zooming.zoom()
        self.set_zoom(1.0 if step == 0 else now + step * zooming.STEP)

    def set_zoom(self, f: float) -> None:
        """Shows everything at zoom f (see zoom.py), saves it as the default and says so in the status line."""
        f = zooming.set_zoom(f)
        apply_theme(self.app, self.s.theme)
        zooming.reapply()
        self.drop.rezoom()
        self._paste_slim = None  # (its one line is taller or shorter now: _fit_paste sizes it again)
        for t in (self.att, self.orders):  # the firms' tags at the new size (_show_orders draws the orders' again)
            t.setIconSize(QSize(*badge_size(z(TAG_HEIGHT))))
        self._mark_attorneys()
        self._place_pairs()
        QTimer.singleShot(0, self._place_pairs)  # again once the new sizes are laid out
        self._fit_minimum()
        self._refresh_jobs()
        # the Who ordered what table is as tall as its rows, which are taller or shorter now
        self._show_orders()
        QTimer.singleShot(0, self, self._show_orders)  # again once the new sizes are laid out
        self._set_opt("zoom", f)
        self._toast(f"Zoom {round(f * 100)} %")

    def _set_theme(self, theme: str) -> None:
        """View -> Theme: as Windows is ("system"), light or dark, at once and for the next time (as Settings ->
        Options -> Theme)."""
        self.s.theme = theme
        apply_theme(self.app, theme)
        self._save_settings()

    def _show_yin(self, on: bool) -> None:
        """View -> Show the yin-yang (as Settings -> Options): the drop zone's picture comes or goes at once."""
        self.s.show_yin = on
        self._save_settings()
        self._job_status()  # (it sets the status again, and with it the picture: _set_status)

    def _avail(self):
        """The part of the window's screen not taken by the taskbar (logical pixels, after Windows scaling)."""
        screen = self.screen() or QApplication.primaryScreen()
        return screen.availableGeometry() if screen else None

    def _fit_minimum(self) -> None:
        """The window's smallest size: 1080 x 720 at 100 % zoom, but never more than the screen has room for.
        At 150 % Windows scaling a 1920 x 1080 screen is 1280 x 690 or so in Qt's pixels: a fixed 720 would put
        the bottom of the window off the screen."""
        avail = self._avail()
        w, h = z(1080), z(720)
        if avail is not None:
            w, h = min(w, int(avail.width() * 0.95)), min(h, int(avail.height() * 0.9))
        self.setMinimumSize(max(480, w), max(360, h))

    def _fit_screen(self) -> None:
        """Brings the window back within its screen: a size saved on a bigger screen (or at another scaling) is
        made smaller, and a window off the screen is moved onto it."""
        avail = self._avail()
        if avail is None or self.isMaximized() or self.isFullScreen():
            return
        frame = self.frameGeometry()
        extra_w, extra_h = frame.width() - self.width(), frame.height() - self.height()
        w = min(self.width(), avail.width() - extra_w)
        h = min(self.height(), avail.height() - extra_h)
        if (w, h) != (self.width(), self.height()):
            self.resize(max(w, self.minimumWidth()), max(h, self.minimumHeight()))
        frame = self.frameGeometry()
        x = min(max(frame.x(), avail.x()), avail.right() - frame.width() + 1)
        y = min(max(frame.y(), avail.y()), avail.bottom() - frame.height() + 1)
        if (x, y) != (frame.x(), frame.y()):
            self.move(max(x, avail.x()), max(y, avail.y()))

    def eventFilter(self, obj, e):
        """Ctrl + the mouse wheel over this window zooms (a text box would otherwise zoom only its own text):
        a step per notch of the wheel (120), as a touchpad sends a notch in many small turns. When the form
        gets wider or narrower (the splitter dragged), its cards' columns are placed again (_place_pairs), and when
        the job list does, its Case / file column (_fit_case_column). Files dragged over a text box are the
        window's to read (_drag_filter); a click on Options ▾ while its box is open only closes it; the paste box
        opens or goes back to one line as it gets or loses the cursor (_fit_paste); and the attorney's address,
        typed under the table, is written into it when its box is left (_att_detail_edited)."""
        t = e.type()
        if t in (QEvent.DragEnter, QEvent.DragMove, QEvent.Drop) and self._drag_filter(obj, e):
            return True
        if t == QEvent.DragLeave and getattr(self, "drop", None) is not None and isinstance(obj, QWidget) \
                and obj.window() is self:
            self.drop._hover(False)  # (left a text box: lit again by the next part of the window it is over)
        if t == QEvent.MouseButtonPress and obj is getattr(self, "options_pop", None) \
                and self.options_btn.rect().contains(self.options_btn.mapFromGlobal(e.globalPosition().toPoint())):
            # a click on Options ▾ while its box is open closes the box, and only that: the click that closes a
            # popup is passed on to what is under it otherwise, and the button would open it again (as
            # QComboBox does)
            obj.setAttribute(Qt.WA_NoMouseReplay)
        if t == QEvent.Resize and obj is getattr(self, "job_list", None):
            self._fit_case_column()
        if t in (QEvent.FocusIn, QEvent.FocusOut) and obj is getattr(self, "paste", None):
            QTimer.singleShot(0, self, self._fit_paste)  # (clicked in: open; left empty: one line again)
        if t == QEvent.FocusOut and obj is getattr(self, "att_edits", {}).get(3) and self._att_typed(3):
            self._att_detail_edited(3)  # (only when typed in: an address the AI filled meanwhile isn't wiped)
        if e.type() == QEvent.Wheel and e.modifiers() & Qt.ControlModifier:
            w = obj if isinstance(obj, QWidget) else None
            if w is not None and w.window() is self:
                dy = e.angleDelta().y()
                if dy and (dy > 0) != (self._wheel > 0):
                    self._wheel = 0  # turned the other way: start counting again
                self._wheel += dy
                steps = int(self._wheel / 120)  # (towards 0: the rest is kept for the next turn)
                if steps:
                    self._wheel -= steps * 120
                    self.set_zoom(zooming.zoom() + steps * zooming.STEP)
                return True
        elif (e.type() == QEvent.Resize and getattr(self, "form_scroll", None) is not None
              and obj is self.form_scroll.viewport()):
            self._place_pairs()
        return super().eventFilter(obj, e)

    def closeEvent(self, e):
        """Asks first while a batch is being made; remembers the window's size and place, and where the divider
        between the job list and the paste box is (_save_splits)."""
        if self.filling and QMessageBox.question(
                self, "Still working", "The files of the batch are still being made. Close anyway?\n\n"
                "(The file being written may be left incomplete.)") != QMessageBox.Yes:
            e.ignore()
            return
        self.s.window_geometry = bytes(self.saveGeometry().toBase64()).decode()
        self._save_splits()
        self._save_settings()  # (a file that can't be written must not keep the window from closing)
        super().closeEvent(e)
