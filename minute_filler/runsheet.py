"""The trial run sheet: which reporter wrote which pages, take by take, kept as an Excel workbook.

When several reporters share a trial, each needs to know how many pages they wrote, for billing. The run
sheet has a row per take, with the columns of the reporters' own Google Sheet:

  Date | Weekday | Reporter | Day's Take | Running Take | Pages written | Starting Page No. |
  Ending Page No. | Witness Start | Witness End | Note

The weekday, the take counts and the page numbers are formulas ("Don't write here"), so a row typed in or
changed by hand is counted too. Witness Start and Witness End get the first and the last running head printed
in the take, as printed ("DIRECT - PLAINTIFF - JANE ROE"), or are left blank for the user
(Settings.runsheet_witness): who was on the stand is not read from them. Filtering the Reporter column shows
that reporter's pages (the total of the rows shown is at the top).

It is there to look back at a trial (the user, 2026-10-26: "to jog my memory of what happened, who worked on what
day, how many pages each person did"), so it has three more tabs: "By Reporter" adds up each reporter's takes and
pages, with their share of the trial's pages and the days they worked; "Witnesses" lists every running head the
takes printed, as printed, with the days, reporters and pages it was on (_witnesses: the heads of each take are
kept on the hidden sheet, so a take of an earlier day still has all of them, not only Witness Start and End);
"Who Ordered" lists, from the records, each invoice made for the case (the firm, its side, the days, the speed,
the number, the amount and whether it is paid) and each firm with a minute agreement but no invoice yet
(order_rows). Generate writes the run sheet first, so Who Ordered is written again once the invoices are made
(refresh_orders): today's invoices are on it too.

There is one run sheet per case, kept in the Trial Runs folder inside the case's own folder (straight in Save to
without a folder for each case, or in the folder Settings -> Options -> Folders gives run sheets). A transcript's
takes are added to the case's run sheet when there is one there, at the top of the case's folder (where 2.13 and
earlier put it), next to the transcript, in a folder the records list for the case's run sheets (a trial's days
can go to different folders) or in Documents\\YinIt Run Sheets (where earlier versions kept them all): see
sheet_folders. One found is added to where it is; it is not moved. It is found by index number or, since one
trial can have several index numbers billed together, by case name; otherwise a new one is started. On a run
sheet of this app, the rows already there stay in their order (with what was typed in them) and each new take
goes before the first later one. A run sheet this app did
not make (downloaded from Google Sheets, say) gets its rows added at the end, or on a row typed in ahead for that
day, with its own formulas copied down; it is backed up first.

A run sheet made before the app was renamed (2.0) is still found and added to: its hidden "DjinnIt" sheet
(OLD_META) is read, and replaced by a "YinIt" one when the run sheet is written again (meta_sheet, _meta).

RunSheets is the output deliver.generate adds the takes through (courthouses.OutputSpec.maker).
"""
from __future__ import annotations

import ast
import json
import os
import re
import shutil
import threading
import time
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import TYPE_CHECKING

from .dates import us_date
from .extract_regex import find_dates, norm_index, strip_line_numbers, title_page_count
from .pdfout import safe_filename, unique_path
from .log import error as log_error
from .models import Attorney, CaseInfo
from .settings import Settings
from .takes import TITLE_ONLY, body_pages, find_takes, fix_marks, reporter_label, title_reporters

if TYPE_CHECKING:
    from .deliver import Making

SHEET, BY_REPORTER, WITNESSES, ORDERS, META = "Run Sheet", "By Reporter", "Witnesses", "Who Ordered", "YinIt"
FORMAT = "YinIt run sheet 2"  # (2: the takes' heads kept on the hidden sheet; 1 is read the same)
BACKUP = "(before YinIt)"  # a run sheet this app didn't make is first copied to "<name> (before YinIt).xlsx"
OLD_FOLDER = "YinIt Run Sheets"  # under Documents: every run sheet's folder before they went with their case's files
OLD_DJINNIT_FOLDER = "DjinnIt Run Sheets"  # and before the rename (2.0), when it couldn't be renamed (settings)
# Run sheets made before the app was renamed (2.0): their hidden sheet, its format and the backup's name
OLD_META, OLD_FORMAT, OLD_BACKUP = "DjinnIt", "DjinnIt run sheet", "(before DjinnIt)"
HEAD_ROW, FIRST = 4, 5  # the column headings, the first take
LAST = 5000             # how far down the totals look
# key, heading, typed by hand ("Write here"), width
COLUMNS = [
    ("date", "Date", True, 11), ("weekday", "Weekday", False, 11), ("reporter", "Reporter", True, 13),
    ("day_take", "Day's Take", False, 9), ("run_take", "Running Take", False, 9),
    ("pages", "Pages written", True, 9), ("start", "Starting Page No.", False, 10),
    ("end", "Ending Page No.", False, 10), ("witness_start", "Witness Start", True, 24),
    ("witness_end", "Witness End", True, 24), ("note", "Note", True, 34),
]
COL = {k: i + 1 for i, (k, *_rest) in enumerate(COLUMNS)}
L = {k: chr(64 + i) for k, i in COL.items()}  # column letters
GREEN, ORANGE, RED, BAND = "93C47D", "F6B26B", "EA9999", "E0F7FA"
NO_RUNSHEET = "a run sheet needs a transcript PDF (its pages and the reporters' initials)"
NOT_FOUND = "reporter's initials not found"
HAND = [k for k, _, hand, _ in COLUMNS if hand]  # the columns typed by hand
_LOCK = threading.Lock()


@dataclass
class Row:
    """One take on the run sheet."""
    day: date | None = None
    reporter: str = ""
    pages: int | None = None
    start: int | None = None
    # Witness Start and End: the first and last running head printed in the take, or "" (Settings.runsheet_witness
    # "blank"); on a row read back from a run sheet, whatever is in those cells
    witness_start: str = ""
    witness_end: str = ""
    note: str = ""
    initials: str = ""  # as on the pages ("ds"); not a column of its own
    extra: dict = field(default_factory=dict)  # cells to the right of the run sheet's columns, kept with the row
    raw: dict = field(default_factory=dict)    # typed cells kept as they were (a formula, "TBD"), by column key
    text: set = field(default_factory=set)     # the keys of raw kept as text, not formulas ("=TBD" typed as text)
    origin: int | None = None                  # the row it was on (its formulas are moved with it)
    # its pages' running heads as printed (takes.Take.heads: "Proceedings, J. Doe - Plaintiff - Direct"): the Who
    # transcribed what card's, not a column of the run sheet
    heads: str = ""
    # the same heads one by one, each once (takes.Take.heads), for the Witnesses tab; on a row read back from a run
    # sheet, those the hidden sheet kept for it (_meta), else its Witness Start and End
    head_list: list = field(default_factory=list)
    # the transcript it is of (batch.Doc.key()) and where its first page is in that transcript's pages
    # (takes.Take.at), for Correct where this take ends; "" and None for a row not read from a transcript
    doc: str = ""
    at: int | None = None
    corrected: bool = False  # the user moved where it or the take before it ends (batch.Doc.fixes)
    case: str = ""  # the case's short caption, for Who transcribed what when it shows several cases (batch.takes_of)

    def end(self) -> int | None:
        """The take's last page number, or None without a first page or page count."""
        return self.start + self.pages - 1 if self.start is not None and self.pages else None

    def same_take(self, other: "Row") -> bool:
        """Already on the run sheet: the same day and pages that overlap, or (without page numbers) the
        same day, reporter and page count."""
        if self.day != other.day:
            return False
        if None not in (self.start, other.start, self.end(), other.end()):
            return self.start <= other.end() and other.start <= self.end()
        return (self.reporter.lower(), self.pages) == (other.reporter.lower(), other.pages)


@dataclass
class RunSheetOpts:
    """The takes to add, where to add them, and (after add_takes) what happened."""
    rows: list[Row] = field(default_factory=list)
    target: str | None = None  # None: as Settings.runsheet_existing says; "": a new run sheet; else that file
    path: Path | None = None   # the run sheet written
    created: bool = False      # a new one
    added: int = 0
    skipped: int = 0           # takes that were on it already (or given twice)
    added_pages: int = 0       # the pages of the takes added
    days: list[CaseInfo] = field(default_factory=list)  # the case of each day of the takes (a batch's several): the
    # attorneys the run sheet keeps for Who Ordered's colours (_kept_attorneys), besides the case written's


def run_sheet_summary(opts: RunSheetOpts) -> str:
    """What generate did to the run sheet, for the 'Saved' message."""
    name = opts.path.name if opts.path else "the run sheet"
    takes = lambda n: f"{n} take{'s' if n != 1 else ''}"
    if opts.created:
        text = f"Started the run sheet {name} with {takes(opts.added)}."
    elif opts.added:
        text = f"Added {takes(opts.added)} to the run sheet {name}."
    else:
        return f"The run sheet {name} already has these takes: nothing was added."
    if opts.skipped:
        text += f" ({takes(opts.skipped)} {'was' if opts.skipped == 1 else 'were'} on it already.)"
    return text


@dataclass
class Found:
    """A run sheet on disk that may be this case's."""
    path: Path
    case_name: str = ""
    index_nos: list[str] = field(default_factory=list)
    ours: bool = True  # made by this app (False: someone else's, added to as it is laid out)
    why: str = ""  # "index" (the same index number) or "name" (the same case name)

    def reason(self) -> str:
        """Why this may be the case's run sheet, for the user ("same index number")."""
        if self.why == "index":
            return "same index number"
        return "same case name" + (f", index {', '.join(self.index_nos)}" if self.index_nos else "")


# ------------------------------------------------------------------ transcripts -> rows

def rows_from(ing, day: date | None, s: Settings, front: str = "", fixes: dict[int, str] | None = None) -> list[Row]:
    """The run sheet rows of one transcript PDF: a row per take. ing: the ingested PDF (its page marks,
    page count and text); day: the day the transcript is of; front: whose the pages before the first initials
    are, as the invoice counts them (batch.Doc.front: Job.front_owner's initials, or "none"; see
    takes.find_takes); fixes: the user's corrections of whose pages are whose (batch.Doc.fixes, see
    takes.fix_marks). The reporters are named as reporter_label says; Witness Start and End are the first and last
    heads printed in the take (Settings.runsheet_witness "headers") or blank. Each row has its take's heads (for
    the card) and where it starts in the marks (Row.at)."""
    if not ing.marks:  # no initials or page numbers to read (a scan): one row for the whole transcript
        return [Row(day, "", transcript_pages(ing), ing.first_page_no, note=NOT_FOUND)]
    title_text, _ = strip_line_numbers(ing.text)
    titles = title_page_count(title_text)
    names = title_reporters("\n".join(title_text.split("\f")[:titles]))
    takes = find_takes(fix_marks(ing.marks, fixes), titles, front)
    heads = s.runsheet_witness == "headers"
    p = s.profile
    rows = []
    for t in takes:
        who = reporter_label(t.initials, names, s.reporter_names(), p.name, p.initials, s.title_case_names)
        note = TITLE_ONLY if t.title_only and len(takes) > 1 else "" if t.initials else NOT_FOUND
        rows.append(Row(day, who, t.pages, t.start, t.heads[0] if heads and t.heads else "",
                        t.last_head if heads else "", note, t.initials, heads=", ".join(t.heads),
                        head_list=list(t.heads), at=t.at))
    return rows


def title_page_total(ing) -> int:
    """How many title pages a transcript begins with (extract_regex.title_page_count), as rows_from counts its
    takes."""
    title_text, _ = strip_line_numbers(ing.text)
    return title_page_count(title_text)


def name_rows(rows: list[Row], known: dict[str, str] | None = None) -> dict[str, str]:
    """Gives every take of a reporter the same name. One transcript's title page may name a reporter where
    another's doesn't, and the run sheet would say "DS" on one day and "Dana" on the next. known: names
    learned before (initials -> name). Changes the rows in place and returns what is known now."""
    def initials(r: Row) -> str:
        return r.initials or (r.reporter.lower() if re.fullmatch(r"[A-Z]{2,3}", r.reporter) else "")
    names = dict(known or {})
    for r in rows:
        if initials(r) and r.reporter and r.reporter != initials(r).upper():
            names.setdefault(initials(r), r.reporter)
    for r in rows:
        if initials(r) and r.reporter == initials(r).upper() and initials(r) in names:
            r.reporter = names[initials(r)]
    return names


def transcript_pages(ing) -> int:
    """The pages of an ingested transcript that are written and billed (not the word index): the count read
    four ways and reconciled (ing.count, takes.count_pages)."""
    return body_pages(ing.marks, ing.page_count, getattr(ing, "count", None))


def to_date(v) -> date | None:
    """A cell's date: a date, a datetime or 'M/D/YYYY' text."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        found = find_dates(v)
        if found:
            m, d, y = (int(x) for x in found[0][2].split("/"))
            try:
                return date(y, m, d)
            except ValueError:
                return None
    return None


# ------------------------------------------------------------------ which run sheet is this case's

# words that don't tell one case from another: party roles, file-name words, places, months
_MONTHS = {date(2000, m, 1).strftime("%B").lower() for m in range(1, 13)} | \
          {date(2000, m, 1).strftime("%b").lower() for m in range(1, 13)} | {"sept"}
_NOISE = {"the", "of", "and", "et", "al", "ano", "inc", "llc", "llp", "pc", "corp", "co", "md", "as", "by", "an",
          "administrator", "administratrix", "executor", "executrix", "estate", "deceased", "individually",
          "guardian", "ad", "litem", "infant", "his", "her", "their", "parent", "natural", "on", "behalf",
          "page", "pages", "sheet", "run", "transcript", "transcripts", "copy", "trial", "minutes", "new", "york",
          "city", "county", "state", "people"} | _MONTHS


def index_numbers(text: str) -> list[str]:
    """Every index number in text, as 'number/year' ('712345-2021' -> '712345/2021')."""
    out = []
    for num, yr in re.findall(r"(?<!\d)(\d{3,7})\s*[-/]\s*(\d{4}|\d{2})(?!\d)", text or ""):
        n = norm_index(num, yr)
        if n and n not in out:
            out.append(n)
    return out


# words of a run sheet's file name, and of a company's name, that are never the party's own
_SHEET_WORDS = {"page", "pages", "sheet", "run", "transcript", "transcripts", "copy", "trial", "minutes"}
_ENTITY = {"inc", "llc", "llp", "pllc", "lp", "pc", "corp", "corporation", "co", "company", "ltd", "md", "do", "dds",
           "esq"}


def _sides(name: str) -> list[str]:
    """One word per side of the 'v.' that tells the case apart: the last name of the first party ('Jane Roe,
    as Administrator ...' -> 'roe'), or the last telling word of a company's name."""
    parts = re.split(r"\s(?:v|vs|versus|against)\.?\s", " " + (name or "").lower().replace("-against-", " v ") + " ",
                     maxsplit=1)
    keys = []
    for p in parts:
        first = re.split(r",|\sas\s|\sand\s|&|\bet\s+al\b", p, maxsplit=1)[0]  # the first party
        words = [w for w in re.findall(r"[a-z][a-z']+", first) if w not in _NOISE and w not in _ENTITY]
        # "The City of New York": nothing but common words, so its last word (not "Run Sheet" of a file name)
        raw = [w for w in re.findall(r"[a-z][a-z']+", first) if w not in _SHEET_WORDS and w not in _MONTHS]
        keys.append(words[-1] if words else raw[-1] if raw else "")
    return keys


def same_name(a: str, b: str) -> bool:
    """'Jane Roe, as Administrator ... v. Sam Poe, M.D., ...' and 'ROE v Poe et al.': the first party on each
    side has the same last name."""
    x, y = _sides(a), _sides(b)
    return len(x) == len(y) == 2 and all(x) and x == y


def runsheets_folder(s: Settings, general: Path | str | None = None) -> Path:
    """Where a new run sheet goes: the run sheets' folder under Settings -> Options -> Folders, else the Trial
    Runs folder inside `general`, the case's own folder (batch.out_dir_for; `general` itself, Save to, without a
    folder for each case: see Settings.folder_for)."""
    return s.folder_for("runsheet", general)


def old_folder() -> Path:
    """Documents\\YinIt Run Sheets, where earlier versions kept every case's run sheet: still looked in, so a
    trial whose run sheet was started there gets the next days' takes on it too."""
    return Path.home() / "Documents" / OLD_FOLDER


def sheet_folders(s: Settings, general: Path | str | None, folders: list[Path],
                  case: CaseInfo | None = None) -> list[Path]:
    """The folders a case's run sheet is looked for in: where a new one would go (runsheets_folder), the case's
    folder itself (`general`: where one went before the Trial Runs folder, 2.13 and earlier), `folders` (the
    transcripts' own), the folders of the case's run sheets the records list (recorded_folders: a trial's
    days can go to different folders) and the old run sheets folders (old_folder, and Documents\\DjinnIt Run
    Sheets when the rename left it there)."""
    recorded = recorded_folders(case, s) if case is not None else []
    case_dir = [Path(general)] if general else []
    return list(dict.fromkeys([runsheets_folder(s, general), *case_dir, *map(Path, folders), *recorded, old_folder(),
                               Path.home() / "Documents" / OLD_DJINNIT_FOLDER]))


def recorded_folders(case: CaseInfo, s: Settings) -> list[Path]:
    """The folders of the run sheets the records list for this case (the same index number, or the same case
    name; newest first): a trial's run sheet is found there whatever folder a later day's files go to (its days'
    transcripts in different folders, a case folder named by date or "(2)"). [] when the records can't be read."""
    from .deliver import ledger_for  # (deliver imports this module)
    try:
        made = ledger_for(s).activity(kind="runsheet")
    except Exception as e:  # (the records locked or damaged: the other folders are still looked in)
        log_error("could not read the records for the case's run sheet", e)
        return []
    out = []
    for a in made:
        if a.file_path and matches(Found(Path(a.file_path), a.case_name, index_numbers(a.index_no)),
                                   case.get("case_name"), case.get("index_no")):
            out.append(Path(a.file_path).parent)
    return list(dict.fromkeys(out))


def meta_sheet(wb) -> str | None:
    """The name of a workbook's hidden sheet about its case (META, or OLD_META in one made before the rename),
    or None."""
    return next((n for n in (META, OLD_META) if n in wb.sheetnames), None)


class SheetUnreadable(OSError):
    """A workbook couldn't be opened just now: held by another program (Dropbox syncing it, an antivirus scan).
    Not the same as "not a run sheet": find_sheets and add_takes raise it rather than start
    "... - Run Sheet (2).xlsx" beside the trial's own run sheet. Its message names the file, for the window."""


def read_info(path: Path, strict: bool = False) -> Found | None:
    """What a workbook says about its case, or None when it isn't a run sheet. A workbook that can't be opened
    just now (OSError: held by another program) is None too (and logged), unless strict: then SheetUnreadable,
    for a caller who must not take "can't read it now" for "no run sheet". One that openpyxl can't open for good
    (damaged, password-protected, no zip at all) is not a run sheet: None, strict or not, so a case's
    "712345-2021 Roe v Poe exhibit list.xlsx" doesn't hold up its run sheet for good."""
    from openpyxl import load_workbook
    # (the log never names the file: its name is often the case's, and the log goes out with problem reports)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except OSError as e:
        if strict:
            raise SheetUnreadable(f"{path.name} could not be read ({e}). If it is being synced or scanned, "
                                  f"try again in a moment.") from e
        log_error("could not open a workbook just now", e)
        return None
    except Exception as e:  # (zipfile.BadZipFile, openpyxl's InvalidFileException, a KeyError inside...)
        log_error("could not open a workbook: not a run sheet", e)
        return None
    try:
        if meta_sheet(wb):
            info = _read_meta(wb)
            # "YinIt run sheet 1", any version, or one made under the old name
            if info.get("format", [""])[0].startswith((FORMAT.rsplit(" ", 1)[0], OLD_FORMAT)):
                return Found(path, (info.get("case") or [""])[0], [i for i in info.get("index", []) if i])
        if {"Setup", "Rates", "Job", "Invoice"} <= set(wb.sheetnames):
            return None  # the invoice spreadsheet: its Page Log looks like a run sheet but isn't one
        for ws in wb.worksheets:  # someone else's run sheet: the columns are there, the case is in the name
            top = [[c for c in row if c is not None] for row in ws.iter_rows(max_row=12, values_only=True)]
            if _header_row(top) is not None:
                # only text: a date cell ("2026-06-02") is not an index number
                text = " ".join([path.stem] + [c for row in top[:3] for c in row if isinstance(c, str)])
                return Found(path, path.stem, index_numbers(text), ours=False)
    except Exception as e:  # (a workbook openpyxl can open but not read through: not a run sheet)
        log_error("could not read through a workbook: not a run sheet", e)
        return None
    finally:
        wb.close()
    return None


def _header_row(rows: list[list]) -> int | None:
    """The row (from 0) of a run sheet's column headings: one with Date, Reporter and Pages in it."""
    for i, row in enumerate(rows):
        cells = " | ".join(str(c).lower() for c in row if c is not None)
        if "reporter" in cells and "page" in cells and "date" in cells:
            return i
    return None


def matches(found: Found, case_name: str, index_no: str) -> str:
    """'index', 'name' or '' - how found is this case's run sheet."""
    if set(index_numbers(index_no)) & set(found.index_nos):
        return "index"
    return "name" if case_name and same_name(case_name, found.case_name) else ""


def find_sheets(case: CaseInfo, folders: list[Path]) -> list[Found]:
    """The run sheets in these folders that may be this case's: same index number first, then this app's own
    before others, then the newest. Excel's lock files, half-saved files and backups are skipped, and so is a
    workbook openpyxl can't open for good (read_info: not a run sheet). One that can't be opened just now (see
    SheetUnreadable) is skipped too, unless its name says it may be the case's
    ("June 2026 712345-2021 Jane Roe v. Sam Poe - Run Sheet.xlsx"): then the error is raised, as adding the takes
    to a new run sheet would split the trial over two."""
    out, seen = [], set()
    for folder in folders:
        try:
            files = sorted(Path(folder).glob("*.xlsx"))
        except OSError:
            continue
        for f in files:
            key = str(f.resolve()).lower()
            if f.name.startswith("~$") or f.name.endswith(".saving.xlsx") or BACKUP in f.name \
                    or OLD_BACKUP in f.name or key in seen:
                continue
            seen.add(key)
            try:
                info = read_info(f, strict=True)
            except SheetUnreadable as e:
                if matches(Found(f, f.stem, index_numbers(f.stem)), case.get("case_name"), case.get("index_no")):
                    raise
                # (its cause, not the error: that one's message starts with the file's name)
                log_error("could not open a workbook just now: skipped, not named for the case", e.__cause__ or e)
                continue
            if info:
                info.why = matches(info, case.get("case_name"), case.get("index_no"))
                if info.why:
                    out.append(info)
    return sorted(out, key=lambda x: (x.why != "index", not x.ours, -x.path.stat().st_mtime))


def choose(case: CaseInfo, s: Settings, folders: list[Path], target: str | None) -> Path | None:
    """The run sheet to add to (None: start a new one). target: what the user chose (see RunSheetOpts);
    without a choice, Settings.runsheet_existing decides - "ask" (nobody to ask here) adds to a run sheet
    with the same index number only."""
    if target is not None:
        return Path(target) if target else None
    if s.runsheet_existing == "new":
        return None
    found = find_sheets(case, folders)
    if s.runsheet_existing != "add":
        found = [f for f in found if f.why == "index"]
    return found[0].path if found else None


# ------------------------------------------------------------------ writing

def new_path(case: CaseInfo, rows: list[Row], s: Settings, out_dir: Path | str | None = None) -> Path:
    """A free file name for a new run sheet in runsheets_folder (out_dir: the case's own folder, so its Trial
    Runs folder), from Settings.runsheet_filename_pattern and the first day of the takes
    ("June 2026 712345-2021 Jane Roe v. Sam Poe - Run Sheet.xlsx")."""
    days = [r.day for r in rows if r.day]
    first = min(days) if days else date.today()
    from .pdfout import short_caption
    try:
        name = s.runsheet_filename_pattern.format(
            month=first.strftime("%B"), year=first.year, case=short_caption(case.get("case_name")) or "Case",
            index=case.get("index_no").replace("/", "-") or "no index", date=us_date(first).replace("/", "-"))
    except (KeyError, IndexError, ValueError, AttributeError, TypeError):  # a bad custom pattern
        name = f"Run Sheet - {case.get('index_no').replace('/', '-')}"
    name = re.sub(r"\s+", " ", name).strip(" -")
    return unique_path(runsheets_folder(s, out_dir) / (safe_filename(name or "Run Sheet") + ".xlsx"))


def add_takes(case: CaseInfo, opts: RunSheetOpts, s: Settings, folders: list[Path],
              out_dir: Path | str | None = None) -> Path:
    """Adds opts.rows to this case's run sheet (see choose) or a new one, and returns its path. out_dir: the
    case's own folder: a new one goes in its Trial Runs folder (unless Settings give run sheets a folder:
    runsheets_folder); folders: where else to look (the transcript's own folder; see sheet_folders). Sets
    opts.path, created, added, skipped and added_pages. ValueError when there are no rows or the chosen file
    isn't a run sheet; SheetUnreadable when the case's run sheet can't be opened just now (nothing is started in
    its place)."""
    if not opts.rows:
        raise ValueError(NO_RUNSHEET)
    with _LOCK:  # one change at a time: the window and a batch may add to the same run sheet
        path = choose(case, s, sheet_folders(s, out_dir, folders, case), opts.target)
        if path and not path.exists():  # chosen, then moved or deleted: a new one in its place
            path = None
        info = read_info(path, strict=True) if path else None
        if path and not info:
            raise ValueError(f"{path.name} is not a run sheet: it has no Date, Reporter and Pages columns, or it "
                             "is damaged or has a password")
        if info and not info.ours:
            _append_theirs(path, opts)
        else:
            path = path or new_path(case, opts.rows, s, out_dir)
            opts.created = info is None
            _write_ours(path, case, opts, exists=info is not None, s=s)
        opts.path = path
        return path


def is_locked(path: Path | str) -> bool:
    """The file is there and another program holds it so that it can't be written (a run sheet open in Excel):
    asked before Generate makes anything, so that nothing is made while the takes can't be added."""
    if not Path(path).is_file():  # (not there, or a folder: Windows refuses to open a folder with PermissionError)
        return False
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True
    except OSError:  # (gone meanwhile: nothing holds it)
        return False


def _save(wb, path: Path) -> None:
    """Saves in one step: a run sheet open in Excel stays as it was (PermissionError). Dropbox or OneDrive may
    hold the file for a moment while they sync it, so a locked file is tried again a few times first."""
    path.parent.mkdir(parents=True, exist_ok=True)
    # each copy of the app and each thread its own: two saving at once must not write or delete one another's
    tmp = path.with_name(f"{path.stem}.saving-{os.getpid()}-{threading.get_ident()}.xlsx")
    try:
        wb.save(tmp)
        for attempt in range(3):
            try:
                os.replace(tmp, path)
                break
            except PermissionError:
                if attempt == 2:
                    raise
                time.sleep(0.3)
    except PermissionError as e:
        raise PermissionError(f"{path.name} is open in another program - close it in Excel and try again") from e
    finally:
        tmp.unlink(missing_ok=True)


def _number(v) -> int | None:
    """A cell's whole number: 12, 12.0, "12" or a sum typed in ("=329-320+1"); None otherwise."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v)
    if isinstance(v, str) and re.fullmatch(r"=[\d\s+\-*()]+", v):
        try:
            return int(_sum(ast.parse(v[1:].strip(), mode="eval").body))
        except (SyntaxError, ValueError, TypeError, RecursionError):
            return None
    return None


def _sum(node) -> int:
    """The value of a sum of whole numbers (+, - and * only)."""
    ops = {ast.Add: lambda a, b: a + b, ast.Sub: lambda a, b: a - b, ast.Mult: lambda a, b: a * b}
    if isinstance(node, ast.Constant) and type(node.value) is int:
        return node.value
    if isinstance(node, ast.BinOp) and type(node.op) in ops:
        return ops[type(node.op)](_sum(node.left), _sum(node.right))
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.USub):
        return -_sum(node.operand)
    raise ValueError("not a sum")


def _unique(rows: list[Row]) -> list[Row]:
    """The rows less any that repeat an earlier one (the same transcript dropped twice)."""
    out: list[Row] = []
    for r in rows:
        if not any(r.same_take(o) for o in out):
            out.append(r)
    return out


def _read_ours(ws) -> list[Row]:
    """The takes on a run sheet of this app, top to bottom; rows with nothing typed in them are skipped."""
    rows, prev_end = [], None
    for r in range(FIRST, ws.max_row + 1):
        v = {k: ws.cell(r, c).value for k, c in COL.items()}
        extra = {c: ws.cell(r, c).value for c in range(len(COLUMNS) + 1, ws.max_column + 1)
                 if ws.cell(r, c).value is not None}
        if all(v[k] in (None, "") for k in HAND) and not extra:
            continue
        pages = _number(v["pages"])
        start = v["start"]
        if _is_formula(start):  # a formula: the page after the row above
            start = prev_end + 1 if prev_end is not None else None
        elif not isinstance(start, (int, float)) or isinstance(start, bool):
            # blank or text ("TBD"): no first page. Counting on from the row above would give the row pages it
            # never had (a scan's take, "not found"), so a later take on them would be skipped as already there.
            start = None
        day = to_date(v["date"])
        # what was typed is written back as it was: a formula, or a date or page count that isn't one ("TBD").
        # (Text that starts with "=", a running head "== DIRECT ==" kept as text, is no formula: _as_text)
        raw = {k: v[k] for k in HAND if _is_formula(v[k]) and ws.cell(r, COL[k]).data_type == "f"}
        if v["date"] not in (None, "") and day is None:
            raw["date"] = v["date"]
        if v["pages"] not in (None, "") and pages is None:
            raw["pages"] = v["pages"]
        text = {k for k in raw if ws.cell(r, COL[k]).data_type != "f"}
        row = Row(day, str(v["reporter"] or ""), pages, int(start) if start is not None else None,
                  str(v["witness_start"] or ""), str(v["witness_end"] or ""), str(v["note"] or ""), extra=extra,
                  raw=raw, text=text, origin=r)
        prev_end = row.end() if row.end() is not None else prev_end
        rows.append(row)
    return rows


def _in_order(old: list[Row], new: list[Row]) -> list[Row]:
    """The rows of the run sheet: the old ones where they were (in the order the user keeps them), each new one
    before the first old row that comes later by date, then first page. Old rows without a date or first page
    are passed over."""
    key = lambda r: (r.day or date.max, r.start if r.start is not None else float("inf"))
    rows = list(old)
    olds = {id(r) for r in old}
    for n in sorted(new, key=key):
        at = next((i for i, o in enumerate(rows) if id(o) in olds and o.day and o.start is not None
                   and key(o) > key(n)), len(rows))
        rows.insert(at, n)
    return rows


def _write_ours(path: Path, case: CaseInfo, opts: RunSheetOpts, exists: bool, s: Settings | None = None) -> None:
    """Writes a run sheet of this app: a new one, or (exists) the one at path rewritten with its old rows and
    the new takes in order, and its tabs (By Reporter, Witnesses, and Who Ordered from the records s names: none
    without s). Sets opts.added, skipped and added_pages."""
    from openpyxl import Workbook, load_workbook
    if exists:
        wb = load_workbook(path)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        old = _read_ours(ws)
        meta = _read_meta(wb)
        _restore_heads(old, _kept_heads(meta))
        case_name = (meta.get("case") or [case.get("case_name")])[0] or case.get("case_name")
        indexes = [i for i in meta.get("index", []) if i]
        known = dict(x.split("=", 1) for x in meta.get("reporter", []) if "=" in x)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        old, case_name, indexes, known, meta = [], case.get("case_name"), [], {}, {}
    new = [r for r in _unique(opts.rows) if not any(r.same_take(o) for o in old)]
    opts.added, opts.skipped = len(new), len(opts.rows) - len(new)
    opts.added_pages = sum(r.pages or 0 for r in new)
    if exists and not new:
        return  # nothing to add: the run sheet stays as it is (it may be open in Excel)
    for i in index_numbers(case.get("index_no")):
        if i not in indexes:
            indexes.append(i)
    rows = _in_order(old, new)
    known = name_rows(rows, known)
    _fill_sheet(ws, rows, case_name, indexes, case)
    _by_reporter(wb, rows, ws.title)
    _witnesses(wb, rows)
    attorneys = _kept_attorneys(meta, case, *opts.days)
    orders = order_rows(case_name, indexes, case, s, _in_turn(attorneys)) if s is not None else None
    if orders is not None:  # (the records unreadable just now: the tab as it was, not "nothing in the records")
        _who_ordered(wb, orders)
    _meta(wb, case_name, indexes, known, rows, attorneys)
    _save(wb, path)


def _moved(value, origin: int | None, row: int, col: int):
    """A formula written on another row than it was on: its references move with it (=A5 on row 5 is =A7 on
    row 7), as when Excel moves a row."""
    if not _is_formula(value) or origin in (None, row):
        return value
    from openpyxl.formula.translate import Translator
    from openpyxl.utils import get_column_letter
    letter = get_column_letter(col)
    try:
        return Translator(value, origin=f"{letter}{origin}").translate_formula(f"{letter}{row}")
    except Exception:  # a formula openpyxl can't read: as it was
        return value


def _fill_sheet(ws, rows: list[Row], case_name: str, indexes: list[str], case: CaseInfo) -> None:
    """Writes the whole Run Sheet tab again: the case at the top, the totals of the rows shown, the headings,
    and a row per take with its formulas."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.properties import PageSetupProperties
    for rng in list(ws.merged_cells.ranges):  # merged cells can't be written: unmerged where the sheet is rewritten
        if rng.min_col <= len(COLUMNS) or rng.max_row >= FIRST:
            ws.unmerge_cells(str(rng))
    for r in range(1, max(ws.max_row, FIRST) + 1):  # start over: the rows are written again in order
        # (the user's own cells to the right of the headings stay; those of the takes move with their rows)
        for c in range(1, (len(COLUMNS) if r < FIRST else max(ws.max_column, len(COLUMNS))) + 1):
            cell = ws.cell(r, c)
            cell.value = None
            cell.style = "Normal"
    fill = lambda color: PatternFill("solid", start_color=color, end_color=color)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    from .pdfout import short_caption
    ws["A1"] = short_caption(case_name, 120) or "Run sheet"
    ws["A1"].font = Font(size=14, bold=True)
    details = [f"Index No. {', '.join(indexes)}" if indexes else ""]
    if case.get("judge"):
        details.append(f"Justice {case.get('judge')}")
    if case.get("part"):
        details.append(f"Part {case.get('part')}")
    ws["A2"] = "   ·   ".join(d for d in details if d)
    ws["A2"].font = Font(color="595959")
    # the totals of the rows shown: filter the Reporter column to see one reporter's pages
    ws[f"{L['witness_start']}2"] = "Pages shown:"
    ws[f"{L['witness_start']}2"].alignment = Alignment(horizontal="right")
    ws[f"{L['witness_end']}2"] = f"=SUBTOTAL(109,{L['pages']}{FIRST}:{L['pages']}{LAST})"
    ws[f"{L['witness_end']}2"].font = Font(bold=True)
    ws[f"{L['witness_end']}2"].alignment = Alignment(horizontal="left")
    ws[f"{L['note']}2"] = f'="Rows shown: "&SUBTOTAL(103,{L["reporter"]}{FIRST}:{L["reporter"]}{LAST})'
    ws[f"{L['note']}2"].font = Font(color="595959")
    for key, heading, hand, width in COLUMNS:
        c = COL[key]
        mark = ws.cell(3, c, "Write here" if hand else "Don't write here")
        mark.fill, mark.font = fill(GREEN if hand else RED), Font(size=6, bold=True, color="FFFFFF" if hand else "000000")
        mark.alignment = Alignment(horizontal="center")
        head = ws.cell(HEAD_ROW, c, heading)
        head.fill, head.font, head.border = fill(GREEN if hand else ORANGE), Font(bold=True), box
        head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L[key]].width = width
    ws.row_dimensions[HEAD_ROW].height = 30

    prev_end, prev_end_row, prev_day = None, None, None
    for i, row in enumerate(rows):
        r = FIRST + i
        a, d, f, g = (f"{L[k]}{r}" for k in ("date", "day_take", "pages", "start"))
        above = r - 1
        values = {
            "date": row.day, "reporter": row.reporter, "pages": row.pages,
            "witness_start": row.witness_start or None, "witness_end": row.witness_end or None, "note": row.note or None,
            "weekday": f'=IF({a}="","",TEXT({a},"dddd"))',
            "day_take": f'=IF({a}="","",IF({L["note"]}{r}="{TITLE_ONLY}",0.5,'
                        f'COUNTIFS({L["date"]}${FIRST}:{a},{a},{L["note"]}${FIRST}:{L["note"]}{r},"<>{TITLE_ONLY}")))',
            "run_take": f'=IF({a}="","",N({L["run_take"]}{above})+IF({a}={L["date"]}{above},{d}-N({L["day_take"]}{above}),{d}))',
            "end": f'=IF(OR({f}="",{g}=""),"",{g}+{f}-1)',
        }
        if row.start is not None and prev_end is not None and row.start == prev_end + 1:
            values["start"] = f"={L['end']}{prev_end_row}+1"  # the row the last page came from (not a blank one)
        else:
            values["start"] = row.start
        for key, value in row.raw.items():
            values[key] = _moved(value, row.origin, r, COL[key])
        first_of_day = row.day != prev_day
        for key, *_ in COLUMNS:
            cell = ws.cell(r, COL[key], values[key])
            # (a formula the user typed, kept in raw, stays a formula; text, "=TBD" kept as text, stays text)
            if key in HAND and (key not in row.raw or key in row.text):
                _as_text(cell)
            cell.border = box
            if i % 2:
                cell.fill = fill(BAND)
            if key in ("date", "start") and first_of_day:
                cell.font = Font(bold=True)
            if key in ("weekday", "day_take", "run_take", "pages", "start", "end"):
                cell.alignment = Alignment(horizontal="center")
        ws.cell(r, COL["date"]).number_format = "m/d/yyyy"
        for c, value in row.extra.items():
            ws.cell(r, c, _moved(value, row.origin, r, c))
        if row.end() is not None:
            prev_end, prev_end_row = row.end(), r
        prev_day = row.day
    ws.freeze_panes = f"A{FIRST}"
    ws.auto_filter.ref = f"A{HEAD_ROW}:{L['note']}{max(FIRST, FIRST + len(rows) - 1)}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = f"{HEAD_ROW}:{HEAD_ROW}"


def _by_reporter(wb, rows: list[Row], sheet: str = SHEET) -> None:
    """The 'By Reporter' sheet: each reporter's takes and pages (formulas over the run sheet, the tab named
    sheet), their share of all the pages, and the days they have takes on (how many, and which: written from the
    rows as they are now, so a row typed in Excel counts once the run sheet is written again)."""
    from openpyxl.styles import Alignment, Font
    if BY_REPORTER in wb.sheetnames:
        del wb[BY_REPORTER]
    ws = wb.create_sheet(BY_REPORTER, 1)
    for c, (heading, width) in enumerate((("Reporter", 18), ("Takes", 9), ("Pages", 9), ("Share", 9), ("Days", 7),
                                          ("Days worked", 60)), 1):
        ws.cell(1, c, heading).font = Font(bold=True)
        ws.column_dimensions[chr(64 + c)].width = width
    tab = "'" + sheet.replace("'", "''") + "'"
    ref = lambda k: f"{tab}!${L[k]}${FIRST}:${L[k]}${LAST}"
    names: dict[str, str] = {}  # COUNTIF and SUMIF don't tell "Dana" from "DANA": one row for both
    days: dict[str, list[date]] = {}
    for r in rows:
        if r.reporter and not _is_formula(r.reporter):
            names.setdefault(r.reporter.casefold(), r.reporter)
            if r.day and r.day not in days.setdefault(r.reporter.casefold(), []):
                days[r.reporter.casefold()].append(r.day)
    total = len(names) + 2
    for i, (key, name) in enumerate(names.items(), 2):
        ws.cell(i, 1, name)
        ws.cell(i, 2, f'=COUNTIFS({ref("reporter")},A{i},{ref("note")},"<>{TITLE_ONLY}")')
        ws.cell(i, 3, f'=SUMIF({ref("reporter")},A{i},{ref("pages")})')
        ws.cell(i, 4, f'=IF(N(C${total})=0,"",C{i}/C${total})').number_format = "0%"
        worked = sorted(days.get(key, []))
        ws.cell(i, 5, len(worked))
        ws.cell(i, 6, ", ".join(us_date(d) for d in worked)).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(total, 1, "Total").font = Font(bold=True)
    # (no reporter named, a scan's takes: 0, not a SUM of the row above that is its own row, a circular reference)
    ws.cell(total, 2, f"=SUM(B2:B{total - 1})" if names else 0).font = Font(bold=True)
    ws.cell(total, 3, f"=SUM(C2:C{total - 1})" if names else 0).font = Font(bold=True)
    every = sorted({r.day for r in rows if r.day})
    ws.cell(total, 5, len(every)).font = Font(bold=True)  # (the trial's days: a day two reporters shared is one)


def heads_of(row: Row) -> list[str]:
    """The running heads of a take, as printed: those read from its pages (Row.head_list), else its Witness Start
    and End cells (a row typed in, or one of a run sheet made before the heads were kept)."""
    return row.head_list or [h for h in dict.fromkeys((row.witness_start, row.witness_end)) if h]


def witness_rows(rows: list[Row]) -> list[tuple[str, list[date], list[str], int]]:
    """Each running head the takes printed, as first printed and once ('PROCEEDINGS' and 'Proceedings' are one,
    as takes.find_takes reads them), in the order of the takes: (the head, the days it was on, the reporters of
    the takes it was on, the pages of those takes)."""
    out: dict[str, list] = {}
    for r in rows:
        for h in heads_of(r):  # ("== DIRECT ==" too: written as text, _as_text)
            entry = out.setdefault(h.casefold(), [h, [], [], 0])
            if r.day and r.day not in entry[1]:
                entry[1].append(r.day)
            if r.reporter and r.reporter.casefold() not in {x.casefold() for x in entry[2]}:
                entry[2].append(r.reporter)
            entry[3] += r.pages or 0
    return [(h, sorted(d), who, pages) for h, d, who, pages in out.values()]


def _witnesses(wb, rows: list[Row]) -> None:
    """The 'Witnesses' sheet: every running head the takes printed (witness_rows), as printed: who was on the
    stand is not read from them, the heads say it in their own words. Always the heads read from the pages, also
    when Settings.runsheet_witness leaves the Witness cells blank for the user."""
    from openpyxl.styles import Alignment, Font
    if WITNESSES in wb.sheetnames:
        del wb[WITNESSES]
    ws = wb.create_sheet(WITNESSES, 2)
    for c, (heading, width) in enumerate((("Running head (as printed)", 48), ("Days", 30), ("Reporters", 24),
                                          ("Pages of its takes", 10)), 1):
        cell = ws.cell(1, c, heading)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[chr(64 + c)].width = width
    found = witness_rows(rows)
    if not found:
        ws.cell(2, 1, "No running heads read from the transcripts yet.").font = Font(italic=True, color="595959")
    for i, (head, days, who, pages) in enumerate(found, 2):
        ws.cell(i, 1, head)
        ws.cell(i, 2, ", ".join(us_date(d) for d in days))
        ws.cell(i, 3, ", ".join(who))
        ws.cell(i, 4, pages)
        for c in (1, 2, 3):
            _as_text(ws.cell(i, c))
            ws.cell(i, c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


ORDER_HEADS = ("", "Firm", "Side", "Attorney", "Days", "Speed", "Invoice No.", "Invoice date", "Amount", "Status")
ORDER_WIDTHS = (2, 30, 10, 26, 30, 14, 14, 12, 26, 22)
_SIDES = {"plaintiff": "Plaintiff", "defense": "Defense", "other": "Other"}
NO_ORDERS = "No minute agreements or invoices for this case in your records yet."


def order_rows(case_name: str, indexes: list[str], case: CaseInfo, s: Settings,
               attorneys: list[Attorney] | None = None) -> list[list] | None:
    """The Who Ordered tab's rows, from the records (Settings.records_dir), for the case of this name or one of
    these index numbers (matches, as the case's run sheet is found; a record with other index numbers is another
    case's, whatever its caption): an invoice each, oldest first ([the firm's colour, firm (or the attorney
    billed, without one), side, the attorney billed (blank when named as the firm), the days, the speed (paid at,
    billed, or those offered, for a firm choosing from its invoice), the number, its date, the amount (paid,
    billed, or the offer), and Open, Paid with its date, or Void, with whose name it is in when made in another
    reporter's]), then each firm with a minute agreement but no invoice yet (its days, the rest blank).
    The side and colour are those of the firm's mark among `attorneys` (the trial's, _kept_attorneys; default the
    case's: firm_marks, as the window shows them), blank for a firm not among them. None when the records can't be
    read (logged)."""
    from .deliver import ledger_for  # (deliver imports this module)
    from .firm_marks import firm_id, marks, shade_of
    from .invoice_calc import fmt, money
    try:
        ledger = ledger_for(s)
        invoices = ledger.invoices()
        agreements = ledger.activity(kind="agreement")
    except Exception as e:  # (the records locked or damaged: the run sheet is still written)
        log_error("could not read the records for the run sheet's Who Ordered tab", e)
        return None
    index = ", ".join(indexes) or case.get("index_no")
    mine = set(index_numbers(index))

    def ours(name: str, idx: str) -> bool:
        """The record is this case's (as matches finds its run sheet), and names none but its index numbers."""
        theirs = index_numbers(idx or "")
        if mine and theirs and not mine & set(theirs):  # ("Garcia v. City of New York" is many cases)
            return False
        return bool(matches(Found(Path(), name or "", theirs), case_name, index))

    given = marks(case.attorneys if attorneys is None else attorneys)

    def money_text(amount: str) -> str:
        """'98' -> '$98.00'; a record's amount as it is when it isn't a sum of money ("" when there is none)."""
        if not str(amount or "").strip():
            return ""
        try:
            return fmt(money(amount))
        except Exception:
            return amount

    def firm_cells(firm: str, attorney: str) -> list:
        """The colour, the firm's name (the attorney's, without a firm) and its side, by its mark among
        `attorneys` (given); no colour and no side for a firm not among them."""
        mark = given.get(firm_id(Attorney(name=attorney, firm=firm)))
        return [shade_of(mark) if mark else None, firm or attorney, _SIDES.get(mark.side, "") if mark else ""]

    out, billed = [], set()
    for inv in sorted((i for i in invoices if ours(i.case_name, i.index_no)), key=lambda i: (i.created, i.invoice_no)):
        billed.add(firm_id(Attorney(name=inv.bill_to, firm=inv.firm)))
        billed.add(firm_id(Attorney(name=inv.bill_to)))  # (an agreement logged with the attorney's name only)
        if inv.status == "paid" and (inv.paid_speed or inv.amount_paid):
            speed = inv.paid_speed or inv.billed_speed
            amount = money_text(inv.amount_paid or inv.amounts.get(speed, ""))
        elif len(inv.amounts) > 1 or inv.billed_speed not in inv.amounts:  # (a firm choosing from its invoice)
            speed, amount = "offered: " + ", ".join(inv.amounts), inv.offered_text()
        else:
            speed, amount = inv.billed_speed, money_text(inv.amounts[inv.billed_speed])
        status = {"paid": f"Paid {us_date(inv.paid_date)}" if inv.paid_date else "Paid", "void": "Void"}.get(
            inv.status, "Open")
        if inv.reporter:
            status += f" (in {inv.reporter.upper()}'s name)"
        out.append(firm_cells(inv.firm, inv.bill_to) + [inv.bill_to if inv.firm else "", inv.dates, speed,
                                                        inv.invoice_no, us_date(inv.created[:10]), amount, status])
    waiting: dict[str, list] = {}  # firms with an agreement and no invoice: their days, each once
    for a in sorted((a for a in agreements if ours(a.case_name, a.index_no)), key=lambda a: (a.ts, a.id)):
        key = firm_id(Attorney(name=a.attorney, firm=a.firm))
        if not key or key in billed:
            continue
        row = waiting.setdefault(key, firm_cells(a.firm, a.attorney) + [a.attorney if a.firm else "", [], "", "",
                                                                         "", "", "Agreement, no invoice yet"])
        for d in [x.strip() for x in (a.dates or "").split(",") if x.strip()]:
            if d not in row[4]:
                row[4].append(d)
    for row in waiting.values():
        row[4] = ", ".join(row[4])
        out.append(row)
    return out


def _who_ordered(wb, rows: list[list]) -> None:
    """The 'Who Ordered' sheet: order_rows, under its headings, each firm's colour in the narrow first column."""
    from openpyxl.styles import Alignment, Font, PatternFill
    at = wb.sheetnames.index(ORDERS) if ORDERS in wb.sheetnames else min(3, len(wb.sheetnames))
    if ORDERS in wb.sheetnames:
        del wb[ORDERS]
    ws = wb.create_sheet(ORDERS, at)
    for c, (heading, width) in enumerate(zip(ORDER_HEADS, ORDER_WIDTHS), 1):
        ws.cell(1, c, heading or None).font = Font(bold=True)
        ws.column_dimensions[chr(64 + c)].width = width
    if not rows:
        ws.cell(2, 2, NO_ORDERS).font = Font(italic=True, color="595959")
    for i, row in enumerate(rows, 2):
        colour, *cells = row
        if colour:
            ws.cell(i, 1).fill = PatternFill("solid", start_color=colour.lstrip("#"), end_color=colour.lstrip("#"))
        for c, value in enumerate(cells, 2):
            cell = ws.cell(i, c, value or None)
            _as_text(cell)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


def _orders_now(wb) -> list[list]:
    """The rows the Who Ordered sheet has now (as _who_ordered wrote them), to tell whether it changed."""
    if ORDERS not in wb.sheetnames:
        return []
    ws = wb[ORDERS]
    if ws.cell(2, 2).value == NO_ORDERS:
        return []
    out = []
    for r in range(2, ws.max_row + 1):
        fill = ws.cell(r, 1).fill
        colour = f"#{fill.start_color.rgb[-6:]}".lower() if fill and fill.fill_type == "solid" else None
        out.append([colour] + [ws.cell(r, c).value or "" for c in range(2, len(ORDER_HEADS) + 1)])
    return out


def refresh_orders(path: Path | str | None, case: CaseInfo, s: Settings, days: list[CaseInfo] = ()) -> bool:
    """Writes the Who Ordered sheet of the run sheet at path again (order_rows), after Generate made the case's
    invoices and agreements: the run sheet is written before them (deliver.generate, batch.fill_jobs). days: the
    case of each day just generated (batch.fill_jobs: a run sheet's days together), whose attorneys are added to
    those the run sheet keeps with the case's (_kept_attorneys), so a firm has the colour the window gives it. Not
    saved when the tab shows the same already and the firms kept are the same, nor for a run sheet this app
    didn't make.
    It is an aside: a run sheet open in Excel, or one that can't be read (or the records), is logged and left as
    it was (False). True when it was saved."""
    from openpyxl import load_workbook
    from .firm_marks import firm_id
    if not path or not Path(path).is_file():
        return False
    try:
        with _LOCK:
            info = read_info(Path(path))
            if not info or not info.ours:
                return False
            wb = load_workbook(path)
            meta = _read_meta(wb)
            kept = _kept_attorneys(meta)
            trial = _kept_attorneys(meta, case, *days)
            name = (meta.get("case") or [case.get("case_name")])[0] or case.get("case_name")
            rows = order_rows(name, [i for i in meta.get("index", []) if i], case, s, _in_turn(trial))
            if rows is None:
                return False
            same = [[r[0].lower() if r[0] else None] + [c or "" for c in r[1:]] for r in rows]
            firms = lambda found: [(day, firm_id(a), a.party) for day, a in found]
            if ORDERS in wb.sheetnames and _orders_now(wb) == same and firms(trial) == firms(kept):
                return False
            _who_ordered(wb, rows)
            _set_attorneys(wb[meta_sheet(wb)], trial)
            _save(wb, Path(path))
            return True
    except Exception as e:  # (open in Excel, or synced just now: the invoices are made all the same)
        log_error("could not write the run sheet's Who Ordered tab again", e)
        return False


def _read_meta(wb) -> dict[str, list[str]]:
    """What the hidden sheet (_meta) says, by key: {"case": [...], "index": [...], "reporter": [...], ...}."""
    meta: dict[str, list[str]] = {}
    if meta_sheet(wb):
        for row in wb[meta_sheet(wb)].iter_rows(values_only=True):
            key, value = (tuple(row) + (None, None))[:2]  # (a sheet cut down to one column by hand: no values)
            if key:
                meta.setdefault(str(key), []).append("" if value is None else str(value))
    return meta


def _take_key(row: Row) -> str:
    """What finds a take's heads again on the hidden sheet: its day, first page and reporter."""
    day = row.day.isoformat() if row.day else ""
    return f"{day}|{'' if row.start is None else row.start}|{row.reporter.casefold()}"


def _json_rows(meta: dict[str, list[str]], key: str, size: int) -> list[list]:
    """The hidden sheet's rows of this key that are JSON lists of `size` items (what _meta wrote), in order; a
    cell changed by hand into something else is passed over."""
    out = []
    for value in meta.get(key, []):
        try:
            item = json.loads(value)
        except ValueError:
            continue
        if isinstance(item, list) and len(item) == size:
            out.append(item)
    return out


def _kept_heads(meta: dict[str, list[str]]) -> list[tuple[str, list[str]]]:
    """The heads the hidden sheet kept, a (_take_key, heads) each, in the order of the takes ([] for a run sheet
    made before they were kept)."""
    return [(k, h) for k, h in _json_rows(meta, "heads", 2)
            if isinstance(k, str) and isinstance(h, list) and h and all(isinstance(x, str) for x in h)]


def _restore_heads(rows: list[Row], kept: list[tuple[str, list[str]]]) -> None:
    """Gives the rows read back from the run sheet the heads kept for their takes (_kept_heads), each kept entry
    to one row: the entry of the row's own day, first page and reporter, else (an earlier take's pages corrected in
    Excel moved its first page, its reporter renamed) one of its day whose first and last heads are those its
    Witness cells name. Only while the cells still say what the heads do: a row whose cells were corrected by
    hand, or one typed in, gets none, and the Witnesses tab reads its cells (heads_of)."""
    def fits(r: Row, key: str, heads: list[str], exact: bool) -> bool:
        """The kept entry (key, heads) is the row's: its cells don't say other heads, and it has the row's key
        (exact), else the row's day (and its reporter when the cells are blank)."""
        printed = {h.casefold() for h in heads}  # (Witness End is the last head printed: it may repeat one)
        if any(c and c.casefold() not in printed for c in (r.witness_start, r.witness_end)):
            return False
        if exact:
            return key == _take_key(r)
        if not key.startswith((r.day.isoformat() if r.day else "") + "|"):
            return False
        # (Witness cells left blank, Settings.runsheet_witness "blank", say nothing: then its reporter's too)
        return bool(r.witness_start or r.witness_end) or key.endswith("|" + r.reporter.casefold())

    left = list(kept)
    for r in rows:
        r.head_list = []
    for exact in (True, False):  # (the exact ones first, so that none is taken by another row of the day)
        for r in (r for r in rows if not r.head_list):
            n = next((n for n, (key, heads) in enumerate(left) if fits(r, key, heads, exact)), None)
            if n is not None:
                r.head_list = list(left.pop(n)[1])


def _case_day(case: CaseInfo) -> str:
    """The case's first day, ISO ('2026-06-03'); "" without a date."""
    days = [d for d in (to_date(x) for _, _, x in find_dates(case.get("dates"))) if d]
    return min(days).isoformat() if days else ""


def _kept_attorneys(meta: dict[str, list[str]], *cases: CaseInfo) -> list[tuple[str, Attorney]]:
    """The attorneys each day of the run sheet named, a (its day, ISO; the attorney) each, earliest day first and
    in that day's own order: those the hidden sheet keeps (as each day was written), each case's in place of those
    kept for its day (_case_day: its attorneys as they are now, a party corrected since), placeholders left out.
    The marks of their firms (firm_marks.marks, each firm by the first of its entries) are then the window's over
    the trial's days (batch.trial_marks), whichever day is generated first."""
    from .firm_marks import firm_id
    days: dict[str, list[Attorney]] = {}
    for d, n, f, p in _json_rows(meta, "attorney", 4):
        days.setdefault(str(d or ""), []).append(Attorney(name=str(n or ""), firm=str(f or ""), party=str(p or "")))
    party: dict[str, str] = {}  # (a party corrected in the window is the firm's on every day: _same_firm_everywhere)
    for c in cases:
        days[_case_day(c)] = [Attorney(name=a.name, firm=a.firm, party=a.party) for a in c.attorneys
                              if not a.is_placeholder()]
        party.update((firm_id(a), a.party) for a in days[_case_day(c)] if a.party)
    for a in (a for group in days.values() for a in group):
        a.party = party.get(firm_id(a), a.party)
    out = []
    for day in sorted(days, key=lambda d: d or "9999"):
        seen = set()  # (one firm once a day)
        for a in days[day]:
            if (key := firm_id(a)) and key not in seen:
                seen.add(key)
                out.append((day, a))
    return out


def _in_turn(found: list[tuple[str, Attorney]]) -> list[Attorney]:
    """The attorneys of _kept_attorneys, earliest day first: firm_marks.marks gives the firms their marks in this
    order, as the window does over the trial's days (batch.trial_marks), whichever day is generated first."""
    return [a for _, a in found]


def _meta(wb, case_name: str, indexes: list[str], reporters: dict[str, str], rows: list[Row] = (),
          attorneys: list[tuple[str, Attorney]] = ()) -> None:
    """The hidden sheet that says what the run sheet is about (so it can be found again), the running heads of
    each take (rows), kept for the Witnesses tab when the run sheet is written again, and the attorneys of its days
    (_kept_attorneys), for the colours of the Who Ordered tab."""
    while meta_sheet(wb):  # (one made before the rename gets the new name)
        del wb[meta_sheet(wb)]
    ws = wb.create_sheet(META)
    ws.append(["format", FORMAT])
    ws.append(["case", case_name])
    for i in indexes or [""]:
        ws.append(["index", i])
    for initials, name in sorted(reporters.items()):
        ws.append(["reporter", f"{initials}={name}"])
    for r in rows:
        if r.head_list:
            ws.append(["heads", json.dumps([_take_key(r), r.head_list])])
    _set_attorneys(ws, attorneys)
    ws.append(["updated", datetime.now().isoformat(timespec="seconds")])
    ws.sheet_state = "hidden"


def _set_attorneys(ws, attorneys: list[tuple[str, Attorney]]) -> None:
    """The attorneys' rows of the hidden sheet (_kept_attorneys reads them), in place of those it had."""
    for r in range(ws.max_row, 0, -1):
        if ws.cell(r, 1).value == "attorney":
            ws.delete_rows(r)
    for day, a in attorneys:
        ws.append(["attorney", json.dumps([day, a.name, a.firm, a.party])])


# ------------------------------------------------------------------ someone else's run sheet

# headings (spaces and line breaks made single spaces, lower case); the first that fits a heading is its column
_THEIR_KEYS = [("date", r"\bdates?\b"), ("weekday", r"weekday|day of"), ("reporter", r"reporter"),
               ("day_take", r"day.?s take"), ("run_take", r"running take"), ("start", r"start(ing)? page"),
               ("end", r"end(ing)? page"), ("pages", r"\bpages?\b"), ("witness_start", r"witness start"),
               ("witness_end", r"witness end"), ("note", r"^notes?$")]


def _their_columns(ws) -> tuple[int, dict[str, int]]:
    """The row of the column headings (from 1) and the column of each heading found, by key (see COLUMNS)."""
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=12)]
    h = _header_row(rows)
    if h is None:
        raise ValueError(f"{ws.title}: no Date, Reporter and Pages columns")
    cols = {}
    for c, value in enumerate(rows[h], 1):
        text = " ".join(str(value or "").split()).lower()
        for key, pattern in _THEIR_KEYS:
            if key not in cols and text and re.search(pattern, text):
                cols[key] = c
                break
    missing = [name for key, name in (("date", "Date"), ("reporter", "Reporter"), ("pages", "Pages")) if key not in cols]
    if missing:
        raise ValueError(f"{ws.title}: no {' or '.join(missing)} column found")
    return h + 1, cols


def _is_formula(v) -> bool:
    """A cell value Excel reads as a formula ("=A5+1")."""
    return isinstance(v, str) and v.startswith("=")


def _as_text(cell) -> None:
    """A cell whose text starts with "=" (a running head printed "== DIRECT EXAMINATION ==") is kept as text:
    openpyxl would write it as a formula, a broken one Excel complains about."""
    if _is_formula(cell.value):
        cell.data_type = "s"


def _append_theirs(path: Path, opts: RunSheetOpts) -> None:
    """Adds the new takes to a run sheet this app did not make, keeping its layout: each on a row typed in
    ahead for its day, else on the first empty row after the takes, with the look and formulas of the row
    above. The file is first copied to "<name> (before YinIt).xlsx" (BACKUP), unless the old app kept a copy
    of it already ("(before DjinnIt)"): the original is kept once. Sets opts.added, skipped and added_pages."""
    from openpyxl import load_workbook
    from openpyxl.formula.translate import Translator
    wb = load_workbook(path)
    cached = load_workbook(path, data_only=True)  # the formulas' last values, as the spreadsheet saved them
    ws = next(w for w in wb.worksheets if _header_row([[c.value for c in r] for r in w.iter_rows(max_row=12)])
              is not None)
    vals = cached[ws.title]
    head, cols = _their_columns(ws)
    get = lambda r, k: ws.cell(r, cols[k]).value if k in cols else None
    num = lambda r, k: (lambda v: v if isinstance(v, (int, float)) and not isinstance(v, bool) else None)(
        vals.cell(r, cols[k]).value if k in cols else None)

    def day_at(r: int) -> date | None:
        v = get(r, "date")
        return to_date(vals.cell(r, cols["date"]).value) if _is_formula(v) else to_date(v)

    def pages_at(r: int) -> int | None:
        # a number, a sum typed in ("=25-1+1") or, for any other formula, the value Excel last saved for it
        p = get(r, "pages")
        n = _number(p)
        if n is None and _is_formula(p):
            n = num(r, "pages")
            n = int(n) if n is not None else None
        return n if n is not None and n > 0 else None

    # what is on it: the rows with a reporter or pages (the page numbers worked out where not saved)
    old, last, prev_end = [], head, None
    ends: dict[int, int] = {}  # the last page of each row that has one
    for r in range(head + 1, ws.max_row + 1):
        pages = pages_at(r)
        if not get(r, "reporter") and pages is None:  # an empty row, or a day typed in ahead
            continue
        start = get(r, "start")
        start = start if isinstance(start, (int, float)) and not _is_formula(start) else num(r, "start")
        if start is None and prev_end is not None and _is_formula(get(r, "start")):  # a formula never saved
            start = prev_end + 1
        row = Row(day_at(r), str(get(r, "reporter") or ""), pages, int(start) if start else None)
        end = num(r, "end")
        end = row.end() if row.end() is not None else int(end) if end is not None else None
        if end is not None:
            ends[r] = prev_end = end
        old.append(row)
        last = r
    new = [x for x in _unique(opts.rows) if not any(x.same_take(o) for o in old)]
    opts.added, opts.skipped = len(new), len(opts.rows) - len(new)
    opts.added_pages = sum(r.pages or 0 for r in new)
    if not new:
        return
    backup = path.with_name(f"{path.stem} {BACKUP}{path.suffix}")
    if not backup.exists() and not path.with_name(f"{path.stem} {OLD_BACKUP}{path.suffix}").exists():
        shutil.copy2(path, backup)
    typed = [cols.get(k) for k in ("date", "reporter", "pages", "witness_start", "witness_end", "note")]
    used: set[int] = set()

    def free(r: int) -> bool:  # no take on it, nor one of this run
        return r not in used and not get(r, "reporter") and pages_at(r) is None

    for row in new:
        # the row typed in ahead for this take's day, else the first empty row after the takes
        r = next((x for x in range(head + 1, ws.max_row + 1) if row.day and free(x) and day_at(x) == row.day), None)
        if r is None:
            r = next(x for x in range(last + 1, ws.max_row + 2 + len(new)) if free(x) and day_at(x) is None)
        for c in range(1, ws.max_column + 1):  # the look of the row above, its formulas copied down
            cell, above = ws.cell(r, c), ws.cell(r - 1, c)
            if above.has_style and not cell.has_style:
                cell._style = copy(above._style)
            if cell.value in (None, "") and _is_formula(above.value) and c not in typed:
                cell.value = Translator(above.value, origin=above.coordinate).translate_formula(cell.coordinate)
        for key, value in (("date", row.day), ("reporter", row.reporter), ("pages", row.pages),
                           ("witness_start", row.witness_start), ("witness_end", row.witness_end),
                           ("note", row.note)):
            # a date formula that already gives this day is kept
            if key in cols and value not in (None, "") and not (key == "date" and _is_formula(get(r, key))
                                                              and day_at(r) == row.day):
                ws.cell(r, cols[key]).value = value
                _as_text(ws.cell(r, cols[key]))
        if "start" in cols:
            cell = ws.cell(r, cols["start"])
            follows = ends.get(r - 1) is not None and row.start == ends[r - 1] + 1
            if row.start is not None and not (follows and _is_formula(cell.value)):
                cell.value = row.start  # a new numbering, no formula to work it out, or nothing above to go on
            elif row.start is None and _is_formula(cell.value) and ends.get(r - 1) is None:
                cell.value = None       # it would point at a row without pages
        if "end" in cols and not _is_formula(ws.cell(r, cols["end"]).value) and row.end() is not None:
            ws.cell(r, cols["end"]).value = row.end()
        if "weekday" in cols and not _is_formula(ws.cell(r, cols["weekday"]).value) and row.day:
            ws.cell(r, cols["weekday"]).value = row.day.strftime("%A")
        if row.end() is not None:
            ends[r] = row.end()
        used.add(r)
        last = max(last, r)
    _save(wb, path)


class RunSheets:
    """The run sheet as an output (courthouses.OutputSpec.maker), made by deliver.generate before any other
    file (group "sheet")."""

    @staticmethod
    def check(m: Making) -> None:
        """ValueError (NO_RUNSHEET) when there are no takes to add, before anything is made."""
        if not (m.runsheet and m.runsheet.rows):
            raise ValueError(NO_RUNSHEET)

    @staticmethod
    def make(m: Making) -> None:
        """Adds the takes to the case's run sheet (add_takes; a new one in runsheets_folder of m.out_dir), recorded
        with the pages of the takes added; a run sheet that had every take already is left as it was, and not
        recorded."""
        path = add_takes(m.case, m.runsheet, m.s, m.folders or [], m.out_dir)
        if m.runsheet.created or m.runsheet.added:
            m.record("runsheet", path, count=m.runsheet.added_pages)  # the pages of the takes added
